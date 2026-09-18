# -*- coding: utf-8 -*-
from __future__ import annotations

"""
GNS - ANÁLISIS DE RECHAZOS EN TRANSCRIPCIONES
Compatible con PyFlow Manager

Entrada:
  - CSV, XLSX o XLS con una fila por conversación.
  - Detecta automáticamente la columna de transcripción (por ejemplo: text).

Salida:
  1) Transcripciones: conserva las columnas originales y agrega el análisis.
  2) Pareto: motivo, casos, participación, acumulado y segmento A/B/C.

Mejoras de precisión:
  - Analiza primero las intervenciones del cliente (external/customer), no todo
    el diálogo; así evita confundir la explicación del asesor con el motivo real.
  - La conclusión de Genesys se usa como respaldo, no como verdad absoluta.
  - Separa llamadas no efectivas, terceros y cierres posiblemente incorrectos.
  - Admite FINANCIAMIENTO, TC o detección AUTO por registro.
  - Valida el XLSX final antes de publicarlo para reducir archivos dañados.

Dependencias:
  pip install pandas openpyxl

Parámetros PyFlow:
  PRODUCT_TYPE                    AUTO | FINANCIAMIENTO | TC. Predeterminado: AUTO.
  SOURCE_SHEET                    Hoja que se procesará del archivo de entrada.
  TRANSCRIPT_COLUMN               Columna que contiene la transcripción.
  EXCLUDE_NON_EFFECTIVE_PARETO    true/false. Predeterminado: false.

El script lee el único CSV/Excel de la carpeta input y guarda el reporte en output.
"""

import argparse
import json
import logging
import os
import re
import sys
import traceback
import unicodedata
import zipfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


SCRIPT_VERSION = "2.3.0"
LOGGER_NAME = "gns_analisis_rechazos_transcripciones"
MAX_EXCEL_TEXT = 32700
SUPPORTED_INPUT_EXTENSIONS = {".csv", ".xlsx", ".xls"}

PYFLOW_PARAMS = {
    "PRODUCT_TYPE": {
        "type": "select",
        "label": "Tipo de análisis: AUTO, FINANCIAMIENTO o TC",
        "required": False,
        "default": "AUTO",
        "options": ["AUTO", "FINANCIAMIENTO", "TC"]
    },
    "SOURCE_SHEET": {
        "type": "excel_sheet",
        "label": "Hoja de origen",
        "required": True,
        "default": ""
    },
    "TRANSCRIPT_COLUMN": {
        "type": "excel_column",
        "label": "Columna de transcripción",
        "required": True,
        "default": ""
    },
    "EXCLUDE_NON_EFFECTIVE_PARETO": {
        "type": "select",
        "label": "Excluir llamadas no efectivas del Pareto: true/false",
        "required": False,
        "default": "false",
        "options": ["false", "true"]
    }
}


@dataclass
class Config:
    input_file_path: str
    output_dir: str
    output_filename: str
    product_type: str
    transcript_column: str
    source_sheet: str
    exclude_non_effective_pareto: bool
    max_rows: int


@dataclass(frozen=True)
class Rule:
    motive: str
    submotive: str
    products: Tuple[str, ...]
    priority: int
    strong: Tuple[str, ...]
    medium: Tuple[str, ...] = ()
    conclusion_tokens: Tuple[str, ...] = ()
    negative: Tuple[str, ...] = ()


@dataclass
class PreparedRule:
    raw: Rule
    strong: Tuple[re.Pattern[str], ...]
    medium: Tuple[re.Pattern[str], ...]
    negative: Tuple[re.Pattern[str], ...]


# =========================================================
# LOGGER Y CONFIGURACIÓN
# =========================================================

def setup_logger() -> logging.Logger:
    logger = logging.getLogger(LOGGER_NAME)
    logger.setLevel(logging.INFO)
    logger.handlers.clear()
    logger.propagate = False
    handler = logging.StreamHandler(sys.stdout)
    handler.setFormatter(logging.Formatter(
        "%(asctime)s | %(levelname)s | %(message)s", "%Y-%m-%d %H:%M:%S"
    ))
    logger.addHandler(handler)
    return logger


def clean_env(value: Any, default: Optional[str] = None) -> Optional[str]:
    if value is None:
        return default
    text = str(value).strip()
    if text == "" or text.lower() in {"null", "none", "undefined"}:
        return default
    return text


def env_str(name: str, default: str = "", required: bool = False) -> str:
    value = clean_env(os.getenv(name), default)
    if required and not value:
        raise ValueError(f"Falta configurar el parámetro requerido: {name}")
    return str(value or "")


def parse_bool(value: Any, default: bool = False) -> bool:
    text = str(value if value is not None else "").strip().lower()
    if text == "":
        return default
    if text in {"1", "true", "t", "yes", "y", "si", "sí", "s"}:
        return True
    if text in {"0", "false", "f", "no", "n"}:
        return False
    raise ValueError(f"Valor booleano inválido: {value}")


def parse_int(value: Any, default: int = 0) -> int:
    text = str(value if value is not None else "").strip()
    if text == "":
        return default
    number = int(float(text))
    if number < 0:
        raise ValueError("MAX_ROWS no puede ser negativo")
    return number


def script_directory() -> Path:
    configured = clean_env(os.getenv("PYFLOW_SCRIPT_DIR"), "")
    if configured:
        return Path(str(configured)).resolve()
    return Path(__file__).resolve().parent


def resolve_input_file() -> Path:
    input_dir = script_directory() / "input"
    if not input_dir.exists():
        raise FileNotFoundError(f"No existe la carpeta de entrada: {input_dir}")

    files = sorted(
        (
            item
            for item in input_dir.iterdir()
            if item.is_file()
            and not item.name.startswith("~$")
            and item.suffix.lower() in SUPPORTED_INPUT_EXTENSIONS
        ),
        key=lambda item: item.name.lower(),
    )

    if not files:
        raise FileNotFoundError(
            f"No se encontró un archivo CSV, XLSX o XLS en la carpeta input: {input_dir}"
        )
    if len(files) > 1:
        names = ", ".join(item.name for item in files)
        raise ValueError(
            "La carpeta input debe contener un solo archivo de entrada. "
            f"Archivos encontrados: {names}"
        )
    return files[0]


def resolve_output_directory() -> Path:
    configured = clean_env(os.getenv("PYFLOW_OUTPUT_DIR"), "")
    if configured:
        return Path(str(configured)).resolve()
    return script_directory() / "output"


def inspect_input_schema(source_sheet: str = "") -> Dict[str, Any]:
    path = resolve_input_file()
    requested_sheet = str(source_sheet or "").strip()

    if path.suffix.lower() == ".csv":
        try:
            frame = pd.read_csv(path, encoding="utf-8-sig", nrows=0)
        except UnicodeDecodeError:
            frame = pd.read_csv(path, encoding="latin-1", nrows=0)
        sheets = ["CSV"]
        selected_sheet = "CSV"
    else:
        with pd.ExcelFile(path) as workbook:
            sheets = [str(name) for name in workbook.sheet_names]
            if not sheets:
                raise ValueError("El archivo Excel no contiene hojas disponibles.")
            selected_sheet = requested_sheet or sheets[0]
            if selected_sheet not in sheets:
                raise ValueError(f"No existe la hoja seleccionada: {selected_sheet}")
            frame = pd.read_excel(workbook, sheet_name=selected_sheet, nrows=0)

    return {
        "fileName": path.name,
        "sheets": sheets,
        "selectedSheet": selected_sheet,
        "columns": [str(column) for column in frame.columns],
    }


def normalize_product(value: Any) -> str:
    text = normalize_text(value).upper()
    aliases = {
        "": "AUTO",
        "AUTO": "AUTO",
        "AUTOMATICO": "AUTO",
        "AUTOMÁTICO": "AUTO",
        "FINANCIAMIENTO": "FINANCIAMIENTO",
        "FINANCIAMIENTOS": "FINANCIAMIENTO",
        "FIN": "FINANCIAMIENTO",
        "TC": "TC",
        "TARJETA": "TC",
        "TARJETA DE CREDITO": "TC",
        "TARJETAS DE CREDITO": "TC",
    }
    if text not in aliases:
        raise ValueError("PRODUCT_TYPE debe ser AUTO, FINANCIAMIENTO o TC")
    return aliases[text]


def load_config() -> Config:
    return Config(
        input_file_path=str(resolve_input_file()),
        output_dir=str(resolve_output_directory()),
        output_filename="",
        product_type=normalize_product(env_str("PRODUCT_TYPE", "AUTO")),
        transcript_column=env_str("TRANSCRIPT_COLUMN", ""),
        source_sheet=env_str("SOURCE_SHEET", ""),
        exclude_non_effective_pareto=parse_bool(
            env_str("EXCLUDE_NON_EFFECTIVE_PARETO", "false"), False
        ),
        max_rows=0,
    )


def log_config(logger: logging.Logger, config: Config) -> None:
    logger.info("Parámetros recibidos:")
    logger.info("- INPUT_FILE_PATH: %s", config.input_file_path)
    logger.info("- OUTPUT_DIR: %s", config.output_dir)
    logger.info("- OUTPUT_FILENAME: %s", config.output_filename or "<automático>")
    logger.info("- PRODUCT_TYPE: %s", config.product_type)
    logger.info("- TRANSCRIPT_COLUMN: %s", config.transcript_column or "<automática>")
    logger.info("- SOURCE_SHEET: %s", config.source_sheet or "<primera hoja>")
    logger.info("- EXCLUDE_NON_EFFECTIVE_PARETO: %s", config.exclude_non_effective_pareto)
    logger.info("- MAX_ROWS: %s", config.max_rows or "todos")


# =========================================================
# NORMALIZACIÓN Y ARCHIVOS
# =========================================================

ILLEGAL_XML_RE = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]")
MARKER_RE = re.compile(r"^\s*([\wáéíóúñü ._-]+?)\s*:\s*(.*)$", re.IGNORECASE)
CUSTOMER_MARKERS = {
    "external", "customer", "cliente", "client", "caller", "usuario", "user",
    "participant external", "participante externo",
}
AGENT_MARKERS = {
    "internal", "agent", "asesor", "agente", "representative", "employee",
    "participant internal", "participante interno",
}


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except Exception:
        pass
    text = str(value).lower()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = ILLEGAL_XML_RE.sub(" ", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def safe_cell(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    value = ILLEGAL_XML_RE.sub(" ", value)
    if len(value) > MAX_EXCEL_TEXT:
        value = value[: MAX_EXCEL_TEXT - 28] + "\n[TRUNCADO PARA EXCEL]"
    if value.startswith("="):
        value = "'" + value
    return value


def sanitize_dataframe_for_excel(df: pd.DataFrame) -> pd.DataFrame:
    result = df.copy()
    for column in result.columns:
        dtype = result[column].dtype
        if pd.api.types.is_object_dtype(dtype) or pd.api.types.is_string_dtype(dtype):
            result[column] = result[column].map(safe_cell)
    return result


def read_input(config: Config, logger: logging.Logger) -> pd.DataFrame:
    path = Path(config.input_file_path)
    if not path.exists():
        raise FileNotFoundError(f"No existe el archivo de entrada: {path}")

    logger.info("Leyendo archivo: %s", path)
    suffix = path.suffix.lower()
    nrows = config.max_rows or None

    if suffix == ".csv":
        try:
            df = pd.read_csv(path, encoding="utf-8-sig", low_memory=False, nrows=nrows)
        except UnicodeDecodeError:
            df = pd.read_csv(path, encoding="latin-1", low_memory=False, nrows=nrows)
    elif suffix in {".xlsx", ".xls"}:
        sheet: Any = config.source_sheet if config.source_sheet else 0
        df = pd.read_excel(path, sheet_name=sheet, nrows=nrows)
    else:
        raise ValueError("Formato no soportado. Utilice CSV, XLSX o XLS.")

    if df.empty:
        raise ValueError("El archivo de entrada no contiene registros.")
    logger.info("Registros cargados: %s", len(df))
    logger.info("Columnas cargadas: %s", len(df.columns))
    return df


def find_column(df: pd.DataFrame, candidates: Sequence[str]) -> Optional[str]:
    normalized = {normalize_text(col): col for col in df.columns}
    for candidate in candidates:
        match = normalized.get(normalize_text(candidate))
        if match is not None:
            return match
    return None


def find_transcript_column(df: pd.DataFrame, configured: str = "") -> str:
    if configured:
        exact = find_column(df, [configured])
        if exact is None:
            raise ValueError(f"No existe la columna TRANSCRIPT_COLUMN={configured}")
        return exact

    candidates = [
        "text", "texto", "transcripcion", "transcripción", "conversation_text",
        "full_text", "transcript", "transcription", "body", "frases", "phrases",
    ]
    exact = find_column(df, candidates)
    if exact:
        return exact

    for col in df.columns:
        norm = normalize_text(col)
        if any(token in norm for token in ("transcrip", "transcript", "text", "frase", "phrase")):
            return col

    raise ValueError(
        "No se encontró la columna de transcripción. Configure TRANSCRIPT_COLUMN "
        "o utilice un nombre como text, transcripcion o transcript."
    )


def extract_customer_lines(transcript: Any) -> Tuple[List[str], str]:
    raw = "" if transcript is None else str(transcript)
    customer_lines: List[str] = []
    unmarked_lines: List[str] = []
    found_markers = False

    for raw_line in raw.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        match = MARKER_RE.match(line)
        if match:
            marker = normalize_text(match.group(1))
            content = match.group(2).strip()
            if marker in CUSTOMER_MARKERS:
                found_markers = True
                if content:
                    customer_lines.append(content)
            elif marker in AGENT_MARKERS:
                found_markers = True
            else:
                unmarked_lines.append(line)
        else:
            if not line.startswith(("[{", "{\"")):
                unmarked_lines.append(line)

    # En los extractores usados por Genesys suelen agregarse, al final, frases
    # destacadas del cliente sin marcador. Se incluyen y se deduplican.
    selected = customer_lines + (unmarked_lines if found_markers else [])
    if not selected and raw.strip():
        selected = [raw.strip()]
        source = "Transcripción completa (sin marcadores)"
    else:
        source = "Intervenciones del cliente" if customer_lines else "Transcripción completa (sin marcadores)"

    unique: List[str] = []
    seen: set[str] = set()
    for line in selected:
        norm = normalize_text(line)
        if len(norm) < 2 or norm in seen:
            continue
        seen.add(norm)
        unique.append(line.strip())
    return unique, source


# =========================================================
# TAXONOMÍA DE RECHAZOS
# =========================================================

def make_rules() -> List[Rule]:
    both = ("FINANCIAMIENTO", "TC", "GENERAL")
    fin = ("FINANCIAMIENTO", "GENERAL")
    tc = ("TC", "GENERAL")

    return [
        Rule(
            "Desea analizarlo / solicita seguimiento",
            "El cliente no rechaza definitivamente y pide tiempo o una llamada posterior",
            both, 98,
            (
                r"\b(vuelva|puede|podria)\s+a?\s*llamar(me)?\b",
                r"\bllam(e|eme|enme)\b.{0,35}\b(despues|manana|tarde|otro dia|proxima semana)\b",
                r"\b(lo|la)\s+(voy a|quiero)\s+(pensar|analizar|revisar|consultar)\b",
                r"\btengo que\s+(pensarlo|analizarlo|consultarlo|hablarlo)\b",
                r"\benvi(e|eme|enme)\b.{0,30}\b(informacion|detalle|propuesta|cotizacion)\b",
            ),
            (r"\bmas adelante\b", r"\bdespues\b", r"\botro dia\b", r"\blo pensare\b"),
            ("VOLVER_LLAMAR", "SEGUIMIENTO", "INTERESADO_PENDIENTE"),
        ),
        Rule(
            "No es el momento / lo evaluaría después",
            "El cliente posterga la decisión sin indicar una objeción específica",
            both, 85,
            (
                r"\bpor (los )?momentos? no\b", r"\bpor ahora no\b", r"\bahorita no\b",
                r"\btodavia no\b", r"\ben este momento no\b", r"\bno es el momento\b",
                r"\bquizas (despues|mas adelante)\b", r"\bde momento no\b",
            ),
            (r"\bpor el momento\b", r"\bahorita\b", r"\btodavia\b", r"\bmas adelante\b"),
            ("NO_ES_EL_MOMENTO", "POR_EL_MOMENTO_NO", "LO_EVALUARA_DESPUES"),
        ),
        Rule(
            "No necesita el dinero / no tiene proyecto",
            "No identifica una necesidad, emergencia o proyecto para usar el financiamiento",
            fin, 96,
            (
                r"\bno necesito\b.{0,35}\b(dinero|financiamiento|prestamo|monto|efectivo|producto)\b",
                r"\bno (lo|la) necesito\b", r"\bno tengo (ningun )?(proyecto|necesidad|emergencia)\b",
                r"\bno tengo en que (usarlo|utilizarlo|gastarlo)\b", r"\btengo suficiente\b",
            ),
            (r"\bno necesito\b", r"\bsin proyecto\b", r"\bninguna necesidad\b"),
            ("NO_NECESITA", "NO_TIENE_PROYECTO", "NO_NECESITA_DINERO"),
        ),
        Rule(
            "Ya tiene préstamo o financiamiento vigente",
            "Ya adquirió o está pagando un financiamiento",
            fin, 94,
            (
                r"\bya tengo\b.{0,25}\b(prestamo|financiamiento|extra financiamiento|extrafinanciamiento)\b",
                r"\bestoy pagando\b.{0,25}\b(prestamo|financiamiento|cuota)\b",
                r"\b(acabo|recien) de (sacar|tomar|aceptar)\b.{0,30}\b(prestamo|financiamiento)\b",
                r"\btengo (uno|un prestamo|un financiamiento) vigente\b",
            ),
            (r"\bya tengo prestamo\b", r"\bfinanciamiento vigente\b"),
            ("YA_TIENE_PRESTAMO", "YA_TIENE_FINANCIAMIENTO", "FINANCIAMIENTO_VIGENTE"),
        ),
        Rule(
            "Tasa o costo financiero alto",
            "Percibe la tasa o los intereses como altos o poco competitivos",
            fin, 93,
            (
                r"\btasa\b.{0,28}\b(alta|elevada|cara|costosa|no me conviene)\b",
                r"\b(interes|intereses)\b.{0,28}\b(alto|altos|elevado|elevados|caro|caros|mucho)\b",
                r"\b(alta|elevada|cara)\b.{0,20}\b(tasa|interes|intereses)\b",
                r"\bpor (la )?(tasa|interes|intereses)\b", r"\bno me conviene\b.{0,25}\b(tasa|interes)\b",
            ),
            (r"\bla tasa\b", r"\blos intereses\b", r"\bcosto financiero\b"),
            ("TASA_ALTA", "TASA_NO_ATRACTIVA", "INTERES_ALTO"),
            (r"\bno tengo (ningun )?interes\b",),
        ),
        Rule(
            "Monto ofrecido no es atractivo",
            "El monto disponible es menor al que el cliente necesita o espera",
            fin, 92,
            (
                r"\bmonto\b.{0,25}\b(bajo|poco|pequeno|insuficiente|no me sirve|no alcanza)\b",
                r"\b(bajo|poco|pequeno|insuficiente)\b.{0,20}\bmonto\b",
                r"\bnecesito (un poco |mucho )?mas\b", r"\bpor ese monto no\b",
                r"\bes muy poquito\b", r"\bno me alcanza\b",
            ),
            (r"\bmonto bajo\b", r"\bmuy poco\b", r"\bnecesito mas\b"),
            ("MONTO_BAJO", "MONTO_NO_ATRACTIVO", "MONTO_INSUFICIENTE"),
        ),
        Rule(
            "Cuota alta o falta de capacidad de pago",
            "La cuota no se ajusta al presupuesto o a la situación económica",
            fin, 91,
            (
                r"\bno (puedo|podria) pagar\b", r"\bcuota\b.{0,25}\b(alta|elevada|mucho|no puedo|no me ajusta)\b",
                r"\bno tengo capacidad\b", r"\bno me ajusta (la )?cuota\b",
                r"\bmi situacion economica\b", r"\bno me da el presupuesto\b",
            ),
            (r"\bcapacidad de pago\b", r"\bpresupuesto\b", r"\bcuota mensual\b"),
            ("CUOTA_ALTA", "SIN_CAPACIDAD_PAGO", "SITUACION_ECONOMICA"),
        ),
        Rule(
            "Exceso de deudas o compromisos actuales",
            "Tiene varias obligaciones y no puede asumir una adicional",
            fin, 90,
            (
                r"\btengo (muchas|demasiadas) deudas\b", r"\bya estoy (muy )?endeudad[oa]\b",
                r"\bmuchos compromisos\b", r"\bdemasiadas cuotas\b", r"\bno puedo con (una|otra|mas) deuda\b",
                r"\btengo varios prestamos\b",
            ),
            (r"\bmuchas deudas\b", r"\bmuchos compromisos\b", r"\bvarios prestamos\b"),
            ("EXCESO_DEUDAS", "MUCHOS_COMPROMISOS", "VARIOS_PRESTAMOS"),
        ),
        Rule(
            "No quiere endeudarse",
            "Rechaza adquirir nuevas deudas o compromisos financieros",
            both, 89,
            (
                r"\bno (me )?quiero endeudar\b", r"\bno quiero (mas )?deudas?\b",
                r"\bno deseo (mas )?deudas?\b", r"\bno quiero comprometerme\b",
                r"\bno quiero (mas )?compromisos\b", r"\bno me gusta deber\b",
            ),
            (r"\bendeudarme\b", r"\bmas deudas\b", r"\bcomprometerme con deudas\b"),
            ("NO_QUIERE_COMPROMETERSE_CON_DEUDAS", "NO_QUIERE_ENDEUDARSE", "NO_MAS_DEUDAS"),
        ),
        Rule(
            "Ya tomó oferta con otra entidad",
            "Ya obtuvo el producto o una mejor oferta en otra institución",
            both, 88,
            (
                r"\bya (lo )?(tome|saque|acepte)\b.{0,35}\b(otro banco|otra entidad|bac|ficohsa|banpais|davivienda)\b",
                r"\bme dieron\b.{0,25}\b(otro banco|otra entidad)\b", r"\btengo una mejor oferta\b",
                r"\bcon otro banco\b.{0,30}\b(mejor|ya tengo|ya saque|ya tome)\b",
            ),
            (r"\botra entidad\b", r"\botro banco\b", r"\bmejor oferta\b"),
            ("OTRA_ENTIDAD", "OFERTA_OTRO_BANCO", "YA_TOMO_OTRA_OFERTA"),
        ),
        Rule(
            "Comisión, seguro u otros cargos",
            "Rechaza costos adicionales distintos de la tasa",
            both, 87,
            (
                r"\bcomision\b.{0,25}\b(alta|elevada|cara|mucho|no me conviene)\b",
                r"\bseguro\b.{0,25}\b(alto|caro|mucho|no quiero)\b",
                r"\b(cargos|costos)\b.{0,25}\b(altos|adicionales|muchos|caros)\b",
                r"\bpor (la )?(comision|seguro|cargos)\b",
            ),
            (r"\bcomision\b", r"\bseguro\b", r"\bcargos adicionales\b"),
            ("COMISION_ALTA", "SEGURO_ALTO", "OTROS_CARGOS"),
        ),
        Rule(
            "Plazo no conveniente",
            "El número de meses o las condiciones del plazo no le convienen",
            fin, 86,
            (
                r"\bplazo\b.{0,25}\b(largo|corto|no me conviene|no me sirve)\b",
                r"\bdemasiados meses\b", r"\bmuy (largo|corto) el plazo\b",
            ),
            (r"\bpor el plazo\b", r"\bplazo no\b"),
            ("PLAZO_NO_CONVENIENTE", "PLAZO_LARGO", "PLAZO_CORTO"),
        ),
        Rule(
            "Afecta el límite disponible de la tarjeta",
            "No desea que el producto reduzca o utilice el límite de su tarjeta",
            fin, 85,
            (
                r"\bafecta\b.{0,25}\blimite\b", r"\b(reduce|baja|consume|ocupa)\b.{0,25}\blimite\b",
                r"\bno quiero usar\b.{0,20}\blimite\b", r"\bme deja sin disponible\b",
            ),
            (r"\blimite disponible\b", r"\busa el limite\b"),
            ("AFECTA_LIMITE", "NO_QUIERE_USAR_LIMITE"),
        ),
        Rule(
            "Modalidad del producto no se ajusta a su necesidad",
            "Prefiere otro desembolso, producto o forma de pago",
            fin, 84,
            (
                r"\bno me sirve (asi|esa modalidad|de esa forma)\b",
                r"\bprefiero (un )?prestamo personal\b", r"\bno quiero que (salga|se cargue) de la tarjeta\b",
                r"\bnecesito (deposito|efectivo|otra modalidad)\b", r"\bese producto no se ajusta\b",
            ),
            (r"\bmodalidad\b", r"\botra forma\b", r"\bprestamo personal\b"),
            ("MODALIDAD_NO_ADECUADA", "PRODUCTO_NO_SE_AJUSTA"),
        ),
        Rule(
            "Prefiere pagar de contado / no usa financiamientos",
            "Evita financiar compras y prefiere pagos de contado",
            fin, 83,
            (
                r"\bprefiero pagar (de )?contado\b", r"\bno uso financiamientos?\b",
                r"\bno me gusta financiar\b", r"\bsolo pago (de )?contado\b",
            ),
            (r"\bde contado\b", r"\bno financio\b"),
            ("PREFIERE_CONTADO", "NO_USA_FINANCIAMIENTO"),
        ),
        Rule(
            "No usa o desea cancelar la tarjeta",
            "No utiliza la tarjeta relacionada o tiene intención de cancelarla",
            both, 82,
            (
                r"\bno uso (la|esa|mi) tarjeta\b", r"\bquiero cancelar (la|mi) tarjeta\b",
                r"\bvoy a cancelar (la|mi) tarjeta\b", r"\bno quiero (la|esa) tarjeta\b",
            ),
            (r"\bcancelar la tarjeta\b", r"\bno uso tarjeta\b"),
            ("NO_USA_TARJETA", "QUIERE_CANCELAR_TARJETA", "CANCELAR_TARJETA"),
        ),
        Rule(
            "No tiene o ya canceló la tarjeta",
            "La tarjeta vinculada al producto ya no existe o está cancelada",
            both, 81,
            (
                r"\bno tengo (esa|la|ninguna) tarjeta\b", r"\bya cancele (esa|la|mi) tarjeta\b",
                r"\bla tarjeta esta cancelada\b", r"\bya no soy cliente\b",
            ),
            (r"\btarjeta cancelada\b", r"\bno tengo tarjeta\b"),
            ("NO_TIENE_TARJETA", "TARJETA_CANCELADA", "YA_NO_ES_CLIENTE"),
        ),
        Rule(
            "Mala experiencia o desconfianza",
            "Rechaza por una experiencia negativa, temor o falta de confianza",
            both, 80,
            (
                r"\bmala experiencia\b", r"\bno confio\b", r"\bdesconfio\b",
                r"\bme (estafaron|enganaron)\b", r"\bproblema con el banco\b",
                r"\bno me resolvieron\b", r"\bmal servicio\b",
            ),
            (r"\bdesconfianza\b", r"\bmiedo\b", r"\bno es seguro\b"),
            ("MALA_EXPERIENCIA", "DESCONFIANZA", "MAL_SERVICIO"),
        ),
        Rule(
            "Saturación de llamadas o rechazo previo",
            "Solicita no recibir más llamadas o ya había rechazado la oferta",
            both, 79,
            (
                r"\bno me llamen (mas|de nuevo)\b", r"\bdejen de llamarme\b",
                r"\bya me han llamado\b", r"\bmuchas llamadas\b", r"\bya (lo )?rechace\b",
            ),
            (r"\bno mas llamadas\b", r"\bya dije que no\b"),
            ("NO_LLAMAR", "SATURACION_LLAMADAS", "RECHAZO_PREVIO"),
        ),
        # Reglas específicas de tarjeta de crédito.
        Rule(
            "Ya posee suficientes tarjetas / producto similar",
            "Ya cuenta con una o varias tarjetas y no desea otra",
            tc, 97,
            (
                r"\bya tengo\b.{0,25}\b(muchas|varias|suficientes) tarjetas\b",
                r"\bya tengo tarjeta\b", r"\bno necesito otra tarjeta\b",
                r"\btengo suficientes tarjetas\b",
            ),
            (r"\botra tarjeta\b", r"\bya tengo una\b"),
            ("YA_TIENE_TARJETA", "YA_TIENE_MUCHAS_TC", "NO_NECESITA_OTRA_TARJETA", "SUFICIENTES_TARJETAS"),
        ),
        Rule(
            "No necesita tarjeta de crédito",
            "No identifica necesidad de adquirir o usar una tarjeta",
            tc, 95,
            (
                r"\bno necesito\b.{0,25}\b(tarjeta|credito)\b", r"\bno me hace falta\b.{0,20}\btarjeta\b",
                r"\bno necesito ese producto\b",
            ),
            (r"\bno necesito tarjeta\b", r"\bno me hace falta\b"),
            ("NO_NECESITA_TARJETA", "NO_NECESITA_PRODUCTO"),
        ),
        Rule(
            "Prefiere débito, efectivo o no usa crédito",
            "Prefiere medios de pago distintos a una tarjeta de crédito",
            tc, 94,
            (
                r"\bprefiero (la )?tarjeta de debito\b", r"\bsolo uso debito\b",
                r"\bprefiero efectivo\b", r"\bno uso credito\b", r"\bno me gusta usar credito\b",
            ),
            (r"\bdebito\b", r"\befectivo\b", r"\bno uso tarjetas de credito\b"),
            ("PREFIERE_DEBITO", "PREFIERE_EFECTIVO", "NO_USA_CREDITO"),
        ),
        Rule(
            "Anualidad, membresía o cargos de la tarjeta",
            "Rechaza el costo anual, membresía o cargos asociados",
            tc, 93,
            (
                r"\b(anualidad|membresia|cuota anual)\b.{0,25}\b(alta|cara|mucho|no quiero|no me conviene)\b",
                r"\bno quiero pagar\b.{0,25}\b(anualidad|membresia|cuota anual)\b",
            ),
            (r"\banualidad\b", r"\bmembresia\b", r"\bcuota anual\b"),
            ("ANUALIDAD_ALTA", "MEMBRESIA_ALTA", "CARGOS_TARJETA"),
        ),
        Rule(
            "Límite de crédito ofrecido no es atractivo",
            "El límite disponible es menor al que el cliente espera",
            tc, 92,
            (
                r"\blimite\b.{0,25}\b(bajo|poco|pequeno|insuficiente|no me sirve)\b",
                r"\bese limite no\b", r"\bnecesito un limite mayor\b",
            ),
            (r"\blimite bajo\b", r"\bpoco limite\b"),
            ("LIMITE_BAJO", "LIMITE_NO_ATRACTIVO"),
        ),
        Rule(
            "Beneficios de la tarjeta no son atractivos",
            "No percibe valor en puntos, descuentos, millas o beneficios",
            tc, 91,
            (
                r"\bbeneficios?\b.{0,30}\b(no me sirven|pocos|malos|no me interesan|no son atractivos)\b",
                r"\bno le veo beneficio\b", r"\bno me sirven (los )?(puntos|millas|descuentos)\b",
                r"\bno acumula\b.{0,20}\b(puntos|millas)\b",
            ),
            (r"\bpuntos\b", r"\bmillas\b", r"\bbeneficios\b"),
            ("BENEFICIOS_NO_ATRACTIVOS", "NO_LE_SIRVEN_PUNTOS"),
        ),
        Rule(
            "No desea realizar el trámite o entregar documentación",
            "Rechaza el proceso, tiempo o documentación requerida",
            tc, 90,
            (
                r"\bno quiero (hacer|realizar) (el )?tramite\b", r"\bmuchos documentos\b",
                r"\bno quiero enviar\b.{0,25}\b(documentos|constancia|informacion)\b",
                r"\bno tengo tiempo\b.{0,25}\b(tramite|documentos|proceso)\b",
            ),
            (r"\btramite\b", r"\bdocumentacion\b", r"\bmuchos requisitos\b"),
            ("NO_DESEA_TRAMITE", "NO_ENTREGA_DOCUMENTOS"),
        ),
        Rule(
            "No cumple requisitos o ingresos",
            "Indica que no cumple los requisitos, ingresos o documentación",
            tc, 89,
            (
                r"\bno cumplo\b.{0,25}\b(requisitos|ingresos|condiciones)\b",
                r"\bno tengo\b.{0,25}\b(constancia|ingresos|trabajo fijo|antiguedad)\b",
                r"\bno soy asalariad[oa]\b",
            ),
            (r"\bno cumplo requisitos\b", r"\bsin ingresos\b"),
            ("NO_CUMPLE_REQUISITOS", "SIN_INGRESOS", "SIN_DOCUMENTACION"),
        ),
        Rule(
            "Producto no le genera valor o beneficio",
            "No percibe utilidad suficiente en la oferta",
            both, 60,
            (
                r"\bno le veo (ningun )?(valor|beneficio|utilidad)\b",
                r"\bno me genera (ningun )?(valor|beneficio)\b", r"\bno me sirve para nada\b",
            ),
            (r"\bno me conviene\b", r"\bno me sirve\b"),
            ("NO_GENERA_VALOR", "NO_LE_VE_BENEFICIO"),
        ),
        Rule(
            "No le interesa, sin explicar el motivo",
            "Rechaza la oferta sin expresar una causa específica",
            both, 20,
            (
                r"\bno me interesa\b", r"\bno tengo (ningun )?interes\b", r"\bno estoy interesad[oa]\b",
                r"\bgracias,? pero no\b", r"\bno,? gracias\b",
            ),
            (r"\bno quiero\b", r"\bno deseo\b", r"\bdefinitivamente no\b"),
            ("NO_LE_INTERESA", "SIN_MOTIVO", "RECHAZO_SIN_MOTIVO"),
        ),
    ]


def compile_rules(rules: Iterable[Rule]) -> List[PreparedRule]:
    prepared: List[PreparedRule] = []
    for rule in rules:
        prepared.append(PreparedRule(
            raw=rule,
            strong=tuple(re.compile(p, re.IGNORECASE) for p in rule.strong),
            medium=tuple(re.compile(p, re.IGNORECASE) for p in rule.medium),
            negative=tuple(re.compile(p, re.IGNORECASE) for p in rule.negative),
        ))
    return prepared


PREPARED_RULES = compile_rules(make_rules())


# =========================================================
# CONTROLES DE CONTACTO Y CONCLUSIONES
# =========================================================

NON_EFFECTIVE_MOTIVE = "Sin motivo identificable / llamada no efectiva"
THIRD_PARTY_MOTIVE = "Cliente no disponible o habló un tercero"
AMBIGUOUS_MOTIVE = "Otros motivos específicos o transcripción ambigua"
WRONG_CLOSE_MOTIVE = "No rechazo definitivo / posible cierre incorrecto"
GENERIC_REJECTION_MOTIVE = "No le interesa, sin explicar el motivo"

NON_EFFECTIVE_CONCLUSIONS = (
    "NO_CONTESTA", "NO_CONTESTAN", "BUZON", "VOICEMAIL", "OCUPADO", "BUSY",
    "LLAMADA_CORTA", "CALLBACK_MACHINE", "NO_CONTACTADO", "SIN_CONTACTO",
    "ABANDONADA", "NUMERO_FUERA_SERVICIO", "FUERA_DE_SERVICIO",
)
THIRD_PARTY_CONCLUSIONS = (
    "TERCERO", "NUMERO_EQUIVOCADO", "NO_ES_EL_CLIENTE", "PERSONA_EQUIVOCADA",
    "FALLECIDO", "FALLECIO", "MENOR_DE_EDAD",
)

INTEREST_PATTERNS = tuple(re.compile(p, re.IGNORECASE) for p in (
    r"\bsi\s+(?:me\s+)?interesa\b",
    r"(?<!no )\bestoy interesad[oa]\b",
    r"(?<!no )\bquiero (tomarlo|aceptarlo|solicitarlo|la tarjeta)\b",
    r"\bcomo (lo|la) solicito\b", r"\bque requisitos\b", r"\bcuanto pagaria\b",
    r"\bprocedamos\b", r"\benvieme la informacion\b",
))

NEGATIVE_INTEREST_PATTERNS = tuple(re.compile(p, re.IGNORECASE) for p in (
    r"\bno\s+(?:estoy|estaria)\s+interesad[oa]\b",
    r"\bno\s+(?:me\s+)?interesa\b",
    r"\bno\s+tengo\s+(?:ningun\s+)?interes\b",
    r"\bno\s+(?:lo\s+|la\s+)?(?:quiero|deseo|acepto|solicito)\b",
    r"\b(?:por ahora|por el momento|de momento|ahorita|ahora|todavia|en este momento)\s+no\b",
    r"\bno\s+(?:ahorita|ahora|por ahora|por el momento|de momento)\b",
    r"\bgracias\s+pero\s+no\b",
    r"\bpara nada\s+(?:estoy\s+)?interesad[oa]\b",
    r"\bno\s*,?\s*gracias\b",
    r"\bdefinitivamente\s+no\b",
))

THIRD_PARTY_TEXT_PATTERNS = tuple(re.compile(p, re.IGNORECASE) for p in (
    r"\bno soy (el|la)\b", r"\bse equivoco de numero\b", r"\bnumero equivocado\b",
    r"\baqui no vive\b", r"\bhabla (su|la|el) (esposa|esposo|mama|madre|padre|hijo|hija|hermano|hermana)\b",
    r"\bno se encuentra\b", r"\bno esta disponible\b",
))


def joined_context(row: pd.Series, columns: Sequence[Optional[str]]) -> str:
    values: List[str] = []
    for col in columns:
        if col and col in row.index:
            value = row.get(col)
            if value is not None and not (isinstance(value, float) and pd.isna(value)):
                values.append(str(value))
    return " | ".join(values)


def conclusion_key(value: Any) -> str:
    text = normalize_text(value).upper()
    return re.sub(r"[^A-Z0-9]+", "_", text).strip("_")


def contains_any_token(text: str, tokens: Sequence[str]) -> bool:
    return any(token in text for token in tokens)


def conclusion_motive(conclusion: str, product: str) -> str:
    key = conclusion_key(conclusion)
    if not key:
        return ""

    mappings: List[Tuple[Tuple[str, ...], str]] = [
        (("VOLVER_LLAMAR", "SEGUIMIENTO", "INTERESADO_PENDIENTE"), "Desea analizarlo / solicita seguimiento"),
        (("NO_ES_EL_MOMENTO", "POR_EL_MOMENTO_NO", "LO_EVALUARA_DESPUES"), "No es el momento / lo evaluaría después"),
        (("NO_NECESITA_DINERO", "NO_TIENE_PROYECTO"), "No necesita el dinero / no tiene proyecto"),
        (("YA_TIENE_PRESTAMO", "YA_TIENE_FINANCIAMIENTO", "FINANCIAMIENTO_VIGENTE"), "Ya tiene préstamo o financiamiento vigente"),
        (("TASA_ALTA", "TASA_NO_ATRACTIVA", "INTERES_ALTO"), "Tasa o costo financiero alto"),
        (("MONTO_BAJO", "MONTO_NO_ATRACTIVO", "MONTO_INSUFICIENTE"), "Monto ofrecido no es atractivo"),
        (("CUOTA_ALTA", "SIN_CAPACIDAD_PAGO", "SITUACION_ECONOMICA"), "Cuota alta o falta de capacidad de pago"),
        (("EXCESO_DEUDAS", "MUCHOS_COMPROMISOS", "VARIOS_PRESTAMOS"), "Exceso de deudas o compromisos actuales"),
        (("NO_QUIERE_COMPROMETERSE_CON_DEUDAS", "NO_QUIERE_ENDEUDARSE", "NO_MAS_DEUDAS"), "No quiere endeudarse"),
        (("OTRA_ENTIDAD", "OFERTA_OTRO_BANCO", "YA_TOMO_OTRA_OFERTA"), "Ya tomó oferta con otra entidad"),
        (("COMISION_ALTA", "SEGURO_ALTO", "OTROS_CARGOS"), "Comisión, seguro u otros cargos"),
        (("PLAZO_NO_CONVENIENTE", "PLAZO_LARGO", "PLAZO_CORTO"), "Plazo no conveniente"),
        (("AFECTA_LIMITE", "NO_QUIERE_USAR_LIMITE"), "Afecta el límite disponible de la tarjeta"),
        (("NO_USA_TARJETA", "QUIERE_CANCELAR_TARJETA", "CANCELAR_TARJETA"), "No usa o desea cancelar la tarjeta"),
        (("NO_TIENE_TARJETA", "TARJETA_CANCELADA", "YA_NO_ES_CLIENTE"), "No tiene o ya canceló la tarjeta"),
        (("MALA_EXPERIENCIA", "DESCONFIANZA", "MAL_SERVICIO"), "Mala experiencia o desconfianza"),
        (("NO_LLAMAR", "SATURACION_LLAMADAS", "RECHAZO_PREVIO"), "Saturación de llamadas o rechazo previo"),
        (("YA_TIENE_TARJETA", "YA_TIENE_MUCHAS_TC", "NO_NECESITA_OTRA_TARJETA", "SUFICIENTES_TARJETAS"), "Ya posee suficientes tarjetas / producto similar"),
        (("NO_NECESITA_TARJETA", "NO_NECESITA_PRODUCTO"), "No necesita tarjeta de crédito"),
        (("PREFIERE_DEBITO", "PREFIERE_EFECTIVO", "NO_USA_CREDITO"), "Prefiere débito, efectivo o no usa crédito"),
        (("ANUALIDAD_ALTA", "MEMBRESIA_ALTA", "CARGOS_TARJETA"), "Anualidad, membresía o cargos de la tarjeta"),
        (("LIMITE_BAJO", "LIMITE_NO_ATRACTIVO"), "Límite de crédito ofrecido no es atractivo"),
        (("BENEFICIOS_NO_ATRACTIVOS", "NO_LE_SIRVEN_PUNTOS"), "Beneficios de la tarjeta no son atractivos"),
        (("NO_DESEA_TRAMITE", "NO_ENTREGA_DOCUMENTOS"), "No desea realizar el trámite o entregar documentación"),
        (("NO_CUMPLE_REQUISITOS", "SIN_INGRESOS", "SIN_DOCUMENTACION"), "No cumple requisitos o ingresos"),
    ]
    for tokens, motive in mappings:
        if contains_any_token(key, tokens):
            return motive

    if "NO_NECESITA" in key:
        return "No necesita tarjeta de crédito" if product == "TC" else "No necesita el dinero / no tiene proyecto"
    if "NO_LE_INTERESA" in key or "SIN_MOTIVO" in key:
        return "No le interesa, sin explicar el motivo"
    return ""


def detect_product(configured: str, row_context: str, client_text: str) -> str:
    if configured != "AUTO":
        return configured

    context = normalize_text(row_context + " " + client_text)
    fin_signals = (
        "extra financiamiento", "extrafinanciamiento", "intra financiamiento",
        "corta cuotas", "corta cuota", "financiamiento", "prestamo", "televentasmp",
        "monto en efectivo", "desembolso",
    )
    tc_signals = (
        "tc cuentas nuevas", "tarjeta de credito nueva", "nueva tarjeta de credito",
        "solicitud de tarjeta", "oferta de tarjeta", "cuentas nuevas tc",
    )
    fin_score = sum(1 for item in fin_signals if item in context)
    tc_score = sum(1 for item in tc_signals if item in context)
    if fin_score > tc_score and fin_score > 0:
        return "FINANCIAMIENTO"
    if tc_score > fin_score and tc_score > 0:
        return "TC"
    return "GENERAL"


def contact_control(
    client_text: str,
    conclusion: str,
    transcript_status: str,
    phrase_count: Any,
) -> Optional[Dict[str, Any]]:
    conclusion_norm = conclusion_key(conclusion)
    status_norm = normalize_text(transcript_status).upper()

    if contains_any_token(conclusion_norm, THIRD_PARTY_CONCLUSIONS) or any(
        pattern.search(client_text) for pattern in THIRD_PARTY_TEXT_PATTERNS
    ):
        return {
            "motive": THIRD_PARTY_MOTIVE,
            "submotive": "No se logró conversar con el titular o el número no corresponde",
            "confidence": "Alta" if contains_any_token(conclusion_norm, THIRD_PARTY_CONCLUSIONS) else "Media",
            "score": 100,
            "source": "Control de contacto",
            "evidence": "Conclusión Genesys: " + conclusion if conclusion else first_evidence_line(client_text),
        }

    bad_status = status_norm not in {"", "OK", "COMPLETADO", "SUCCESS"}
    no_phrases = False
    try:
        no_phrases = int(float(phrase_count)) <= 0
    except Exception:
        pass

    if contains_any_token(conclusion_norm, NON_EFFECTIVE_CONCLUSIONS) or bad_status or no_phrases or len(client_text) < 3:
        reason = "No hubo conversación suficiente con el cliente para identificar un rechazo"
        evidence = ""
        if conclusion:
            evidence = "Conclusión Genesys: " + conclusion
        elif transcript_status:
            evidence = "Estado transcripción: " + transcript_status
        return {
            "motive": NON_EFFECTIVE_MOTIVE,
            "submotive": reason,
            "confidence": "Alta" if contains_any_token(conclusion_norm, NON_EFFECTIVE_CONCLUSIONS) else "Media",
            "score": 100,
            "source": "Control de contacto",
            "evidence": evidence,
        }
    return None


def first_evidence_line(text: str) -> str:
    for line in str(text).splitlines():
        if line.strip():
            return line.strip()[:350]
    return ""


def evidence_from_lines(lines: Sequence[str], patterns: Sequence[re.Pattern[str]]) -> str:
    for pattern in patterns:
        for line in lines:
            if pattern.search(normalize_text(line)):
                return safe_cell(line.strip())[:500]
    return ""


def final_interest_state(text: str) -> str:
    """Devuelve la última intención explícita: positive, negative o vacío."""
    events: List[Tuple[int, int, str]] = []

    for pattern in INTEREST_PATTERNS:
        for match in pattern.finditer(text):
            events.append((match.end(), 0, "positive"))

    for pattern in NEGATIVE_INTEREST_PATTERNS:
        for match in pattern.finditer(text):
            # En un empate, la negación debe prevalecer.
            events.append((match.end(), 1, "negative"))

    if not events:
        return ""
    return max(events, key=lambda event: (event[0], event[1]))[2]


# =========================================================
# CLASIFICACIÓN
# =========================================================

def classify_row(
    row: pd.Series,
    transcript_col: str,
    configured_product: str,
    conclusion_col: Optional[str],
    queue_col: Optional[str],
    campaign_col: Optional[str],
    status_col: Optional[str],
    phrase_count_col: Optional[str],
) -> Dict[str, Any]:
    transcript = row.get(transcript_col, "")
    lines, text_source = extract_customer_lines(transcript)
    client_original = "\n".join(lines)
    client_norm = normalize_text(client_original)

    conclusion = str(row.get(conclusion_col, "") or "") if conclusion_col else ""
    transcript_status = str(row.get(status_col, "") or "") if status_col else ""
    phrase_count = row.get(phrase_count_col, None) if phrase_count_col else None
    row_context = joined_context(row, [queue_col, campaign_col, conclusion_col])
    product = detect_product(configured_product, row_context, client_original)

    conclusion_suggested = conclusion_motive(conclusion, product)
    control = contact_control(client_norm, conclusion, transcript_status, phrase_count)
    if control:
        motive = control["motive"]
        return {
            "Motivo principal identificado": motive,
            "Submotivo identificado": control["submotive"],
            "Evidencia del cliente": control["evidence"],
            "Confianza de clasificación": control["confidence"],
            "Puntaje clasificación": control["score"],
            "Fuente de clasificación": control["source"],
            "Tipo de producto detectado": product,
            "Motivo sugerido por conclusión": conclusion_suggested,
            "Coincide con conclusión": compare_conclusion(motive, conclusion_suggested),
            "Requiere revisión manual": "Sí" if motive in {NON_EFFECTIVE_MOTIVE, THIRD_PARTY_MOTIVE} else "No",
            "Fuente de texto analizado": text_source,
            "Intervenciones del cliente": len(lines),
            "Versión clasificador": SCRIPT_VERSION,
        }

    # Solo se marca cierre incorrecto cuando la última intención explícita es positiva.
    interest = final_interest_state(client_norm) == "positive"
    conclusion_no_interest = "NO_LE_INTERESA" in conclusion_key(conclusion)
    if interest and conclusion_no_interest:
        motive = WRONG_CLOSE_MOTIVE
        return {
            "Motivo principal identificado": motive,
            "Submotivo identificado": "El cliente mostró interés o solicitó condiciones, pero la conclusión indica rechazo",
            "Evidencia del cliente": evidence_from_lines(lines, INTEREST_PATTERNS),
            "Confianza de clasificación": "Alta",
            "Puntaje clasificación": 100,
            "Fuente de clasificación": "Texto del cliente + validación de conclusión",
            "Tipo de producto detectado": product,
            "Motivo sugerido por conclusión": conclusion_suggested,
            "Coincide con conclusión": "No",
            "Requiere revisión manual": "Sí",
            "Fuente de texto analizado": text_source,
            "Intervenciones del cliente": len(lines),
            "Versión clasificador": SCRIPT_VERSION,
        }

    candidates: List[Dict[str, Any]] = []
    conclusion_key_value = conclusion_key(conclusion)

    for prepared in PREPARED_RULES:
        rule = prepared.raw
        if product == "GENERAL":
            if "GENERAL" not in rule.products:
                continue
        elif product not in rule.products:
            continue

        if any(pattern.search(client_norm) for pattern in prepared.negative):
            continue

        strong_matches = [p for p in prepared.strong if p.search(client_norm)]
        medium_matches = [p for p in prepared.medium if p.search(client_norm)]
        conclusion_hit = any(token in conclusion_key_value for token in rule.conclusion_tokens)

        if not strong_matches and not medium_matches and not conclusion_hit:
            continue

        # El texto del cliente domina. La conclusión solo refuerza o sirve de fallback.
        score = min(len(strong_matches), 3) * 20 + min(len(medium_matches), 3) * 7
        if conclusion_hit:
            score += 5 if (strong_matches or medium_matches) else 12

        # Una sola palabra genérica no debe crear una clasificación fuerte.
        if not strong_matches and len(medium_matches) == 1 and not conclusion_hit:
            score -= 2

        candidates.append({
            "rule": rule,
            "strong": strong_matches,
            "medium": medium_matches,
            "conclusion_hit": conclusion_hit,
            "score": score,
        })

    if not candidates:
        if conclusion_suggested:
            motive = conclusion_suggested
            return {
                "Motivo principal identificado": motive,
                "Submotivo identificado": "Clasificación basada únicamente en la conclusión de Genesys; validar manualmente",
                "Evidencia del cliente": "Conclusión Genesys: " + conclusion,
                "Confianza de clasificación": "Baja",
                "Puntaje clasificación": 12,
                "Fuente de clasificación": "Conclusión de Genesys (fallback)",
                "Tipo de producto detectado": product,
                "Motivo sugerido por conclusión": conclusion_suggested,
                "Coincide con conclusión": "Sí",
                "Requiere revisión manual": "Sí",
                "Fuente de texto analizado": text_source,
                "Intervenciones del cliente": len(lines),
                "Versión clasificador": SCRIPT_VERSION,
            }

        return {
            "Motivo principal identificado": AMBIGUOUS_MOTIVE,
            "Submotivo identificado": "La transcripción no contiene evidencia suficiente para asignar una causa específica",
            "Evidencia del cliente": first_evidence_line(client_original),
            "Confianza de clasificación": "Baja",
            "Puntaje clasificación": 0,
            "Fuente de clasificación": "Sin evidencia suficiente",
            "Tipo de producto detectado": product,
            "Motivo sugerido por conclusión": "",
            "Coincide con conclusión": "No aplica",
            "Requiere revisión manual": "Sí",
            "Fuente de texto analizado": text_source,
            "Intervenciones del cliente": len(lines),
            "Versión clasificador": SCRIPT_VERSION,
        }

    # Jerarquia de evidencia:
    # 1) motivo especifico expresado por el cliente;
    # 2) conclusion especifica de Genesys cuando el texto solo dice "no me interesa";
    # 3) rechazo generico del cliente.
    used_specific_conclusion = False
    specific_text_candidates = [
        item for item in candidates
        if item["rule"].motive != GENERIC_REJECTION_MOTIVE and item["strong"]
    ]
    generic_text_present = any(
        item["rule"].motive == GENERIC_REJECTION_MOTIVE
        and (item["strong"] or item["medium"])
        for item in candidates
    )
    specific_conclusion_candidates = [
        item for item in candidates
        if conclusion_suggested
        and conclusion_suggested != GENERIC_REJECTION_MOTIVE
        and item["rule"].motive == conclusion_suggested
        and item["conclusion_hit"]
    ]

    if specific_text_candidates:
        candidate_pool = specific_text_candidates
    elif generic_text_present and specific_conclusion_candidates:
        candidate_pool = specific_conclusion_candidates
        used_specific_conclusion = True
    else:
        candidate_pool = candidates

    candidate_pool.sort(
        key=lambda item: (
            item["score"],
            len(item["strong"]),
            item["rule"].priority,
        ),
        reverse=True,
    )
    winner = candidate_pool[0]
    second_score = candidate_pool[1]["score"] if len(candidate_pool) > 1 else -999
    margin = winner["score"] - second_score
    rule = winner["rule"]

    if winner["strong"] and winner["score"] >= 20 and margin >= 7:
        confidence = "Alta"
    elif winner["score"] >= 12:
        confidence = "Media"
    else:
        confidence = "Baja"

    patterns_for_evidence: Sequence[re.Pattern[str]] = winner["strong"] or winner["medium"]
    evidence = evidence_from_lines(lines, patterns_for_evidence)
    if not evidence and winner["conclusion_hit"]:
        evidence = "Conclusión Genesys: " + conclusion

    if used_specific_conclusion:
        source = "Rechazo general + conclusión específica"
    elif winner["conclusion_hit"] and (winner["strong"] or winner["medium"]):
        source = "Texto del cliente + conclusión"
    elif winner["strong"] or winner["medium"]:
        source = "Texto del cliente"
    else:
        source = "Conclusión de Genesys (fallback)"

    motive = rule.motive
    review = (
        confidence == "Baja"
        or margin < 4
        or used_specific_conclusion
        or motive in {AMBIGUOUS_MOTIVE, WRONG_CLOSE_MOTIVE}
    )

    return {
        "Motivo principal identificado": motive,
        "Submotivo identificado": rule.submotive,
        "Evidencia del cliente": evidence,
        "Confianza de clasificación": confidence,
        "Puntaje clasificación": int(winner["score"]),
        "Fuente de clasificación": source,
        "Tipo de producto detectado": product,
        "Motivo sugerido por conclusión": conclusion_suggested,
        "Coincide con conclusión": compare_conclusion(motive, conclusion_suggested),
        "Requiere revisión manual": "Sí" if review else "No",
        "Fuente de texto analizado": text_source,
        "Intervenciones del cliente": len(lines),
        "Versión clasificador": SCRIPT_VERSION,
    }


def compare_conclusion(motive: str, conclusion_suggested: str) -> str:
    if not conclusion_suggested:
        return "No aplica"
    return "Sí" if motive == conclusion_suggested else "No"


def analyze_dataframe(df: pd.DataFrame, config: Config, logger: logging.Logger) -> pd.DataFrame:
    transcript_col = find_transcript_column(df, config.transcript_column)
    conclusion_col = find_column(df, [
        "conclusion_ultima", "conclusión_ultima", "conclusion_original", "conclusión_original",
        "wrapup", "wrapUpName", "wrapup_name", "conclusion",
    ])
    queue_col = find_column(df, ["queueName", "queue_name", "cola", "nombre_cola"])
    campaign_col = find_column(df, ["CampaignName", "campaignName", "campaign_name", "campaña", "campana"])
    status_col = find_column(df, ["transcript_estado", "transcript_status", "estado_transcripcion"])
    phrase_count_col = find_column(df, ["phrases_count", "phrase_count", "cantidad_frases"])

    logger.info("Columna de transcripción: %s", transcript_col)
    logger.info("Columna de conclusión: %s", conclusion_col or "<no encontrada>")
    logger.info("Columna de cola: %s", queue_col or "<no encontrada>")
    logger.info("Columna de campaña: %s", campaign_col or "<no encontrada>")

    results: List[Dict[str, Any]] = []
    total = len(df)
    for position, (_, row) in enumerate(df.iterrows(), start=1):
        results.append(classify_row(
            row=row,
            transcript_col=transcript_col,
            configured_product=config.product_type,
            conclusion_col=conclusion_col,
            queue_col=queue_col,
            campaign_col=campaign_col,
            status_col=status_col,
            phrase_count_col=phrase_count_col,
        ))
        if position % 2500 == 0 or position == total:
            logger.info("Clasificación: %s/%s registros", position, total)

    analysis = pd.DataFrame(results, index=df.index)
    output = df.copy()
    for column in analysis.columns:
        output[column] = analysis[column]
    return output


# =========================================================
# PARETO Y EXCEL
# =========================================================

def pareto_segment(participation: float, cumulative: float) -> str:
    previous = cumulative - participation
    if previous < 0.80:
        return "A"
    if previous < 0.95:
        return "B"
    return "C"


def build_pareto(base: pd.DataFrame, exclude_non_effective: bool) -> pd.DataFrame:
    working = base.copy()
    if exclude_non_effective:
        working = working[~working["Motivo principal identificado"].isin({
            NON_EFFECTIVE_MOTIVE, THIRD_PARTY_MOTIVE,
        })]

    counts = (
        working.groupby("Motivo principal identificado", dropna=False)
        .size()
        .reset_index(name="Casos")
        .rename(columns={"Motivo principal identificado": "Motivo de rechazo"})
        .sort_values(["Casos", "Motivo de rechazo"], ascending=[False, True])
        .reset_index(drop=True)
    )
    total = int(counts["Casos"].sum())
    counts.insert(0, "Rango", range(1, len(counts) + 1))
    counts["Participación"] = counts["Casos"] / total if total else 0.0
    counts["Casos acumulados"] = counts["Casos"].cumsum()
    counts["Participación acumulada"] = counts["Participación"].cumsum()
    counts["Segmento Pareto"] = [
        pareto_segment(float(part), float(cum))
        for part, cum in zip(counts["Participación"], counts["Participación acumulada"])
    ]

    total_row = pd.DataFrame([{
        "Rango": "",
        "Motivo de rechazo": "TOTAL",
        "Casos": total,
        "Participación": 1.0 if total else 0.0,
        "Casos acumulados": total,
        "Participación acumulada": 1.0 if total else 0.0,
        "Segmento Pareto": "",
    }])
    return pd.concat([counts, total_row], ignore_index=True)


def default_output_name(config: Config) -> str:
    if config.output_filename.strip():
        filename = config.output_filename.strip()
    else:
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"Analisis_Rechazos_{stamp}.xlsx"
    if not filename.lower().endswith(".xlsx"):
        filename += ".xlsx"
    return filename


def style_workbook(path: Path) -> None:
    wb = load_workbook(path)
    blue = "1F4E78"
    white = "FFFFFF"
    light_blue = "D9EAF7"
    orange = "F4B183"

    ws = wb["Transcripciones"]
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 32
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=blue)
        cell.font = Font(color=white, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    analysis_headers = {
        "Motivo principal identificado", "Submotivo identificado", "Evidencia del cliente",
        "Confianza de clasificación", "Puntaje clasificación", "Fuente de clasificación",
        "Tipo de producto detectado", "Motivo sugerido por conclusión", "Coincide con conclusión",
        "Requiere revisión manual", "Fuente de texto analizado", "Intervenciones del cliente",
        "Versión clasificador",
    }
    for col_idx, cell in enumerate(ws[1], start=1):
        header = str(cell.value or "")
        if header in analysis_headers:
            cell.fill = PatternFill("solid", fgColor="2F75B5")

        norm = normalize_text(header)
        width = 18
        if "conversationid" in norm or norm.endswith("id") or "communication" in norm:
            width = 38
        elif "fecha" in norm or "date" in norm or "start" in norm or "end" in norm:
            width = 21
        elif header == "Motivo principal identificado":
            width = 44
        elif header == "Submotivo identificado":
            width = 55
        elif header == "Evidencia del cliente":
            width = 55
        elif header in {"Fuente de clasificación", "Motivo sugerido por conclusión", "Fuente de texto analizado"}:
            width = 36
        elif "text" in norm or "transcrip" in norm:
            width = 80
        elif "conclusion" in norm or "queue" in norm or "campaign" in norm:
            width = 35
        ws.column_dimensions[get_column_letter(col_idx)].width = min(width, 80)

    # Solo las columnas textuales del análisis se envuelven; aplicar bordes a 45 mil
    # filas generaría libros pesados y aumenta el riesgo de corrupción.
    header_index = {str(c.value or ""): c.column for c in ws[1]}
    for name in ("Motivo principal identificado", "Submotivo identificado", "Evidencia del cliente"):
        col = header_index.get(name)
        if col:
            for cells in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=col, max_col=col):
                cells[0].alignment = Alignment(vertical="top", wrap_text=True)

    pws = wb["Pareto"]
    pws.freeze_panes = "A2"
    pws.auto_filter.ref = f"A1:G{max(1, pws.max_row - 1)}"
    pws.row_dimensions[1].height = 28
    for cell in pws[1]:
        cell.fill = PatternFill("solid", fgColor=blue)
        cell.font = Font(color=white, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    pws.column_dimensions["A"].width = 10
    pws.column_dimensions["B"].width = 55
    pws.column_dimensions["C"].width = 16
    pws.column_dimensions["D"].width = 18
    pws.column_dimensions["E"].width = 20
    pws.column_dimensions["F"].width = 24
    pws.column_dimensions["G"].width = 18
    pws["D2"].number_format = "0.00%"
    pws["F2"].number_format = "0.00%"
    for row in range(2, pws.max_row + 1):
        pws.cell(row, 4).number_format = "0.00%"
        pws.cell(row, 6).number_format = "0.00%"
        segment = pws.cell(row, 7).value
        if segment == "A":
            pws.cell(row, 7).fill = PatternFill("solid", fgColor=light_blue)
        elif segment == "B":
            pws.cell(row, 7).fill = PatternFill("solid", fgColor="FFF2CC")
        elif segment == "C":
            pws.cell(row, 7).fill = PatternFill("solid", fgColor=orange)

    if pws.max_row >= 2:
        for cell in pws[pws.max_row]:
            cell.fill = PatternFill("solid", fgColor=blue)
            cell.font = Font(color=white, bold=True)

    wb.save(path)
    wb.close()


def validate_xlsx(path: Path, expected_rows: int) -> None:
    with zipfile.ZipFile(path, "r") as archive:
        broken = archive.testzip()
        if broken:
            raise ValueError(f"El XLSX contiene una entrada dañada: {broken}")

    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        required = {"Transcripciones", "Pareto"}
        missing = required.difference(wb.sheetnames)
        if missing:
            raise ValueError(f"Faltan hojas en el XLSX final: {sorted(missing)}")
        rows = wb["Transcripciones"].max_row - 1
        if rows != expected_rows:
            raise ValueError(f"Validación de filas falló: esperadas={expected_rows}, generadas={rows}")
    finally:
        wb.close()


def write_output(base: pd.DataFrame, pareto: pd.DataFrame, config: Config, logger: logging.Logger) -> Path:
    output_dir = Path(config.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    final_path = output_dir / default_output_name(config)
    temp_path = output_dir / f".{final_path.stem}.tmp.xlsx"

    if temp_path.exists():
        temp_path.unlink()

    logger.info("Preparando datos seguros para Excel...")
    safe_base = sanitize_dataframe_for_excel(base)
    safe_pareto = sanitize_dataframe_for_excel(pareto)

    logger.info("Escribiendo Excel temporal: %s", temp_path)
    with pd.ExcelWriter(temp_path, engine="openpyxl") as writer:
        safe_base.to_excel(writer, sheet_name="Transcripciones", index=False)
        safe_pareto.to_excel(writer, sheet_name="Pareto", index=False)

    logger.info("Aplicando formato ligero y estable...")
    style_workbook(temp_path)

    logger.info("Validando estructura interna del XLSX...")
    validate_xlsx(temp_path, len(base))

    if final_path.exists():
        final_path.unlink()
    os.replace(temp_path, final_path)
    return final_path


# =========================================================
# PROCESO PRINCIPAL
# =========================================================

def run(config: Config, logger: logging.Logger) -> Path:
    df = read_input(config, logger)
    base = analyze_dataframe(df, config, logger)
    logger.info("Generando Pareto...")
    pareto = build_pareto(base, config.exclude_non_effective_pareto)
    output = write_output(base, pareto, config, logger)

    review_count = int((base["Requiere revisión manual"] == "Sí").sum())
    high_count = int((base["Confianza de clasificación"] == "Alta").sum())
    medium_count = int((base["Confianza de clasificación"] == "Media").sum())
    low_count = int((base["Confianza de clasificación"] == "Baja").sum())
    logger.info("Registros: %s", len(base))
    logger.info("Confianza alta/media/baja: %s / %s / %s", high_count, medium_count, low_count)
    logger.info("Revisión manual sugerida: %s", review_count)
    logger.info("Archivo validado: %s", output)
    return output


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Análisis de rechazos en transcripciones para PyFlow Manager")
    parser.add_argument("--product-type", "--PRODUCT_TYPE", dest="product_type", default=env_str("PRODUCT_TYPE", "AUTO"))
    parser.add_argument("--source-sheet", "--SOURCE_SHEET", dest="source_sheet", default=env_str("SOURCE_SHEET", ""))
    parser.add_argument("--transcript-column", "--TRANSCRIPT_COLUMN", dest="transcript_column", default=env_str("TRANSCRIPT_COLUMN", ""))
    parser.add_argument(
        "--exclude-non-effective-pareto", "--EXCLUDE_NON_EFFECTIVE_PARETO",
        dest="exclude_non_effective_pareto", default=env_str("EXCLUDE_NON_EFFECTIVE_PARETO", "false"),
    )
    parser.add_argument(
        "--inspect-input-schema",
        action="store_true",
        help=argparse.SUPPRESS,
    )
    return parser


def config_from_args(args: argparse.Namespace) -> Config:
    return Config(
        input_file_path=str(resolve_input_file()),
        output_dir=str(resolve_output_directory()),
        output_filename="",
        product_type=normalize_product(args.product_type),
        transcript_column=str(args.transcript_column or ""),
        source_sheet=str(args.source_sheet or ""),
        exclude_non_effective_pareto=parse_bool(args.exclude_non_effective_pareto, False),
        max_rows=0,
    )


def main() -> int:
    logger = setup_logger()
    started = datetime.now()
    try:
        args = build_parser().parse_args()
        if args.inspect_input_schema:
            try:
                schema = inspect_input_schema(args.source_sheet)
                print("PYFLOW_INPUT_SCHEMA=" + json.dumps(schema, ensure_ascii=False))
                return 0
            except Exception as exc:
                print(f"PYFLOW_INPUT_SCHEMA_ERROR={exc}")
                return 2
        config = config_from_args(args)
        logger.info("=" * 88)
        logger.info("INICIO ANÁLISIS DE RECHAZOS EN TRANSCRIPCIONES | versión %s", SCRIPT_VERSION)
        log_config(logger, config)
        logger.info("=" * 88)
        output = run(config, logger)
        duration = (datetime.now() - started).total_seconds()
        logger.info("=" * 88)
        logger.info("PROCESO FINALIZADO CORRECTAMENTE")
        logger.info("Archivo generado: %s", output)
        logger.info("Duración total: %.2f segundos", duration)
        logger.info("=" * 88)
        return 0
    except Exception as exc:
        logger.exception("El proceso terminó con error: %s", exc)
        logger.error(traceback.format_exc())
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
