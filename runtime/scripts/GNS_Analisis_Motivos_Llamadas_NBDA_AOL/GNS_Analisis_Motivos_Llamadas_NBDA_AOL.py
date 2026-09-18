# -*- coding: utf-8 -*-
"""
GNS_Analisis_Motivos_Llamadas_PyFlow.py

Script para PyFlow Manager:
- Lee transcripciones de llamadas desde CSV o Excel.
- Clasifica cada llamada en un motivo raíz y submotivo raíz.
- Genera un resumen tipo Pareto como el análisis mostrado en Excel:
  Motivo raíz | Submotivo raíz | Llamadas | % | % acumulado
- Exporta un Excel con resumen y detalle, y un HTML ejecutivo.

Modo de análisis:
1) rules  : Clasificación local por reglas/palabras clave. No usa internet ni API externa.
2) llm    : Clasificación con endpoint OpenAI-compatible. Requiere LLM_API_KEY.
3) hybrid : Usa reglas primero; si la confianza es baja, consulta LLM. Requiere LLM_API_KEY.

Autor: generado para PyFlow Manager
"""

from __future__ import annotations

import argparse
import base64
import datetime as _dt
import hashlib
import html
import json
import mimetypes
import os
import re
import sys
import time
import unicodedata
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

try:
    import pandas as pd
except ImportError as exc:
    raise SystemExit("[ERROR] Falta instalar pandas. Instale con: pip install pandas openpyxl") from exc

try:
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    load_workbook = None


# -----------------------------------------------------------------------------
# Parámetros para PyFlow Manager
# -----------------------------------------------------------------------------
PYFLOW_PARAMS = {
    "DATE_FROM": {
        "label": "Fecha inicial",
        "type": "date",
        "required": False,
        "default": "",
        "description": "Filtra llamadas desde esta fecha. Vacio = sin filtro inicial.",
    },
    "DATE_TO": {
        "label": "Fecha final",
        "type": "date",
        "required": False,
        "default": "",
        "description": "Filtra llamadas hasta esta fecha inclusive. Vacio = sin filtro final.",
    },
    "EMAIL_TO": {
        "type": "tags",
        "label": "Para",
        "required": False,
        "description": "Destinatarios del correo. Si queda vacio, no se envia correo.",
    },
    "EMAIL_CC": {
        "type": "tags",
        "label": "CC",
        "required": False,
        "description": "Copias del correo.",
    },
    "EMAIL_SUBJECT": {
        "type": "text",
        "label": "Asunto",
        "required": False,
        "default": "Analisis de motivos de llamadas NBDA AOL",
        "description": "Asunto del correo. Puede modificarse antes de ejecutar.",
    },
    "GRAPH_TENANT_ID": {
        "type": "global",
        "global_key": "GRAPH_TENANT_ID",
        "label": "Microsoft Graph Tenant ID",
        "required": False,
    },
    "GRAPH_CLIENT_ID": {
        "type": "global",
        "global_key": "GRAPH_CLIENT_ID",
        "label": "Microsoft Graph Client ID",
        "required": False,
    },
    "GRAPH_CLIENT_SECRET": {
        "type": "global",
        "global_key": "GRAPH_CLIENT_SECRET",
        "label": "Microsoft Graph Client Secret",
        "required": False,
        "secret": True,
    },
    "GRAPH_SENDER_EMAIL": {
        "type": "global",
        "global_key": "GRAPH_SENDER_EMAIL",
        "label": "Correo remitente Graph",
        "required": False,
    },
}

HIDDEN_CLI_PARAMS = [
    "INPUT_FILE",
    "OUTPUT_DIR",
    "SHEET_NAME",
    "TEXT_COLUMN",
    "CALL_ID_COLUMN",
    "DATE_COLUMN",
    "CLIENT_COLUMN",
    "CONCLUSION_COLUMN",
    "GROUP_BY_CALL_ID",
    "ANALYSIS_MODE",
    "MIN_CONFIDENCE_LLM",
    "LLM_API_KEY",
    "LLM_BASE_URL",
    "LLM_MODEL",
    "MAX_CHARS_LLM",
]


# -----------------------------------------------------------------------------
# Taxonomía de motivos raíz
# Ajuste aquí si más adelante desea agregar/quitar motivos.
# -----------------------------------------------------------------------------
TAXONOMY: List[Dict[str, Any]] = [
    {
        "motivo": "Migración a Atlántida HN / nueva app",
        "submotivo": "Credenciales o flujo anterior no funcionan en nueva plataforma",
        "priority": 80,
        "patterns": [
            (r"nueva\s+(aplicaci[o??]n|app)", 80),
            (r"va\s+a\s+reemplazar|reemplaza(r|ndo)?", 28),
            (r"nueva\s+plataforma", 18),
            (r"banca\s+digital\s+nueva|nueva\s+banca\s+digital", 14),
            (r"nueva\s+banca\s+en\s+l[i??]nea|banca\s+en\s+l[i??]nea\s+nueva", 18),
            (r"\bmigraci[oó]n\b", 6),
            (r"\bmigrar\b|\bmigrando\b|\bmigrado\b", 5),
            (r"atl[aá]ntida\s*hn", 7),
            (r"nueva\s+app|aplicaci[oó]n\s+nueva|app\s+nueva", 6),
            (r"banca\s+digital\s+nueva|nueva\s+banca\s+digital", 5),
            (r"credenciales\s+anteriores|usuario\s+anterior|contrase[nñ]a\s+anterior", 5),
            (r"no\s+(me\s+)?funciona(n)?\s+(mis\s+)?credenciales", 5),
            (r"antes\s+ingresaba|antes\s+entraba|ya\s+no\s+puedo\s+ingresar", 4),
            (r"cambio\s+de\s+plataforma|nueva\s+plataforma", 5),
        ],
    },
    {
        "motivo": "No recibe código OTP / SMS / Token",
        "submotivo": "El código de verificación no llega o falla el token",
        "priority": 95,
        "patterns": [
            (r"no\s+((me|le|les|nos)\s+)?llega\s+(el\s+)?c[oó]digo", 10),
            (r"no\s+.*llega.*c[oó]digo", 8),
            (r"no\s+(he\s+)?recibido\s+(el\s+)?c[oó]digo", 9),
            (r"c[oó]digo\s+(no\s+)?llega", 8),
            (r"no\s+recibe\s+c[oó]digo", 9),
            (r"reenviar\s+(c[o??]digo|contrase[n??]a|clave)|enviar\s+(c[o??]digo|contrase[n??]a|clave)", 9),
            (r"(c[o??]digo|clave|contrase[n??]a)\s+.*(correo|tel[e??]fono|celular)", 7),
            (r"correo\s+.*(c[o??]digo|clave|contrase[n??]a)", 7),
            (r"c[o??]digo\s+de\s+seguridad", 7),
            (r"otp|token|sms|mensaje\s+de\s+texto", 5),
            (r"c[oó]digo\s+de\s+verificaci[oó]n", 6),
            (r"c[oó]digo\s+inv[aá]lido|token\s+inv[aá]lido|otp\s+inv[aá]lido", 7),
            (r"reenviar\s+c[oó]digo|enviar\s+otro\s+c[oó]digo", 5),
            (r"correo\s+.*c[oó]digo|c[oó]digo\s+.*correo", 4),
            (r"celular\s+.*c[oó]digo|tel[eé]fono\s+.*c[oó]digo", 4),
        ],
    },
    {
        "motivo": "Actualización de datos / teléfono o correo asociado",
        "submotivo": "Datos registrados no permiten validar identidad o recuperar acceso",
        "priority": 90,
        "patterns": [
            (r"actualizar\s+(mis\s+)?datos|actualizaci[oó]n\s+de\s+datos", 8),
            (r"actualizar\s+(mi\s+)?(tel[eé]fono|numero|n[uú]mero|correo)", 8),
            (r"(tel[eé]fono|numero|n[uú]mero|correo)\s+.*(asociado|incorrecto|desactualizado)", 7),
            (r"datos\s+(des)?actualizados|datos\s+incorrectos", 7),
            (r"correo\s+(des)?actualizado|correo\s+incorrecto|cambiar\s+correo", 7),
            (r"tel[eé]fono\s+(des)?actualizado|n[uú]mero\s+(des)?actualizado|cambiar\s+(mi\s+)?n[uú]mero", 7),
            (r"no\s+coincide(n)?\s+(mis\s+)?datos", 6),
            (r"validar\s+(mi\s+)?identidad|validaci[oó]n\s+de\s+identidad", 5),
            (r"recuperar\s+acceso\s+.*datos|datos\s+.*recuperar\s+acceso", 5),
            (r"tel[eé]fono\s+asociado|correo\s+asociado", 6),
        ],
    },
    {
        "motivo": "Cambio de dispositivo / reinstalación",
        "submotivo": "El cliente cambió celular, reinstaló la app o perdió acceso desde otro equipo",
        "priority": 75,
        "patterns": [
            (r"cambi[eé]\s+(de\s+)?celular|cambio\s+de\s+celular|cambio\s+de\s+tel[eé]fono", 8),
            (r"nuevo\s+celular|nuevo\s+tel[eé]fono|nuevo\s+dispositivo", 7),
            (r"cambio\s+de\s+dispositivo", 8),
            (r"reinstal[eé]\s+la\s+app|reinstalaci[oó]n|instal[eé]\s+de\s+nuevo", 7),
            (r"perd[ií]\s+(mi\s+)?celular|me\s+robaron\s+(el\s+)?celular", 7),
            (r"otro\s+equipo|otro\s+dispositivo", 5),
            (r"vincular\s+(mi\s+)?dispositivo|desvincular\s+dispositivo", 6),
        ],
    },
    {
        "motivo": "Error de app/web o autenticación",
        "submotivo": "La app/web muestra error, no permite avanzar o rechaza credenciales",
        "priority": 70,
        "patterns": [
            (r"error\s+(en\s+)?(la\s+)?app|app\s+.*error", 8),
            (r"error\s+(en\s+)?(la\s+)?p[aá]gina|web\s+.*error|p[aá]gina\s+web\s+.*error", 8),
            (r"no\s+(me\s+)?deja\s+(avanzar|continuar|ingresar|entrar)", 7),
            (r"no\s+puedo\s+(ingresar|entrar|acceder)", 5),
            (r"rechaza\s+(mis\s+)?credenciales|credenciales\s+inv[aá]lidas", 7),
            (r"usuario\s+o\s+contrase[nñ]a\s+incorrect", 6),
            (r"autenticaci[oó]n|autenticar|autenticando", 5),
            (r"pantalla\s+en\s+blanco|se\s+queda\s+cargando|no\s+carga", 6),
            (r"app\s+no\s+funciona|p[aá]gina\s+no\s+funciona|web\s+no\s+funciona", 8),
        ],
    },
    {
        "motivo": "Usuario bloqueado / requiere desbloqueo",
        "submotivo": "El usuario está bloqueado o el cliente solicita desbloqueo",
        "priority": 92,
        "patterns": [
            (r"usuario\s+bloqueado|bloque[oó]\s+de\s+usuario", 9),
            (r"me\s+bloque[oó]|se\s+bloque[oó]\s+(mi\s+)?usuario", 8),
            (r"desbloquear\s+(mi\s+)?usuario|desbloqueo\s+de\s+usuario", 9),
            (r"cuenta\s+bloqueada|acceso\s+bloqueado", 7),
            (r"intentos\s+fallidos|demasiados\s+intentos", 6),
        ],
    },
    {
        "motivo": "Contraseña temporal vencida / recuperación incompleta",
        "submotivo": "Clave temporal o enlace de recuperación expiró antes de completar el proceso",
        "priority": 85,
        "patterns": [
            (r"contrase[nñ]a\s+temporal|clave\s+temporal", 1),
            (r"clave\s+vencida|contrase[nñ]a\s+vencida", 8),
            (r"temporal\s+venci[oó]|temporal\s+expir[oó]", 8),
            (r"enlace\s+(de\s+)?recuperaci[oó]n\s+(vencido|expirado)", 8),
            (r"link\s+(vencido|expirado)", 7),
            (r"recuperaci[oó]n\s+incompleta|no\s+complet[eó]\s+la\s+recuperaci[oó]n", 6),
            (r"restablecer\s+contrase[nñ]a|recuperar\s+contrase[nñ]a", 1),
        ],
    },
    {
        "motivo": "Consulta operativa / validación general NBDA",
        "submotivo": "La llamada no es claramente cambio de contraseña; es soporte general de banca digital",
        "priority": 10,
        "patterns": [
            (r"consulta\s+general|soporte\s+general", 5),
            (r"banca\s+digital", 3),
            (r"nbda", 4),
            (r"validaci[oó]n\s+general", 4),
            (r"informaci[oó]n\s+general", 4),
        ],
    },
]

OTHER_MOTIVO = "Otros / no clasificado"
OTHER_SUBMOTIVO = "No se encontró una señal suficiente en la transcripción"

VALID_MOTIVES = {(t["motivo"], t["submotivo"]) for t in TAXONOMY}


# -----------------------------------------------------------------------------
# Utilidades generales
# -----------------------------------------------------------------------------
def log(msg: str) -> None:
    print(f"[{_dt.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}] {msg}", flush=True)


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and pd.isna(value):
        return ""
    text = str(value)
    text = text.replace("\u00a0", " ")
    text = re.sub(r"\s+", " ", text).strip()
    return text


def strip_accents(text: str) -> str:
    nfkd = unicodedata.normalize("NFKD", text)
    return "".join(ch for ch in nfkd if not unicodedata.combining(ch))


def text_for_matching(text: str) -> str:
    text = normalize_text(text).lower()
    text = strip_accents(text)
    return text


def parse_bool(value: Any, default: bool = False) -> bool:
    if value is None:
        return default
    txt = str(value).strip().lower()
    if txt in {"1", "true", "t", "yes", "y", "si", "sí", "s"}:
        return True
    if txt in {"0", "false", "f", "no", "n"}:
        return False
    return default


def parse_float(value: Any, default: float) -> float:
    try:
        if value is None or str(value).strip() == "":
            return default
        return float(str(value).replace(",", "."))
    except Exception:
        return default


def parse_int(value: Any, default: int) -> int:
    try:
        if value is None or str(value).strip() == "":
            return default
        return int(float(str(value).replace(",", ".")))
    except Exception:
        return default


def split_list_value(value: Any) -> List[str]:
    if value is None:
        return []
    if isinstance(value, list):
        raw_items = value
    else:
        text = str(value).strip()
        if not text:
            return []
        try:
            parsed = json.loads(text)
            raw_items = parsed if isinstance(parsed, list) else [text]
        except Exception:
            raw_items = re.split(r"[;,]", text)

    result: List[str] = []
    seen = set()
    for item in raw_items:
        email = normalize_text(item).strip()
        if email and email.lower() not in seen:
            seen.add(email.lower())
            result.append(email)
    return result


def parse_cli_args() -> Dict[str, str]:
    """
    Soporta parámetros tipo:
      --INPUT_FILE "C:/archivo.csv"
      --input_file "C:/archivo.csv"
      INPUT_FILE="C:/archivo.csv"
    """
    parser = argparse.ArgumentParser(add_help=True)
    visible_names = list(PYFLOW_PARAMS.keys()) if isinstance(PYFLOW_PARAMS, dict) else [p["name"] for p in PYFLOW_PARAMS]
    known_names = list(dict.fromkeys(visible_names + HIDDEN_CLI_PARAMS))
    for name in known_names:
        parser.add_argument(f"--{name}", dest=name, required=False)
        parser.add_argument(f"--{name.lower()}", dest=name, required=False)

    parsed, unknown = parser.parse_known_args()
    result = {k: v for k, v in vars(parsed).items() if v is not None}

    # Soporte simple KEY=VALUE si PyFlow o ejecución manual lo envía así.
    for item in unknown:
        if "=" in item and not item.startswith("--"):
            key, val = item.split("=", 1)
            key = key.strip().upper()
            if key in known_names:
                result[key] = val.strip().strip('"')
    return result


def get_param(name: str, cli: Dict[str, str], default: Any = None) -> Any:
    # Orden de prioridad: CLI -> ENV exacto -> ENV con PYFLOW_ -> default del PYFLOW_PARAMS -> default recibido.
    if name in cli and cli[name] not in (None, ""):
        return cli[name]
    if name in os.environ and os.environ[name] != "":
        return os.environ[name]
    prefixed = f"PYFLOW_{name}"
    if prefixed in os.environ and os.environ[prefixed] != "":
        return os.environ[prefixed]
    if isinstance(PYFLOW_PARAMS, dict):
        item_default = PYFLOW_PARAMS.get(name, {}).get("default", None)
        if item_default not in (None, ""):
            return item_default
    else:
        for item in PYFLOW_PARAMS:
            if item.get("name") == name:
                item_default = item.get("default", None)
                if item_default not in (None, ""):
                    return item_default
    return default


def norm_col_name(col: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "", strip_accents(str(col)).lower())


def detect_column(df: pd.DataFrame, requested: str, candidates: List[str], purpose: str, required: bool) -> Optional[str]:
    if requested and str(requested).strip() and str(requested).strip().upper() != "AUTO":
        req = str(requested).strip()
        if req in df.columns:
            return req
        # comparación flexible
        req_norm = norm_col_name(req)
        for col in df.columns:
            if norm_col_name(col) == req_norm:
                return col
        raise ValueError(f"No se encontró la columna configurada para {purpose}: {req}")

    candidate_norms = [norm_col_name(c) for c in candidates]
    for target in candidate_norms:
        for col in df.columns:
            if norm_col_name(col) == target:
                return col

    # búsqueda contiene
    for col in df.columns:
        ncol = norm_col_name(col)
        if any(target in ncol or ncol in target for target in candidate_norms):
            return col

    if required:
        raise ValueError(
            f"No pude detectar automáticamente la columna para {purpose}. "
            f"Columnas disponibles: {list(df.columns)}"
        )
    return None


def read_input_file(path: Path, sheet_name: str = "") -> pd.DataFrame:
    if not path.exists():
        raise FileNotFoundError(f"No existe el archivo de entrada: {path}")

    suffix = path.suffix.lower()
    log(f"Leyendo archivo de entrada: {path}")

    if suffix in {".xlsx", ".xlsm", ".xls"}:
        if sheet_name and sheet_name.strip():
            df = pd.read_excel(path, sheet_name=sheet_name.strip(), dtype=str)
        else:
            df = pd.read_excel(path, sheet_name=0, dtype=str)
        return df

    if suffix in {".csv", ".txt"}:
        encodings = ["utf-8-sig", "utf-8", "latin1", "cp1252"]
        last_error: Optional[Exception] = None
        for enc in encodings:
            try:
                return pd.read_csv(path, sep=None, engine="python", dtype=str, encoding=enc)
            except Exception as exc:
                last_error = exc
        raise ValueError(f"No pude leer el CSV/TXT con codificaciones comunes. Último error: {last_error}")

    raise ValueError(f"Formato no soportado: {suffix}. Use CSV o Excel.")


def resolve_input_file(configured_path: Path) -> Path:
    """Usa INPUT_FILE si viene informado; si no, toma el archivo mas reciente de input."""
    if configured_path and str(configured_path) not in {".", ""}:
        return configured_path

    script_dir = Path(__file__).resolve().parent
    input_dir = script_dir / "input"
    supported = {".xlsx", ".xlsm", ".xls", ".csv", ".txt"}

    if not input_dir.exists():
        raise FileNotFoundError(f"No existe la carpeta input esperada: {input_dir}")

    candidates = [
        p for p in input_dir.iterdir()
        if p.is_file() and p.suffix.lower() in supported and not p.name.startswith("~$")
    ]
    if not candidates:
        raise FileNotFoundError(
            f"No se encontraron archivos CSV/XLSX en la carpeta input: {input_dir}"
        )

    selected = max(candidates, key=lambda p: p.stat().st_mtime)
    log(f"INPUT_FILE vacio. Se usara automaticamente el archivo mas reciente de input: {selected}")
    return selected


def resolve_output_dir(configured_path: Path) -> Path:
    """Usa OUTPUT_DIR si viene informado; si no, guarda en output junto al script."""
    if configured_path and str(configured_path) not in {".", ""}:
        return configured_path
    return Path(__file__).resolve().parent / "output"


def apply_date_filter(df: pd.DataFrame, date_col: Optional[str], date_from: str, date_to: str) -> pd.DataFrame:
    if not date_col:
        if date_from or date_to:
            log("[WARN] Se indicó rango de fechas, pero no se detectó columna de fecha. No se aplicará filtro.")
        return df

    if not date_from and not date_to:
        return df

    work = df.copy()
    work["__fecha_parseada"] = pd.to_datetime(work[date_col], errors="coerce", dayfirst=False)
    before = len(work)

    if date_from:
        start = pd.to_datetime(date_from, errors="coerce")
        if pd.isna(start):
            raise ValueError(f"DATE_FROM inválida: {date_from}. Use YYYY-MM-DD.")
        work = work[work["__fecha_parseada"] >= start]

    if date_to:
        end = pd.to_datetime(date_to, errors="coerce")
        if pd.isna(end):
            raise ValueError(f"DATE_TO inválida: {date_to}. Use YYYY-MM-DD.")
        # inclusive hasta final del día
        work = work[work["__fecha_parseada"] < (end + pd.Timedelta(days=1))]

    after = len(work)
    log(f"Filtro de fechas aplicado sobre {date_col}: {before:,} -> {after:,} filas")
    work = work.drop(columns=["__fecha_parseada"])
    return work


def aggregate_by_call(
    df: pd.DataFrame,
    text_col: str,
    call_id_col: Optional[str],
    date_col: Optional[str],
    client_col: Optional[str],
    conclusion_col: Optional[str],
    group_by_call_id: bool,
) -> pd.DataFrame:
    work = df.copy()
    work[text_col] = work[text_col].fillna("").astype(str)

    def first_non_empty(values: Any) -> str:
        iterable = values if hasattr(values, "__iter__") and not isinstance(values, str) else [values]
        for value in iterable:
            text = normalize_text(value)
            if text and text.lower() not in {"nan", "none", "null"}:
                return text
        return ""

    if not group_by_call_id or not call_id_col:
        out = pd.DataFrame()
        out["Call ID"] = work[call_id_col] if call_id_col else range(1, len(work) + 1)
        out["Fecha"] = work[date_col] if date_col else ""
        out["Cliente"] = work[client_col] if client_col else ""
        out["Conclusión"] = work[conclusion_col].map(first_non_empty) if conclusion_col else ""
        out["Transcripción"] = work[text_col]
        return out

    log(f"Agrupando segmentos por ID de llamada/conversación: {call_id_col}")
    work[call_id_col] = work[call_id_col].fillna("").astype(str).map(normalize_text)
    # Si vienen IDs vacíos, no se deben juntar todos en una sola llamada.
    blank_ids = work[call_id_col].eq("") | work[call_id_col].str.lower().isin({"nan", "none", "null"})
    if blank_ids.any():
        work.loc[blank_ids, call_id_col] = [f"SIN_ID_{i+1}" for i in range(blank_ids.sum())]
    work["__row_order"] = range(len(work))

    agg_dict: Dict[str, Any] = {
        text_col: lambda s: "\n".join([normalize_text(x) for x in s if normalize_text(x)]),
        "__row_order": "min",
    }
    if date_col:
        agg_dict[date_col] = "min"
    if client_col:
        agg_dict[client_col] = "first"
    if conclusion_col:
        agg_dict[conclusion_col] = first_non_empty

    grouped = work.groupby(call_id_col, dropna=False).agg(agg_dict).reset_index()
    grouped = grouped.sort_values("__row_order")

    out = pd.DataFrame()
    out["Call ID"] = grouped[call_id_col]
    out["Fecha"] = grouped[date_col] if date_col else ""
    out["Cliente"] = grouped[client_col] if client_col else ""
    out["Conclusión"] = grouped[conclusion_col] if conclusion_col else ""
    out["Transcripción"] = grouped[text_col]

    log(f"Llamadas únicas después de agrupar: {len(out):,}")
    return out


# -----------------------------------------------------------------------------
# Clasificación local por reglas
# -----------------------------------------------------------------------------
def classify_rules(transcript: str) -> Dict[str, Any]:
    clean = text_for_matching(transcript)
    if not clean or len(clean) < 8:
        return {
            "motivo": OTHER_MOTIVO,
            "submotivo": OTHER_SUBMOTIVO,
            "confidence": 0.0,
            "mode": "rules",
            "signals": "texto vacío o demasiado corto",
        }

    results: List[Tuple[float, int, Dict[str, Any], List[str]]] = []
    for tax in TAXONOMY:
        score = 0.0
        signals: List[str] = []
        for pattern, weight in tax["patterns"]:
            # Los patrones tienen acentos; comparamos en texto sin acentos, por eso normalizamos también el patrón.
            pattern_norm = strip_accents(pattern.lower())
            matches = re.findall(pattern_norm, clean, flags=re.IGNORECASE)
            if matches:
                # Evita inflar demasiado por repetición; máximo 3 repeticiones por patrón.
                add = weight * min(len(matches), 3)
                score += add
                signals.append(pattern)
        if score > 0:
            results.append((score, int(tax.get("priority", 0)), tax, signals))

    if not results:
        return {
            "motivo": OTHER_MOTIVO,
            "submotivo": OTHER_SUBMOTIVO,
            "confidence": 0.25,
            "mode": "rules",
            "signals": "sin coincidencias fuertes",
        }

    results.sort(key=lambda x: (x[0], x[1]), reverse=True)
    best_score, best_priority, best_tax, best_signals = results[0]

    # Penalización si hay empate muy cercano entre categorías.
    second_score = results[1][0] if len(results) > 1 else 0.0
    margin = best_score - second_score
    confidence = 0.45 + min(best_score / 25.0, 0.45) + min(max(margin, 0) / 25.0, 0.10)
    confidence = max(0.0, min(confidence, 0.98))

    return {
        "motivo": best_tax["motivo"],
        "submotivo": best_tax["submotivo"],
        "confidence": round(confidence, 4),
        "mode": "rules",
        "signals": "; ".join(best_signals[:8]),
    }


# -----------------------------------------------------------------------------
# Clasificación opcional con LLM OpenAI-compatible
# -----------------------------------------------------------------------------
def build_llm_prompt(transcript: str) -> List[Dict[str, str]]:
    taxonomy_text = "\n".join(
        f"- Motivo raíz: {t['motivo']} | Submotivo raíz: {t['submotivo']}" for t in TAXONOMY
    )
    system = (
        "Eres un analista de contact center bancario. "
        "Tu tarea es clasificar transcripciones de llamadas en una taxonomía fija. "
        "No inventes categorías fuera de la taxonomía. "
        "Devuelve solo JSON válido."
    )
    user = f"""
Clasifica la siguiente transcripción en el motivo raíz y submotivo raíz más probable.

Taxonomía permitida:
{taxonomy_text}
- Motivo raíz: {OTHER_MOTIVO} | Submotivo raíz: {OTHER_SUBMOTIVO}

Criterios:
- Usa "No recibe código OTP / SMS / Token" si el problema central es que no llega o falla el código/token/SMS/correo de verificación.
- Usa "Migración a Atlántida HN / nueva app" si el problema central es la nueva app/plataforma, migración o credenciales anteriores que ya no funcionan.
- Usa "Actualización de datos / teléfono o correo asociado" si el problema central son datos, correo o teléfono desactualizados/incorrectos.
- Usa "Error de app/web o autenticación" si la app/web muestra error, no permite avanzar o rechaza credenciales sin que sea claramente OTP, datos o migración.
- Usa "Otros / no clasificado" solo si no hay señal suficiente.

Devuelve este JSON exacto:
{{
  "motivo_raiz": "...",
  "submotivo_raiz": "...",
  "confianza": 0.0,
  "justificacion_breve": "máximo 20 palabras"
}}

Transcripción:
--- INICIO TRANSCRIPCIÓN ---
{transcript}
--- FIN TRANSCRIPCIÓN ---
"""
    return [
        {"role": "system", "content": system},
        {"role": "user", "content": user},
    ]


def call_llm_classification(
    transcript: str,
    api_key: str,
    base_url: str,
    model: str,
    max_chars: int,
    timeout_seconds: int = 60,
    max_retries: int = 3,
) -> Dict[str, Any]:
    try:
        import requests
    except ImportError as exc:
        raise RuntimeError("Falta instalar requests para usar ANALYSIS_MODE=llm/hybrid: pip install requests") from exc

    text = normalize_text(transcript)
    if len(text) > max_chars:
        text = text[:max_chars] + "\n[TRUNCADO]"

    url = base_url.rstrip("/") + "/chat/completions"
    headers = {
        "Authorization": f"Bearer {api_key}",
        "Content-Type": "application/json",
    }
    payload = {
        "model": model,
        "messages": build_llm_prompt(text),
        "temperature": 0,
        "response_format": {"type": "json_object"},
    }

    last_error: Optional[Exception] = None
    for attempt in range(1, max_retries + 1):
        try:
            response = requests.post(url, headers=headers, json=payload, timeout=timeout_seconds)
            if response.status_code >= 400:
                raise RuntimeError(f"HTTP {response.status_code}: {response.text[:500]}")
            data = response.json()
            content = data["choices"][0]["message"]["content"]
            parsed = json.loads(content)

            motivo = normalize_text(parsed.get("motivo_raiz", ""))
            submotivo = normalize_text(parsed.get("submotivo_raiz", ""))
            confianza = parse_float(parsed.get("confianza", 0.0), 0.0)
            justificacion = normalize_text(parsed.get("justificacion_breve", ""))

            if (motivo, submotivo) not in VALID_MOTIVES and motivo != OTHER_MOTIVO:
                # Si el LLM devuelve una categoría inválida, se fuerza a otros.
                motivo, submotivo, confianza = OTHER_MOTIVO, OTHER_SUBMOTIVO, min(confianza, 0.40)
                justificacion = "Categoría devuelta fuera de taxonomía"

            if motivo == OTHER_MOTIVO:
                submotivo = OTHER_SUBMOTIVO

            return {
                "motivo": motivo or OTHER_MOTIVO,
                "submotivo": submotivo or OTHER_SUBMOTIVO,
                "confidence": round(max(0.0, min(confianza, 1.0)), 4),
                "mode": "llm",
                "signals": justificacion,
            }
        except Exception as exc:
            last_error = exc
            wait = min(2 ** attempt, 10)
            log(f"[WARN] Error LLM intento {attempt}/{max_retries}: {exc}. Reintentando en {wait}s...")
            time.sleep(wait)

    raise RuntimeError(f"No fue posible clasificar con LLM. Último error: {last_error}")


def classify_transcript(
    transcript: str,
    mode: str,
    min_confidence_llm: float,
    llm_config: Dict[str, Any],
    cache: Dict[str, Dict[str, Any]],
) -> Dict[str, Any]:
    mode = (mode or "rules").strip().lower()
    rules_result = classify_rules(transcript)

    if mode == "rules":
        return rules_result

    api_key = normalize_text(llm_config.get("api_key", ""))
    if not api_key:
        log("[WARN] ANALYSIS_MODE requiere LLM_API_KEY, pero no se encontró. Se usará rules.")
        return rules_result

    should_call_llm = mode == "llm" or (mode == "hybrid" and rules_result["confidence"] < min_confidence_llm)
    if not should_call_llm:
        rules_result["mode"] = "hybrid_rules"
        return rules_result

    digest = hashlib.sha256(normalize_text(transcript).encode("utf-8", errors="ignore")).hexdigest()
    if digest in cache:
        cached = dict(cache[digest])
        cached["mode"] = cached.get("mode", "llm") + "_cache"
        return cached

    try:
        result = call_llm_classification(
            transcript=transcript,
            api_key=api_key,
            base_url=llm_config.get("base_url", "https://api.openai.com/v1"),
            model=llm_config.get("model", "gpt-4o-mini"),
            max_chars=parse_int(llm_config.get("max_chars", 4500), 4500),
        )
        cache[digest] = result
        return result
    except Exception as exc:
        log(f"[WARN] Falló LLM; se usará resultado por reglas. Detalle: {exc}")
        rules_result["mode"] = "rules_fallback"
        return rules_result


# -----------------------------------------------------------------------------
# Salidas Excel / HTML
# -----------------------------------------------------------------------------
def build_summary(detail: pd.DataFrame) -> pd.DataFrame:
    summary = (
        detail.groupby(["Motivo raíz", "Submotivo raíz"], dropna=False)
        .agg(Llamadas=("Call ID", "count"))
        .reset_index()
        .sort_values("Llamadas", ascending=False)
    )
    total = int(summary["Llamadas"].sum())
    if total == 0:
        summary["%"] = 0.0
        summary["% acumulado"] = 0.0
    else:
        summary["%"] = summary["Llamadas"] / total
        summary["% acumulado"] = summary["%"].cumsum()

    total_row = pd.DataFrame(
        [{
            "Motivo raíz": "Total general",
            "Submotivo raíz": "",
            "Llamadas": total,
            "%": 1.0 if total > 0 else 0.0,
            "% acumulado": 1.0 if total > 0 else 0.0,
        }]
    )
    return pd.concat([summary, total_row], ignore_index=True)


def build_conclusion_summary(detail: pd.DataFrame) -> pd.DataFrame:
    work = detail.copy()
    if "Conclusión" not in work.columns:
        work["Conclusión"] = ""
    work["Conclusión"] = work["Conclusión"].fillna("").astype(str).map(normalize_text)
    work.loc[work["Conclusión"].eq(""), "Conclusión"] = "Sin conclusión"

    summary = (
        work.groupby(["Conclusión", "Motivo raíz", "Submotivo raíz"], dropna=False)
        .agg(Llamadas=("Call ID", "count"))
        .reset_index()
    )
    total = int(summary["Llamadas"].sum())
    if total == 0:
        summary["% conclusión"] = 0.0
        summary["% total"] = 0.0
    else:
        total_by_conclusion = summary.groupby("Conclusión")["Llamadas"].transform("sum")
        summary["% conclusión"] = summary["Llamadas"] / total_by_conclusion
        summary["% total"] = summary["Llamadas"] / total

    summary = summary.sort_values(
        ["Conclusión", "Llamadas", "Motivo raíz", "Submotivo raíz"],
        ascending=[True, False, True, True],
    )

    total_row = pd.DataFrame(
        [{
            "Conclusión": "Total general",
            "Motivo raíz": "",
            "Submotivo raíz": "",
            "Llamadas": total,
            "% conclusión": 1.0 if total > 0 else 0.0,
            "% total": 1.0 if total > 0 else 0.0,
        }]
    )
    return pd.concat([summary, total_row], ignore_index=True)


def format_excel(path: Path) -> None:
    if load_workbook is None:
        log("[WARN] openpyxl no está instalado; se generó Excel sin formato avanzado.")
        return

    wb = load_workbook(path)
    red = "ED1C24"
    light_blue = "DDEBF7"
    white = "FFFFFF"
    gray = "F2F2F2"
    border_color = "B7C9E2"

    thin = Side(style="thin", color=border_color)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        # Header
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor=red)
            cell.font = Font(color=white, bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border

        # Body
        for row in ws.iter_rows(min_row=2):
            is_total = str(row[0].value).strip().lower() == "total general" if row and row[0].value else False
            for idx, cell in enumerate(row, start=1):
                cell.border = border
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if is_total:
                    cell.fill = PatternFill("solid", fgColor=red)
                    cell.font = Font(color=white, bold=True)
                elif ws.title == "Pareto" and idx in (1, 2):
                    cell.fill = PatternFill("solid", fgColor=light_blue)
                    if idx == 1:
                        cell.font = Font(bold=True)
                elif ws.title != "Pareto" and row[0].row % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor=gray)

        # Formatos porcentaje / confianza
        headers = [cell.value for cell in ws[1]]
        for col_idx, header in enumerate(headers, start=1):
            if header in {"%", "% acumulado", "% conclusión", "% total"}:
                for cell in ws.iter_cols(min_col=col_idx, max_col=col_idx, min_row=2, max_row=ws.max_row):
                    for c in cell:
                        c.number_format = "0.00%"
            if header == "Confianza":
                for cell in ws.iter_cols(min_col=col_idx, max_col=col_idx, min_row=2, max_row=ws.max_row):
                    for c in cell:
                        c.number_format = "0.00%"

        # Ancho columnas
        for col_idx, col in enumerate(ws.columns, start=1):
            max_len = 0
            col_letter = get_column_letter(col_idx)
            for cell in col:
                value = "" if cell.value is None else str(cell.value)
                max_len = max(max_len, min(len(value), 80))
            width = max(10, min(max_len + 2, 70))
            if ws.title == "Base Clasificada" and headers[col_idx - 1] == "Transcripción":
                width = 80
            ws.column_dimensions[col_letter].width = width

    wb.save(path)


def write_outputs(
    summary: pd.DataFrame,
    summary_by_conclusion: pd.DataFrame,
    detail: pd.DataFrame,
    output_dir: Path,
) -> Tuple[Path, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    timestamp = _dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    excel_path = output_dir / f"Analisis_Motivos_Llamadas_{timestamp}.xlsx"
    html_path = output_dir / f"Analisis_Motivos_Llamadas_{timestamp}.html"

    taxonomy_df = pd.DataFrame(
        [{"Motivo raíz": t["motivo"], "Submotivo raíz": t["submotivo"], "Prioridad": t.get("priority", "")} for t in TAXONOMY]
    )

    with pd.ExcelWriter(excel_path, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Pareto", index=False)
        summary_by_conclusion.to_excel(writer, sheet_name="Resumen Conclusiones", index=False)
        detail.to_excel(writer, sheet_name="Base Clasificada", index=False)

    format_excel(excel_path)
    write_html(summary, detail, html_path)
    return excel_path, html_path


def pct(value: Any) -> str:
    try:
        return f"{float(value) * 100:,.2f}%"
    except Exception:
        return "0.00%"


def write_html_legacy(summary: pd.DataFrame, detail: pd.DataFrame, path: Path) -> None:
    total = int(detail.shape[0])
    rows_no_total = summary[summary["Motivo raíz"] != "Total general"].copy()
    top = rows_no_total.iloc[0].to_dict() if not rows_no_total.empty else {}
    top3_pct = float(rows_no_total.head(3)["%"].sum()) if not rows_no_total.empty else 0.0
    generated_at = _dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    def render_table(df: pd.DataFrame) -> str:
        table_rows = []
        for _, r in df.iterrows():
            is_total = str(r.get("Motivo raíz", "")) == "Total general"
            cls = " class='total-row'" if is_total else ""
            table_rows.append(
                f"<tr{cls}>"
                f"<td>{html.escape(str(r.get('Motivo raíz', '')))}</td>"
                f"<td>{html.escape(str(r.get('Submotivo raíz', '')))}</td>"
                f"<td class='num'>{int(r.get('Llamadas', 0)):,}</td>"
                f"<td class='num'>{pct(r.get('%', 0))}</td>"
                f"<td class='num'>{pct(r.get('% acumulado', 0))}</td>"
                "</tr>"
            )
        return "\n".join(table_rows)

    html_text = f"""<!doctype html>
<html lang="es">
<head>
<meta charset="utf-8">
<title>Análisis de Motivos de Llamadas</title>
<style>
    body {{ font-family: Arial, sans-serif; margin: 28px; color: #1f2937; background: #f8fafc; }}
    h1 {{ margin-bottom: 4px; color: #111827; }}
    .subtitle {{ color: #6b7280; margin-top: 0; }}
    .cards {{ display: flex; gap: 16px; flex-wrap: wrap; margin: 22px 0; }}
    .card {{ background: #ffffff; border-radius: 14px; padding: 18px 20px; box-shadow: 0 2px 10px rgba(0,0,0,.07); min-width: 220px; }}
    .label {{ font-size: 12px; text-transform: uppercase; color: #6b7280; letter-spacing: .05em; }}
    .value {{ font-size: 28px; font-weight: 700; margin-top: 8px; color: #111827; }}
    .small {{ font-size: 13px; color: #4b5563; margin-top: 8px; }}
    table {{ width: 100%; border-collapse: collapse; background: white; box-shadow: 0 2px 10px rgba(0,0,0,.07); }}
    th {{ background: #ed1c24; color: white; padding: 9px 8px; text-align: center; font-size: 14px; }}
    td {{ border: 1px solid #b7c9e2; padding: 7px 8px; vertical-align: top; }}
    td:first-child {{ background: #ddebf7; font-weight: 700; }}
    .num {{ text-align: right; white-space: nowrap; }}
    .total-row td {{ background: #ed1c24 !important; color: white; font-weight: 700; }}
    .footer {{ margin-top: 18px; color: #6b7280; font-size: 12px; }}
</style>
</head>
<body>
<h1>Análisis de motivos raíz de llamadas</h1>
<p class="subtitle">Generado: {html.escape(generated_at)}</p>

<div class="cards">
    <div class="card">
        <div class="label">Llamadas analizadas</div>
        <div class="value">{total:,}</div>
    </div>
    <div class="card">
        <div class="label">Motivo principal</div>
        <div class="value">{pct(top.get('%', 0))}</div>
        <div class="small">{html.escape(str(top.get('Motivo raíz', 'N/D')))}</div>
    </div>
    <div class="card">
        <div class="label">Top 3 acumulado</div>
        <div class="value">{pct(top3_pct)}</div>
        <div class="small">Concentración de los tres principales motivos.</div>
    </div>
</div>

<table>
<thead>
<tr>
    <th>Motivo raíz</th>
    <th>Submotivo raíz</th>
    <th>Llamadas</th>
    <th>%</th>
    <th>% acumulado</th>
</tr>
</thead>
<tbody>
{render_table(summary)}
</tbody>
</table>

<p class="footer">Fuente: transcripciones procesadas por PyFlow Manager. Clasificación según taxonomía configurada en el script.</p>
</body>
</html>
"""
    path.write_text(html_text, encoding="utf-8")


def write_html(summary: pd.DataFrame, detail: pd.DataFrame, path: Path) -> None:
    total = int(detail.shape[0])
    rows_no_total = summary[summary["Motivo raíz"] != "Total general"].copy()
    top = rows_no_total.iloc[0].to_dict() if not rows_no_total.empty else {}
    top3_pct = float(rows_no_total.head(3)["%"].sum()) if not rows_no_total.empty else 0.0

    fecha_inicio = "sin filtro inicial"
    fecha_fin = "sin filtro final"
    if "Fecha" in detail.columns:
        parsed_dates = pd.to_datetime(detail["Fecha"], errors="coerce")
        if parsed_dates.notna().any():
            fecha_inicio = parsed_dates.min().strftime("%Y-%m-%d")
            fecha_fin = parsed_dates.max().strftime("%Y-%m-%d")

    def render_rows(df: pd.DataFrame) -> str:
        table_rows = []
        for _, r in df.iterrows():
            motivo = str(r.get("Motivo raíz", ""))
            submotivo = str(r.get("Submotivo raíz", ""))
            is_total = motivo == "Total general"
            if is_total:
                table_rows.append(
                    "<tr style=\"background-color: #ed1c24;\">"
                    "<td style=\"padding: 12px; font-size: 12px; color: #ffffff; font-family: Arial, sans-serif;\"><strong>TOTAL GENERAL</strong></td>"
                    f"<td align=\"right\" style=\"padding: 12px; font-size: 13px; font-weight: bold; color: #ffffff; font-family: Arial, sans-serif;\">{int(r.get('Llamadas', 0)):,}</td>"
                    f"<td align=\"right\" style=\"padding: 12px; font-size: 13px; font-weight: bold; color: #ffffff; font-family: Arial, sans-serif;\">{pct(r.get('%', 0))}</td>"
                    "<td align=\"right\" style=\"padding: 12px; font-size: 13px; font-weight: bold; color: #ffffff; font-family: Arial, sans-serif;\">-</td>"
                    "</tr>"
                )
                continue
            table_rows.append(
                "<tr>"
                "<td style=\"padding: 10px 12px; border-bottom: 1px solid #e2e8f0; font-size: 12px; color: #1e293b; font-family: Arial, sans-serif;\">"
                f"<strong style=\"color: #0f172a;\">{html.escape(motivo)}</strong>"
                f"<div style=\"font-size: 11px; color: #64748b; margin-top: 2px;\">{html.escape(submotivo)}</div>"
                "</td>"
                f"<td align=\"right\" style=\"padding: 10px 12px; border-bottom: 1px solid #e2e8f0; font-size: 12px; font-weight: bold; color: #1e293b; font-family: Arial, sans-serif;\">{int(r.get('Llamadas', 0)):,}</td>"
                f"<td align=\"right\" style=\"padding: 10px 12px; border-bottom: 1px solid #e2e8f0; font-size: 12px; color: #475569; font-family: Arial, sans-serif;\">{pct(r.get('%', 0))}</td>"
                f"<td align=\"right\" style=\"padding: 10px 12px; border-bottom: 1px solid #e2e8f0; font-size: 12px; color: #475569; font-family: Arial, sans-serif;\">{pct(r.get('% acumulado', 0))}</td>"
                "</tr>"
            )
        return "\n".join(table_rows)

    html_text = f"""<!DOCTYPE html PUBLIC "-//W3C//DTD XHTML 1.0 Transitional//EN" "http://www.w3.org/TR/xhtml1/DTD/xhtml1-transitional.dtd">
<html xmlns="http://www.w3.org/1999/xhtml" lang="es">
<head>
  <meta http-equiv="Content-Type" content="text/html; charset=utf-8" />
  <meta name="viewport" content="width=device-width, initial-scale=1.0" />
  <title>Análisis de Motivos de Llamadas NBDA y AOL</title>
  <style type="text/css">
    body {{ margin: 0; padding: 0; min-width: 100%; width: 100% !important; background-color: #f1f5f9; font-family: Arial, Helvetica, sans-serif; }}
    table {{ border-collapse: collapse !important; }}
    @media screen and (max-width: 600px) {{
      .container {{ width: 100% !important; max-width: 100% !important; }}
      .col-3 {{ display: block !important; width: 100% !important; box-sizing: border-box; margin-bottom: 12px !important; }}
      .col-3-last {{ margin-bottom: 0 !important; }}
      .responsive-padding {{ padding: 15px !important; }}
      .mobile-scroll {{ display: block !important; width: 100% !important; overflow-x: auto !important; -webkit-overflow-scrolling: touch; }}
    }}
  </style>
</head>
<body style="margin: 0; padding: 0; background-color: #f1f5f9; -webkit-text-size-adjust: 100%; -ms-text-size-adjust: 100%;">
  <table border="0" cellpadding="0" cellspacing="0" width="100%" style="background-color: #f1f5f9; padding: 20px 0;">
    <tr><td align="center" valign="top">
      <table class="container" border="0" cellpadding="0" cellspacing="0" width="600" style="background-color: #ffffff; border-radius: 8px; overflow: hidden; box-shadow: 0 4px 6px rgba(0, 0, 0, 0.05); border: 1px solid #e2e8f0;">
        <tr>
          <td bgcolor="#ed1c24" align="left" style="padding: 24px 30px; background-color: #ed1c24;">
            <h1 style="color: #ffffff; font-size: 20px; font-weight: bold; margin: 0; font-family: Arial, sans-serif;">Análisis de Motivos de Llamadas NBDA y AOL</h1>
            <p style="color: #ffcccc; font-size: 12px; margin: 4px 0 0 0; font-family: Arial, sans-serif;">Periodo: {html.escape(fecha_inicio)} al {html.escape(fecha_fin)}</p>
          </td>
        </tr>
        <tr><td class="responsive-padding" style="padding: 30px; background-color: #ffffff;">
          <p style="color: #334155; font-size: 14px; line-height: 1.5; margin: 0 0 20px 0;">
            Buen día,<br /><br />
            Comparto el análisis correspondiente a las conclusiones de cambio de contraseña AOL, NBDA y Soporte técnico NBDA para el periodo comprendido desde el <strong>{html.escape(fecha_inicio)}</strong> hasta <strong>{html.escape(fecha_fin)}</strong>.
          </p>
          <table border="0" cellpadding="0" cellspacing="0" width="100%" style="margin-bottom: 25px;"><tr>
            <td class="col-3" valign="top" width="31%" style="background-color: #f8fafc; border: 1px solid #e2e8f0; border-radius: 6px; padding: 12px; text-align: left;">
              <span style="color: #64748b; font-size: 10px; font-weight: bold; text-transform: uppercase; letter-spacing: 0.5px; display: block; font-family: Arial, sans-serif;">Llamadas Analizadas</span>
              <strong style="color: #0f172a; font-size: 22px; font-weight: bold; display: block; margin-top: 6px; font-family: Arial, sans-serif;">{total:,}</strong>
              <span style="color: #64748b; font-size: 11px; display: block; margin-top: 4px; font-family: Arial, sans-serif;">Muestra total procesada</span>
            </td>
            <td class="col-3" width="3%" style="font-size: 0; line-height: 0;">&nbsp;</td>
            <td class="col-3" valign="top" width="31%" style="background-color: #f8fafc; border: 1px solid #e2e8f0; border-radius: 6px; padding: 12px; text-align: left;">
              <span style="color: #64748b; font-size: 10px; font-weight: bold; text-transform: uppercase; letter-spacing: 0.5px; display: block; font-family: Arial, sans-serif;">Motivo Principal</span>
              <strong style="color: #ed1c24; font-size: 22px; font-weight: bold; display: block; margin-top: 6px; font-family: Arial, sans-serif;">{pct(top.get('%', 0))}</strong>
              <span style="color: #475569; font-size: 11px; display: block; margin-top: 4px; font-family: Arial, sans-serif; line-height: 1.2;">{html.escape(str(top.get('Motivo raíz', 'N/D')))}</span>
            </td>
            <td class="col-3" width="3%" style="font-size: 0; line-height: 0;">&nbsp;</td>
            <td class="col-3 col-3-last" valign="top" width="31%" style="background-color: #f8fafc; border: 1px solid #e2e8f0; border-radius: 6px; padding: 12px; text-align: left;">
              <span style="color: #64748b; font-size: 10px; font-weight: bold; text-transform: uppercase; letter-spacing: 0.5px; display: block; font-family: Arial, sans-serif;">Top 3 Acumulado</span>
              <strong style="color: #0f172a; font-size: 22px; font-weight: bold; display: block; margin-top: 6px; font-family: Arial, sans-serif;">{pct(top3_pct)}</strong>
              <span style="color: #64748b; font-size: 11px; display: block; margin-top: 4px; font-family: Arial, sans-serif; line-height: 1.2;">Concentrado en 3 motivos</span>
            </td>
          </tr></table>
          <div class="mobile-scroll">
            <table border="0" cellpadding="0" cellspacing="0" width="100%" style="border-collapse: collapse; min-width: 480px; border: 1px solid #e2e8f0; border-radius: 6px; overflow: hidden;">
              <thead><tr style="background-color: #f8fafc;">
                <th align="left" style="padding: 10px 12px; border-bottom: 2px solid #cbd5e1; font-size: 12px; font-weight: bold; color: #475569; font-family: Arial, sans-serif; width: 40%;">Motivo Raíz / Submotivo</th>
                <th align="right" style="padding: 10px 12px; border-bottom: 2px solid #cbd5e1; font-size: 12px; font-weight: bold; color: #475569; font-family: Arial, sans-serif; width: 20%;">Llamadas</th>
                <th align="right" style="padding: 10px 12px; border-bottom: 2px solid #cbd5e1; font-size: 12px; font-weight: bold; color: #475569; font-family: Arial, sans-serif; width: 20%;">%</th>
                <th align="right" style="padding: 10px 12px; border-bottom: 2px solid #cbd5e1; font-size: 12px; font-weight: bold; color: #475569; font-family: Arial, sans-serif; width: 20%;">% Acum.</th>
              </tr></thead>
              <tbody>{render_rows(summary)}</tbody>
            </table>
          </div>
        </td></tr>
        <tr><td align="center" bgcolor="#1e293b" style="padding: 20px; background-color: #1e293b;">
          <p style="color: #94a3b8; font-size: 11px; margin: 0; font-family: Arial, sans-serif;">Estrategia y Gestión - Banco Atlántida S.A.</p>
          <p style="color: #64748b; font-size: 10px; margin: 6px 0 0 0; font-family: Arial, sans-serif;">Clasificación de transacciones de soporte y accesos de canales.</p>
        </td></tr>
      </table>
    </td></tr>
  </table>
</body>
</html>
"""
    path.write_text(html_text, encoding="utf-8")


def get_graph_access_token(tenant_id: str, client_id: str, client_secret: str) -> str:
    try:
        import requests
    except ImportError as exc:
        raise RuntimeError("Falta instalar requests para enviar correo por Microsoft Graph.") from exc

    url = f"https://login.microsoftonline.com/{tenant_id}/oauth2/v2.0/token"
    data = {
        "grant_type": "client_credentials",
        "client_id": client_id,
        "client_secret": client_secret,
        "scope": "https://graph.microsoft.com/.default",
    }
    response = requests.post(url, data=data, timeout=30)
    if response.status_code >= 400:
        raise RuntimeError(f"Error obteniendo token Graph {response.status_code}: {response.text[:800]}")
    return response.json()["access_token"]


def build_file_attachment(path: Path) -> Dict[str, Any]:
    mime_type = mimetypes.guess_type(str(path))[0] or "application/octet-stream"
    content = base64.b64encode(path.read_bytes()).decode("ascii")
    return {
        "@odata.type": "#microsoft.graph.fileAttachment",
        "name": path.name,
        "contentType": mime_type,
        "contentBytes": content,
    }


def send_report_email(
    to_recipients: List[str],
    cc_recipients: List[str],
    subject: str,
    sender_email: str,
    tenant_id: str,
    client_id: str,
    client_secret: str,
    excel_path: Path,
    html_path: Path,
    total_calls: int,
) -> bool:
    if not to_recipients:
        log("Correo no enviado: no hay destinatarios configurados en EMAIL_TO.")
        return False

    missing = [
        name for name, value in {
            "GRAPH_TENANT_ID": tenant_id,
            "GRAPH_CLIENT_ID": client_id,
            "GRAPH_CLIENT_SECRET": client_secret,
            "GRAPH_SENDER_EMAIL": sender_email,
        }.items()
        if not normalize_text(value)
    ]
    if missing:
        log(f"[WARN] Correo no enviado: faltan variables Graph: {', '.join(missing)}")
        return False

    try:
        import requests
    except ImportError as exc:
        log(f"[WARN] Correo no enviado: falta requests. Detalle: {exc}")
        return False

    body = html_path.read_text(encoding="utf-8")

    payload = {
        "message": {
            "subject": subject or "Analisis de motivos de llamadas NBDA AOL",
            "body": {"contentType": "HTML", "content": body},
            "toRecipients": [{"emailAddress": {"address": email}} for email in to_recipients],
            "ccRecipients": [{"emailAddress": {"address": email}} for email in cc_recipients],
            "attachments": [build_file_attachment(excel_path), build_file_attachment(html_path)],
        },
        "saveToSentItems": "true",
    }

    try:
        token = get_graph_access_token(tenant_id, client_id, client_secret)
        url = f"https://graph.microsoft.com/v1.0/users/{sender_email}/sendMail"
        response = requests.post(
            url,
            json=payload,
            headers={"Authorization": f"Bearer {token}", "Content-Type": "application/json"},
            timeout=60,
        )
        if response.status_code >= 400:
            raise RuntimeError(f"Error enviando correo Graph {response.status_code}: {response.text[:1000]}")
        log(f"Correo enviado a: {', '.join(to_recipients + cc_recipients)}")
        return True
    except Exception as exc:
        log(f"[WARN] No se pudo enviar el correo: {exc}")
        return False


# -----------------------------------------------------------------------------
# Main
# -----------------------------------------------------------------------------
def main() -> int:
    start_time = time.time()
    cli = parse_cli_args()

    input_file = Path(str(get_param("INPUT_FILE", cli, "")).strip().strip('"'))
    output_dir = resolve_output_dir(Path(str(get_param("OUTPUT_DIR", cli, "")).strip().strip('"')))
    sheet_name = str(get_param("SHEET_NAME", cli, "")).strip()
    text_column_param = str(get_param("TEXT_COLUMN", cli, "AUTO")).strip()
    call_id_column_param = str(get_param("CALL_ID_COLUMN", cli, "AUTO")).strip()
    date_column_param = str(get_param("DATE_COLUMN", cli, "AUTO")).strip()
    client_column_param = str(get_param("CLIENT_COLUMN", cli, "AUTO")).strip()
    conclusion_column_param = str(get_param("CONCLUSION_COLUMN", cli, "AUTO")).strip()
    date_from = str(get_param("DATE_FROM", cli, "")).strip()
    date_to = str(get_param("DATE_TO", cli, "")).strip()
    group_by_call_id = parse_bool(get_param("GROUP_BY_CALL_ID", cli, "true"), True)
    analysis_mode = str(get_param("ANALYSIS_MODE", cli, "rules")).strip().lower()
    min_confidence_llm = parse_float(get_param("MIN_CONFIDENCE_LLM", cli, "0.72"), 0.72)
    email_to = split_list_value(get_param("EMAIL_TO", cli, ""))
    email_cc = split_list_value(get_param("EMAIL_CC", cli, ""))
    email_subject = str(
        get_param("EMAIL_SUBJECT", cli, "Analisis de motivos de llamadas NBDA AOL")
    ).strip()

    llm_config = {
        "api_key": str(get_param("LLM_API_KEY", cli, "")).strip(),
        "base_url": str(get_param("LLM_BASE_URL", cli, "https://api.openai.com/v1")).strip(),
        "model": str(get_param("LLM_MODEL", cli, "gpt-4o-mini")).strip(),
        "max_chars": parse_int(get_param("MAX_CHARS_LLM", cli, "4500"), 4500),
    }

    graph_config = {
        "tenant_id": str(get_param("GRAPH_TENANT_ID", cli, "")).strip(),
        "client_id": str(get_param("GRAPH_CLIENT_ID", cli, "")).strip(),
        "client_secret": str(get_param("GRAPH_CLIENT_SECRET", cli, "")).strip(),
        "sender_email": str(get_param("GRAPH_SENDER_EMAIL", cli, "")).strip(),
    }

    input_file = resolve_input_file(input_file)

    if analysis_mode not in {"rules", "hybrid", "llm"}:
        raise ValueError("ANALYSIS_MODE debe ser: rules, hybrid o llm")

    log("Iniciando análisis de motivos raíz de llamadas")
    log(f"Modo de análisis: {analysis_mode}")

    df = read_input_file(input_file, sheet_name)
    original_rows = len(df)
    log(f"Filas leídas: {original_rows:,}")

    # Limpieza básica de nombres de columnas vacíos
    df.columns = [str(c).strip() if str(c).strip() else f"Columna_{i+1}" for i, c in enumerate(df.columns)]

    text_col = detect_column(
        df,
        text_column_param,
        candidates=[
            "Transcripción", "Transcripcion", "transcript", "transcripts", "texto", "texto llamada",
            "call transcript", "conversation transcript", "body", "contenido", "mensaje", "transcription",
            "Transcript Text", "full_text", "text",
        ],
        purpose="transcripción",
        required=True,
    )
    call_id_col = detect_column(
        df,
        call_id_column_param,
        candidates=[
            "conversationId", "conversation_id", "id_conversation", "id conversación", "id conversacion",
            "call_id", "callId", "id llamada", "id_llamada", "conversation", "id", "interaccion", "interaction_id",
        ],
        purpose="ID de llamada/conversación",
        required=False,
    )
    date_col = detect_column(
        df,
        date_column_param,
        candidates=[
            "fecha", "date", "startTime", "start_time", "inicio", "fecha inicio", "fecha_inicio",
            "conversationStart", "conversation_start", "created_at", "createdAt", "aniomesdia",
        ],
        purpose="fecha",
        required=False,
    )
    client_col = detect_column(
        df,
        client_column_param,
        candidates=[
            "cliente", "dni", "identidad", "documento", "telefono", "teléfono", "ani", "numero", "número",
            "caller", "customer", "customer_id", "id_cliente", "etiqueta_externa", "external_tag",
        ],
        purpose="cliente/DNI/teléfono",
        required=False,
    )

    conclusion_col = detect_column(
        df,
        conclusion_column_param,
        candidates=[
            "conclusion_ultima", "conclusion_original", "conclusion", "conclusión",
            "wrapUpCodeName", "wrapupcodename", "wrapUpCode", "wrapupcode",
            "wrapUpCodeId_ultima", "wrapUpCodeId", "wrapupcodeid",
        ],
        purpose="conclusión",
        required=False,
    )

    log(f"Columna transcripción: {text_col}")
    log(f"Columna ID llamada: {call_id_col or 'No detectada'}")
    log(f"Columna fecha: {date_col or 'No detectada'}")
    log(f"Columna cliente: {client_col or 'No detectada'}")
    log(f"Columna conclusión: {conclusion_col or 'No detectada'}")

    df = apply_date_filter(df, date_col, date_from, date_to)
    calls_df = aggregate_by_call(df, text_col, call_id_col, date_col, client_col, conclusion_col, group_by_call_id)

    # Elimina llamadas sin texto útil
    before_text_filter = len(calls_df)
    calls_df["Transcripción"] = calls_df["Transcripción"].fillna("").astype(str).map(normalize_text)
    calls_df = calls_df[calls_df["Transcripción"].str.len() > 0].copy()
    log(f"Filtro llamadas con transcripción vacía: {before_text_filter:,} -> {len(calls_df):,}")

    if calls_df.empty:
        raise ValueError("No quedaron llamadas con transcripción para analizar.")

    cache: Dict[str, Dict[str, Any]] = {}
    classifications: List[Dict[str, Any]] = []

    total_calls = len(calls_df)
    for idx, transcript in enumerate(calls_df["Transcripción"].tolist(), start=1):
        if idx == 1 or idx % 50 == 0 or idx == total_calls:
            log(f"Clasificando llamada {idx:,}/{total_calls:,}")
        result = classify_transcript(
            transcript=transcript,
            mode=analysis_mode,
            min_confidence_llm=min_confidence_llm,
            llm_config=llm_config,
            cache=cache,
        )
        classifications.append(result)

    class_df = pd.DataFrame(classifications)
    detail = calls_df.reset_index(drop=True).copy()
    detail["Motivo raíz"] = class_df["motivo"]
    detail["Submotivo raíz"] = class_df["submotivo"]
    detail["Confianza"] = class_df["confidence"].astype(float)
    detail["Modo clasificación"] = class_df["mode"]
    detail["Señales / justificación"] = class_df["signals"]

    # Orden de columnas para que sea cómodo revisar
    detail = detail[
        [
            "Call ID", "Fecha", "Cliente", "Conclusión", "Motivo raíz", "Submotivo raíz", "Confianza",
            "Modo clasificación", "Señales / justificación", "Transcripción",
        ]
    ]

    summary = build_summary(detail)
    summary_by_conclusion = build_conclusion_summary(detail)
    excel_path, html_path = write_outputs(summary, summary_by_conclusion, detail, output_dir)

    send_report_email(
        to_recipients=email_to,
        cc_recipients=email_cc,
        subject=email_subject,
        sender_email=graph_config["sender_email"],
        tenant_id=graph_config["tenant_id"],
        client_id=graph_config["client_id"],
        client_secret=graph_config["client_secret"],
        excel_path=excel_path,
        html_path=html_path,
        total_calls=len(detail),
    )

    elapsed = time.time() - start_time
    log("Proceso finalizado correctamente")
    log(f"Excel generado: {excel_path}")
    log(f"HTML generado: {html_path}")
    log(f"Tiempo total: {elapsed:,.1f} segundos")
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Exception as exc:
        log(f"[ERROR] {exc}")
        sys.exit(1)
