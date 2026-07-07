# -*- coding: utf-8 -*-
from __future__ import annotations

import argparse
import csv
import html
import json
import logging
import math
import os
import re
import sys
import unicodedata
from collections import Counter
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

try:
    import pandas as pd
except ImportError as exc:
    raise SystemExit("Falta instalar pandas/openpyxl. Ejecuta: pip install pandas openpyxl") from exc


PYFLOW_PARAMS = {
    "LIMITE_FILAS": {
        "type": "text",
        "label": "Limite de filas para prueba, 0 = todas",
        "default": "0",
        "required": False
    },
    "REANALIZAR_EXISTENTES": {
        "type": "select",
        "label": "Reanalizar existentes",
        "options": ["NO", "SI"],
        "default": "NO",
        "required": False
    }
}


BASE_DIR = Path(__file__).resolve().parent
CONFIG_DIR = BASE_DIR / "config"
INPUT_DIR = BASE_DIR / "input"
IMPUT_DIR = BASE_DIR / "imput"
OUTPUT_DIR = BASE_DIR / "output"


COLUMNAS_CANDIDATAS_DEFAULT = {
    "conversation_id": ["conversation_id", "conversationId", "CONVERSATION_ID", "id_conversacion"],
    "fecha_hora": ["fecha_hora", "fecha", "FECHA_HORA", "FECHA", "conversationStart", "conversationStart_local"],
    "operador": ["operador", "agent", "agente", "usuario", "USER_NAME", "agentName", "nombre_agente"],
    "cola": ["cola", "queue", "QUEUE_NAME", "queueName", "nombre_cola"],
    "wrapup": ["conclusion_ultima", "conclusion_original", "wrapup", "conclusion", "CONCLUSION", "WRAPUP", "wrapUpCodeName", "wrapUpCodeId_ultima", "wrapUpCodeId"],
    "cliente_id": ["cliente_id", "ani", "telefono", "numero_cliente", "externalTag", "external_tag", "remote"],
    "duracion": ["duracion", "duracion_segundos", "duration_seconds", "DURACION_SEGUNDOS"],
    "temas_genesys": ["temas_genesys", "topics", "temas", "categorias"],
    "transcripcion": ["text", "transcripcion", "TRANSCRIPCION", "texto_llamada", "conversation_text", "transcript"]
}


FRASES_NO_CONTACTO = [
    "no contesta", "no contesto", "buzon", "voicemail", "busy", "ocupado", "abandono",
    "abandonada", "timeout", "sin contacto", "llamada muda", "sin audio", "no se escucha",
    "cliente no responde", "llamada caida", "se corta la llamada"
]

FRASES_RESPUESTA_UTIL = [
    "le informo", "le indico", "con gusto le indico", "puede realizarlo", "puede hacerlo",
    "los requisitos son", "el horario es", "su saldo", "saldo disponible", "fecha de pago",
    "pago minimo", "monto a pagar", "la cuota", "la tasa", "debe ingresar", "debe seleccionar",
    "le explico", "le explicare", "los pasos son", "queda registrado", "numero de gestion",
    "numero de caso", "numero de reclamo", "le voy a generar", "vamos a validar", "voy a escalar",
    "vamos a bloquear", "se bloqueo", "bloqueo preventivo", "se estara validando", "plazo de respuesta",
    "tiempo de respuesta", "seguimiento", "notificacion", "debe presentarse con", "puede consultar",
    "me aparece", "le aparece", "aparece que", "verificando", "validando", "validar",
    "correcto", "transferencia electronica", "transferencia de su banca en linea",
    "fue realizado", "se realizo", "se hizo", "movimiento", "transaccion",
    "puede verificar", "en su banca en linea", "en el sistema", "permítame un momento",
    "permitame un momento", "vamos a realizar la verificacion", "realizar la verificacion"
]

FRASES_RIESGO_NO_RESOLUCION = [
    "no sabria decirle", "no puedo ayudarle", "no le puedo ayudar", "llame despues",
    "intente mas tarde", "no tengo informacion", "no tenemos sistema", "no puedo hacer nada",
    "eso no se atiende aqui", "aqui no vemos eso", "no tengo acceso", "no es mi problema"
]

FRASES_REDIRECCION = [
    "vaya a agencia", "debe ir a agencia", "tiene que presentarse", "en sucursal",
    "oficina", "agencia mas cercana", "presencial", "debe avocarse"
]

FRASES_EXPLICACION_REDIRECCION = [
    "requisitos", "debe llevar", "presentar", "con su identidad", "horario", "ubicacion",
    "pasos", "motivo", "porque", "para que le", "alli le", "en agencia podran"
]

FRASES_CONSULTA_CLIENTE = [
    "quiero saber", "consulta", "me puede indicar", "necesito saber", "cual es", "cuanto",
    "cuando", "donde", "como puedo", "requisitos", "horario", "saldo", "movimientos",
    "estado de cuenta", "fecha de pago", "monto", "cuota"
]

FRASES_PROBLEMA_CLIENTE = [
    "problema", "no puedo", "no me deja", "no funciona", "reclamo", "queja", "no reconozco",
    "me debitaron", "fraude", "bloquear", "perdi", "robaron", "no me llega", "error",
    "no me resolvieron", "no me solucionan", "me urge", "necesito solucion"
]


@dataclass
class RunConfig:
    limite_filas: int
    reanalizar_existentes: bool
    min_chars: int
    max_chars: int
    batch_size: int


def pyflow_progress(value: int) -> None:
    print(f"PYFLOW_PROGRESS={max(0, min(100, int(value)))}", flush=True)


def get_param(name: str, default: Any = None) -> Any:
    value = os.environ.get(name)
    if value is None or str(value).strip() == "":
        return default
    return value


def env_int(name: str, default: int) -> int:
    try:
        return int(float(str(get_param(name, default)).strip()))
    except Exception:
        return default


def env_bool(name: str, default: bool = False) -> bool:
    value = str(get_param(name, "SI" if default else "NO")).strip().upper()
    return value in {"SI", "S", "TRUE", "1", "YES", "Y"}


def setup_logger() -> logging.Logger:
    logger = logging.getLogger("p2_8_local")
    logger.handlers.clear()
    logger.setLevel(logging.INFO)
    handler = logging.StreamHandler(sys.stdout)
    handler.setFormatter(logging.Formatter("%(asctime)s | %(levelname)s | %(message)s", "%Y-%m-%d %H:%M:%S"))
    logger.addHandler(handler)
    return logger


def normalizar_texto(texto: Any) -> str:
    if texto is None:
        return ""
    if isinstance(texto, float) and math.isnan(texto):
        return ""
    text = str(texto).replace("\r", "\n")
    text = "".join(ch for ch in text if unicodedata.category(ch)[0] != "C" or ch in "\n\t")
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = text.lower()
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def normalizar_columna(nombre: Any) -> str:
    text = normalizar_texto(str(nombre or "").replace("\ufeff", ""))
    text = re.sub(r"[^a-z0-9]+", "_", text)
    return text.strip("_")


def cargar_json_obligatorio(path: Path) -> Dict[str, Any]:
    if not path.exists():
        raise FileNotFoundError(f"Falta archivo de configuracion requerido: {path}")
    with path.open("r", encoding="utf-8-sig") as fh:
        return json.load(fh)


def cargar_config() -> Tuple[Dict[str, Any], Dict[str, Any], Dict[str, Any]]:
    parametros = cargar_json_obligatorio(CONFIG_DIR / "parametros_p2_8.json")
    prompt = cargar_json_obligatorio(CONFIG_DIR / "prompt_ia_p2_8.json")
    matriz = cargar_json_obligatorio(CONFIG_DIR / "matriz_evaluacion_p2_8.json")
    return parametros, prompt, matriz


def load_run_config(parametros: Dict[str, Any]) -> RunConfig:
    return RunConfig(
        limite_filas=env_int("LIMITE_FILAS", 0),
        reanalizar_existentes=env_bool("REANALIZAR_EXISTENTES", False),
        min_chars=int(parametros.get("min_caracteres_transcripcion_evaluable", 120)),
        max_chars=int(parametros.get("max_caracteres_transcripcion", 18000)),
        batch_size=int(parametros.get("batch_size", 5000))
    )


def input_folder() -> Path:
    INPUT_DIR.mkdir(parents=True, exist_ok=True)
    if any(INPUT_DIR.iterdir()):
        return INPUT_DIR
    if IMPUT_DIR.exists() and any(IMPUT_DIR.iterdir()):
        return IMPUT_DIR
    return INPUT_DIR


def listar_archivos_input() -> List[Path]:
    folder = input_folder()
    files = sorted(p for p in folder.iterdir() if p.is_file() and p.suffix.lower() in {".csv", ".txt", ".xlsx", ".xlsm", ".xls"})
    if not files:
        extra = " Tambien revise que la carpeta se llame input, no imput." if folder == INPUT_DIR else ""
        raise FileNotFoundError(f"No encontre CSV/Excel en {folder}.{extra}")
    return files


def detectar_separador(path: Path) -> str:
    sample = path.read_text(encoding="utf-8-sig", errors="ignore")[:4096]
    candidates = [",", ";", "\t"]
    counts = {sep: sample.count(sep) for sep in candidates}
    return max(counts, key=counts.get) if max(counts.values()) > 0 else ","


def leer_archivo(path: Path) -> pd.DataFrame:
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm", ".xls"}:
        df = pd.read_excel(path, dtype=str)
    else:
        sep = detectar_separador(path)
        try:
            df = pd.read_csv(path, sep=sep, dtype=str, low_memory=False, encoding="utf-8-sig")
        except UnicodeDecodeError:
            df = pd.read_csv(path, sep=sep, dtype=str, low_memory=False, encoding="latin-1")
    df["__source_file"] = path.name
    return df


def leer_todos_los_archivos(logger: logging.Logger) -> pd.DataFrame:
    files = listar_archivos_input()
    logger.info("Archivos detectados en input/imput: %s", len(files))
    frames = []
    for idx, path in enumerate(files, start=1):
        logger.info("Leyendo archivo %s/%s: %s", idx, len(files), path.name)
        frames.append(leer_archivo(path))
    return pd.concat(frames, ignore_index=True)


def encontrar_columna(df: pd.DataFrame, candidatos: Iterable[str], campo: str = "") -> Optional[str]:
    normalized = {normalizar_columna(col): col for col in df.columns}
    blocked_by_field = {
        "transcripcion": {
            "transcript_estado", "transcript_error", "transcript_intentos",
            "phrases_count", "transcript_first_marker", "transcript_last_marker"
        },
        "wrapup": {"wrapupcodeid", "wrapupcodeid_ultima"},
    }
    blocked = {normalizar_columna(x) for x in blocked_by_field.get(campo, set())}
    for cand in candidatos:
        key = normalizar_columna(cand)
        if key in normalized and key not in blocked:
            return normalized[key]
    for cand in candidatos:
        key = normalizar_columna(cand)
        if campo == "transcripcion" and key in {"transcript", "transcripcion"}:
            continue
        for col_norm, original in normalized.items():
            if col_norm in blocked:
                continue
            if key and (key in col_norm or col_norm in key):
                return original
    return None


def resolver_columnas(df: pd.DataFrame, parametros: Dict[str, Any]) -> Dict[str, Optional[str]]:
    candidatos = {campo: list(lista) for campo, lista in COLUMNAS_CANDIDATAS_DEFAULT.items()}
    for campo, lista in (parametros.get("columnas_candidatas", {}) or {}).items():
        base = candidatos.get(campo, [])
        merged = list(lista or []) + [x for x in base if x not in (lista or [])]
        candidatos[campo] = merged
    columnas = {campo: encontrar_columna(df, lista, campo) for campo, lista in candidatos.items()}
    if not columnas.get("transcripcion"):
        raise ValueError("No encontre columna de transcripcion. Use transcripcion/transcript/text/texto_llamada o ajuste config/parametros_p2_8.json.")
    if columnas.get("transcripcion") in {"transcript_estado", "transcript_error", "transcript_intentos"}:
        raise ValueError(f"Columna de transcripcion invalida detectada: {columnas.get('transcripcion')}. Debe ser text/transcripcion/texto_llamada.")
    return columnas


def valor(row: pd.Series, col: Optional[str], default: str = "") -> str:
    if not col or col not in row:
        return default
    value = row.get(col, default)
    if pd.isna(value):
        return default
    return str(value)


def limpiar_transcripcion(texto: Any, max_chars: int) -> str:
    text = str(texto or "").replace("\x00", " ")
    text = re.sub(r"\s+", " ", text).strip()
    if len(text) > max_chars:
        half = max_chars // 2
        text = text[:half] + " ...[TRANSCRIPCION RECORTADA]... " + text[-half:]
    return text


def contiene(texto: str, frases: Iterable[str]) -> List[str]:
    t = normalizar_texto(texto)
    hits = []
    for frase in frases:
        f = normalizar_texto(frase)
        if f and f in t:
            hits.append(frase)
    return hits


def separar_hablantes(transcripcion: str) -> Tuple[str, str, str]:
    transcripcion = (transcripcion or "").replace("\\n", "\n")
    cliente_parts: List[str] = []
    operador_parts: List[str] = []
    tag_pattern = re.compile(r"(?i)\b(cliente|customer|usuario|caller|external|operador|agente|agent|representative|asesor|internal)\s*:")
    matches = list(tag_pattern.finditer(transcripcion))
    for index, match in enumerate(matches):
        speaker = normalizar_texto(match.group(1))
        start = match.end()
        end = matches[index + 1].start() if index + 1 < len(matches) else len(transcripcion)
        chunk = transcripcion[start:end].strip()
        if speaker in {"cliente", "customer", "usuario", "caller", "external"}:
            cliente_parts.append(chunk)
        else:
            operador_parts.append(chunk)
    general = limpiar_transcripcion(transcripcion, 50000)
    cliente = " ".join(cliente_parts).strip() or general
    operador = " ".join(operador_parts).strip() or general
    return cliente, operador, general


def evidencia(texto: str, frases: Iterable[str], limit: int = 220) -> str:
    t_norm = normalizar_texto(texto)
    for frase in frases:
        f = normalizar_texto(frase)
        pos = t_norm.find(f)
        if f and pos >= 0:
            start = max(0, pos - 60)
            return texto[start:start + limit].strip()
    return limpiar_transcripcion(texto, limit)[:limit]


def es_no_evaluable(row_data: Dict[str, str], transcripcion: str, parametros: Dict[str, Any]) -> Tuple[bool, str]:
    wrapup = row_data.get("wrapup", "")
    wrapups_na = parametros.get("wrapups_no_evaluables", [])
    if contiene(wrapup, wrapups_na + FRASES_NO_CONTACTO):
        return True, "Wrap-up indica no contacto o llamada no evaluable."
    if len(normalizar_texto(transcripcion)) < int(parametros.get("min_caracteres_transcripcion_evaluable", 120)):
        return True, "Transcripcion insuficiente para evaluar."
    if contiene(transcripcion, FRASES_NO_CONTACTO) and len(normalizar_texto(transcripcion)) < 700:
        return True, "La transcripcion sugiere no contacto o audio insuficiente."
    return False, ""


def clasificar_tipo_interaccion(texto_cliente: str, texto_general: str) -> str:
    if contiene(texto_cliente, FRASES_PROBLEMA_CLIENTE):
        return "Reclamo o incidente"
    if contiene(texto_cliente, FRASES_CONSULTA_CLIENTE):
        return "Consulta informativa"
    if contiene(texto_general, ["bloquear", "generar", "registrar", "escalar", "reclamo", "caso"]):
        return "Gestion operativa"
    return "Otro"


def evaluar_local(row_data: Dict[str, str], transcripcion: str, parametros: Dict[str, Any]) -> Dict[str, Any]:
    no_eval, razon = es_no_evaluable(row_data, transcripcion, parametros)
    if no_eval:
        return {
            "p2_8_resuelve_eficientemente": "N/A",
            "justificacion_breve": razon,
            "motivo_cliente": "No evidenciado",
            "evidencia_cliente": "No evidenciado",
            "evidencia_operador": "No evidenciado",
            "requiere_revision_supervisor": "NO",
            "tipo_interaccion": "No contactado o no evaluable",
            "confianza": "Media"
        }

    texto_cliente, texto_operador, texto_general = separar_hablantes(transcripcion)
    utiles = contiene(texto_operador, parametros.get("frases_indican_respuesta_util", []) + FRASES_RESPUESTA_UTIL)
    riesgo_frases = [
        frase for frase in (parametros.get("frases_riesgo_no_resolucion", []) + FRASES_RIESGO_NO_RESOLUCION)
        if normalizar_texto(frase) not in {"no se", "no s"}
    ]
    riesgos = contiene(texto_operador, riesgo_frases)
    redireccion = contiene(texto_operador, FRASES_REDIRECCION)
    explicacion_redireccion = contiene(texto_operador, FRASES_EXPLICACION_REDIRECCION)
    cliente_tiene_motivo = bool(contiene(texto_cliente, FRASES_CONSULTA_CLIENTE + FRASES_PROBLEMA_CLIENTE)) or len(normalizar_texto(texto_cliente)) > 80
    consulta_informativa_atendida = (
        clasificar_tipo_interaccion(texto_cliente, texto_general) == "Consulta informativa"
        and len(normalizar_texto(texto_operador)) > 120
        and bool(contiene(texto_operador, [
            "me aparece", "le aparece", "aparece que", "correcto", "verificando",
            "validando", "transferencia", "movimiento", "transaccion", "se realizo",
            "fue realizado", "puede verificar", "banca en linea", "en el sistema"
        ]))
    )
    tipo = clasificar_tipo_interaccion(texto_cliente, texto_general)

    if redireccion and not explicacion_redireccion and not utiles:
        result = "No"
        just = "Se detecta redireccion sin explicacion suficiente ni pasos claros para el cliente."
        revision = "SI"
    elif riesgos and not utiles:
        result = "No"
        just = "Se detectan frases de riesgo sin evidencia clara de solucion, orientacion o proximo paso util."
        revision = "SI"
    elif utiles or consulta_informativa_atendida:
        result = "Si"
        just = "Hay evidencia de respuesta util, orientacion, gestion, escalamiento o pasos relacionados con el motivo."
        revision = "NO" if not riesgos else "SI"
    elif not cliente_tiene_motivo:
        result = "N/A"
        just = "No se identifica una consulta o problema real del cliente con evidencia suficiente."
        revision = "NO"
    else:
        result = "No"
        just = "El cliente presenta un motivo evaluable, pero no se detecta orientacion, gestion o respuesta util suficiente."
        revision = "SI"

    motivo_hits = contiene(texto_cliente, FRASES_CONSULTA_CLIENTE + FRASES_PROBLEMA_CLIENTE)
    evidencia_cliente = evidencia(texto_cliente, motivo_hits or FRASES_CONSULTA_CLIENTE + FRASES_PROBLEMA_CLIENTE)
    evidencia_operador = evidencia(texto_operador, utiles + riesgos + redireccion + explicacion_redireccion)
    motivo_cliente = evidencia_cliente[:180] if evidencia_cliente else "Motivo inferido por conversacion, revisar detalle."

    return {
        "p2_8_resuelve_eficientemente": result,
        "justificacion_breve": just[:300],
        "motivo_cliente": motivo_cliente,
        "evidencia_cliente": evidencia_cliente or "No evidenciado",
        "evidencia_operador": evidencia_operador or "No evidenciado",
        "requiere_revision_supervisor": revision,
        "tipo_interaccion": tipo,
        "confianza": "Media" if revision == "SI" else "Alta"
    }


def checkpoint_paths(output_dir: Path) -> Tuple[Path, Path]:
    return output_dir / "checkpoint_p2_8.csv", output_dir / "checkpoint_p2_8.jsonl"


def cargar_checkpoint(output_dir: Path) -> set[str]:
    checkpoint_csv, checkpoint_jsonl = checkpoint_paths(output_dir)
    path = checkpoint_csv if checkpoint_csv.exists() else checkpoint_jsonl
    procesados: set[str] = set()
    if not path.exists():
        return procesados
    try:
        if path.suffix.lower() == ".jsonl":
            with path.open("r", encoding="utf-8") as fh:
                for line in fh:
                    if line.strip():
                        item = json.loads(line)
                        if item.get("conversation_id"):
                            procesados.add(str(item["conversation_id"]))
        else:
            with path.open("r", encoding="utf-8-sig", newline="") as fh:
                for row in csv.DictReader(fh):
                    if row.get("conversation_id"):
                        procesados.add(str(row["conversation_id"]))
    except Exception:
        return set()
    return procesados


def append_checkpoint(output_dir: Path, registro: Dict[str, Any]) -> None:
    checkpoint_csv, checkpoint_jsonl = checkpoint_paths(output_dir)
    liviano = {
        "conversation_id": registro.get("conversation_id", ""),
        "estado": registro.get("estado_proceso", ""),
        "fecha_hora": registro.get("fecha_hora", ""),
        "operador": registro.get("operador", ""),
        "wrapup": registro.get("wrapup", ""),
        "p2_8_resuelve_eficientemente": registro.get("p2_8_resuelve_eficientemente", ""),
        "justificacion_breve": registro.get("justificacion_breve", ""),
        "error": registro.get("error_proceso", "")
    }
    exists = checkpoint_csv.exists()
    with checkpoint_csv.open("a", newline="", encoding="utf-8-sig") as fh:
        writer = csv.DictWriter(fh, fieldnames=list(liviano.keys()))
        if not exists:
            writer.writeheader()
        writer.writerow(liviano)
    with checkpoint_jsonl.open("a", encoding="utf-8") as fh:
        fh.write(json.dumps(liviano, ensure_ascii=False) + "\n")


def pct(part: int, total: int) -> float:
    return round((part / total) * 100, 1) if total else 0.0


def generar_html(df: pd.DataFrame, output_dir: Path, ts: str) -> Path:
    total = len(df)
    counts = df["p2_8_resuelve_eficientemente"].value_counts().to_dict()
    top_ops = (
        df[df["p2_8_resuelve_eficientemente"] == "No"]
        .groupby("operador", dropna=False)
        .size()
        .sort_values(ascending=False)
        .head(10)
    )
    top_wrap = (
        df[df["p2_8_resuelve_eficientemente"] == "No"]
        .groupby("wrapup", dropna=False)
        .size()
        .sort_values(ascending=False)
        .head(10)
    )
    sample_no = df[df["p2_8_resuelve_eficientemente"] == "No"].head(20)

    def rows_from_series(series: pd.Series) -> str:
        return "".join(f"<tr><td>{html.escape(str(k or 'Sin dato'))}</td><td>{int(v):,}</td></tr>" for k, v in series.items())

    sample_rows = "".join(
        "<tr>"
        f"<td>{html.escape(str(r.get('conversation_id', '')))}</td>"
        f"<td>{html.escape(str(r.get('operador', '')))}</td>"
        f"<td>{html.escape(str(r.get('wrapup', '')))}</td>"
        f"<td>{html.escape(str(r.get('justificacion_breve', '')))}</td>"
        "</tr>"
        for _, r in sample_no.iterrows()
    )

    content = f"""<!doctype html>
<html><head><meta charset="utf-8"><title>Resumen P2.8</title>
<style>
body{{font-family:Arial,sans-serif;background:#f7f7f7;color:#333;margin:0;padding:24px}}
.wrap{{max-width:1100px;margin:auto;background:white;border:1px solid #ddd}}
.head{{background:#DA282D;color:white;padding:22px}}
.cards{{display:grid;grid-template-columns:repeat(4,1fr);gap:12px;padding:18px}}
.card{{background:#fafafa;border:1px solid #e5e5e5;padding:14px}}
.card b{{display:block;font-size:25px;color:#DA282D;margin-top:6px}}
.grid{{display:grid;grid-template-columns:1fr 1fr;gap:14px;padding:18px}}
.panel{{border:1px solid #e5e5e5;background:#fff;padding:14px}}
h3{{margin:0 0 10px;border-bottom:2px solid #DA282D;padding-bottom:6px}}
table{{width:100%;border-collapse:collapse;font-size:13px}}td,th{{border-bottom:1px solid #eee;padding:7px;text-align:left}}
.full{{grid-column:1 / -1}}
</style></head><body><div class="wrap">
<div class="head"><h1>Pregunta 2.8 - Resuelve eficientemente el problema del cliente</h1><p>Evaluacion local por reglas, sin IA pagada.</p></div>
<div class="cards">
<div class="card">Total<b>{total:,}</b></div>
<div class="card">Si<b>{counts.get('Si', 0):,} ({pct(counts.get('Si', 0), total)}%)</b></div>
<div class="card">No<b>{counts.get('No', 0):,} ({pct(counts.get('No', 0), total)}%)</b></div>
<div class="card">N/A<b>{counts.get('N/A', 0):,} ({pct(counts.get('N/A', 0), total)}%)</b></div>
</div>
<div class="grid">
<div class="panel"><h3>Top operadores con No</h3><table>{rows_from_series(top_ops)}</table></div>
<div class="panel"><h3>Top conclusiones con No</h3><table>{rows_from_series(top_wrap)}</table></div>
<div class="panel full"><h3>Muestra de llamadas con No</h3><table><tr><th>Conversation ID</th><th>Operador</th><th>Wrapup</th><th>Justificacion</th></tr>{sample_rows}</table></div>
</div>
</div></body></html>"""
    path = output_dir / f"Resumen_P2_8_{ts}.html"
    path.write_text(content, encoding="utf-8")
    return path


def generar_excel(df: pd.DataFrame, output_dir: Path, ts: str) -> Path:
    xlsx = output_dir / f"Evaluacion_P2_8_{ts}.xlsx"
    resumen = pd.DataFrame([
        {"Indicador": "Total llamadas analizadas", "Valor": len(df)},
        {"Indicador": "Si", "Valor": int((df["p2_8_resuelve_eficientemente"] == "Si").sum())},
        {"Indicador": "No", "Valor": int((df["p2_8_resuelve_eficientemente"] == "No").sum())},
        {"Indicador": "N/A", "Valor": int((df["p2_8_resuelve_eficientemente"] == "N/A").sum())},
        {"Indicador": "Requiere revision supervisor", "Valor": int((df["requiere_revision_supervisor"] == "SI").sum())}
    ])
    with pd.ExcelWriter(xlsx, engine="openpyxl") as writer:
        resumen.to_excel(writer, sheet_name="Resumen", index=False)
        df.to_excel(writer, sheet_name="Detalle P2.8", index=False)
    try:
        from openpyxl import load_workbook
        from openpyxl.styles import Font, PatternFill
        wb = load_workbook(xlsx)
        for ws in wb.worksheets:
            if ws.max_row:
                ws.freeze_panes = "A2"
                ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.fill = PatternFill("solid", fgColor="DA282D")
                cell.font = Font(color="FFFFFF", bold=True)
            for col in ws.columns:
                width = min(60, max(10, max(len(str(c.value or "")) for c in col) + 2))
                ws.column_dimensions[col[0].column_letter].width = width
        wb.save(xlsx)
    except Exception:
        pass
    return xlsx


def build_row_data(row: pd.Series, columnas: Dict[str, Optional[str]], idx: int, max_chars: int) -> Tuple[Dict[str, str], str]:
    row_data = {
        "conversation_id": valor(row, columnas.get("conversation_id"), f"row_{idx + 1}"),
        "fecha_hora": valor(row, columnas.get("fecha_hora")),
        "operador": valor(row, columnas.get("operador")),
        "cola": valor(row, columnas.get("cola")),
        "wrapup": valor(row, columnas.get("wrapup")),
        "cliente_id": valor(row, columnas.get("cliente_id")),
        "duracion": valor(row, columnas.get("duracion")),
        "temas_genesys": valor(row, columnas.get("temas_genesys"))
    }
    transcripcion = limpiar_transcripcion(valor(row, columnas.get("transcripcion")), max_chars)
    return row_data, transcripcion


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Evaluacion local P2.8 sin IA pagada.")
    parser.add_argument("--limite-filas", type=int, default=None, help="Procesar solo N filas para prueba.")
    parser.add_argument("--no-reanudar", action="store_true", help="Ignorar checkpoint y reprocesar todo.")
    parser.add_argument("--sin-ia", action="store_true", help="Compatibilidad: este script siempre corre sin IA pagada.")
    return parser.parse_args()


def main() -> int:
    logger = setup_logger()
    args = parse_args()
    pyflow_progress(1)
    try:
        parametros, prompt_cfg, matriz_cfg = cargar_config()
        run_cfg = load_run_config(parametros)
        if args.limite_filas is not None:
            run_cfg.limite_filas = args.limite_filas
        if args.no_reanudar:
            run_cfg.reanalizar_existentes = True

        OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
        logger.info("Inicio evaluacion local P2.8. No usa OpenAI ni APIs pagadas.")
        logger.info("Config cargada: %s | %s | %s", prompt_cfg.get("version", ""), matriz_cfg.get("version", ""), parametros.get("version", ""))

        files = listar_archivos_input()
        logger.info("Archivos detectados en input/imput: %s", len(files))
        procesados = set() if run_cfg.reanalizar_existentes else cargar_checkpoint(OUTPUT_DIR)
        resultados: List[Dict[str, Any]] = []
        evaluadas = 0
        omitidas = 0
        total_leidas = 0

        for file_idx, path in enumerate(files, start=1):
            if run_cfg.limite_filas > 0 and evaluadas >= run_cfg.limite_filas:
                break

            logger.info("Leyendo archivo %s/%s: %s", file_idx, len(files), path.name)
            df_raw = leer_archivo(path)
            columnas = resolver_columnas(df_raw, parametros)
            total_leidas += len(df_raw)

            if run_cfg.limite_filas > 0:
                faltantes = max(run_cfg.limite_filas - evaluadas, 0)
                df_raw = df_raw.head(faltantes)

            total_archivo = len(df_raw)
            logger.info("Filas detectadas en %s: %s", path.name, total_archivo)

            for idx, (_, row) in enumerate(df_raw.iterrows(), start=1):
                row_data, transcripcion = build_row_data(row, columnas, evaluadas, run_cfg.max_chars)
                cid = row_data["conversation_id"]
                if cid in procesados:
                    omitidas += 1
                    continue
                try:
                    evaluacion = evaluar_local(row_data, transcripcion, parametros)
                    registro = {
                        **row_data,
                        "p2_8_resuelve_eficientemente": evaluacion["p2_8_resuelve_eficientemente"],
                        "justificacion_breve": evaluacion["justificacion_breve"],
                        "motivo_cliente": evaluacion["motivo_cliente"],
                        "evidencia_cliente": evaluacion["evidencia_cliente"],
                        "evidencia_operador": evaluacion["evidencia_operador"],
                        "requiere_revision_supervisor": evaluacion["requiere_revision_supervisor"],
                        "tipo_interaccion": evaluacion["tipo_interaccion"],
                        "confianza": evaluacion["confianza"],
                        "estado_proceso": "OK",
                        "error_proceso": ""
                    }
                except Exception as exc:
                    registro = {
                        **row_data,
                        "p2_8_resuelve_eficientemente": "N/A",
                        "justificacion_breve": "Error de procesamiento. Revisar detalle tecnico.",
                        "motivo_cliente": "No evidenciado",
                        "evidencia_cliente": "No evidenciado",
                        "evidencia_operador": "No evidenciado",
                        "requiere_revision_supervisor": "SI",
                        "tipo_interaccion": "No contactado o no evaluable",
                        "confianza": "Baja",
                        "estado_proceso": "ERROR",
                        "error_proceso": str(exc)[:500]
                    }
                resultados.append(registro)
                append_checkpoint(OUTPUT_DIR, registro)
                procesados.add(cid)
                evaluadas += 1

                if evaluadas == 1 or evaluadas % max(run_cfg.batch_size, 1) == 0 or idx == total_archivo:
                    if run_cfg.limite_filas > 0:
                        progress = 5 + int((evaluadas / max(run_cfg.limite_filas, 1)) * 85)
                    else:
                        base = (file_idx - 1) / max(len(files), 1)
                        part = (idx / max(total_archivo, 1)) / max(len(files), 1)
                        progress = 5 + int((base + part) * 85)
                    logger.info(
                        "Avance: evaluadas=%s | omitidas=%s | archivo=%s/%s | nuevas=%s",
                        evaluadas,
                        omitidas,
                        file_idx,
                        len(files),
                        len(resultados),
                    )
                    pyflow_progress(min(progress, 95))

                if run_cfg.limite_filas > 0 and evaluadas >= run_cfg.limite_filas:
                    break

            del df_raw

        logger.info("Filas leidas desde archivos: %s", total_leidas)
        logger.info("Filas evaluadas en esta ejecucion: %s", evaluadas)
        logger.info("Filas omitidas por checkpoint: %s", omitidas)

        resultado_df = pd.DataFrame(resultados)
        if resultado_df.empty:
            logger.info("No hubo registros nuevos; todos estaban en checkpoint.")
            return 0

        detalle_cols = [
            "conversation_id", "fecha_hora", "operador", "cola", "wrapup", "cliente_id",
            "p2_8_resuelve_eficientemente", "justificacion_breve", "motivo_cliente",
            "evidencia_cliente", "evidencia_operador", "requiere_revision_supervisor",
            "tipo_interaccion", "confianza", "estado_proceso", "error_proceso"
        ]
        for col in detalle_cols:
            if col not in resultado_df.columns:
                resultado_df[col] = ""
        resultado_df = resultado_df[detalle_cols]

        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        out_csv = OUTPUT_DIR / f"Evaluacion_P2_8_{ts}.csv"
        resultado_df.to_csv(out_csv, index=False, encoding="utf-8-sig")
        out_xlsx = generar_excel(resultado_df, OUTPUT_DIR, ts)
        out_html = generar_html(resultado_df, OUTPUT_DIR, ts)
        pyflow_progress(100)

        logger.info("Listo. CSV: %s", out_csv)
        logger.info("Excel: %s", out_xlsx)
        logger.info("HTML: %s", out_html)
        logger.info("Checkpoint liviano: %s", checkpoint_paths(OUTPUT_DIR)[0])
        return 0
    except Exception as exc:
        logger.exception("Error general: %s", exc)
        pyflow_progress(100)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
