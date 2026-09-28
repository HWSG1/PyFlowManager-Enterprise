# -*- coding: utf-8 -*-
"""Genesys Cloud -> atenciones omnicanal -> Excel / SAP HANA.

Dependencias: requests, python-dotenv (opcional), tzdata en Windows,
openpyxl (Excel), hdbcli (HANA). No depende de otros scripts ni del Excel modelo.

Columnas y tipos incorporados desde Modelo_Genesys_HANA_Tablas.xlsx.
Los filtros seleccionan CONVERSACIONES: se exporta su secuencia completa de
atenciones, incluso en otros días/medios, para preservar transferencias y la
clave conversationId + sequence al reprocesar. No se recortan las atenciones
al intervalo del job. No se extraen transcripciones, grabaciones ni datos IVR.

Una atención = episodio interactivo de una sesión de agente hasta su conclusión.
Hold/delay no abren episodios. Si no hay agente, una atención automática por
medio resume la interacción; nunca una fila por cada flow ni por ACD/IVR.

--self-test ejecuta únicamente pruebas sintéticas sin red, credenciales ni HANA.
Sin --self-test, el modo de salida por defecto es SAP HANA; --dry-run impide
cualquier escritura HANA (las lecturas de metadatos/catálogos sí son posibles).
"""
from __future__ import annotations

import argparse
import json
import logging
import os
import re
import sys
import time
import math
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from email.utils import parsedate_to_datetime
from pathlib import Path
from typing import Any, Iterator
from urllib.parse import quote
from zoneinfo import ZoneInfo

try:
    from dotenv import load_dotenv
except ImportError:
    def load_dotenv(*args: Any, **kwargs: Any) -> bool:
        return False

PYFLOW_PARAMS = {'GENESYS_CLIENT_ID': {'type': 'global',
                       'global_key': 'GENESYS_CLIENT_ID',
                       'label': 'Genesys Client ID',
                       'required': True},
 'GENESYS_CLIENT_SECRET': {'type': 'global',
                           'global_key': 'GENESYS_CLIENT_SECRET',
                           'label': 'Genesys Client Secret',
                           'required': True,
                           'secret': True},
 'GENESYS_REGION': {'type': 'global',
                    'global_key': 'GENESYS_REGION',
                    'label': 'Genesys Region / Domain',
                    'required': True},
 'START_DATE': {'type': 'date', 'label': 'Fecha inicial local', 'required': False},
 'END_DATE': {'type': 'date', 'label': 'Fecha final local', 'required': False},
 'DAYS_BACK': {'type': 'number',
               'label': 'Días hacia atrás si no se indican fechas',
               'required': False,
               'default': '1'},
 'GENESYS_TIMEZONE': {'type': 'text',
                      'label': 'Zona horaria Genesys',
                      'required': True,
                      'default': 'America/Tegucigalpa'},
 'ORIGINAL_DIRECTION': {'type': 'select',
                        'label': 'Dirección original',
                        'required': False,
                        'options': ['ambas', 'inbound', 'outbound'],
                        'default': 'ambas'},
 'FLOW_SELECTION_ID': {'type': 'genesys_flow', 'label': 'Nombre de flujo (buscar y seleccionar)', 'required': False},
 'FLOW_ID': {'type': 'tags', 'label': 'Flow ID (prioridad sobre selector)', 'required': False},
 'MEDIA_TYPE': {'type': 'select',
                'label': 'Tipo de medio (voice = voz, message = mensaje)',
                'required': False,
                'options': ['todos',
                            'voice',
                            'message',
                            'email',
                            'chat',
                            'callback',
                            'cobrowse',
                            'internalmessage',
                            'screenmonitoring',
                            'screenshare',
                            'video',
                            'unknown'],
                'default': 'todos'},
 'PARTICIPANT_PURPOSE': {'type': 'select',
                         'label': 'Segmento / propósito del participante',
                         'required': False,
                         'options': ['todos',
                                     'acd',
                                     'agent',
                                     'api',
                                     'botflow',
                                     'campaign',
                                     'customer',
                                     'dialer',
                                     'external',
                                     'fax',
                                     'group',
                                     'inbound',
                                     'ivr',
                                     'manual',
                                     'outbound',
                                     'station',
                                     'user',
                                     'voicemail',
                                     'voicesurveyflow',
                                     'workflow'],
                         'default': 'todos'},
 'CONVERSATION_ID': {'type': 'tags', 'label': 'Conversation ID específico', 'required': False},
 'USER_ID': {'type': 'tags', 'label': 'User ID del agente', 'required': False},
 'USER_NAME': {'type': 'genesys_users', 'label': 'Nombre/correo del agente', 'required': False},
 'QUEUE_ID': {'type': 'tags', 'label': 'Queue ID', 'required': False},
 'QUEUE_NAME': {'type': 'genesys_queues', 'label': 'Nombre de cola', 'required': False},
 'CAMPAIGN_ID': {'type': 'tags', 'label': 'Campaign ID', 'required': False},
 'CAMPAIGN_NAME': {'type': 'genesys_campaigns', 'label': 'Nombre de campaña', 'required': False},
 'CONTACT_LIST_ID': {'type': 'tags', 'label': 'Contact List ID', 'required': False},
 'CONTACT_LIST_NAME': {'type': 'genesys_contactlists', 'label': 'Nombre lista de contacto', 'required': False},
 'WRAPUP_CODE_ID': {'type': 'tags', 'label': 'WrapUpCode ID opcional', 'required': False},
 'WRAPUP_CODE_NAME': {'type': 'genesys_wrapupcodes', 'label': 'Nombre de conclusión opcional', 'required': False},
 'MAX_CONVERSATIONS': {'type': 'number', 'label': 'Máximo conversaciones; vacío = todas', 'required': False},
 'HPR_HOST': {'type': 'global', 'global_key': 'HPR_HOST', 'label': 'HPR_HOST', 'required': False},
 'HPR_HOST_ESPEJO': {'type': 'global', 'global_key': 'HPR_HOST_ESPEJO', 'label': 'HPR_HOST_ESPEJO', 'required': False},
 'HPR_PORT': {'type': 'global', 'global_key': 'HPR_PORT', 'label': 'HPR_PORT', 'required': False},
 'HPR_USER': {'type': 'global', 'global_key': 'HPR_USER', 'label': 'HPR_USER', 'required': False},
 'HPR_PASSWORD': {'type': 'global',
                  'global_key': 'HPR_PASSWORD',
                  'label': 'HPR_PASSWORD',
                  'required': False,
                  'secret': True},
 'HANA_SCHEMA': {'type': 'text', 'label': 'HANA_SCHEMA', 'required': False, 'default': 'BI_SS'},
 'HANA_MAIN_TABLE': {'type': 'text', 'label': 'HANA_MAIN_TABLE', 'required': False, 'default': 'GNS_API_INTERACCIONES'},
 'HANA_VOICE_TABLE': {'type': 'text',
                      'label': 'HANA_VOICE_TABLE',
                      'required': False,
                      'default': 'GNS_API_INTERACCIONES_VOICE'},
 'HANA_DIGITAL_TABLE': {'type': 'text',
                        'label': 'HANA_DIGITAL_TABLE',
                        'required': False,
                        'default': 'GNS_API_INTERACCIONES_DIGITAL'},
 'HANA_EMAIL_TABLE': {'type': 'text',
                      'label': 'HANA_EMAIL_TABLE',
                      'required': False,
                      'default': 'GNS_API_INTERACCIONES_EMAIL'},
 'HANA_USERS_TABLE': {'type': 'text', 'label': 'HANA_USERS_TABLE', 'required': False, 'default': 'GNS_API_USUARIOS'},
 'HANA_QUEUES_TABLE': {'type': 'text', 'label': 'HANA_QUEUES_TABLE', 'required': False, 'default': 'GNS_API_COLAS'},
 'HANA_WRAPUPS_TABLE': {'type': 'text',
                        'label': 'HANA_WRAPUPS_TABLE',
                        'required': False,
                        'default': 'GNS_API_CAT_CONCLUSIONES'},
 'OUTPUT_DESTINATION': {'type': 'select',
                        'label': 'OUTPUT_DESTINATION',
                        'options': ['Excel', 'SAP HANA', 'Excel + SAP HANA'],
                        'default': 'SAP HANA',
                        'required': False},
 'DIRECTION': {'type': 'select',
               'label': 'DIRECTION',
               'options': ['ambas', 'inbound', 'outbound'],
               'default': 'ambas',
               'required': False},
 'ROUTING_TYPE': {'type': 'select',
                  'label': 'ROUTING_TYPE',
                  'options': ['todos', 'Predictive', 'Standard', 'Bullseye', 'Last', 'Preferred'],
                  'default': 'todos',
                  'required': False},
 'TRANSFER_FILTER': {'type': 'select',
                     'label': 'TRANSFER_FILTER',
                     'options': ['todas', 'solo_transferidas', 'no_transferidas'],
                     'default': 'todas',
                     'required': False},
 'OUTPUT_DIR': {'type': 'text', 'label': 'Carpeta de salida Excel', 'required': False},
 'DRY_RUN': {'type': 'boolean', 'label': 'Simular (sin escrituras HANA)', 'default': 'false', 'required': False},
 'MAX_IDENTIFICATION_REQUESTS': {'type': 'number',
                                 'label': 'Máximo consultas adicionales de identificación por ejecución',
                                 'default': '1000',
                                 'required': False},
 'HANA_BATCH_SIZE': {'type': 'number',
                     'label': 'Conversaciones por transacción HANA',
                     'default': '100',
                     'required': False}}

# (nombre exacto, tipo HANA, nullable, fuente documentada).
SCHEMAS = {'INTERACCIONES': [('conversationId', 'NVARCHAR(36)', False, 'conversation.conversationId'),
                   ('conversationStart', 'TIMESTAMP', False, 'conversation.conversationStart'),
                   ('conversationEnd', 'TIMESTAMP', True, 'conversation.conversationEnd'),
                   ('originatingDirection', 'NVARCHAR(30)', True, 'conversation.originatingDirection'),
                   ('conversationInitiator', 'NVARCHAR(50)', True, 'conversation.conversationInitiator'),
                   ('customerParticipation', 'BOOLEAN', True, 'conversation.customerParticipation'),
                   ('externalTag', 'NVARCHAR(1000)', True, 'conversation.externalTag'),
                   ('externalContactId', 'NVARCHAR(36)', True, 'participants[].externalContactId'),
                   ('mediaType', 'NVARCHAR(30)', False, 'participants[].sessions[].mediaType'),
                   ('direction', 'NVARCHAR(30)', True, 'participants[].sessions[].direction'),
                   ('provider', 'NVARCHAR(100)', True, 'participants[].sessions[].provider'),
                   ('participantId', 'NVARCHAR(36)', True, 'participants[].participantId'),
                   ('participantName', 'NVARCHAR(250)', True, 'participants[].participantName'),
                   ('purpose', 'NVARCHAR(50)', True, 'participants[].purpose'),
                   ('userId', 'NVARCHAR(36)', True, 'participants[].userId'),
                   ('selectedAgentId', 'NVARCHAR(36)', True, 'participants[].sessions[].selectedAgentId'),
                   ('sessionId', 'NVARCHAR(36)', True, 'participants[].sessions[].sessionId'),
                   ('peerId', 'NVARCHAR(36)', True, 'participants[].sessions[].peerId'),
                   ('queueId', 'NVARCHAR(36)', True, 'participants[].sessions[].segments[].queueId'),
                   ('segmentStart', 'TIMESTAMP', True, 'participants[].sessions[].segments[].segmentStart'),
                   ('segmentEnd', 'TIMESTAMP', True, 'participants[].sessions[].segments[].segmentEnd'),
                   ('segmentType', 'NVARCHAR(1000)', True, 'participants[].sessions[].segments[].segmentType'),
                   ('disconnectType', 'NVARCHAR(500)', True, 'participants[].sessions[].segments[].disconnectType'),
                   ('wrapUpCode', 'NVARCHAR(36)', True, 'participants[].sessions[].segments[].wrapUpCode'),
                   ('wrapUpNote', 'NCLOB', True, 'participants[].sessions[].segments[].wrapUpNote'),
                   ('flowId', 'NVARCHAR(5000)', True, 'participants[].sessions[].flow.flowId'),
                   ('flowName', 'NVARCHAR(5000)', True, 'participants[].sessions[].flow.flowName'),
                   ('flowType', 'NVARCHAR(2000)', True, 'participants[].sessions[].flow.flowType'),
                   ('flowVersion', 'NVARCHAR(2000)', True, 'participants[].sessions[].flow.flowVersion'),
                   ('exitReason', 'NVARCHAR(2000)', True, 'participants[].sessions[].flow.exitReason'),
                   ('transferType', 'NVARCHAR(1000)', True, 'participants[].sessions[].flow.transferType'),
                   ('transferTargetAddress', 'NVARCHAR(5000)', True, 'participants[].sessions[].flow.transferTargetAddress'),
                   ('transferTargetName', 'NVARCHAR(5000)', True, 'participants[].sessions[].flow.transferTargetName'),
                   ('SPD_IDENTIFICACION', 'NVARCHAR(30)', True, 'participants[].attributes.SPD_IDENTIFICACION'),
                   ('Identificacion', 'NVARCHAR(30)', True, 'participants[].attributes.Identificacion'),
                   ('ScripterIdentificacion', 'NVARCHAR(30)', True, 'participants[].attributes.ScripterIdentificacion'),
                   ('spd_hasRecentContact', 'NVARCHAR(20)', True, 'participants[].attributes.spd_hasRecentContact'),
                   ('SPD_Gestion', 'NVARCHAR(1000)', True, 'participants[].attributes.SPD_Gestion'),
                   ('SPD_Gestion_FH', 'NVARCHAR(100)', True, 'participants[].attributes.SPD_Gestion_FH'),
                   ('tAnswered', 'BIGINT', True, "participants[].sessions[].metrics[name='tAnswered'].value"),
                   ('tAcd', 'BIGINT', True, "participants[].sessions[].metrics[name='tAcd'].value"),
                   ('tTalkComplete', 'BIGINT', True, "participants[].sessions[].metrics[name='tTalkComplete'].value"),
                   ('tHeldComplete', 'BIGINT', True, "participants[].sessions[].metrics[name='tHeldComplete'].value"),
                   ('tAcw', 'BIGINT', True, "participants[].sessions[].metrics[name='tAcw'].value"),
                   ('tHandle', 'BIGINT', True, "participants[].sessions[].metrics[name='tHandle'].value"),
                   ('tFlow', 'BIGINT', True, "participants[].sessions[].metrics[name='tFlow'].value"),
                   ('divisionIds', 'NVARCHAR(5000)', True, 'conversation.divisionIds'),
                   ('teamId', 'NVARCHAR(36)', True, 'participants[].teamId'),
                   ('requestedRoutings', 'NVARCHAR(1000)', True, 'participants[].sessions[].requestedRoutings'),
                   ('usedRouting', 'NVARCHAR(100)', True, 'participants[].sessions[].usedRouting'),
                   ('eligibleAgentCounts', 'NVARCHAR(2000)', True, 'participants[].sessions[].eligibleAgentCounts'),
                   ('remote', 'NVARCHAR(500)', True, 'participants[].sessions[].remote'),
                   ('recording', 'BOOLEAN', True, 'participants[].sessions[].recording'),
                   ('requestedRoutingUserIds', 'NVARCHAR(5000)', True, 'participants[].sessions[].segments[].requestedRoutingUserIds'),
                   ('errorCode', 'NVARCHAR(500)', True, 'participants[].sessions[].segments[].errorCode'),
                   ('recognitionFailureReason', 'NVARCHAR(500)', True, 'participants[].sessions[].flow.recognitionFailureReason'),
                   ('nConnected', 'BIGINT', True, "participants[].sessions[].metrics[name='nConnected'].value"),
                   ('nOffered', 'BIGINT', True, "participants[].sessions[].metrics[name='nOffered'].value"),
                   ('tConnected', 'BIGINT', True, "participants[].sessions[].metrics[name='tConnected'].value"),
                   ('tAgentResponseTime', 'BIGINT', True, "participants[].sessions[].metrics[name='tAgentResponseTime'].value"),
                   ('tAlert', 'BIGINT', True, "participants[].sessions[].metrics[name='tAlert'].value"),
                   ('nOverSla', 'BIGINT', True, "participants[].sessions[].metrics[name='nOverSla'].value"),
                   ('nBlindTransferred', 'BIGINT', True, "participants[].sessions[].metrics[name='nBlindTransferred'].value"),
                   ('nTransferred', 'BIGINT', True, "participants[].sessions[].metrics[name='nTransferred'].value"),
                   ('nBotInteractions', 'BIGINT', True, "participants[].sessions[].metrics[name='nBotInteractions'].value"),
                   ('nError', 'BIGINT', True, "participants[].sessions[].metrics[name='nError'].value"),
                   ('tAbandon', 'BIGINT', True, "participants[].sessions[].metrics[name='tAbandon'].value"),
                   ('SPD_ANI', 'NVARCHAR(100)', True, 'participants[].attributes.SPD_ANI'),
                   ('SPD_Comentario', 'NCLOB', True, 'participants[].attributes.SPD_Comentario'),
                   ('ScripterComentario', 'NCLOB', True, 'participants[].attributes.ScripterComentario'),
                   ('ScripterAutenticacion', 'NVARCHAR(250)', True, 'participants[].attributes.ScripterAutenticacion'),
                   ('ScripterNombre', 'NVARCHAR(500)', True, 'participants[].attributes.ScripterNombre'),
                   ('ScripterCorreo', 'NVARCHAR(500)', True, 'participants[].attributes.ScripterCorreo'),
                   ('knowledgeBaseIds', 'NVARCHAR(5000)', True, 'conversation.knowledgeBaseIds'),
                   ('selfServed', 'BOOLEAN', True, 'conversation.selfServed'),
                   ('screenRecording', 'BOOLEAN', True, 'participants[].screenRecording'),
                   ('remoteNameDisplayable', 'NVARCHAR(500)', True, 'participants[].sessions[].remoteNameDisplayable'),
                   ('flowInType', 'NVARCHAR(100)', True, 'participants[].sessions[].flowInType'),
                   ('flowOutType', 'NVARCHAR(100)', True, 'participants[].sessions[].flowOutType'),
                   ('journeyCustomerId', 'NVARCHAR(250)', True, 'participants[].sessions[].journeyCustomerId'),
                   ('journeyCustomerIdType', 'NVARCHAR(100)', True, 'participants[].sessions[].journeyCustomerIdType'),
                   ('journeyCustomerSessionId', 'NVARCHAR(250)', True, 'participants[].sessions[].journeyCustomerSessionId'),
                   ('journeyCustomerSessionIdType', 'NVARCHAR(100)', True, 'participants[].sessions[].journeyCustomerSessionIdType'),
                   ('agentBullseyeRing', 'INTEGER', True, 'participants[].sessions[].agentBullseyeRing'),
                   ('routingRing', 'INTEGER', True, 'participants[].sessions[].routingRing'),
                   ('routingRule', 'NVARCHAR(250)', True, 'participants[].sessions[].routingRule'),
                   ('routingRuleType', 'NVARCHAR(100)', True, 'participants[].sessions[].routingRuleType'),
                   ('selectedAgentRank', 'INTEGER', True, 'participants[].sessions[].selectedAgentRank'),
                   ('proposedAgents', 'NCLOB', True, 'participants[].sessions[].proposedAgents'),
                   ('waitingInteractionCounts', 'NVARCHAR(2000)', True, 'participants[].sessions[].waitingInteractionCounts'),
                   ('startingLanguage', 'NVARCHAR(50)', True, 'participants[].sessions[].flow.startingLanguage'),
                   ('endingLanguage', 'NVARCHAR(50)', True, 'participants[].sessions[].flow.endingLanguage'),
                   ('entryReason', 'NVARCHAR(1000)', True, 'participants[].sessions[].flow.entryReason'),
                   ('entryType', 'NVARCHAR(100)', True, 'participants[].sessions[].flow.entryType'),
                   ('outcomes', 'NCLOB', True, 'participants[].sessions[].flow.outcomes'),
                   ('nFlow', 'BIGINT', True, "participants[].sessions[].metrics[name='nFlow'].value"),
                   ('nFlowOutcome', 'BIGINT', True, "participants[].sessions[].metrics[name='nFlowOutcome'].value"),
                   ('nFlowOutcomeFailed', 'BIGINT', True, "participants[].sessions[].metrics[name='nFlowOutcomeFailed'].value"),
                   ('tFlowOutcome', 'BIGINT', True, "participants[].sessions[].metrics[name='tFlowOutcome'].value"),
                   ('tFlowExit', 'BIGINT', True, "participants[].sessions[].metrics[name='tFlowExit'].value"),
                   ('tFlowDisconnect', 'BIGINT', True, "participants[].sessions[].metrics[name='tFlowDisconnect'].value"),
                   ('sequence', 'INTEGER', False, 'DERIVADO ETL'),
                   ('transferSequence', 'INTEGER', True, 'DERIVADO ETL'),
                   ('queueName', 'NVARCHAR(250)', True, 'GNS_API_COLAS'),
                   ('wrapUpCodeName', 'NVARCHAR(500)', True, 'GNS_API_CAT_CONCLUSIONES'),
                   ('userName', 'NVARCHAR(250)', True, 'GNS_API_USUARIOS'),
                   ('selectedAgentName', 'NVARCHAR(250)', True, 'GNS_API_USUARIOS'),
                   ('transferToUserName', 'NVARCHAR(250)', True, 'GNS_API_USUARIOS'),
                   ('transferToQueueName', 'NVARCHAR(250)', True, 'GNS_API_COLAS'),
                   ('fechaCarga', 'TIMESTAMP', False, 'ETL')],
 'INTERACCIONES_VOICE': [('conversationId', 'NVARCHAR(36)', False, 'conversation.conversationId'),
                         ('sequence', 'INTEGER', False, 'DERIVADO ETL'),
                         ('ani', 'NVARCHAR(100)', True, 'participants[].sessions[].ani'),
                         ('dnis', 'NVARCHAR(100)', True, 'participants[].sessions[].dnis'),
                         ('outboundCampaignId', 'NVARCHAR(36)', True, 'participants[].sessions[].outboundCampaignId'),
                         ('outboundContactId', 'NVARCHAR(250)', True, 'participants[].sessions[].outboundContactId'),
                         ('outboundContactListId', 'NVARCHAR(36)', True, 'participants[].sessions[].outboundContactListId'),
                         ('callbackNumbers', 'NVARCHAR(2000)', True, 'participants[].sessions[].callbackNumbers'),
                         ('callbackScheduledTime', 'TIMESTAMP', True, 'participants[].sessions[].callbackScheduledTime'),
                         ('destinationAddresses', 'NVARCHAR(5000)', True, 'participants[].sessions[].destinationAddresses'),
                         ('dialerCampaignId', 'NVARCHAR(36)', True, 'participants[].attributes.dialerCampaignId'),
                         ('dialerContactListId', 'NVARCHAR(36)', True, 'participants[].attributes.dialerContactListId'),
                         ('dialerContactId', 'NVARCHAR(250)', True, 'participants[].attributes.dialerContactId'),
                         ('mediaStatsMinConversationMos', 'DECIMAL(12,6)', True, 'conversation.mediaStatsMinConversationMos'),
                         ('mediaStatsMinConversationRFactor', 'DECIMAL(12,6)', True, 'conversation.mediaStatsMinConversationRFactor'),
                         ('edgeId', 'NVARCHAR(36)', True, 'participants[].sessions[].edgeId'),
                         ('protocolCallId', 'NVARCHAR(500)', True, 'participants[].sessions[].protocolCallId'),
                         ('sessionDnis', 'NVARCHAR(1000)', True, 'participants[].sessions[].sessionDnis'),
                         ('callbackUserName', 'NVARCHAR(500)', True, 'participants[].sessions[].callbackUserName'),
                         ('scriptId', 'NVARCHAR(36)', True, 'participants[].sessions[].scriptId'),
                         ('timeoutSeconds', 'INTEGER', True, 'participants[].sessions[].timeoutSeconds'),
                         ('skipEnabled', 'BOOLEAN', True, 'participants[].sessions[].skipEnabled'),
                         ('sipResponseCodes', 'NVARCHAR(1000)', True, 'participants[].sessions[].segments[].sipResponseCodes'),
                         ('q850ResponseCodes', 'NVARCHAR(1000)', True, 'participants[].sessions[].segments[].q850ResponseCodes'),
                         ('issuedCallback', 'BOOLEAN', True, 'participants[].sessions[].flow.issuedCallback'),
                         ('tActiveCallback', 'BIGINT', True, "participants[].sessions[].metrics[name='tActiveCallback'].value"),
                         ('tActiveCallbackComplete', 'BIGINT', True, "participants[].sessions[].metrics[name='tActiveCallbackComplete'].value"),
                         ('nOutbound', 'BIGINT', True, "participants[].sessions[].metrics[name='nOutbound'].value"),
                         ('tContacting', 'BIGINT', True, "participants[].sessions[].metrics[name='tContacting'].value"),
                         ('tDialing', 'BIGINT', True, "participants[].sessions[].metrics[name='tDialing'].value"),
                         ('nOutboundAttempted', 'BIGINT', True, "participants[].sessions[].metrics[name='nOutboundAttempted'].value"),
                         ('nOutboundConnected', 'BIGINT', True, "participants[].sessions[].metrics[name='nOutboundConnected'].value"),
                         ('tFirstConnect', 'BIGINT', True, "participants[].sessions[].metrics[name='tFirstConnect'].value"),
                         ('tFirstDial', 'BIGINT', True, "participants[].sessions[].metrics[name='tFirstDial'].value"),
                         ('fechaCarga', 'TIMESTAMP', False, 'ETL')],
 'INTERACCIONES_DIGITAL': [('conversationId', 'NVARCHAR(36)', False, 'conversation.conversationId'),
                           ('sequence', 'INTEGER', False, 'DERIVADO ETL'),
                           ('messageType', 'NVARCHAR(100)', True, 'participants[].sessions[].messageType'),
                           ('addressFrom', 'NVARCHAR(1000)', True, 'participants[].sessions[].addressFrom'),
                           ('addressTo', 'NVARCHAR(1000)', True, 'participants[].sessions[].addressTo'),
                           ('oMessageCount', 'BIGINT', True, "participants[].sessions[].metrics[name='oMessageCount'].value"),
                           ('addressOther', 'NVARCHAR(1000)', True, 'participants[].sessions[].addressOther'),
                           ('addressSelf', 'NVARCHAR(1000)', True, 'participants[].sessions[].addressSelf'),
                           ('SPD_ABANDONO', 'NVARCHAR(100)', True, 'participants[].attributes.SPD_ABANDONO'),
                           ('SPD_AUTOGESTION', 'NVARCHAR(100)', True, 'participants[].attributes.SPD_AUTOGESTION'),
                           ('SPD_tipoDeCanal', 'NVARCHAR(250)', True, 'participants[].attributes.SPD_tipoDeCanal'),
                           ('SPD_QUEUE', 'NVARCHAR(250)', True, 'participants[].attributes.SPD_QUEUE'),
                           ('SPD_redSocial', 'NVARCHAR(250)', True, 'participants[].attributes.SPD_redSocial'),
                           ('channel', 'NVARCHAR(250)', True, 'participants[].attributes.channel'),
                           ('itxds_queue', 'NVARCHAR(250)', True, 'participants[].attributes.itxds_queue'),
                           ('itxds_channel', 'NVARCHAR(250)', True, 'participants[].attributes.itxds_channel'),
                           ('mediaCount', 'INTEGER', True, 'participants[].sessions[].mediaCount'),
                           ('fechaCarga', 'TIMESTAMP', False, 'ETL')],
 'INTERACCIONES_EMAIL': [('conversationId', 'NVARCHAR(36)', False, 'conversation.conversationId'),
                         ('sequence', 'INTEGER', False, 'DERIVADO ETL'),
                         ('addressFrom', 'NVARCHAR(1000)', True, 'participants[].sessions[].addressFrom'),
                         ('addressTo', 'NVARCHAR(1000)', True, 'participants[].sessions[].addressTo'),
                         ('subject', 'NVARCHAR(2000)', True, 'participants[].sessions[].segments[].subject'),
                         ('addressOther', 'NVARCHAR(1000)', True, 'participants[].sessions[].addressOther'),
                         ('addressSelf', 'NVARCHAR(1000)', True, 'participants[].sessions[].addressSelf'),
                         ('SDP_Asunto_Correo', 'NVARCHAR(1000)', True, 'participants[].attributes.SDP_Asunto_Correo'),
                         ('cc', 'NVARCHAR(5000)', True, 'participants[].sessions[].cc'),
                         ('fechaCarga', 'TIMESTAMP', False, 'ETL')]}

LOG = logging.getLogger("gns_interacciones_omnicanal")
UTC = timezone.utc
FILTER_SELECTORS = {
    "USER_ID": ("USER_NAME", "/api/v2/users", "users"),
    "QUEUE_ID": ("QUEUE_NAME", "/api/v2/routing/queues", "queues"),
    "WRAPUP_CODE_ID": ("WRAPUP_CODE_NAME", "/api/v2/routing/wrapupcodes", "wrapups"),
    "FLOW_ID": ("FLOW_SELECTION_ID", "/api/v2/flows", "flows"),
    "CAMPAIGN_ID": ("CAMPAIGN_NAME", "/api/v2/outbound/campaigns", "campaigns"),
    "CONTACT_LIST_ID": ("CONTACT_LIST_NAME", "/api/v2/outbound/contactlists", "contactlists"),
}
TABLE_KEYS = dict(zip(SCHEMAS, ("HANA_MAIN_TABLE", "HANA_VOICE_TABLE", "HANA_DIGITAL_TABLE", "HANA_EMAIL_TABLE")))
IDENTIFICATION_KEYS = ("SPD_IDENTIFICACION", "Identificacion", "ScripterIdentificacion")


def env_str(name: str, default: str = "") -> str:
    value = str(os.environ.get(name, default)).strip()
    return default if value.lower() in ("", "none", "null", "undefined") else value


def split_filter_values(value: Any) -> list[str]:
    if value is None:
        return []
    return list(dict.fromkeys(part.strip() for part in re.split(r"[;,\r\n]+", str(value)) if part.strip()))


def join_filter_values(values: list[str]) -> str:
    return ";".join(dict.fromkeys(values))


def parse_selection(value: str) -> list[str] | None:
    """None indica un nombre antiguo; JSON inválido nunca amplía silenciosamente el filtro."""
    if not value.strip().startswith(("[", "{")):
        return None
    items = json.loads(value)
    items = [items] if isinstance(items, dict) else items
    if not isinstance(items, list) or any(not isinstance(item, dict) or not isinstance(item.get("id"), str)
                                          or not item["id"].strip() for item in items):
        raise ValueError("Selector JSON inválido: se esperaba una lista de objetos con id.")
    return list(dict.fromkeys(item["id"].strip() for item in items))


def identifier(value: str) -> str:
    if not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", value):
        raise ValueError("Identificador SQL no válido.")
    return '"' + value + '"'


def parse_date(value: str) -> date:
    for fmt in ("%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(value, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"Fecha inválida: {value}. Use YYYY-MM-DD o DD/MM/YYYY.")


def timestamp(value: Any) -> datetime | None:
    if value in (None, ""):
        return None
    result = value if isinstance(value, datetime) else datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    if result.tzinfo is None:
        raise ValueError("Timestamp Genesys sin zona horaria.")
    return result.astimezone(UTC)


def utc_text(value: datetime) -> str:
    return value.astimezone(UTC).isoformat(timespec="milliseconds").replace("+00:00", "Z")


def day_interval(day: date, zone: str) -> tuple[datetime, datetime]:
    tz = ZoneInfo(zone)
    start = datetime.combine(day, datetime.min.time(), tzinfo=tz)
    end = datetime.combine(day + timedelta(days=1), datetime.min.time(), tzinfo=tz)
    return start.astimezone(UTC), end.astimezone(UTC)


@dataclass
class Config:
    values: dict[str, str]
    first_day: date
    last_day: date
    zone: str = "America/Tegucigalpa"
    output: str = "hana"
    dry_run: bool = False
    filters: dict[str, list[str]] = field(default_factory=dict)
    timeout: int = 60
    retries: int = 5
    poll_seconds: float = 5
    max_polls: int = 720
    page_size: int = 500
    api_sleep: float = 0.1
    max_conversations: int = 0
    max_identification_requests: int = 1000
    batch_size: int = 100

    def get(self, key: str, default: str = "") -> str:
        return self.values.get(key, default)


def load_config(argv: list[str] | None = None) -> Config:
    load_dotenv()
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    for name in ("start-date", "end-date", "days-back", "media-type", "original-direction", "direction",
                 "participant-purpose", "user-id", "queue-id", "wrapup-code-id", "flow-id", "campaign-id",
                 "contact-list-id", "conversation-id", "routing-type", "transfer-filter", "output-dir",
                 "max-conversations"):
        parser.add_argument("--" + name)
    parser.add_argument("--output", choices=("excel", "hana", "both"))
    parser.add_argument("--dry-run", action="store_true", default=None)
    parser.add_argument("--self-test", action="store_true", help="Pruebas sin conexiones externas")
    args = parser.parse_args(argv)
    values = {key: env_str(key, str(meta.get("default", ""))) for key, meta in PYFLOW_PARAMS.items()}
    # Días relativos explícitos reemplazan un rango guardado en el entorno.
    if args.days_back is not None and args.start_date is None and args.end_date is None:
        values.update(START_DATE="", END_DATE="")
    for key, value in vars(args).items():
        if value is not None and key not in ("output", "self_test"):
            values[key.upper()] = str(value)
    zone = values["GENESYS_TIMEZONE"] or "America/Tegucigalpa"
    ZoneInfo(zone)
    if values["START_DATE"] or values["END_DATE"]:
        if not (values["START_DATE"] and values["END_DATE"]):
            raise ValueError("START_DATE y END_DATE deben informarse juntos.")
        first, last = parse_date(values["START_DATE"]), parse_date(values["END_DATE"])
    else:
        days = int(values["DAYS_BACK"] or "1")
        if days < 1:
            raise ValueError("DAYS_BACK debe ser mayor que cero.")
        last = datetime.now(ZoneInfo(zone)).date() - timedelta(days=1)
        first = last - timedelta(days=days - 1)
    if last < first:
        raise ValueError("END_DATE no puede ser menor que START_DATE.")
    for key in ("MEDIA_TYPE", "ORIGINAL_DIRECTION", "DIRECTION", "PARTICIPANT_PURPOSE", "ROUTING_TYPE", "TRANSFER_FILTER"):
        options = PYFLOW_PARAMS[key]["options"]
        value = values[key] or PYFLOW_PARAMS[key]["default"]
        match = next((opt for opt in options if opt.lower() == value.lower()), None)
        if match is None:
            raise ValueError(f"Opción inválida para {key}: {value}")
        values[key] = match
    destinations = {"Excel": "excel", "SAP HANA": "hana", "Excel + SAP HANA": "both"}
    output = args.output or destinations.get(values["OUTPUT_DESTINATION"])
    if output is None:
        raise ValueError("OUTPUT_DESTINATION no válido.")
    for key in ("GENESYS_CLIENT_ID", "GENESYS_CLIENT_SECRET", "GENESYS_REGION"):
        if not values[key]:
            raise ValueError(f"Falta parámetro requerido: {key}")
    dry = values["DRY_RUN"].lower() in ("true", "1", "yes", "si", "sí")
    if output in ("hana", "both"):
        for key in ("HPR_HOST", "HPR_PORT", "HPR_USER", "HPR_PASSWORD"):
            if not values[key] and not dry:
                raise ValueError(f"Falta parámetro HANA: {key}")
    for key in TABLE_KEYS.values():
        identifier(values[key])
    for key in ("HANA_SCHEMA", "HANA_USERS_TABLE", "HANA_QUEUES_TABLE", "HANA_WRAPUPS_TABLE"):
        identifier(values[key])
    config = Config(values, first, last, zone, output, dry,
                    timeout=int(env_str("REQUEST_TIMEOUT", "60")), retries=int(env_str("MAX_RETRIES", "5")),
                    poll_seconds=float(env_str("JOB_POLL_SECONDS", "5")), max_polls=int(env_str("JOB_MAX_POLLS", "720")),
                    page_size=int(env_str("PAGE_SIZE", "500")), api_sleep=float(env_str("API_SLEEP_SECONDS", "0.1")),
                    max_conversations=int(values["MAX_CONVERSATIONS"] or 0),
                    max_identification_requests=int(values["MAX_IDENTIFICATION_REQUESTS"] or 1000),
                    batch_size=int(values["HANA_BATCH_SIZE"] or 100))
    if min(config.timeout, config.max_polls, config.page_size, config.batch_size) < 1:
        raise ValueError("Timeout, páginas, lotes y máximo de sondeos deben ser positivos.")
    if min(config.retries, config.poll_seconds, config.api_sleep, config.max_conversations, config.max_identification_requests) < 0:
        raise ValueError("Los límites y tiempos no pueden ser negativos.")
    if config.page_size > 10000:
        raise ValueError("PAGE_SIZE no puede superar 10000.")
    return config


def retry_delay(response: Any, attempt: int) -> float:
    header = response.headers.get("Retry-After", "") if response is not None else ""
    if header:
        try:
            return max(0.0, float(header))
        except ValueError:
            try:
                return max(0.0, (parsedate_to_datetime(header) - datetime.now(UTC)).total_seconds())
            except (ValueError, TypeError):
                pass
    return min(60.0, 2.0 ** attempt)


class GenesysClient:
    """Cliente secuencial: una sesión HTTP, renovación OAuth y backoff acotado."""
    def __init__(self, config: Config):
        import requests
        self.requests = requests
        self.http = requests.Session()
        self.config = config
        domain = re.sub(r"^https?://", "", config.get("GENESYS_REGION").lower()).strip("/")
        self.domain = re.sub(r"^(api|login|apps)\.", "", domain)
        if not re.fullmatch(r"[a-z0-9.-]+", self.domain):
            raise ValueError("GENESYS_REGION debe ser un dominio Genesys.")
        self.token = ""
        self.expires = 0.0
        self.catalogs: dict[str, list[dict[str, Any]]] = {}

    def close(self) -> None:
        self.http.close()

    def authenticate(self) -> None:
        c = self.config
        for attempt in range(c.retries + 1):
            response = None
            try:
                response = self.http.post(f"https://login.{self.domain}/oauth/token",
                    auth=(c.get("GENESYS_CLIENT_ID"), c.get("GENESYS_CLIENT_SECRET")),
                    data={"grant_type": "client_credentials"}, timeout=c.timeout)
            except (self.requests.Timeout, self.requests.ConnectionError):
                if attempt == c.retries:
                    raise RuntimeError("No se pudo conectar al servicio OAuth.") from None
            if response is not None and response.ok:
                payload = response.json()
                self.token = payload.get("access_token", "")
                if not self.token:
                    raise RuntimeError("OAuth no devolvió access_token.")
                self.expires = time.monotonic() + max(1, float(payload.get("expires_in", 3600)) - 60)
                return
            if response is not None and response.status_code != 429 and response.status_code < 500:
                raise RuntimeError(f"OAuth rechazado: HTTP {response.status_code}")
            if attempt == c.retries:
                raise RuntimeError("OAuth agotó los reintentos.")
            time.sleep(retry_delay(response, attempt))

    def request(self, method: str, path: str, *, optional: bool = False, **kwargs: Any) -> Any:
        if not path.startswith("/api/v2/"):
            raise ValueError("Ruta API no válida.")
        refreshed = False
        for attempt in range(self.config.retries + 2):
            if not self.token or time.monotonic() >= self.expires:
                self.authenticate()
            response = None
            try:
                time.sleep(self.config.api_sleep)
                response = self.http.request(method, f"https://api.{self.domain}{path}",
                    headers={"Authorization": f"Bearer {self.token}", "Content-Type": "application/json"},
                    timeout=self.config.timeout, **kwargs)
            except (self.requests.Timeout, self.requests.ConnectionError):
                pass
            if response is not None:
                if response.ok:
                    return response.json() if response.content else {}
                if response.status_code == 401 and not refreshed:
                    self.token = ""
                    refreshed = True
                    continue
                if optional and response.status_code in (403, 404):
                    LOG.warning("Fallback no disponible | ruta=%s | HTTP=%s", path, response.status_code)
                    return None
                if response.status_code != 429 and response.status_code < 500:
                    raise RuntimeError(f"Genesys HTTP {response.status_code} | {method} {path}")
            if attempt >= self.config.retries:
                raise RuntimeError(f"Genesys agotó reintentos | {method} {path}")
            delay = retry_delay(response, attempt)
            LOG.warning("Reintento API | ruta=%s | HTTP=%s | espera=%.1fs", path,
                        response.status_code if response is not None else "conexión/timeout", delay)
            time.sleep(delay)
        raise RuntimeError("No se pudo completar la solicitud Genesys.")

    def entities(self, path: str) -> list[dict[str, Any]]:
        if path in self.catalogs:
            return self.catalogs[path]
        entities: dict[str, dict[str, Any]] = {}
        for page in range(1, 10001):
            data = self.request("GET", path, params={"pageSize": 100, "pageNumber": page})
            items = data.get("entities") or []
            for item in items:
                if item.get("id"):
                    entities[str(item["id"])] = item
            if not items or (data.get("pageCount") is not None and page >= int(data["pageCount"])) or (
                    data.get("pageCount") is None and not data.get("nextUri") and len(items) < 100):
                result = list(entities.values())
                self.catalogs[path] = result
                return result
        raise RuntimeError(f"Catálogo excede el límite de paginación: {path}")


def hana_connect(config: Config, write: bool = False) -> Any:
    from hdbcli import dbapi
    host = config.get("HPR_HOST") if write else (config.get("HPR_HOST_ESPEJO") or config.get("HPR_HOST"))
    if not host or not config.get("HPR_USER") or not config.get("HPR_PASSWORD"):
        raise RuntimeError("Credenciales HANA incompletas.")
    return dbapi.connect(address=host, port=int(config.get("HPR_PORT") or "30015"),
                         user=config.get("HPR_USER"), password=config.get("HPR_PASSWORD"),
                         connectTimeout=15000)


def load_catalogs(config: Config, api: GenesysClient) -> dict[str, dict[str, str]]:
    definitions = {
        "users": ("HANA_USERS_TABLE", "ID", "NAME", "/api/v2/users"),
        "queues": ("HANA_QUEUES_TABLE", "ID", "QUEUE_NAME", "/api/v2/routing/queues"),
        "wrapups": ("HANA_WRAPUPS_TABLE", "wrap_upcode", "wrap_upcode_name", "/api/v2/routing/wrapupcodes"),
    }
    catalogs: dict[str, dict[str, str]] = {}
    conn = None
    try:
        conn = hana_connect(config)
    except Exception:
        if config.output != "excel" and not config.dry_run:
            raise RuntimeError("No fue posible abrir HANA para los catálogos; revise globales y hdbcli.") from None
        LOG.warning("HANA no disponible; catálogos mediante API Genesys una sola vez.")
    try:
        for name, (table_key, id_col, name_col, endpoint) in definitions.items():
            rows = None
            if conn is not None:
                cursor = conn.cursor()
                try:
                    if name == "wrapups":
                        # HANA conserva minúsculas solo cuando se crean entre comillas.
                        cursor.execute('SELECT COLUMN_NAME FROM SYS.TABLE_COLUMNS WHERE SCHEMA_NAME = ? AND TABLE_NAME = ?',
                                       (config.get("HANA_SCHEMA"), config.get(table_key)))
                        columns = [row[0] for row in cursor.fetchall()]
                        resolved = []
                        for expected in (id_col, name_col):
                            matches = [col for col in columns if col.lower() == expected]
                            if len(matches) != 1:
                                raise ValueError(f"Columna ausente o ambigua: {expected}")
                            resolved.append(matches[0])
                        id_col, name_col = resolved
                    cursor.execute(f'SELECT DISTINCT {identifier(id_col)}, {identifier(name_col)} '
                                   f'FROM {identifier(config.get("HANA_SCHEMA"))}.{identifier(config.get(table_key))}')
                    rows = cursor.fetchall()
                except Exception:
                    # Los catálogos se consultan sin crear ni alterar tablas.
                    if config.output != "excel" and not config.dry_run:
                        raise RuntimeError(f"No se pudo leer {config.get(table_key)} ({id_col}, {name_col}).") from None
                    LOG.warning("Catálogo %s no disponible en HANA; se usará Genesys.", name)
                finally:
                    cursor.close()
            if rows is None:
                rows = [(item.get("id"), item.get("name")) for item in api.entities(endpoint)]
            catalog: dict[str, str] = {}
            for entity_id, entity_name in rows:
                if entity_id is None or entity_name is None:
                    continue
                key, value = str(entity_id), str(entity_name)
                if key in catalog and catalog[key] != value:
                    # No elegir aleatoriamente entre nombres históricos.
                    LOG.warning("Catálogo %s: ID con nombres distintos; se conserva elección determinista.", name)
                    value = min(catalog[key], value)
                catalog[key] = value
            catalogs[name] = catalog
            LOG.info("Catálogo %s cargado: %s", name, len(catalog))
    finally:
        if conn is not None:
            conn.close()
    return catalogs


def resolve_filters(config: Config, api: GenesysClient, catalogs: dict[str, dict[str, str]]) -> None:
    for id_key, (selection_key, endpoint, catalog_name) in FILTER_SELECTORS.items():
        manual, selected = config.get(id_key), config.get(selection_key)
        if manual:
            config.filters[id_key] = split_filter_values(manual)
            continue
        if not selected:
            config.filters[id_key] = []
            continue
        ids = parse_selection(selected)
        if ids is None and selection_key == "FLOW_SELECTION_ID":
            ids = split_filter_values(selected)  # El selector de flujo existente entrega UUID directamente.
        if ids is None:
            entries = [{"id": key, "name": value} for key, value in catalogs.get(catalog_name, {}).items()]
            wanted = split_filter_values(selected)
            ids = []
            for name in wanted:
                matches = [item for item in entries if str(item.get("name", "")).casefold() == name.casefold()]
                if len(matches) != 1:
                    entries = api.entities(endpoint)
                    matches = [item for item in entries if name.casefold() in (
                        str(item.get("name", "")).casefold(), str(item.get("email", "")).casefold())]
                if len(matches) != 1:
                    raise ValueError(f"{selection_key}: nombre no encontrado o ambiguo; use el selector con ID.")
                ids.append(str(matches[0]["id"]))
        config.filters[id_key] = list(dict.fromkeys(ids))
    config.filters["CONVERSATION_ID"] = split_filter_values(config.get("CONVERSATION_ID"))
    active = [key for key, values in config.filters.items() if values]
    active += [key for key, default in (("MEDIA_TYPE", "todos"), ("PARTICIPANT_PURPOSE", "todos"),
               ("ORIGINAL_DIRECTION", "ambas"), ("DIRECTION", "ambas"), ("ROUTING_TYPE", "todos"),
               ("TRANSFER_FILTER", "todas")) if config.get(key) != default]
    if active:
        LOG.info("Filtros activos: %s", ", ".join(active))
    else:
        LOG.warning("No hay filtros adicionales aparte del rango de fechas.")


def dimension_filter(dimension: str, values: list[str]) -> dict[str, Any]:
    return {"type": "or", "predicates": [{"type": "dimension", "dimension": dimension,
            "operator": "matches", "value": value} for value in values]}


def job_body(config: Config, start: datetime, end: datetime) -> dict[str, Any]:
    body: dict[str, Any] = {"interval": f"{utc_text(start)}/{utc_text(end)}", "order": "asc", "orderBy": "conversationStart"}
    # Sólo dimensiones cuya ubicación nativa es fiable. Campaña/lista también pueden
    # existir sólo en attributes de Dialer: se validan después para no perder casos.
    filters = []
    for key, dimension in (("USER_ID", "userId"), ("QUEUE_ID", "queueId"), ("WRAPUP_CODE_ID", "wrapUpCode"),
                           ("FLOW_ID", "flowId")):
        if config.filters.get(key):
            filters.append(dimension_filter(dimension, config.filters[key]))
    for key, dimension in (("MEDIA_TYPE", "mediaType"), ("PARTICIPANT_PURPOSE", "purpose")):
        if config.get(key) != "todos":
            filters.append(dimension_filter(dimension, [config.get(key)]))
    # Cada bloque es independiente a nivel conversación, no un AND en una sola sesión.
    if filters:
        body["segmentFilters"] = filters
    if config.get("ORIGINAL_DIRECTION") != "ambas":
        body["conversationFilters"] = [dimension_filter("originatingDirection", [config.get("ORIGINAL_DIRECTION")])]
    return body


def daily_conversations(api: GenesysClient, config: Config, day: date) -> Iterator[dict[str, Any]]:
    start, end = day_interval(day, config.zone)
    LOG.info("Fecha local: %s | intervalo UTC: %s/%s", day, utc_text(start), utc_text(end))
    base = "/api/v2/analytics/conversations/details/jobs"
    data = api.request("POST", base, json=job_body(config, start, end))
    job_id = data.get("jobId") or data.get("id")
    if not job_id:
        raise RuntimeError("Genesys no devolvió jobId.")
    path = base + "/" + quote(str(job_id), safe="")
    LOG.info("Job creado: %s", job_id)
    for _ in range(config.max_polls):
        state = str(api.request("GET", path).get("state", "")).upper()
        LOG.info("Job %s | estado=%s", job_id, state)
        if state == "FULFILLED":
            break
        if state in ("FAILED", "CANCELLED", "CANCELED", "EXPIRED"):
            raise RuntimeError(f"Job {job_id} terminó en {state}.")
        time.sleep(config.poll_seconds)
    else:
        raise RuntimeError(f"Job {job_id} excedió el tiempo de espera; no se cargará un resultado parcial.")
    cursor = None
    cursors: set[str] = set()
    page = total = 0
    while True:
        params: dict[str, Any] = {"pageSize": config.page_size}
        if cursor:
            params["cursor"] = cursor
        data = api.request("GET", path + "/results", params=params)
        items = data.get("conversations") or []
        page += 1
        total += len(items)
        LOG.info("Job %s | página=%s | conversaciones descargadas=%s", job_id, page, total)
        yield from items
        cursor = data.get("cursor")
        if not cursor:
            return
        if cursor in cursors:
            raise RuntimeError("Cursor repetido en resultados Genesys.")
        cursors.add(cursor)


def sessions_of(conversation: dict[str, Any]) -> list[tuple[dict[str, Any], dict[str, Any]]]:
    return [(p, s) for p in conversation.get("participants") or [] for s in p.get("sessions") or []]


def unique(values: list[Any]) -> list[Any]:
    result, seen = [], set()
    for value in values:
        if value is None or value == "":
            continue
        key = json.dumps(value, sort_keys=True, ensure_ascii=False, default=str)
        if key not in seen:
            result.append(value)
            seen.add(key)
    return result


def attributes_of(conversation: dict[str, Any]) -> dict[str, Any]:
    """Sólo atributos presentes en el modelo; nunca recoger navegación IVR."""
    allowed = {name for schema in SCHEMAS.values() for name, _, _, source in schema if ".attributes." in source}
    result: dict[str, Any] = {}
    for participant in conversation.get("participants") or []:
        for key, value in (participant.get("attributes") or {}).items():
            if key in allowed and value not in (None, "") and key not in result:
                result[key] = str(value)
    return result


def session_bounds(session: dict[str, Any]) -> tuple[datetime | None, datetime | None]:
    segments = session.get("segments") or []
    starts = [timestamp(s.get("segmentStart")) for s in segments if s.get("segmentStart")]
    ends = [timestamp(s.get("segmentEnd")) for s in segments if s.get("segmentEnd")]
    return (min(starts) if starts else None, max(ends) if ends else None)


def segment_order(segment: dict[str, Any]) -> tuple[Any, ...]:
    return (timestamp(segment.get("segmentStart")) or datetime.min.replace(tzinfo=UTC),
            timestamp(segment.get("segmentEnd")) or datetime.max.replace(tzinfo=UTC),
            str(segment.get("segmentType", "")), str(segment.get("wrapUpCode", "")))


@dataclass
class Attention:
    participant: dict[str, Any]
    session: dict[str, Any]
    segments: list[dict[str, Any]]
    contexts: list[tuple[dict[str, Any], dict[str, Any]]] = field(default_factory=list)
    metrics: dict[str, Any] = field(default_factory=dict)
    transfer_sequence: int | None = None
    transfer_to_user: str | None = None
    transfer_to_queue: str | None = None
    native_transfer: bool = False

    @property
    def start(self) -> datetime | None:
        return session_bounds({"segments": self.segments})[0]

    @property
    def end(self) -> datetime | None:
        return session_bounds({"segments": self.segments})[1]

    @property
    def queue(self) -> str | None:
        # Prioridad a la cola donde hubo interacción o conclusión, no una alerta posterior.
        values = [s.get("queueId") for s in self.segments if s.get("queueId") and
                  (s.get("segmentType") == "interact" or s.get("wrapUpCode"))]
        return values[-1] if values else next((s.get("queueId") for s in self.segments if s.get("queueId")), None)


def agent_episodes(participant: dict[str, Any], session: dict[str, Any]) -> list[Attention]:
    """Agrupa segmentos técnicos; una conclusión cierra el episodio, no cada hold."""
    result: list[Attention] = []
    pending: list[dict[str, Any]] = []
    closed = False
    queue = None

    def finish() -> None:
        if pending and any(s.get("segmentType") in ("interact", "contacting", "dialing") or s.get("wrapUpCode") for s in pending):
            result.append(Attention(participant, session, list(pending)))

    for segment in sorted(session.get("segments") or [], key=segment_order):
        meaningful = segment.get("segmentType") == "interact" or bool(segment.get("wrapUpCode"))
        new_queue = segment.get("queueId")
        different_wrapup = closed and segment.get("wrapUpCode") and any(
            s.get("wrapUpCode") and s["wrapUpCode"] != segment["wrapUpCode"] for s in pending)
        if pending and (different_wrapup or (closed and segment.get("segmentType") == "interact") or
                        (meaningful and queue and new_queue and queue != new_queue)):
            finish()
            pending, closed, queue = [], False, None
        pending.append(segment)
        if meaningful and new_queue:
            queue = new_queue
        if segment.get("segmentType") == "wrapup" or segment.get("wrapUpCode"):
            closed = True
    finish()
    # Callback pendiente puede no tener todavía interact, pero sí una acción real programada.
    if not result and session.get("mediaType") == "callback" and session.get("callbackScheduledTime"):
        result.append(Attention(participant, session, list(session.get("segments") or [])))
    return result


def metric_totals(metrics: list[dict[str, Any]]) -> dict[str, Any]:
    result: dict[str, Any] = {}
    for metric in metrics:
        key, value = metric.get("name"), metric.get("value")
        if key and isinstance(value, (float, int)) and not isinstance(value, bool) and math.isfinite(value):
            result[key] = result.get(key, 0) + value
    return result


def assign_metrics(attentions: list[Attention]) -> None:
    """Cada evento de métrica de una sesión se asigna una sola vez por emitDate."""
    groups: dict[int, list[Attention]] = defaultdict(list)
    for attention in attentions:
        groups[id(attention.session)].append(attention)
    for group in groups.values():
        group.sort(key=lambda a: a.start or datetime.min.replace(tzinfo=UTC))
        for metric in group[0].session.get("metrics") or []:
            emitted = timestamp(metric.get("emitDate"))
            if emitted is None:
                # Evento sin fecha: asignar una sola vez; no multiplicar por conclusiones.
                target = group[0]
            else:
                candidates = [a for a in group if (a.start is None or a.start <= emitted) and
                              (a.end is None or emitted <= a.end)]
                target = candidates[-1] if candidates else min(group, key=lambda a: abs(
                    ((a.end or a.start or emitted) - emitted).total_seconds()))
            for key, value in metric_totals([metric]).items():
                target.metrics[key] = target.metrics.get(key, 0) + value


def assign_contexts(attentions: list[Attention], conversation: dict[str, Any]) -> None:
    """Una sesión técnica se asocia a una atención, nunca se replica en todas."""
    primary_ids = {id(a.session) for a in attentions}
    for participant, session in sessions_of(conversation):
        if participant.get("purpose") == "ivr" or id(session) in primary_ids:
            continue
        if participant.get("purpose") in ("agent", "user"):
            continue
        candidates = [a for a in attentions if a.session.get("mediaType") == session.get("mediaType")]
        if not candidates:
            continue
        start, end = session_bounds(session)
        anchor = start if participant.get("purpose") in ("customer", "external") else (end or start)
        related = [a for a in candidates if session.get("sessionId") and (
            session.get("sessionId") == a.session.get("peerId") or
            (session.get("peerId") and session.get("peerId") == a.session.get("sessionId")))]
        if related:
            candidates = related
        if anchor:
            later = [a for a in candidates if a.start is not None and a.start >= anchor]
            overlapping = [a for a in candidates if a.start is not None and a.start <= anchor and
                           (a.end is None or a.end >= anchor)]
            if later:
                target = min(later, key=lambda a: a.start)
            elif overlapping:
                target = min(overlapping, key=lambda a: abs((a.start - (start or anchor)).total_seconds()))
            else:
                continue  # No asignar un bot posterior a una atención ya cerrada.
        elif len(candidates) == 1:
            target = candidates[0]
        else:
            continue  # Sin evidencia temporal no inventar asociación.
        target.contexts.append((participant, session))
    for attention in attentions:
        attention.contexts.sort(key=lambda pair: session_bounds(pair[1])[0] or datetime.min.replace(tzinfo=UTC))
        # Las métricas nativas del agente prevalecen; complementar sólo métricas ausentes.
        totals: dict[str, Any] = {}
        for p, s in attention.contexts:
            for key, value in metric_totals(s.get("metrics") or []).items():
                # No sumar duraciones del cliente sobre métricas de atención.
                if p.get("purpose") in ("acd", "botflow", "workflow", "dialer", "campaign"):
                    totals[key] = totals.get(key, 0) + value
        for key, value in totals.items():
            if key not in attention.metrics:
                attention.metrics[key] = value


def native_transfer(session: dict[str, Any], segments: list[dict[str, Any]] | None = None,
                    metrics: dict[str, Any] | None = None) -> bool:
    values = metric_totals(session.get("metrics") or []) if metrics is None else metrics
    return any(s.get("disconnectType") == "transfer" for s in (session.get("segments") or [] if segments is None else segments)) or any(
        (values.get(key) or 0) > 0 for key in ("nTransferred", "nBlindTransferred"))


def link_transfers(attentions: list[Attention]) -> bool:
    transfer_index = 0
    transferred = False
    for attention in attentions:
        transferred = transferred or attention.native_transfer
    for index, current in enumerate(attentions):
        if not current.participant.get("userId"):
            continue
        # Los segmentos wrapup pueden superponerse con el agente receptor.
        interactions = [s for s in current.segments if s.get("segmentType") == "interact"]
        end = max((timestamp(s["segmentEnd"]) for s in interactions if s.get("segmentEnd")), default=None)
        if end is None:
            continue
        candidates = [a for a in attentions[index + 1:] if a.participant.get("userId") and
                      a.session.get("mediaType") == current.session.get("mediaType") and a.start is not None and
                      (a.participant.get("userId") != current.participant.get("userId") or a.queue != current.queue) and
                      a.start >= end]
        if not candidates:
            continue
        nxt = candidates[0]
        if len([a for a in candidates if a.start == nxt.start]) != 1:
            continue
        gap = (nxt.start - end).total_seconds()
        bridge = any(p.get("purpose") == "acd" and any(seg.get("queueId") == nxt.queue for seg in s.get("segments") or [])
                     and session_bounds(s)[0] is not None and session_bounds(s)[0] >= end
                     for p, s in nxt.contexts)
        peer_link = bool(current.session.get("sessionId") and (
            current.session.get("sessionId") == nxt.session.get("peerId") or
            (current.session.get("peerId") and current.session.get("peerId") == nxt.session.get("sessionId"))))
        # Evidencia + continuidad: un email retomado días después no implica transferencia.
        verified = peer_link or (current.native_transfer and (gap <= 120 or bridge)) or (bridge and gap <= 120)
        if not verified:
            continue
        transferred = True
        transfer_index += 1
        nxt.transfer_sequence = transfer_index
        current.transfer_to_user = nxt.participant.get("userId")
        current.transfer_to_queue = nxt.queue
    return transferred


def build_attentions(conversation: dict[str, Any]) -> tuple[list[Attention], bool]:
    attentions: list[Attention] = []
    pairs = sessions_of(conversation)
    for participant, session in pairs:
        if participant.get("purpose") in ("agent", "user"):
            attentions.extend(agent_episodes(participant, session))
    handled_media = {a.session.get("mediaType") for a in attentions}
    by_media: dict[str, list[tuple[dict[str, Any], dict[str, Any]]]] = defaultdict(list)
    for p, s in pairs:
        media = s.get("mediaType")
        if media and media not in handled_media and p.get("purpose") not in ("ivr", "acd", "agent", "user"):
            if any(seg.get("segmentType") == "interact" for seg in s.get("segments") or []) or s.get("callbackScheduledTime"):
                by_media[media].append((p, s))
    priorities = {"customer": 0, "external": 1, "botflow": 2, "workflow": 3}
    for pairs_for_media in by_media.values():
        p, s = min(pairs_for_media, key=lambda pair: (
            priorities.get(pair[0].get("purpose"), 4), session_bounds(pair[1])[0] or datetime.min.replace(tzinfo=UTC),
            str(pair[0].get("participantId", "")), str(pair[1].get("sessionId", ""))))
        attentions.append(Attention(p, s, sorted(s.get("segments") or [], key=segment_order)))
    attentions.sort(key=lambda a: (a.start or timestamp(conversation.get("conversationStart")) or datetime.min.replace(tzinfo=UTC),
                                  str(a.participant.get("participantId", "")), str(a.session.get("sessionId", "")),
                                  str(a.queue or "")))
    assign_metrics(attentions)
    for attention in attentions:
        attention.native_transfer = native_transfer(attention.session, attention.segments, attention.metrics)
    assign_contexts(attentions, conversation)
    transferred = link_transfers(attentions)
    # También aceptar evidencia nativa en un bot o ACD, aunque el receptor sea desconocido.
    transferred = transferred or any(native_transfer(s) for p, s in pairs if p.get("purpose") != "ivr")
    return attentions, transferred


def conversation_matches(conversation: dict[str, Any], config: Config, transferred: bool) -> bool:
    cid = str(conversation.get("conversationId") or "")
    if config.filters.get("CONVERSATION_ID") and cid not in config.filters["CONVERSATION_ID"]:
        return False
    if config.get("ORIGINAL_DIRECTION") != "ambas" and conversation.get("originatingDirection") != config.get("ORIGINAL_DIRECTION"):
        return False
    pairs = sessions_of(conversation)
    # Medio y dirección deben coexistir en la misma sesión relevante.
    relevant = [(p, s) for p, s in pairs if p.get("purpose") != "ivr" and
                (config.get("MEDIA_TYPE") == "todos" or s.get("mediaType") == config.get("MEDIA_TYPE")) and
                (config.get("DIRECTION") == "ambas" or s.get("direction") == config.get("DIRECTION"))]
    if not relevant:
        return False
    # Propósito/usuario/flujo pueden corresponder a participantes distintos en la misma conversación.
    if config.get("PARTICIPANT_PURPOSE") != "todos" and not any(
            p.get("purpose") == config.get("PARTICIPANT_PURPOSE") for p, _ in relevant):
        return False
    if config.get("ROUTING_TYPE") != "todos" and not any(s.get("usedRouting") == config.get("ROUTING_TYPE") for _, s in relevant):
        return False
    attributes = attributes_of(conversation)
    observed = {
        "USER_ID": [p.get("userId") for p, _ in relevant],
        "QUEUE_ID": [g.get("queueId") for _, s in relevant for g in s.get("segments") or []],
        "WRAPUP_CODE_ID": [g.get("wrapUpCode") for _, s in relevant for g in s.get("segments") or []],
        "FLOW_ID": [(s.get("flow") or {}).get("flowId") for _, s in relevant],
        "CAMPAIGN_ID": [s.get("outboundCampaignId") for _, s in relevant] + [(p.get("attributes") or {}).get("dialerCampaignId") for p in conversation.get("participants") or []],
        "CONTACT_LIST_ID": [s.get("outboundContactListId") for _, s in relevant] + [(p.get("attributes") or {}).get("dialerContactListId") for p in conversation.get("participants") or []],
    }
    for key, values in observed.items():
        if config.filters.get(key) and not set(config.filters[key]).intersection(str(v) for v in values if v):
            return False
    mode = config.get("TRANSFER_FILTER")
    return not ((mode == "solo_transferidas" and not transferred) or (mode == "no_transferidas" and transferred))


class IdentificationLookup:
    def __init__(self, config: Config, api: GenesysClient):
        self.config, self.api = config, api
        self.cache: dict[str, dict[str, Any]] = {}
        self.requests = 0
        self.counts: Counter = Counter()
        self.limit_logged = False

    def enrich(self, conversation: dict[str, Any]) -> dict[str, Any]:
        attributes = attributes_of(conversation)
        if attributes.get("SPD_IDENTIFICACION"):
            self.counts["analytics"] += 1
            return attributes
        cid = conversation["conversationId"]
        if cid not in self.cache:
            found: dict[str, Any] = {}
            paths = {"message": "messages", "voice": "calls", "callback": "calls", "email": "emails"}
            endpoints = unique([paths[s["mediaType"]] for _, s in sessions_of(conversation) if s.get("mediaType") in paths])
            for endpoint in endpoints:
                if self.requests >= self.config.max_identification_requests:
                    if not self.limit_logged:
                        LOG.warning("Límite de fallback de identificación alcanzado: %s", self.requests)
                        self.limit_logged = True
                    break
                self.requests += 1
                data = self.api.request("GET", f"/api/v2/conversations/{endpoint}/{quote(cid, safe='')}", optional=True)
                if data:
                    found.update({k: v for k, v in attributes_of(data).items() if k in IDENTIFICATION_KEYS})
                if found.get("SPD_IDENTIFICACION"):
                    break
            self.cache[cid] = found
        for key, value in self.cache[cid].items():
            if not attributes.get(key):
                attributes[key] = str(value)
        self.counts["fallback" if attributes.get("SPD_IDENTIFICACION") else "sin_identificacion"] += 1
        return attributes


def combine(values: list[Any], json_array: bool = False) -> Any:
    values = unique(values)
    if not values:
        return None
    if json_array:
        result = []
        for value in values:
            result.extend(value if isinstance(value, list) else [value])
        return unique(result)
    if len(values) == 1:
        return values[0]
    return ";".join(str(value) for value in values)


def source_value(name: str, source: str, conversation: dict[str, Any], attention: Attention,
                 attributes: dict[str, Any]) -> Any:
    if source.startswith("conversation."):
        return conversation.get(name)
    if ".attributes." in source:
        own = (attention.participant.get("attributes") or {}).get(name)
        if own not in (None, ""):
            return str(own)
        local = [(p.get("attributes") or {}).get(name) for p, _ in attention.contexts]
        local = [v for v in local if v not in (None, "")]
        return str(local[-1]) if local else attributes.get(name)
    if ".metrics[" in source:
        return attention.metrics.get(name)
    if ".flow." in source:
        sessions = [s for p, s in attention.contexts + [(attention.participant, attention.session)] if p.get("purpose") != "ivr"]
        return combine([(s.get("flow") or {}).get(name) for s in sessions], json_array=name == "outcomes")
    if ".segments[]." in source:
        values = [segment.get(name) for segment in attention.segments if segment.get(name) is not None]
        if not values and name == "requestedRoutingUserIds":
            values = [segment[name] for p, session in attention.contexts if p.get("purpose") == "acd"
                      for segment in session.get("segments") or [] if segment.get(name) is not None]
        if name == "queueId":
            return attention.queue
        if name == "segmentStart":
            return attention.start
        if name == "segmentEnd":
            return attention.end
        if name in ("wrapUpCode", "wrapUpNote", "subject"):
            return values[-1] if values else None
        return combine(values, json_array=name in ("requestedRoutingUserIds", "sipResponseCodes", "q850ResponseCodes"))
    if source.startswith("participants[].sessions[]."):
        value = attention.session.get(name)
        if value is not None:
            return value
        # Completar routing sólo desde el ACD asociado; nunca derivar usedRouting de requestedRoutings.
        routing = {"requestedRoutings", "usedRouting", "routingRule", "routingRuleType", "selectedAgentId",
                   "selectedAgentRank", "proposedAgents", "eligibleAgentCounts", "waitingInteractionCounts",
                   "agentBullseyeRing", "routingRing"}
        permitted = [(p, s) for p, s in attention.contexts if p.get("purpose") == "acd"] if name in routing else attention.contexts
        found = [s.get(name) for _, s in permitted if s.get(name) is not None]
        # No transferir IDs técnicos de otra sesión a la atención actual.
        if name in ("sessionId", "peerId", "direction", "provider", "edgeId", "protocolCallId"):
            return None
        if found:
            return found[-1]
        # Datos del extremo remoto pueden estar sólo en la sesión del cliente.
        # Utilizar el dato común sólo si es inequívoco para ese medio, sin inventar IDs.
        if name in ("ani", "dnis", "addressFrom", "addressTo", "addressSelf", "addressOther", "messageType",
                    "outboundCampaignId", "outboundContactListId", "outboundContactId", "callbackNumbers"):
            candidates = unique([s.get(name) for p, s in sessions_of(conversation)
                                 if p.get("purpose") != "ivr" and s.get("mediaType") == attention.session.get("mediaType")])
            if len(candidates) == 1:
                return candidates[0]
        return None
    if source.startswith("participants[]."):
        value = attention.participant.get(name)
        if value is None and name == "externalContactId":
            candidates = unique([p.get(name) for p in conversation.get("participants") or [] if p.get("purpose") != "ivr"])
            if len(candidates) == 1:
                value = candidates[0]
        return value
    return None


def convert_value(name: str, kind: str, nullable: bool, value: Any, config: Config) -> Any:
    if value is None or value == "":
        if not nullable:
            raise ValueError(f"Columna obligatoria sin valor: {name}")
        return None
    if kind == "TIMESTAMP":
        if name == "fechaCarga" and isinstance(value, datetime) and value.tzinfo is None:
            return value
        return timestamp(value).astimezone(ZoneInfo(config.zone)).replace(tzinfo=None)
    if kind in ("INTEGER", "BIGINT"):
        number = Decimal(str(value))
        if not number.is_finite() or number != number.to_integral_value():
            raise ValueError(f"Valor no entero en {name}")
        return int(number)
    if kind.startswith("DECIMAL"):
        number = Decimal(str(value))
        if not number.is_finite():
            raise ValueError(f"Decimal no finito en {name}")
        return number
    if kind == "BOOLEAN":
        if isinstance(value, bool):
            return value
        if str(value).lower() in ("true", "1"):
            return True
        if str(value).lower() in ("false", "0"):
            return False
        # issuedCallback puede aparecer en varios flows; no tratar la cadena como True.
        if isinstance(value, str) and ";" in value:
            parts = value.lower().split(";")
            if all(p in ("true", "false") for p in parts):
                return "true" in parts
        raise ValueError(f"Booleano no válido en {name}")
    text = json.dumps(value, ensure_ascii=False, separators=(",", ":"), default=str) if isinstance(value, (list, dict)) else str(value)
    # La longitud depende del destino: Excel valida su límite de celda y HANA
    # valida la capacidad real consultada en SYS.TABLE_COLUMNS, no el modelo inicial.
    return text


def normalize(conversation: dict[str, Any], attentions: list[Attention], attributes: dict[str, Any],
              catalogs: dict[str, dict[str, str]], config: Config) -> dict[str, list[dict[str, Any]]]:
    result = {sheet: [] for sheet in SCHEMAS}
    loaded_at = datetime.now(ZoneInfo(config.zone)).replace(tzinfo=None)
    for sequence, attention in enumerate(attentions, 1):
        if attention.participant.get("purpose") == "ivr":
            raise AssertionError("IVR nunca debe convertirse en atención.")
        enriched = {
            "sequence": sequence, "fechaCarga": loaded_at, "transferSequence": attention.transfer_sequence,
            "transferToUserName": catalogs["users"].get(attention.transfer_to_user),
            "transferToQueueName": catalogs["queues"].get(attention.transfer_to_queue),
        }
        sheets = ["INTERACCIONES"]
        media = attention.session.get("mediaType")
        if media in ("voice", "callback"):
            sheets.append("INTERACCIONES_VOICE")
        elif media == "message":
            sheets.append("INTERACCIONES_DIGITAL")
        elif media == "email":
            sheets.append("INTERACCIONES_EMAIL")
        for sheet in sheets:
            row = {name: source_value(name, source, conversation, attention, attributes)
                   for name, _, _, source in SCHEMAS[sheet]}
            row.update({key: value for key, value in enriched.items() if key in row})
            for id_key, name_key, catalog in (("userId", "userName", "users"), ("queueId", "queueName", "queues"),
                        ("wrapUpCode", "wrapUpCodeName", "wrapups"), ("selectedAgentId", "selectedAgentName", "users")):
                if name_key in row:
                    row[name_key] = catalogs[catalog].get(row.get(id_key))
            result[sheet].append({name: convert_value(name, kind, nullable, row[name], config)
                                  for name, kind, nullable, _ in SCHEMAS[sheet]})
    return result


def hana_column(name: str) -> str:
    """Nombres físicos verificados contra TABLE_COLUMNS_202609272148.csv.

    Mantener los nombres API en normalización y Excel; traducir solo al escribir HANA.
    """
    if name == "sequence":
        return "SECUENCIA"
    name = re.sub(r"([A-Z]+)([A-Z][a-z])", r"\1_\2", name)
    return re.sub(r"([a-z0-9])([A-Z])", r"\1_\2", name).upper()


def merge_statement(config: Config, sheet: str) -> str:
    columns = [column[0] for column in SCHEMAS[sheet]]
    table = identifier(config.get("HANA_SCHEMA")) + "." + identifier(config.get(TABLE_KEYS[sheet]))
    quoted = [identifier(hana_column(col)) for col in columns]
    source = ", ".join(f"? AS {col}" for col in quoted)
    updates = []
    for col in columns:
        if col in ("conversationId", "sequence"):
            continue
        q = identifier(hana_column(col))
        # No borrar identificación ya conocida si una consulta posterior no la devuelve.
        expression = f"COALESCE(S.{q}, T.{q})" if col in IDENTIFICATION_KEYS else f"S.{q}"
        updates.append(f"T.{q} = {expression}")
    return (f"MERGE INTO {table} AS T USING (SELECT {source} FROM DUMMY) AS S "
            'ON T."CONVERSATION_ID" = S."CONVERSATION_ID" AND T."SECUENCIA" = S."SECUENCIA" '
            f'WHEN MATCHED THEN UPDATE SET {", ".join(updates)} '
            f'WHEN NOT MATCHED THEN INSERT ({", ".join(quoted)}) VALUES ({", ".join("S." + col for col in quoted)})')


def cleanup_batches(config: Config, sheet: str, conversations: list) -> Iterator[tuple[str, tuple]]:
    """Agrupar predicados completos; nunca partir el conjunto de secuencias de un ID."""
    table = identifier(config.get("HANA_SCHEMA")) + "." + identifier(config.get(TABLE_KEYS[sheet]))
    clauses, parameters = [], []
    for conversation in conversations:
        primary = conversation["INTERACCIONES"]
        if not primary:
            continue
        cid = primary[0]["conversationId"]
        sequences = sorted({row["sequence"] for row in conversation[sheet]})
        values = (cid, *sequences)
        if clauses and (len(clauses) >= 100 or len(parameters) + len(values) > 2000):
            yield f'DELETE FROM {table} WHERE ' + ' OR '.join(clauses), tuple(parameters)
            clauses, parameters = [], []
        predicate = '"CONVERSATION_ID" = ?'
        if sequences:
            predicate += ' AND "SECUENCIA" NOT IN (' + ','.join('?' for _ in sequences) + ')'
        clauses.append('(' + predicate + ')')
        parameters.extend(values)
    if clauses:
        yield f'DELETE FROM {table} WHERE ' + ' OR '.join(clauses), tuple(parameters)


class HanaWriter:
    """Una transacción para las cuatro tablas. Sin DDL; MERGE por claves del modelo."""
    def __init__(self, config: Config):
        self.config = config
        self.conn = None
        self.buffer: list[dict[str, list[dict[str, Any]]]] = []
        self.counts: Counter = Counter()
        if not config.dry_run:
            self.conn = hana_connect(config, write=True)
            self.conn.setautocommit(False)
            try:
                self.validate_schema()
            except Exception:
                self.conn.close()
                self.conn = None
                raise

    def validate_schema(self) -> None:
        self.text_limits = {}
        cur = self.conn.cursor()
        try:
            for sheet, table_key in TABLE_KEYS.items():
                cur.execute("SELECT COLUMN_NAME, DATA_TYPE_NAME, LENGTH, SCALE, IS_NULLABLE FROM SYS.TABLE_COLUMNS "
                            "WHERE SCHEMA_NAME = ? AND TABLE_NAME = ? ORDER BY POSITION",
                            (self.config.get("HANA_SCHEMA"), self.config.get(table_key)))
                actual = cur.fetchall()
                by_name = {str(row[0]): row for row in actual}
                expected_names = {hana_column(col[0]) for col in SCHEMAS[sheet]}
                if set(by_name) != expected_names:
                    missing = sorted(expected_names - set(by_name))
                    extra = sorted(set(by_name) - expected_names)
                    raise RuntimeError(f"Columnas de {self.config.get(table_key)} incompatibles: faltan={missing}, adicionales={extra}; no se escribirá.")
                for expected in SCHEMAS[sheet]:
                    row = by_name[hana_column(expected[0])]
                    name, kind, nullable, _ = expected
                    base = kind.split("(")[0]
                    if str(row[1]).upper() != base:
                        raise RuntimeError(f"Tipo HANA incompatible: {sheet}.{name}")
                    if base == "NVARCHAR" and int(row[2]) < int(re.search(r"\((\d+)\)", kind).group(1)):
                        raise RuntimeError(f"Longitud HANA inferior al modelo: {sheet}.{name}")
                    if base == "NVARCHAR":
                        self.text_limits[(sheet, name)] = int(row[2])
                    if base == "DECIMAL":
                        precision, scale = map(int, re.search(r"\((\d+),(\d+)\)", kind).groups())
                        if int(row[3]) < scale or int(row[2]) - int(row[3]) < precision - scale:
                            raise RuntimeError(f"Precisión HANA inferior al modelo: {sheet}.{name}")
                cur.execute("SELECT COLUMN_NAME FROM SYS.CONSTRAINTS WHERE SCHEMA_NAME = ? AND TABLE_NAME = ? "
                            "AND IS_PRIMARY_KEY = 'TRUE' ORDER BY POSITION",
                            (self.config.get("HANA_SCHEMA"), self.config.get(table_key)))
                keys = {str(row[0]) for row in cur.fetchall()}
                if keys != {"CONVERSATION_ID", "SECUENCIA"}:
                    raise RuntimeError(f"Clave primaria incompatible en {self.config.get(table_key)}.")
        finally:
            cur.close()

    def add(self, rows: dict[str, list[dict[str, Any]]]) -> None:
        if self.config.dry_run:
            return
        for sheet, items in rows.items():
            for row in items:
                for column, value in row.items():
                    limit = self.text_limits.get((sheet, column))
                    if limit is not None and isinstance(value, str):
                        length = len(value.encode("utf-16-le")) // 2
                        if length > limit:
                            raise ValueError(
                                f'{self.config.get(TABLE_KEYS[sheet])}.{hana_column(column)}: '
                                f'longitud={length}, capacidad HANA={limit}; '
                                f'conversationId={row.get("conversationId")}, sequence={row.get("sequence")}. '
                                'No se truncará información.')
        self.buffer.append(rows)
        if len(self.buffer) >= self.config.batch_size:
            self.flush()

    def flush(self) -> None:
        if self.config.dry_run or not self.buffer:
            return
        if self.conn is None:
            raise RuntimeError("Conexión de escritura no disponible.")
        cur = self.conn.cursor()
        batch_counts: Counter = Counter()
        started = time.monotonic()
        cleanup_queries = 0
        try:
            # Primero MERGE de principal y extensiones. Parámetros enlazados, sin SQL por dato.
            for sheet in SCHEMAS:
                rows = [row for conversation in self.buffer for row in conversation[sheet]]
                columns = [col[0] for col in SCHEMAS[sheet]]
                sql = merge_statement(self.config, sheet)
                for offset in range(0, len(rows), 500):
                    cur.executemany(sql,
                                    [tuple(row[col] for col in columns) for row in rows[offset:offset + 500]])
                batch_counts[sheet] += len(rows)
            # El snapshot completo puede cambiar sus secuencias/medio al cerrar un email.
            # Quitar únicamente claves obsoletas de las conversaciones de ESTE lote.
            written_at = time.monotonic()
            for sheet in reversed(list(SCHEMAS)):
                for sql, parameters in cleanup_batches(self.config, sheet, self.buffer):
                    cur.execute(sql, parameters)
                    cleanup_queries += 1
            cleaned_at = time.monotonic()
            self.conn.commit()
        except Exception:
            self.conn.rollback()
            raise RuntimeError("Falló el lote HANA: rollback de las cuatro tablas; los lotes anteriores permanecen confirmados.") from None
        finally:
            cur.close()
        self.counts.update(batch_counts)
        LOG.info("Tiempos lote HANA (%s conversaciones): escritura=%.2fs | limpieza=%.2fs (%s consultas) | commit=%.2fs",
                 len(self.buffer), written_at - started, cleaned_at - written_at,
                 cleanup_queries, time.monotonic() - cleaned_at)
        LOG.info("Registros escritos HANA por tabla: %s", dict(self.counts))
        self.buffer.clear()

    def close(self) -> None:
        if self.conn is not None:
            self.conn.close()


class ExcelWriter:
    """Cuatro hojas exactas, en streaming; rechaza exceso de filas/texto en vez de truncar."""
    def __init__(self, config: Config):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.utils import get_column_letter
        directory = Path(config.get("OUTPUT_DIR") or Path(__file__).resolve().parent / "output")
        directory.mkdir(parents=True, exist_ok=True)
        stamp = datetime.now(ZoneInfo(config.zone)).strftime("%Y%m%d_%H%M%S")
        self.path = directory / f"GNS_Interacciones_Omnicanal_{stamp}.xlsx"
        if self.path.exists():
            raise FileExistsError(f"Ya existe la salida: {self.path}")
        self.temporary = self.path.with_suffix(".part.xlsx")
        self.book = Workbook(write_only=True)
        self.counts: Counter = Counter()
        self.completed = False
        for name, schema in SCHEMAS.items():
            sheet = self.book.create_sheet(name)
            sheet.freeze_panes = "A2"
            for index, (column, kind, _, _) in enumerate(schema, 1):
                sheet.column_dimensions[get_column_letter(index)].width = 22 if kind == "TIMESTAMP" else min(42, max(16, len(column) + 2))
            header = []
            for column, _, _, _ in schema:
                cell = WriteOnlyCell(sheet, column)
                cell.fill = PatternFill("solid", fgColor="17365D")
                cell.font = Font(bold=True, color="FFFFFF")
                header.append(cell)
            sheet.append(header)

    def add(self, rows: dict[str, list[dict[str, Any]]]) -> None:
        from openpyxl.cell import WriteOnlyCell
        for name, items in rows.items():
            sheet = self.book[name]
            for row in items:
                if self.counts[name] >= 1048575:
                    raise ValueError(f"{name} excede el máximo de filas de Excel. Use SAP HANA o un rango menor.")
                cells = []
                for column, _, _, _ in SCHEMAS[name]:
                    value = row[column]
                    if isinstance(value, str) and len(value.encode("utf-16-le")) // 2 > 32767:
                        raise ValueError(f"{name}.{column} supera 32767 caracteres de Excel. Use SAP HANA; no se truncó.")
                    cell = WriteOnlyCell(sheet, value)
                    if isinstance(value, str):
                        cell.data_type = "s"  # DNI con ceros iniciales y texto que comience con =.
                        cell.number_format = "@"
                    elif isinstance(value, datetime):
                        cell.number_format = "yyyy-mm-dd hh:mm:ss"
                    cells.append(cell)
                sheet.append(cells)
                self.counts[name] += 1

    def finish(self) -> Path:
        from openpyxl.utils import get_column_letter
        for name, schema in SCHEMAS.items():
            self.book[name].auto_filter.ref = f"A1:{get_column_letter(len(schema))}{self.counts[name] + 1}"
        self.book.save(self.temporary)
        self.temporary.replace(self.path)
        self.completed = True
        LOG.info("Archivo Excel generado: %s", self.path)
        return self.path

    def close(self) -> None:
        # Cerrar streams incluso en error, sin publicar un Excel incompleto.
        if not self.completed:
            for sheet in self.book:
                try:
                    sheet.close()
                except Exception:
                    pass
            try:
                self.book.save(self.temporary)
            except Exception:
                pass
            if self.temporary.exists():
                self.temporary.unlink()
        self.book.close()


def overlaps_range(conversation: dict[str, Any], config: Config) -> bool:
    start, _ = day_interval(config.first_day, config.zone)
    _, end = day_interval(config.last_day, config.zone)
    conversation_start = timestamp(conversation.get("conversationStart"))
    conversation_end = timestamp(conversation.get("conversationEnd"))
    return bool(conversation_start and conversation_start < end and (conversation_end is None or conversation_end > start))


def run(config: Config) -> None:
    api = GenesysClient(config)
    hana = excel = None
    totals: Counter = Counter({name: 0 for name in SCHEMAS})
    routing: Counter = Counter()
    lookup = IdentificationLookup(config, api)
    started = time.monotonic()
    try:
        LOG.info("Rango local %s al %s inclusive | zona=%s | salida=%s | dry_run=%s",
                 config.first_day, config.last_day, config.zone, config.output, config.dry_run)
        api.authenticate()
        print("PYFLOW_PROGRESS=2", flush=True)
        catalogs = load_catalogs(config, api)
        resolve_filters(config, api, catalogs)
        print("PYFLOW_PROGRESS=10", flush=True)
        if config.output in ("hana", "both"):
            hana = HanaWriter(config)
        if config.output in ("excel", "both"):
            excel = ExcelWriter(config)
        seen: set[str] = set()
        days = (config.last_day - config.first_day).days + 1
        ids = config.filters.get("CONVERSATION_ID", [])

        def consume(conversations: Iterator[dict[str, Any]]) -> bool:
            for conversation in conversations:
                totals["descargadas"] += 1
                cid = conversation.get("conversationId")
                if not cid:
                    raise ValueError("Genesys devolvió una conversación sin conversationId.")
                if cid in seen:
                    continue
                seen.add(cid)
                if not overlaps_range(conversation, config):
                    continue
                attentions, transferred = build_attentions(conversation)
                if not attentions or not conversation_matches(conversation, config, transferred):
                    continue
                attrs = lookup.enrich(conversation)
                rows = normalize(conversation, attentions, attrs, catalogs, config)
                totals["filtradas"] += 1
                totals["transferidas"] += int(transferred)
                for sheet, items in rows.items():
                    totals[sheet] += len(items)
                routing.update(row["usedRouting"] or "sin_usedRouting" for row in rows["INTERACCIONES"])
                if excel is not None:
                    excel.add(rows)
                if hana is not None:
                    hana.add(rows)
                if totals["filtradas"] == 1 or totals["filtradas"] % 100 == 0:
                    LOG.info("Conversaciones después de filtros=%s | filas=%s | transferidas=%s",
                             totals["filtradas"], totals["INTERACCIONES"], totals["transferidas"])
                if config.max_conversations and totals["filtradas"] >= config.max_conversations:
                    return True
            return False

        if ids:
            LOG.info("Consulta directa de %s conversationId; se validarán rango y todos los filtros.", len(ids))
            consume(api.request("GET", f"/api/v2/analytics/conversations/{quote(cid, safe='')}/details") for cid in ids)
        else:
            for offset in range(days):
                day = config.first_day + timedelta(days=offset)
                stop = consume(daily_conversations(api, config, day))
                if hana is not None:
                    hana.flush()
                print(f"PYFLOW_PROGRESS={20 + int(65 * (offset + 1) / days)}", flush=True)
                if stop:
                    break
        print("PYFLOW_PROGRESS=90", flush=True)
        if hana is not None:
            hana.flush()
        if excel is not None:
            excel.finish()
        LOG.info("Resumen extracción: %s", dict(totals))
        LOG.info("SPD_IDENTIFICACION Analytics=%s | fallback=%s | sin identificación=%s | requests fallback=%s",
                 lookup.counts["analytics"], lookup.counts["fallback"], lookup.counts["sin_identificacion"], lookup.requests)
        LOG.info("Routing: %s", {key: routing[key] for key in ("Predictive", "Standard", "Bullseye", "Last", "Preferred", "sin_usedRouting")})
        LOG.info("Registros HANA: %s", dict(hana.counts) if hana else {})
        if config.dry_run:
            LOG.info("DRY_RUN: no se ejecutaron escrituras en SAP HANA.")
        print("PYFLOW_PROGRESS=100", flush=True)
    finally:
        if excel is not None:
            excel.close()
        if hana is not None:
            hana.close()
        api.close()
        LOG.info("Duración total: %.2f segundos", time.monotonic() - started)


def self_test() -> int:
    """Validaciones internas reproducibles. No usa credenciales ni abre conexiones."""
    import copy
    import io
    import tempfile
    import unittest
    from contextlib import redirect_stdout
    from unittest.mock import Mock, patch

    class OmnichannelTests(unittest.TestCase):
        def setUp(self) -> None:
            values = {key: str(meta.get("default", "")) for key, meta in PYFLOW_PARAMS.items()}
            self.c = Config(values, date(2026, 9, 20), date(2026, 9, 20), output="excel", api_sleep=0, poll_seconds=0)
            self.catalogs = {"users": {"u1": "Agente 1", "u2": "Agente 2", "u3": "Agente 3"},
                             "queues": {"q1": "Cola 1", "q2": "Cola 2"}, "wrapups": {"w1": "Resuelto"}}
            self.api = Mock()
            self.api.entities.return_value = []
            resolve_filters(self.c, self.api, self.catalogs)

        def seg(self, start: int, end: int | None, kind: str = "interact", **kwargs: Any) -> dict[str, Any]:
            base = datetime(2026, 9, 20, 6, tzinfo=UTC)
            row = {"segmentStart": utc_text(base + timedelta(seconds=start)), "segmentType": kind, **kwargs}
            if end is not None:
                row["segmentEnd"] = utc_text(base + timedelta(seconds=end))
            return row

        def conv(self, media: str = "voice", count: int = 1) -> dict[str, Any]:
            participants = []
            for i in range(count):
                start = i * 100
                segments = [self.seg(start, start + 60, queueId=f"q{min(i + 1, 2)}",
                                     disconnectType="transfer" if i < count - 1 else "client"),
                            self.seg(start + 60, start + 70, "wrapup", wrapUpCode="w1", queueId=f"q{min(i + 1, 2)}")]
                participants.append({"participantId": f"p{i + 1}", "purpose": "agent", "userId": f"u{i + 1}",
                    "sessions": [{"sessionId": f"s{i + 1}", "mediaType": media, "direction": "inbound", "segments": segments,
                                  "metrics": [{"name": "tHandle", "value": 70000, "emitDate": segments[-1]["segmentEnd"]}]}]})
            return {"conversationId": "test-conversation", "conversationStart": "2026-09-20T06:00:00.000Z",
                    "conversationEnd": "2026-09-20T07:00:00.000Z", "originatingDirection": "inbound", "participants": participants}

        def normalized(self, conv: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
            attentions, _ = build_attentions(conv)
            return normalize(conv, attentions, attributes_of(conv), self.catalogs, self.c)

        def test_01_honduras_day(self) -> None:
            a, b = day_interval(date(2026, 9, 20), "America/Tegucigalpa")
            self.assertEqual(utc_text(a), "2026-09-20T06:00:00.000Z")
            self.assertEqual(utc_text(b), "2026-09-21T06:00:00.000Z")

        def test_02_technical_segments_do_not_multiply(self) -> None:
            conv = self.conv()
            session = conv["participants"][0]["sessions"][0]
            session["segments"] += [self.seg(10, 15, "hold"), self.seg(15, 17, "delay"), self.seg(17, 18, "system")]
            conv["participants"] += [{"participantId": "tech-" + purpose, "purpose": purpose,
                                      "sessions": [{"mediaType": "voice", "segments": [self.seg(0, 5)], "flow": {"flowName": "IVR"} if purpose == "ivr" else {}}]}
                                     for purpose in ("ivr", "acd")]
            rows = self.normalized(conv)["INTERACCIONES"]
            self.assertEqual(len(rows), 1)
            self.assertNotEqual(rows[0]["purpose"], "ivr")
            self.assertNotIn("IVR", rows[0]["flowName"] or "")

        def test_03_agent_sequences_and_transfer(self) -> None:
            rows = self.normalized(self.conv(count=3))["INTERACCIONES"]
            self.assertEqual([r["sequence"] for r in rows], [1, 2, 3])
            self.assertEqual([r["transferSequence"] for r in rows], [None, 1, 2])
            self.assertEqual(rows[0]["transferToUserName"], "Agente 2")
            self.assertEqual(rows[0]["transferToQueueName"], "Cola 2")

        def test_04_transfer_unknown_destination(self) -> None:
            conv = self.conv()
            conv["participants"][0]["sessions"][0]["metrics"].append({"name": "nTransferred", "value": 1})
            attentions, transferred = build_attentions(conv)
            self.assertTrue(transferred)
            rows = self.normalized(conv)["INTERACCIONES"]
            self.assertIsNone(rows[0]["transferToUserName"])
            self.assertIsNone(rows[0]["transferToQueueName"])

        def test_05_identification_is_text(self) -> None:
            conv = self.conv()
            conv["participants"][0]["attributes"] = {"SPD_IDENTIFICACION": "0012345678901", "Identificacion": "00012", "ScripterIdentificacion": "0013"}
            row = self.normalized(conv)["INTERACCIONES"][0]
            self.assertEqual(row["SPD_IDENTIFICACION"], "0012345678901")
            self.assertEqual(row["Identificacion"], "00012")

        def test_06_email_five_agents_and_missing_dni(self) -> None:
            rows = self.normalized(self.conv("email", 5))
            self.assertEqual([r["sequence"] for r in rows["INTERACCIONES"]], [1, 2, 3, 4, 5])
            self.assertEqual(len(rows["INTERACCIONES_EMAIL"]), 5)
            self.assertIsNone(rows["INTERACCIONES"][0]["SPD_IDENTIFICACION"])

        def test_07_whatsapp_flow_without_extra_rows(self) -> None:
            conv = self.conv("message")
            conv["participants"].insert(0, {"purpose": "botflow", "participantId": "bot", "sessions": [{
                "sessionId": "bot-session", "mediaType": "message", "direction": "inbound",
                "segments": [self.seg(-30, 0)], "flow": {"flowId": "bot-id", "flowName": "BANCOATLANTIDA ABI WhatsApp PROD", "flowType": "DIGITALBOT"}}]})
            rows = self.normalized(conv)
            self.assertEqual(len(rows["INTERACCIONES"]), 1)
            self.assertEqual(rows["INTERACCIONES"][0]["flowName"], "BANCOATLANTIDA ABI WhatsApp PROD")
            self.c.filters["FLOW_ID"] = ["bot-id"]
            self.c.values["PARTICIPANT_PURPOSE"] = "botflow"
            self.assertTrue(conversation_matches(conv, self.c, False))

        def test_08_voice_outbound_callback(self) -> None:
            for media in ("voice", "callback"):
                conv = self.conv(media)
                session = conv["participants"][0]["sessions"][0]
                session.update(ani="tel:+50400123", dnis="tel:+50400456", outboundCampaignId="campaign-id",
                               outboundContactListId="list-id", outboundContactId="contact-id",
                               callbackNumbers=["00123"], callbackScheduledTime="2026-09-20T12:00:00Z")
                rows = self.normalized(conv)
                row = rows["INTERACCIONES_VOICE"][0]
                self.assertEqual(row["ani"], "tel:+50400123")
                self.assertEqual(row["outboundContactId"], "contact-id")
                self.assertEqual(row["callbackScheduledTime"], datetime(2026, 9, 20, 6))
                self.assertEqual(rows["INTERACCIONES"][0]["mediaType"], media)

        def test_09_predictive_and_standard(self) -> None:
            conv = self.conv()
            session = conv["participants"][0]["sessions"][0]
            fields = {"usedRouting": "Predictive", "requestedRoutings": ["Predictive"], "routingRule": "rule",
                      "routingRuleType": "Predictive", "selectedAgentId": "u1", "selectedAgentRank": 1,
                      "proposedAgents": [{"userId": "u1", "rank": 1}], "eligibleAgentCounts": [3],
                      "waitingInteractionCounts": [5], "agentBullseyeRing": 2, "routingRing": 3}
            session.update(fields)
            session["segments"][0]["requestedRoutingUserIds"] = ["u1"]
            row = self.normalized(conv)["INTERACCIONES"][0]
            for key, value in fields.items():
                self.assertEqual(json.loads(row[key]) if isinstance(value, list) else row[key], value)
            self.assertEqual(json.loads(row["requestedRoutingUserIds"]), ["u1"])
            session["usedRouting"] = "Standard"
            self.assertEqual(self.normalized(conv)["INTERACCIONES"][0]["usedRouting"], "Standard")
            session.pop("usedRouting")
            self.assertIsNone(self.normalized(conv)["INTERACCIONES"][0]["usedRouting"])
            self.c.values["ROUTING_TYPE"] = "Predictive"
            self.assertFalse(conversation_matches(conv, self.c, False))

        def test_10_multiple_selectors_and_manual_priority(self) -> None:
            for key, (selector, _, _) in FILTER_SELECTORS.items():
                self.c.values[key] = ""
                self.c.values[selector] = '[{"id":"id-a","name":"Nombre; ambiguo"},{"id":"id-b","name":"Nombre"}]'
                resolve_filters(self.c, self.api, self.catalogs)
                self.assertEqual(self.c.filters[key], ["id-a", "id-b"])
                self.c.values[key] = "manual;manual,otro\notro"
                self.c.values[selector] = "[malformed ignored because manual wins"
                resolve_filters(self.c, self.api, self.catalogs)
                self.assertEqual(self.c.filters[key], ["manual", "otro"])

        def test_11_direction_and_combined_filters(self) -> None:
            conv = self.conv("message")
            conv["participants"][0]["sessions"][0]["direction"] = "outbound"
            self.c.values.update(ORIGINAL_DIRECTION="inbound", DIRECTION="outbound", MEDIA_TYPE="message")
            self.c.filters.update(USER_ID=["u1"], QUEUE_ID=["q1"], WRAPUP_CODE_ID=["w1"])
            self.assertTrue(conversation_matches(conv, self.c, False))
            self.c.values["DIRECTION"] = "inbound"
            self.assertFalse(conversation_matches(conv, self.c, False))

        def test_12_bot_only(self) -> None:
            conv = self.conv("message")
            conv["participants"][0].update(purpose="customer", userId=None)
            for number in range(3):
                conv["participants"].append({"purpose": "botflow", "participantId": f"bot{number}", "sessions": [{
                    "mediaType": "message", "sessionId": f"bot-s{number}", "segments": [self.seg(number, number + 1)],
                    "flow": {"flowId": f"b{number}", "flowName": f"Bot {number}"}}]})
            rows = self.normalized(conv)["INTERACCIONES"]
            self.assertEqual(len(rows), 1)
            self.assertIn("Bot 2", rows[0]["flowName"])

        def test_13_fallback_cache_and_no_dni_inference(self) -> None:
            conv = self.conv("message")
            api = Mock()
            api.request.return_value = {"participants": [{"attributes": {"SPD_IDENTIFICACION": "000321"}}]}
            lookup = IdentificationLookup(self.c, api)
            for _ in range(2):
                self.assertEqual(lookup.enrich(conv)["SPD_IDENTIFICACION"], "000321")
            self.assertEqual(api.request.call_count, 1)
            conv["conversationId"] = "other"
            api.request.return_value = {"participants": [{"addressFrom": "00123@example.com"}]}
            self.assertNotIn("SPD_IDENTIFICACION", lookup.enrich(conv))

        def test_14_excel_exact_schema_and_text(self) -> None:
            from openpyxl import load_workbook
            with tempfile.TemporaryDirectory() as directory:
                self.c.values["OUTPUT_DIR"] = directory
                conv = self.conv()
                conv["participants"][0]["attributes"] = {"SPD_IDENTIFICACION": "000123", "SPD_Comentario": "=1+1"}
                for participant in conv["participants"]:
                    for session in participant.get("sessions", []):
                        for segment in session.get("segments", []):
                            if segment.get("wrapUpCode"):
                                segment["wrapUpCode"] = "LONG-CODE-" * 30
                writer = ExcelWriter(self.c)
                try:
                    writer.add(self.normalized(conv))
                    path = writer.finish()
                finally:
                    writer.close()
                book = load_workbook(path)
                self.assertEqual(book.sheetnames, list(SCHEMAS))
                for name, schema in SCHEMAS.items():
                    self.assertEqual([c.value for c in book[name][1]], [c[0] for c in schema])
                    self.assertEqual(book[name].freeze_panes, "A2")
                    self.assertTrue(book[name].auto_filter.ref)
                columns = [c[0] for c in SCHEMAS["INTERACCIONES"]]
                sheet = book["INTERACCIONES"]
                self.assertEqual(sheet.cell(2, columns.index("SPD_IDENTIFICACION") + 1).value, "000123")
                self.assertEqual(sheet.cell(2, columns.index("SPD_Comentario") + 1).data_type, "s")
                self.assertEqual(sheet.cell(2, columns.index("wrapUpCode") + 1).value, "LONG-CODE-" * 30)
                book.close()

        def test_hana_actual_text_capacity(self) -> None:
            writer = HanaWriter(self.c.__class__(**{**self.c.__dict__, "dry_run": True}))
            writer.config = copy.deepcopy(self.c)
            writer.config.dry_run = False
            writer.config.batch_size = 100
            writer.conn = Mock()
            writer.text_limits = {("INTERACCIONES", "wrapUpCode"): 250}
            rows = self.normalized(self.conv())
            rows["INTERACCIONES"][0]["wrapUpCode"] = "x" * 250
            writer.add(rows)
            self.assertEqual(len(writer.buffer), 1)
            rows["INTERACCIONES"][0]["wrapUpCode"] = "x" * 251
            with self.assertRaisesRegex(ValueError, "longitud=251, capacidad HANA=250; conversationId="):
                writer.add(rows)
            self.assertEqual(len(writer.buffer), 1)
            writer.conn.cursor.assert_not_called()

        def test_15_idempotent_hana_mock_and_rollback(self) -> None:
            rows = self.normalized(self.conv())
            writer = HanaWriter(self.c.__class__(**{**self.c.__dict__, "dry_run": True}))
            writer.config = copy.deepcopy(self.c)
            writer.text_limits = {}
            writer.conn = Mock()
            stored: dict[str, dict[tuple[Any, Any], tuple[Any, ...]]] = defaultdict(dict)
            cursor = writer.conn.cursor.return_value

            def execute_many(sql_text: str, values: list[tuple[Any, ...]]) -> None:
                sheet = next(s for s in SCHEMAS if f'."{self.c.get(TABLE_KEYS[s])}" AS T' in sql_text)
                columns = [c[0] for c in SCHEMAS[sheet]]
                for row in values:
                    stored[sheet][(row[columns.index("conversationId")], row[columns.index("sequence")])] = row

            cursor.executemany.side_effect = execute_many
            for _ in range(2):
                writer.add(rows)
                writer.flush()
            self.assertEqual(len(stored["INTERACCIONES"]), 1)
            self.assertEqual(writer.conn.commit.call_count, 2)
            rows["INTERACCIONES"][0]["wrapUpCode"] = "updated"
            writer.add(rows)
            writer.flush()
            cols = [c[0] for c in SCHEMAS["INTERACCIONES"]]
            self.assertEqual(next(iter(stored["INTERACCIONES"].values()))[cols.index("wrapUpCode")], "updated")
            cursor.executemany.side_effect = RuntimeError("simulated")
            writer.add(rows)
            with self.assertRaises(RuntimeError):
                writer.flush()
            writer.conn.rollback.assert_called_once()

        def test_cleanup_batches_preserve_exact_keys(self) -> None:
            import sqlite3
            conn = sqlite3.connect(":memory:")
            conn.execute('ATTACH DATABASE ":memory:" AS BI_SS')
            sheet = "INTERACCIONES_EMAIL"
            table = identifier(self.c.get("HANA_SCHEMA")) + "." + identifier(self.c.get(TABLE_KEYS[sheet]))
            conn.execute(f'CREATE TABLE {table} (CONVERSATION_ID TEXT, SECUENCIA INTEGER)')
            snapshots = []
            expected = {("outside", 9)}
            initial = [("outside", 9)]
            for index in range(205):
                cid = f"cid-{index}"
                keep = [1, 3] if index % 2 else []
                snapshots.append({"INTERACCIONES": [{"conversationId": cid}],
                                  sheet: [{"conversationId": cid, "sequence": seq} for seq in keep]})
                initial.extend((cid, seq) for seq in (1, 2, 3, 4))
                expected.update((cid, seq) for seq in keep)
            conn.executemany(f'INSERT INTO {table} VALUES (?, ?)', initial)
            queries = list(cleanup_batches(self.c, sheet, snapshots))
            self.assertEqual(len(queries), 3)
            for sql, params in queries:
                conn.execute(sql, params)
            self.assertEqual(set(conn.execute(f'SELECT * FROM {table}')), expected)
            conn.close()

        def test_cleanup_failure_rolls_back_whole_batch(self) -> None:
            writer = HanaWriter(self.c.__class__(**{**self.c.__dict__, "dry_run": True}))
            writer.config = copy.deepcopy(self.c)
            writer.text_limits = {}
            writer.conn = Mock()
            writer.conn.cursor.return_value.execute.side_effect = RuntimeError("cleanup failure")
            writer.add(self.normalized(self.conv()))
            with self.assertRaises(RuntimeError):
                writer.flush()
            writer.conn.commit.assert_not_called()
            writer.conn.rollback.assert_called_once()
            self.assertEqual(dict(writer.counts), {})
            self.assertEqual(len(writer.buffer), 1)

        def test_16_dry_run_never_connects_writer(self) -> None:
            self.c.dry_run = True
            with patch(__name__ + ".hana_connect", side_effect=AssertionError("No connection")):
                writer = HanaWriter(self.c)
                writer.add(self.normalized(self.conv()))
                writer.flush()
                writer.close()

        def test_17_pagination_and_jobs(self) -> None:
            api = Mock()
            api.request.side_effect = [{"jobId": "job"}, {"state": "QUEUED"}, {"state": "FULFILLED"},
                                       {"conversations": [self.conv()], "cursor": "second"}, {"conversations": [self.conv()]}]
            rows = list(daily_conversations(api, self.c, self.c.first_day))
            self.assertEqual(len(rows), 2)
            self.assertEqual(api.request.call_args.kwargs["params"]["cursor"], "second")
            self.assertEqual(api.request.call_args_list[0].kwargs["json"]["interval"],
                             "2026-09-20T06:00:00.000Z/2026-09-21T06:00:00.000Z")

        def test_18_same_session_repeated_attention_no_duplicate_metrics(self) -> None:
            conv = self.conv()
            session = conv["participants"][0]["sessions"][0]
            session["segments"] += [self.seg(100, 120), self.seg(120, 130, "wrapup", wrapUpCode="w1")]
            rows = self.normalized(conv)["INTERACCIONES"]
            self.assertEqual(len(rows), 2)
            self.assertEqual(sum(r["tHandle"] or 0 for r in rows), 70000)

        def test_19_cli_precedence_and_dates(self) -> None:
            environment = {"GENESYS_CLIENT_ID": "test", "GENESYS_CLIENT_SECRET": "test", "GENESYS_REGION": "mypurecloud.com",
                           "DATE": "2026-01-01", "OUTPUT_DESTINATION": "Excel"}
            with patch.dict(os.environ, environment, clear=True), patch(__name__ + ".load_dotenv"):
                config = load_config(["--start-date", "20/09/2026", "--end-date", "23/09/2026", "--dry-run"])
                self.assertEqual(config.first_day, date(2026, 9, 20))
                self.assertEqual(config.last_day, date(2026, 9, 23))
                self.assertTrue(config.dry_run)
                with self.assertRaises(ValueError):
                    load_config(["--start-date", "2026-09-20"])

        def test_20_schema_sizes_and_ivr_exclusion(self) -> None:
            self.assertEqual([len(schema) for schema in SCHEMAS.values()], [110, 35, 18, 10])
            self.assertFalse(any("SPD_IVR" in column[0] for schema in SCHEMAS.values() for column in schema))
            conv = self.conv()
            conv["participants"][0]["purpose"] = "ivr"
            self.assertEqual(build_attentions(conv)[0], [])

        def test_21_no_transfer_from_many_segments_or_later_email(self) -> None:
            conv = self.conv("email", 2)
            first = conv["participants"][0]["sessions"][0]
            first["segments"][0]["disconnectType"] = "client"
            self.assertFalse(build_attentions(conv)[1])
            self.assertTrue(all(r["transferSequence"] is None for r in self.normalized(conv)["INTERACCIONES"]))

        def test_22_retry_after_and_token_refresh(self) -> None:
            self.c.values.update(GENESYS_REGION="mypurecloud.com", GENESYS_CLIENT_ID="test", GENESYS_CLIENT_SECRET="test")
            import types
            request_module = types.SimpleNamespace(Session=Mock, Timeout=TimeoutError, ConnectionError=ConnectionError)
            with patch.dict(sys.modules, {"requests": request_module}):
                api = GenesysClient(self.c)
            responses = []
            for status in (401, 429, 200):
                response = Mock(ok=status == 200, status_code=status, content=b"{}", headers={"Retry-After": "0"})
                response.json.return_value = {"result": True}
                responses.append(response)
            api.http = Mock()
            api.http.request.side_effect = responses
            auth = Mock(ok=True)
            auth.json.return_value = {"access_token": "test", "expires_in": 3600}
            api.http.post.return_value = auth
            with patch("time.sleep"):
                self.assertEqual(api.request("GET", "/api/v2/test"), {"result": True})
            self.assertEqual(api.http.post.call_count, 2)
            self.assertEqual(api.http.request.call_count, 3)

        def test_23_max_conversations_after_filters(self) -> None:
            self.c.max_conversations = 1
            self.c.output = "hana"
            self.c.dry_run = True
            self.c.values["USER_ID"] = "u2"
            rejected = self.conv()
            accepted = self.conv()
            accepted["conversationId"] = "accepted"
            accepted["participants"][0]["userId"] = "u2"
            accepted["participants"][0]["attributes"] = {"SPD_IDENTIFICACION": "0001"}
            extra = copy.deepcopy(accepted)
            extra["conversationId"] = "extra"
            sink = Mock()
            sink.counts = Counter()
            with patch(__name__ + ".GenesysClient", return_value=Mock()), \
                 patch(__name__ + ".load_catalogs", return_value=self.catalogs), \
                 patch(__name__ + ".HanaWriter", return_value=sink), \
                 patch(__name__ + ".daily_conversations", return_value=iter([rejected, accepted, extra])), redirect_stdout(io.StringIO()):
                run(self.c)
            self.assertEqual(sink.add.call_count, 1)
            self.assertEqual(sink.add.call_args.args[0]["INTERACCIONES"][0]["conversationId"], "accepted")

        def test_24_daily_jobs_and_cross_day_deduplication(self) -> None:
            self.c.last_day = date(2026, 9, 21)
            self.c.output = "hana"
            self.c.dry_run = True
            conv = self.conv("email")
            conv["conversationEnd"] = "2026-09-22T00:00:00Z"
            conv["participants"][0]["attributes"] = {"SPD_IDENTIFICACION": "0001"}
            sink = Mock()
            sink.counts = Counter()
            with patch(__name__ + ".GenesysClient", return_value=Mock()), \
                 patch(__name__ + ".load_catalogs", return_value=self.catalogs), \
                 patch(__name__ + ".HanaWriter", return_value=sink), \
                 patch(__name__ + ".daily_conversations", side_effect=lambda *args: iter([conv])) as daily, redirect_stdout(io.StringIO()):
                run(self.c)
            self.assertEqual([call.args[2] for call in daily.call_args_list], [date(2026, 9, 20), date(2026, 9, 21)])
            self.assertEqual(sink.add.call_count, 1)

        def test_hana_physical_column_names(self) -> None:
            for source, target in {"conversationId": "CONVERSATION_ID", "sequence": "SECUENCIA",
                                   "mediaStatsMinConversationRFactor": "MEDIA_STATS_MIN_CONVERSATION_R_FACTOR",
                                   "SPD_Gestion_FH": "SPD_GESTION_FH", "SPD_tipoDeCanal": "SPD_TIPO_DE_CANAL"}.items():
                self.assertEqual(hana_column(source), target)
            for sheet in SCHEMAS:
                sql = merge_statement(self.c, sheet)
                self.assertIn('T."CONVERSATION_ID" = S."CONVERSATION_ID"', sql)
                self.assertIn('T."SECUENCIA" = S."SECUENCIA"', sql)
                self.assertNotIn('"conversationId"', sql)
                self.assertNotIn('"sequence"', sql)

        def test_25_legacy_selectors_and_catalog_queries(self) -> None:
            self.c.values["USER_NAME"] = "Agente 1"
            resolve_filters(self.c, self.api, self.catalogs)
            self.assertEqual(self.c.filters["USER_ID"], ["u1"])
            connection = Mock()
            cursor = connection.cursor.return_value
            cursor.fetchall.side_effect = [[("u1", "Agente 1")], [("q1", "Cola 1")],
                                          [("wrap_upcode",), ("wrap_upcode_name",), ("fecha_carga",)], [("w1", "Resuelto")]]
            with patch(__name__ + ".hana_connect", return_value=connection):
                loaded = load_catalogs(self.c, self.api)
            self.assertEqual(loaded["wrapups"]["w1"], "Resuelto")
            sql = " ".join(call.args[0] for call in cursor.execute.call_args_list)
            self.assertIn('"BI_SS"."GNS_API_CAT_CONCLUSIONES"', sql)
            self.assertIn('"wrap_upcode", "wrap_upcode_name"', sql)
            self.assertIn('"QUEUE_NAME"', sql)
            self.api.entities.assert_not_called()

        def test_26_flow_metrics_not_duplicated_across_agents(self) -> None:
            conv = self.conv("message", 2)
            conv["participants"].append({"purpose": "botflow", "sessions": [{"mediaType": "message",
                "segments": [self.seg(-40, -1)], "flow": {"flowId": "bot"},
                "metrics": [{"name": "tFlow", "value": 39000}]}]})
            rows = self.normalized(conv)["INTERACCIONES"]
            self.assertEqual(sum(r["tFlow"] or 0 for r in rows), 39000)
            self.assertEqual(rows[0]["flowId"], "bot")
            self.assertIsNone(rows[1]["flowId"])

        def test_27_transfer_and_campaign_post_filters(self) -> None:
            conv = self.conv()
            conv["participants"][0]["attributes"] = {"dialerCampaignId": "campaign", "dialerContactListId": "list"}
            self.c.filters.update(CAMPAIGN_ID=["campaign"], CONTACT_LIST_ID=["list"])
            self.c.values["TRANSFER_FILTER"] = "solo_transferidas"
            self.assertFalse(conversation_matches(conv, self.c, False))
            self.assertTrue(conversation_matches(conv, self.c, True))
            self.c.filters["CONTACT_LIST_ID"] = ["different"]
            self.assertFalse(conversation_matches(conv, self.c, True))

    # Garantía adicional: una prueba nunca puede tocar la red aunque se omita un mock.
    suite = unittest.defaultTestLoader.loadTestsFromTestCase(OmnichannelTests)
    with patch("socket.socket.connect", side_effect=AssertionError("Red prohibida en --self-test")):
        outcome = unittest.TextTestRunner(verbosity=2).run(suite)
    return 0 if outcome.wasSuccessful() else 1


def main(argv: list[str] | None = None) -> int:
    argv = sys.argv[1:] if argv is None else argv
    logging.basicConfig(level=logging.INFO, format="%(asctime)s | %(levelname)s | %(message)s", datefmt="%Y-%m-%d %H:%M:%S", stream=sys.stdout)
    if "--self-test" in argv:
        return self_test()
    try:
        config = load_config(argv)
        run(config)
        return 0
    except KeyboardInterrupt:
        LOG.error("Proceso interrumpido; no se publicará un Excel incompleto.")
        return 130
    except Exception as exc:
        # Sin traceback de librerías HTTP/HANA que pueda contener credenciales.
        message = str(exc)
        for key in ("GENESYS_CLIENT_SECRET", "HPR_PASSWORD"):
            secret = env_str(key)
            if secret:
                message = message.replace(secret, "********")
        LOG.error("Proceso fallido (%s): %s", type(exc).__name__, message)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
