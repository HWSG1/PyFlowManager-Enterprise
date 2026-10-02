# =========================================================
# GNS EXTRACTOR DE TRANSCRIPCIONES - Genesys Cloud -> SAP HANA + Excel/CSV
# =========================================================
#
# Objetivo:
#   Extraer transcripciones disponibles de Genesys Cloud en texto plano.
#   Una fila principal por CONVERSATION_ID; detalle por transcript en tablas hijas.
#   Principal V3: 18 columnas; dirección y medio como valores únicos separados por /.
#   Solo se cargan/exportan interacciones con TEXT no vacío.
#
# Modos:
#   1) solo_transcript:
#      Extrae conversación, comunicación y texto transcrito.
#
#   2) transcript_campania:
#      Además de la transcripción, intenta enriquecer con datos de campaña
#      usando ContactId y ContactListId del participante Dialer.
#
# Ejemplos:
#   py .\GNS_Extractor_Transcripciones_PyFlow.py --start-date 2026-06-01 --end-date 2026-06-03
#   py .\GNS_Extractor_Transcripciones_PyFlow.py --date 2026-06-01
##
# Notas:
#   - Carga las cinco tablas BI_SS.GNS_API_TRANSCRIPCIONES* ya creadas con el SQL adjunto.
#   - HPR_HOST/HPR_PORT/HPR_USER/HPR_PASSWORD del entorno; no escribe al espejo.
#   - Dependencias: requests, hdbcli, openpyxl; python-dotenv y tzdata según entorno.
#   - TEXT íntegro en HANA; si Excel no admite una celda, se genera .completo.csv.
#   - Enriquecimiento dinámico de campaña en .campania.jsonl, sin columnas extras HANA.
#   - Indicadores decimales opcionales incompatibles: NULL con advertencia y JSON
#     íntegro en normalization_warnings; no se inventan valores numéricos.
#   - --dry-run mantiene la consulta Genesys pero no conecta HANA ni genera Excel/CSV.
#   - El enriquecimiento de campaña se hace contra Genesys Outbound Contact Lists.
#   - Si no se indican fechas, usa DAYS_BACK para extraer últimos N días cerrados.
#   - No limita el rango máximo; usar filtros para controlar volumen.
#   - Permite filtros por campaña, lista de contacto, conversationId, usuario, cola y conclusión.
#   - Permite seleccionar flujo, medio, dirección original y propósito del participante.
#   - FLOW_ID manual tiene prioridad sobre FLOW_SELECTION_ID del selector por nombre.
#   - Sin filtro de propósito incluye conversaciones que no pasaron por agent.
#   - Requiere que Speech and Text Analytics tenga transcripción disponible.
#   - Siempre devuelve la conclusión original en columnas wrapUpCodeId y conclusion_original.
# =========================================================

import os
import sys
import csv
import json
import time
import math
import argparse
import logging
import traceback
import hashlib
import threading
import re
from decimal import Decimal, InvalidOperation, localcontext
from pathlib import Path
from urllib.parse import quote, urlparse
from concurrent.futures import wait, FIRST_COMPLETED
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass
from datetime import datetime, date, timedelta, timezone
from zoneinfo import ZoneInfo
from typing import Any, Dict, List, Optional, Tuple

import requests

try:
    from dotenv import load_dotenv
except Exception:
    def load_dotenv(*args, **kwargs):
        return False


PYFLOW_PARAMS = {
    "GENESYS_CLIENT_ID": {"type": "global", "global_key": "GENESYS_CLIENT_ID", "label": "Genesys Client ID", "required": True},
    "GENESYS_CLIENT_SECRET": {"type": "global", "global_key": "GENESYS_CLIENT_SECRET", "label": "Genesys Client Secret", "required": True, "secret": True},
    "GENESYS_REGION": {"type": "global", "global_key": "GENESYS_REGION", "label": "Genesys Region / Domain", "required": True},

    "DATE": {"type": "date", "label": "Fecha específica local", "required": False},
    "START_DATE": {"type": "date", "label": "Fecha inicial local", "required": False},
    "END_DATE": {"type": "date", "label": "Fecha final local", "required": False},
    "DAYS_BACK": {"type": "number", "label": "Días hacia atrás si no se indican fechas", "required": False, "default": "30"},
    "GENESYS_TIMEZONE": {"type": "text", "label": "Zona horaria Genesys", "required": True, "default": "America/Tegucigalpa"},

    "OUTPUT_MODE": {
        "type": "select",
        "label": "Salida requerida",
        "required": True,
        "options": ["solo_transcript", "transcript_campania"],
        "default": "solo_transcript"
    },
    "ORIGINAL_DIRECTION": {
        "type": "select",
        "label": "Dirección original",
        "required": False,
        "options": ["ambas", "inbound", "outbound"],
        "default": "ambas"
    },
    "FLOW_SELECTION_ID": {"type": "genesys_flow", "label": "Nombre de flujo (buscar y seleccionar)", "required": False},
    "FLOW_ID": {"type": "text", "label": "ID del flujo (prioridad sobre el nombre seleccionado)", "required": False},
    "MEDIA_TYPE": {"type": "select", "label": "Tipo de medio (voice = voz, message = mensaje)", "required": False, "options": ["todos", "voice", "message", "email", "chat", "callback", "cobrowse", "internalmessage", "screenmonitoring", "screenshare", "video", "unknown"], "default": "todos"},
    "PARTICIPANT_PURPOSE": {"type": "select", "label": "Segmento / propósito del participante", "required": False, "options": ["todos", "acd", "agent", "api", "botflow", "campaign", "customer", "dialer", "external", "fax", "group", "inbound", "ivr", "manual", "outbound", "station", "user", "voicemail", "voicesurveyflow", "workflow"], "default": "todos"},
    "CONVERSATION_ID": {"type": "tags", "label": "Conversation ID específico", "required": False},
    "USER_ID": {"type": "tags", "label": "User ID del agente", "required": False},
    "USER_NAME": {"type": "genesys_users", "label": "Nombre/correo del agente", "required": False},
    "QUEUE_ID": {"type": "tags", "label": "Queue ID", "required": False},
    "QUEUE_NAME": {"type": "genesys_queues", "label": "Nombre de cola", "required": False},
    "CAMPAIGN_ID": {"type": "tags", "label": "Campaign ID", "required": False},
    "CAMPAIGN_NAME": {"type": "genesys_campaigns", "label": "Nombre de campaña", "required": False},
    "CONTACT_LIST_ID": {"type": "tags", "label": "Contact List ID", "required": False},
    "CONTACT_LIST_NAME": {"type": "genesys_contactlists", "label": "Nombre lista de contacto", "required": False},
    "WRAPUP_CODE_ID": {"type": "tags", "label": "WrapUpCode ID opcional", "required": False},
    "WRAPUP_CODE_NAME": {"type": "genesys_wrapupcodes", "label": "Nombre de conclusión opcional", "required": False},
    "MAX_CONVERSATIONS": {"type": "number", "label": "Máximo conversaciones; vacío = todas", "required": False},
    "MAX_TRANSCRIPT_WORKERS": {"type": "number", "label": "Consultas paralelas de transcript", "required": False, "default": "6"},
    "OUTPUT_FORMAT": {"type": "select", "label": "Formato salida", "required": False, "options": ["xlsx", "csv"], "default": "xlsx"},
    "OUTPUT_CSV": {"type": "text", "label": "Ruta de salida opcional", "required": False},
    "JSON_OUTPUT_DIR": {"type": "text", "label": "Carpeta JSON opcional", "required": False}
}

LOGGER_NAME = "gns_extractor_transcripciones_pyflow"
TOKEN_LOCK = threading.Lock()
LATEST_TOKEN = ""


class TranscriptNotFound(Exception):
    """La comunicación no tiene transcripción disponible en Genesys."""
    pass


def _clean_env_value(value: Any, default: Optional[str] = None) -> Optional[str]:
    if value is None:
        return default
    text = str(value).strip()
    if text == "" or text.lower() in ("null", "none", "undefined"):
        return default
    return text


def env_str(name: str, default: Optional[str] = None, required: bool = False) -> str:
    value = _clean_env_value(os.getenv(name), default)
    if required and not value:
        raise ValueError(f"Falta configurar variable/parámetro requerido: {name}")
    return "" if value is None else str(value)


def env_int(name: str, default: int, required: bool = False) -> int:
    value = env_str(name, str(default), required=required)
    try:
        return int(value)
    except Exception:
        raise ValueError(f"El parámetro {name} debe ser numérico. Valor recibido: {value!r}")


def env_float(name: str, default: float, required: bool = False) -> float:
    value = env_str(name, str(default), required=required)
    try:
        return float(value)
    except Exception:
        raise ValueError(f"El parámetro {name} debe ser numérico. Valor recibido: {value!r}")


def env_bool(name: str, default: bool = False) -> bool:
    value = env_str(name, "", required=False).lower()
    if not value:
        return default
    return value in ("1", "true", "yes", "y", "si", "sí")


def split_filter_values(value: Any) -> List[str]:
    """Acepta valores separados por ;, coma o salto de linea y elimina duplicados."""
    if value is None:
        return []

    raw = str(value).replace("\r", "\n").replace(",", ";").replace("\n", ";")
    result: List[str] = []

    for item in raw.split(";"):
        text = item.strip()
        if text and text not in result:
            result.append(text)

    return result


def join_filter_values(values: List[str]) -> str:
    return ";".join([str(value).strip() for value in values if str(value).strip()])


def filter_contains(value: str, allowed_values: str) -> bool:
    items = split_filter_values(allowed_values)
    if not items:
        return True
    return str(value or "") in items


def normalize_genesys_domain(value: str) -> str:
    value = str(value or "mypurecloud.com").strip()
    value = value.replace("https://", "").replace("http://", "").strip("/")
    if value.startswith("api."):
        value = value[4:]
    if value.startswith("login."):
        value = value[6:]
    if value.startswith("apps."):
        value = value[5:]
    return value


def genesys_api_url_from_region(region: str) -> str:
    return f"https://api.{normalize_genesys_domain(region)}"


def genesys_login_url_from_region(region: str) -> str:
    return f"https://login.{normalize_genesys_domain(region)}/oauth/token"


class CompactConsoleFilter(logging.Filter):
    """Acota mensajes repetitivos entre workers; conserva errores y un balance final."""
    PREFIXES = (
        "HTTP %s sin reintento", "HTTP 429", "HTTP %s | intento", "Error request",
        "Job conversaciones estado", "Conversaciones recuperadas página",
        "Catálogo conclusiones página", "Consultando conversación específica",
        "Normalización numérica", "Enriquecimiento campaña", "Transcript |",
        "Fuente íntegra del transcript", "HANA | %s:", "Sin transcripción disponible |",
    )

    def __init__(self):
        super().__init__()
        self.lock = threading.Lock()
        self.stats = {}

    def filter(self, record):
        if record.levelno >= logging.ERROR:
            return True
        category = next((p for p in self.PREFIXES if str(record.msg).startswith(p)), None)
        if category is None:
            return True
        now = time.monotonic()
        with self.lock:
            stat = self.stats.setdefault(category, {"total": 0, "hidden": 0, "last": now})
            stat["total"] += 1
            if stat["total"] <= 3 or now - stat["last"] >= 60:
                stat["last"] = now
                return True
            stat["hidden"] += 1
            return False

    def summarize(self, logger):
        with self.lock:
            summary = [(key, value.copy()) for key, value in self.stats.items()]
        for category, stat in summary:
            if stat["hidden"]:
                logger.info("Resumen consola | %s | eventos=%s | líneas repetitivas omitidas=%s",
                            category, stat["total"], stat["hidden"])


def setup_logger() -> logging.Logger:
    logger = logging.getLogger(LOGGER_NAME)
    logger.setLevel(logging.INFO)
    logger.handlers.clear()
    logger.propagate = False
    formatter = logging.Formatter("%(asctime)s | %(levelname)s | %(message)s", "%Y-%m-%d %H:%M:%S")
    console = logging.StreamHandler(sys.stdout)
    console.setFormatter(formatter)
    console.addFilter(CompactConsoleFilter())
    logger.addHandler(console)
    logger.info("Consola compacta: hasta 3 ejemplos por categoría repetitiva y luego uno por minuto; errores fatales siempre visibles.")
    return logger


@dataclass
class Config:
    genesys_client_id: str
    genesys_client_secret: str
    genesys_api_url: str
    genesys_login_url: str
    timezone_name: str
    days_back: int
    output_mode: str
    original_direction: str
    conversation_id: str
    user_id: str
    user_name: str
    queue_id: str
    queue_name: str
    campaign_id: str
    campaign_name: str
    contact_list_id: str
    contact_list_name: str
    wrapup_code_id: str
    wrapup_code_name: str
    max_conversations: int
    max_transcript_workers: int
    page_size: int
    request_timeout: int
    max_retries: int
    api_sleep_seconds: float
    job_poll_seconds: int
    job_max_polls: int
    output_csv: str
    output_format: str
    save_transcript_json: bool
    json_output_dir: str
    log_every_n: int
    transcript_debug: bool
    dry_run: bool
    flow_id: str = ""
    media_type: str = "todos"
    participant_purpose: str = "todos"


def load_config() -> Config:
    load_dotenv()
    region = env_str("GENESYS_REGION", "mypurecloud.com", required=True)
    output_mode = env_str("OUTPUT_MODE", "solo_transcript")
    if output_mode not in ("solo_transcript", "transcript_campania"):
        raise ValueError("OUTPUT_MODE debe ser solo_transcript o transcript_campania")
    original_direction = env_str("ORIGINAL_DIRECTION", "ambas").lower()
    if original_direction not in ("ambas", "inbound", "outbound"):
        raise ValueError("ORIGINAL_DIRECTION debe ser ambas, inbound u outbound")
    media_type = env_str("MEDIA_TYPE", "todos").lower()
    participant_purpose = env_str("PARTICIPANT_PURPOSE", "todos").lower()
    for key, value in (("MEDIA_TYPE", media_type), ("PARTICIPANT_PURPOSE", participant_purpose)):
        if value not in PYFLOW_PARAMS[key]["options"]:
            raise ValueError(f"{key} no válido: {value}")
    output_format = env_str("OUTPUT_FORMAT", "xlsx").lower()
    if output_format not in ("xlsx", "csv"):
        raise ValueError("OUTPUT_FORMAT debe ser xlsx o csv")
    return Config(
        genesys_client_id=env_str("GENESYS_CLIENT_ID", required=True),
        genesys_client_secret=env_str("GENESYS_CLIENT_SECRET", required=True),
        genesys_api_url=genesys_api_url_from_region(region),
        genesys_login_url=genesys_login_url_from_region(region),
        timezone_name=env_str("GENESYS_TIMEZONE", "America/Tegucigalpa"),
        days_back=env_int("DAYS_BACK", 30),
        output_mode=output_mode,
        original_direction=original_direction,
        flow_id=env_str("FLOW_ID", "") or env_str("FLOW_SELECTION_ID", ""),
        media_type=media_type,
        participant_purpose=participant_purpose,
        conversation_id=env_str("CONVERSATION_ID", ""),
        user_id=env_str("USER_ID", ""),
        user_name=env_str("USER_NAME", ""),
        queue_id=env_str("QUEUE_ID", ""),
        queue_name=env_str("QUEUE_NAME", ""),
        campaign_id=env_str("CAMPAIGN_ID", ""),
        campaign_name=env_str("CAMPAIGN_NAME", ""),
        contact_list_id=env_str("CONTACT_LIST_ID", ""),
        contact_list_name=env_str("CONTACT_LIST_NAME", ""),
        wrapup_code_id=env_str("WRAPUP_CODE_ID", ""),
        wrapup_code_name=env_str("WRAPUP_CODE_NAME", ""),
        max_conversations=env_int("MAX_CONVERSATIONS", 0),
        max_transcript_workers=max(1, env_int("MAX_TRANSCRIPT_WORKERS", 6)),
        page_size=env_int("PAGE_SIZE", 500),
        request_timeout=env_int("REQUEST_TIMEOUT", 120),
        max_retries=env_int("MAX_RETRIES", 5),
        api_sleep_seconds=env_float("API_SLEEP_SECONDS", 0.0),
        job_poll_seconds=env_int("JOB_POLL_SECONDS", 5),
        job_max_polls=env_int("JOB_MAX_POLLS", 120),
        output_csv=env_str("OUTPUT_CSV", ""),
        output_format=output_format,
        save_transcript_json=env_bool("SAVE_TRANSCRIPT_JSON", False),
        json_output_dir=env_str("JSON_OUTPUT_DIR", ""),
        log_every_n=max(1, env_int("LOG_EVERY_N", 500)),
        transcript_debug=env_bool("TRANSCRIPT_DEBUG", False),
        dry_run=env_bool("DRY_RUN", False),
    )


def log_params(logger: logging.Logger, names: List[str]) -> None:
    secret_words = ("SECRET", "PASSWORD", "TOKEN", "KEY")
    logger.info("Parámetros recibidos:")
    for name in names:
        value = env_str(name, "")
        if any(w in name.upper() for w in secret_words) and value:
            value = "********"
        logger.info("- %s: %s", name, value if value else "<vacío>")


def parse_local_date(value: str) -> date:
    value = str(value or "").strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(value, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"Fecha inválida: {value!r}. Use YYYY-MM-DD o DD/MM/YYYY.")


def to_utc_z(dt: datetime) -> str:
    return dt.astimezone(timezone.utc).strftime("%Y-%m-%dT%H:%M:%S.000Z")


def parse_genesys_datetime(value: Any) -> Optional[datetime]:
    text = str(value or "").strip()
    if not text:
        return None

    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00"))
    except ValueError:
        return None


def format_local_datetime(value: Any, tz_name: str) -> str:
    dt = parse_genesys_datetime(value)
    if not dt:
        return ""
    return dt.astimezone(ZoneInfo(tz_name)).strftime("%Y-%m-%d %H:%M:%S")


def format_local_date(value: Any, tz_name: str) -> str:
    dt = parse_genesys_datetime(value)
    if not dt:
        return ""
    return dt.astimezone(ZoneInfo(tz_name)).strftime("%Y-%m-%d")


def duration_seconds(start_value: Any, end_value: Any) -> int:
    start = parse_genesys_datetime(start_value)
    end = parse_genesys_datetime(end_value)
    if not start or not end:
        return 0
    return max(0, int((end - start).total_seconds()))


def format_duration(seconds: int) -> str:
    seconds = max(0, int(seconds or 0))
    hours, remainder = divmod(seconds, 3600)
    minutes, secs = divmod(remainder, 60)
    if hours:
        return f"{hours}h {minutes}m {secs}s"
    if minutes:
        return f"{minutes}m {secs}s"
    return f"{secs}s"




def parse_dates(args: argparse.Namespace, tz_name: str) -> Tuple[str, str, str]:
    """
    Prioridad de fechas:
    1) DATE: extrae un solo día.
    2) START_DATE + END_DATE: rango local, END_DATE inclusivo.
    3) DAYS_BACK: si no se indican fechas, toma los últimos N días cerrados.

    Ejemplo DAYS_BACK=30:
    desde hoy 00:00 menos 30 días hasta hoy 00:00, en la zona horaria configurada.
    """
    tz = ZoneInfo(tz_name)

    if args.date:
        d = parse_local_date(args.date)
        start_local_dt = datetime(d.year, d.month, d.day, 0, 0, 0, tzinfo=tz)
        end_local_dt = start_local_dt + timedelta(days=1)
        return to_utc_z(start_local_dt), to_utc_z(end_local_dt), f"Día local {d.isoformat()}"

    if args.start_date or args.end_date:
        if not (args.start_date and args.end_date):
            raise ValueError("Debe informar START_DATE y END_DATE juntos, o dejar ambos vacíos para usar DAYS_BACK.")

        d1 = parse_local_date(args.start_date)
        d2 = parse_local_date(args.end_date)
        if d2 < d1:
            raise ValueError("La fecha final no puede ser menor que la fecha inicial.")

        start_local_dt = datetime(d1.year, d1.month, d1.day, 0, 0, 0, tzinfo=tz)
        end_local_dt = datetime(d2.year, d2.month, d2.day, 0, 0, 0, tzinfo=tz) + timedelta(days=1)
        return to_utc_z(start_local_dt), to_utc_z(end_local_dt), f"Rango local {d1.isoformat()} al {d2.isoformat()}"

    days_back = env_int("DAYS_BACK", 30)
    if days_back <= 0:
        raise ValueError("DAYS_BACK debe ser mayor que cero.")

    today = datetime.now(tz).replace(hour=0, minute=0, second=0, microsecond=0)
    start_local_dt = today - timedelta(days=days_back)
    end_local_dt = today

    return (
        to_utc_z(start_local_dt),
        to_utc_z(end_local_dt),
        f"Automático últimos {days_back} días cerrados"
    )

def request_with_retry(method: str, url: str, config: Config, logger: logging.Logger, **kwargs) -> requests.Response:
    """Request HTTP con reintentos.

    Para errores no recuperables, como 404 al pedir un transcript que no existe,
    se puede pasar no_retry_statuses={404} para fallar rápido sin esperar 5 reintentos.
    """
    no_retry_statuses = set(kwargs.pop("no_retry_statuses", set()) or set())
    # Solo los códigos explícitamente esperados durante una búsqueda alternativa.
    quiet_statuses = set(kwargs.pop("quiet_statuses", set()) or set()) & no_retry_statuses
    last_error = None
    global LATEST_TOKEN
    is_api = url.startswith(config.genesys_api_url + "/api/")

    for attempt in range(1, config.max_retries + 1):
        try:
            if is_api and LATEST_TOKEN:
                kwargs["headers"] = {**kwargs.get("headers", {}), "Authorization": "Bearer " + LATEST_TOKEN}
            response = requests.request(method, url, timeout=config.request_timeout, **kwargs)

            if is_api and response.status_code == 401 and attempt < config.max_retries:
                old_token = kwargs.get("headers", {}).get("Authorization", "")
                with TOKEN_LOCK:
                    if not LATEST_TOKEN or old_token == "Bearer " + LATEST_TOKEN:
                        LATEST_TOKEN = get_access_token(config, logger)
                continue

            if response.status_code in no_retry_statuses:
                if response.status_code >= 400 and response.status_code not in quiet_statuses:
                    logger.info("HTTP %s sin reintento | %s", response.status_code, response.text[:500])
                response.raise_for_status()
                return response

            if response.status_code == 429:
                retry_after = response.headers.get("Retry-After")
                wait = int(retry_after) if retry_after and retry_after.isdigit() else min(60, 5 * attempt)
                logger.warning("HTTP 429 | intento %s/%s | esperando %ss", attempt, config.max_retries, wait)
                time.sleep(wait)
                continue

            if response.status_code >= 500:
                wait = min(60, 5 * attempt)
                logger.warning("HTTP %s | intento %s/%s | esperando %ss", response.status_code, attempt, config.max_retries, wait)
                time.sleep(wait)
                continue

            if response.status_code >= 400:
                logger.error("Error HTTP %s | %s", response.status_code, response.text[:1000])

            response.raise_for_status()
            return response

        except requests.HTTPError as exc:
            last_error = exc
            status = getattr(exc.response, "status_code", None)
            if status in no_retry_statuses:
                raise
            if attempt >= config.max_retries:
                break
            wait = min(60, 5 * attempt)
            logger.warning("Error request | intento %s/%s | %s | esperando %ss", attempt, config.max_retries, exc, wait)
            time.sleep(wait)
        except Exception as exc:
            last_error = exc
            if attempt >= config.max_retries:
                break
            wait = min(60, 5 * attempt)
            logger.warning("Error request | intento %s/%s | %s | esperando %ss", attempt, config.max_retries, exc, wait)
            time.sleep(wait)

    raise RuntimeError(f"No se pudo completar request después de {config.max_retries} intentos: {last_error}")


def get_access_token(config: Config, logger: logging.Logger) -> str:
    logger.info("Solicitando token OAuth en Genesys Cloud...")
    response = request_with_retry(
        "POST",
        config.genesys_login_url,
        config,
        logger,
        data={
            "grant_type": "client_credentials",
            "client_id": config.genesys_client_id,
            "client_secret": config.genesys_client_secret,
        },
        headers={"Content-Type": "application/x-www-form-urlencoded"},
    )
    token = response.json().get("access_token")
    if not token:
        raise RuntimeError("No se recibió access_token desde Genesys.")
    logger.info("Token obtenido correctamente.")
    return token


def genesys_headers(token: str) -> Dict[str, str]:
    return {"Authorization": f"Bearer {token}", "Content-Type": "application/json"}



def normalize_text(value: str) -> str:
    return " ".join(str(value or "").strip().lower().split())


def find_first_by_name(items: List[Dict[str, Any]], wanted: str, name_keys: Tuple[str, ...] = ("name",)) -> Optional[Dict[str, Any]]:
    wanted_n = normalize_text(wanted)
    if not wanted_n:
        return None
    exact = []
    partial = []
    for item in items:
        for key in name_keys:
            name = normalize_text(str(item.get(key) or ""))
            if not name:
                continue
            if name == wanted_n:
                exact.append(item)
            elif wanted_n in name:
                partial.append(item)
    return exact[0] if exact else (partial[0] if partial else None)


def paged_get_entities(config: Config, token: str, logger: logging.Logger, path: str, page_size: int = 100, extra_query: str = "") -> List[Dict[str, Any]]:
    items: List[Dict[str, Any]] = []
    page = 1
    while True:
        sep = "&" if "?" in path else "?"
        url = f"{config.genesys_api_url}{path}{sep}pageSize={page_size}&pageNumber={page}"
        if extra_query:
            url += "&" + extra_query.lstrip("&")
        response = request_with_retry("GET", url, config, logger, headers=genesys_headers(token))
        data = response.json() if response.text else {}
        entities = data.get("entities") or []
        items.extend(entities)
        page_count = data.get("pageCount") or data.get("page_count") or page
        if not entities or page >= int(page_count):
            break
        page += 1
        time.sleep(config.api_sleep_seconds)
    return items


def resolve_queue_id(config: Config, token: str, logger: logging.Logger) -> str:
    if config.queue_id:
        return join_filter_values(split_filter_values(config.queue_id))
    if not config.queue_name:
        return ""

    names = split_filter_values(config.queue_name)
    logger.info("Resolviendo Queue por nombre: %s", ", ".join(names))
    queues = paged_get_entities(config, token, logger, "/api/v2/routing/queues")
    resolved: List[str] = []

    for name in names:
        found = find_first_by_name(queues, name)
        if not found:
            raise ValueError(f"No se encontró cola con nombre: {name}")
        logger.info("Queue resuelta: %s -> %s", found.get("name"), found.get("id"))
        resolved.append(str(found.get("id") or ""))

    return join_filter_values(resolved)


def resolve_user_id(config: Config, token: str, logger: logging.Logger) -> str:
    if config.user_id:
        return join_filter_values(split_filter_values(config.user_id))
    if not config.user_name:
        return ""

    resolved: List[str] = []

    for name in split_filter_values(config.user_name):
        logger.info("Resolviendo usuario por nombre/correo: %s", name)
        # La búsqueda de usuarios soporta q por nombre/correo en Genesys.
        users = paged_get_entities(config, token, logger, "/api/v2/users", extra_query=f"q={requests.utils.quote(name)}")
        found = find_first_by_name(users, name, ("name", "email", "username"))
        if not found and users:
            found = users[0]
        if not found:
            raise ValueError(f"No se encontró usuario con: {name}")
        logger.info("Usuario resuelto: %s -> %s", found.get("name") or found.get("email"), found.get("id"))
        resolved.append(str(found.get("id") or ""))

    return join_filter_values(resolved)


def resolve_campaign_id(config: Config, token: str, logger: logging.Logger) -> str:
    if config.campaign_id:
        return join_filter_values(split_filter_values(config.campaign_id))
    if not config.campaign_name:
        return ""

    names = split_filter_values(config.campaign_name)
    logger.info("Resolviendo campaña por nombre: %s", ", ".join(names))
    campaigns = paged_get_entities(config, token, logger, "/api/v2/outbound/campaigns")
    resolved: List[str] = []

    for name in names:
        found = find_first_by_name(campaigns, name)
        if not found:
            raise ValueError(f"No se encontró campaña con nombre: {name}")
        logger.info("Campaña resuelta: %s -> %s", found.get("name"), found.get("id"))
        resolved.append(str(found.get("id") or ""))

    return join_filter_values(resolved)


def resolve_contact_list_id(config: Config, token: str, logger: logging.Logger) -> str:
    if config.contact_list_id:
        return join_filter_values(split_filter_values(config.contact_list_id))
    if not config.contact_list_name:
        return ""

    names = split_filter_values(config.contact_list_name)
    logger.info("Resolviendo lista de contacto por nombre: %s", ", ".join(names))
    lists = paged_get_entities(config, token, logger, "/api/v2/outbound/contactlists")
    resolved: List[str] = []

    for name in names:
        found = find_first_by_name(lists, name)
        if not found:
            raise ValueError(f"No se encontró lista de contacto con nombre: {name}")
        logger.info("Lista de contacto resuelta: %s -> %s", found.get("name"), found.get("id"))
        resolved.append(str(found.get("id") or ""))

    return join_filter_values(resolved)


def resolve_wrapup_code_id(config: Config, token: str, logger: logging.Logger) -> str:
    if config.wrapup_code_id:
        return join_filter_values(split_filter_values(config.wrapup_code_id))
    if not config.wrapup_code_name:
        return ""

    names = split_filter_values(config.wrapup_code_name)
    logger.info("Resolviendo conclusiones por nombre: %s", ", ".join(names))
    wrapup_codes = paged_get_entities(config, token, logger, "/api/v2/routing/wrapupcodes")
    resolved: List[str] = []

    for name in names:
        found = find_first_by_name(wrapup_codes, name)
        if not found:
            raise ValueError(f"No se encontró conclusión con nombre: {name}")
        logger.info("Conclusión resuelta: %s -> %s", found.get("name"), found.get("id"))
        resolved.append(str(found.get("id") or ""))

    return join_filter_values(resolved)


def apply_resolved_filters(config: Config, token: str, logger: logging.Logger) -> Config:
    """Resuelve filtros por nombre a sus IDs y los coloca en el mismo config."""
    # Los selectores guardan ID y nombre; los nombres antiguos siguen siendo compatibles.
    for field in ("user", "queue", "campaign", "contact_list", "wrapup_code"):
        raw = getattr(config, field + "_name")
        if raw.lstrip().startswith("["):
            selected = json.loads(raw)
            if not isinstance(selected, list) or any(
                not isinstance(item, dict) or not isinstance(item.get("id"), str) or not item["id"].strip()
                for item in selected
            ):
                raise ValueError(f"Selección no válida: {field}")
            if not getattr(config, field + "_id"):
                setattr(config, field + "_id", join_filter_values([item["id"] for item in selected]))
            setattr(config, field + "_name", "")
    config.queue_id = resolve_queue_id(config, token, logger)
    config.user_id = resolve_user_id(config, token, logger)
    config.campaign_id = resolve_campaign_id(config, token, logger)
    config.contact_list_id = resolve_contact_list_id(config, token, logger)
    config.wrapup_code_id = resolve_wrapup_code_id(config, token, logger)
    return config


def validate_filter_safety(config: Config, logger: logging.Logger) -> None:
    filters = {
        "FLOW_ID/FLOW_SELECTION_ID": config.flow_id,
        "MEDIA_TYPE": "" if config.media_type == "todos" else config.media_type,
        "PARTICIPANT_PURPOSE": "" if config.participant_purpose == "todos" else config.participant_purpose,
        "ORIGINAL_DIRECTION": "" if config.original_direction == "ambas" else config.original_direction,
        "CONVERSATION_ID": config.conversation_id,
        "USER_ID/USER_NAME": config.user_id or config.user_name,
        "QUEUE_ID/QUEUE_NAME": config.queue_id or config.queue_name,
        "CAMPAIGN_ID/CAMPAIGN_NAME": config.campaign_id or config.campaign_name,
        "CONTACT_LIST_ID/CONTACT_LIST_NAME": config.contact_list_id or config.contact_list_name,
        "WRAPUP_CODE_ID/WRAPUP_CODE_NAME": config.wrapup_code_id or config.wrapup_code_name,
        "MAX_CONVERSATIONS": str(config.max_conversations) if config.max_conversations else "",
    }
    active = [k for k, v in filters.items() if str(v or "").strip()]
    if active:
        logger.info("Filtros activos: %s", ", ".join(active))
    else:
        logger.warning("No hay filtros adicionales aparte del rango de fechas.")


def build_predicate(dimension: str, value: str, operator: str = "matches") -> Dict[str, str]:
    pred = {"dimension": dimension, "value": value}
    if operator:
        pred["operator"] = operator
    return pred


def build_dimension_filter(dimension: str, values: str, operator: str = "matches") -> Optional[Dict[str, Any]]:
    items = split_filter_values(values)
    if not items:
        return None

    return {
        "type": "or" if len(items) > 1 else "and",
        "predicates": [build_predicate(dimension, item, operator) for item in items],
    }


def build_details_job_body(start_utc: str, end_utc: str, config: Config) -> Dict[str, Any]:
    """Construye el job de Analytics Details con filtros seguros.

    Nota: algunas dimensiones dependen de cómo Genesys registra los segmentos.
    Además del filtro en el job, se hace una validación posterior por atributos
    para campaña/lista cuando esos datos vienen en participantes Dialer.
    """
    segment_filters: List[Dict[str, Any]] = []

    for dimension, values in (
        ("mediaType", "" if config.media_type == "todos" else config.media_type),
        ("purpose", "" if config.participant_purpose == "todos" else config.participant_purpose),
        ("flowId", config.flow_id),
        ("wrapUpCode", config.wrapup_code_id),
        ("queueId", config.queue_id),
        ("userId", config.user_id),
        ("outboundCampaignId", config.campaign_id),
        ("outboundContactListId", config.contact_list_id),
    ):
        dimension_filter = build_dimension_filter(dimension, values)
        if dimension_filter:
            segment_filters.append(dimension_filter)

    body = {
        "order": "asc",
        "orderBy": "conversationStart",
        "interval": f"{start_utc}/{end_utc}",
    }
    if segment_filters:
        body["segmentFilters"] = segment_filters
    if config.original_direction in ("inbound", "outbound"):
        body["conversationFilters"] = [build_dimension_filter("originatingDirection", config.original_direction)]
    return body


def conversation_matches_post_filters(conversation: Dict[str, Any], config: Config) -> bool:
    """Filtro posterior para datos que pueden venir como atributos Dialer."""
    if config.conversation_id:
        cid = conversation.get("conversationId") or conversation.get("id")
        if not filter_contains(str(cid or ""), config.conversation_id):
            return False

    direction = str(conversation.get("originatingDirection") or "").lower()
    if config.original_direction in ("inbound", "outbound") and direction != config.original_direction:
        return False

    dialer = extract_dialer_attributes(conversation)
    if config.campaign_id and not filter_contains(dialer.get("CampaignId", ""), config.campaign_id):
        return False
    if config.contact_list_id and not filter_contains(dialer.get("ContactListId", ""), config.contact_list_id):
        return False

    participants = conversation.get("participants") or []
    sessions = [s for p in participants for s in p.get("sessions") or []]
    # Flujo y agente pueden estar en sesiones diferentes de la conversación.
    if config.participant_purpose != "todos" and not any(
        p.get("purpose") == config.participant_purpose for p in participants
    ):
        return False
    if config.media_type != "todos" and not any(s.get("mediaType") == config.media_type for s in sessions):
        return False
    if config.flow_id and not any(
        filter_contains(str((s.get("flow") or {}).get("flowId") or ""), config.flow_id) for s in sessions
    ):
        return False

    return True


def create_conversation_details_job(config: Config, token: str, start_utc: str, end_utc: str, logger: logging.Logger) -> str:
    url = f"{config.genesys_api_url}/api/v2/analytics/conversations/details/jobs"
    body = build_details_job_body(start_utc, end_utc, config)
    logger.info("Creando job de conversaciones detalle...")
    response = request_with_retry("POST", url, config, logger, headers=genesys_headers(token), json=body)
    data = response.json()
    job_id = data.get("jobId") or data.get("id") or data.get("job", {}).get("id")
    if not job_id:
        raise RuntimeError(f"No se recibió jobId. Respuesta: {data}")
    logger.info("Job conversaciones creado: %s", job_id)
    return job_id


def wait_details_job(config: Config, token: str, job_id: str, logger: logging.Logger) -> None:
    url = f"{config.genesys_api_url}/api/v2/analytics/conversations/details/jobs/{job_id}"
    done_status = {"FULFILLED", "Succeeded", "Complete", "Completed", "SUCCESS", "COMPLETED"}
    fail_status = {"FAILED", "Failed", "Canceled", "Cancelled", "CANCELED", "CANCELLED"}
    for attempt in range(1, config.job_max_polls + 1):
        response = request_with_retry("GET", url, config, logger, headers=genesys_headers(token))
        data = response.json()
        status = data.get("state") or data.get("status") or data.get("job", {}).get("status")
        if status in done_status:
            logger.info("Job conversaciones finalizado: %s", status)
            return
        if status in fail_status:
            raise RuntimeError(f"Job conversaciones terminó con estado {status}. Respuesta: {data}")
        logger.info("Job conversaciones estado %s | intento %s/%s", status, attempt, config.job_max_polls)
        time.sleep(config.job_poll_seconds)
    raise TimeoutError(f"Job conversaciones {job_id} no finalizó.")


def fetch_details_job_results(config: Config, token: str, job_id: str, logger: logging.Logger) -> List[Dict[str, Any]]:
    conversations: List[Dict[str, Any]] = []
    cursor = ""
    while True:
        url = f"{config.genesys_api_url}/api/v2/analytics/conversations/details/jobs/{job_id}/results?pageSize={config.page_size}"
        if cursor:
            url += f"&cursor={cursor}"
        response = request_with_retry("GET", url, config, logger, headers=genesys_headers(token))
        data = response.json()
        items = data.get("conversations") or data.get("entities") or []
        conversations.extend(items)
        logger.info("Conversaciones recuperadas página: %s | acumulado: %s", len(items), len(conversations))
        if config.max_conversations and len(conversations) >= config.max_conversations:
            return conversations[: config.max_conversations]
        cursor = data.get("cursor") or data.get("nextCursor") or ""
        if not cursor or not items:
            break
        time.sleep(config.api_sleep_seconds)
    return conversations



def fetch_conversation_details_by_id(config: Config, token: str, conversation_id: str, logger: logging.Logger) -> Dict[str, Any]:
    """Obtiene el detalle de una conversación específica."""
    url = f"{config.genesys_api_url}/api/v2/analytics/conversations/{conversation_id}/details"
    logger.info("Consultando conversación específica: %s", conversation_id)
    response = request_with_retry("GET", url, config, logger, headers=genesys_headers(token))
    data = response.json() if response.text else {}
    return data


def get_wrapup_catalog(config: Config, token: str, logger: logging.Logger) -> Dict[str, str]:
    """Consulta el catálogo de conclusiones para traducir wrapUpCode ID a nombre."""
    logger.info("Consultando catálogo de conclusiones para enriquecer la data...")
    catalog: Dict[str, str] = {}
    page = 1

    while True:
        url = f"{config.genesys_api_url}/api/v2/routing/wrapupcodes?pageSize=100&pageNumber={page}"
        response = request_with_retry("GET", url, config, logger, headers=genesys_headers(token))
        data = response.json() if response.text else {}
        entities = data.get("entities") or []

        for item in entities:
            wrapup_id = str(item.get("id") or "")
            wrapup_name = str(item.get("name") or "")
            if wrapup_id:
                catalog[wrapup_id] = wrapup_name

        page_count = int(data.get("pageCount") or page)
        logger.info("Catálogo conclusiones página %s/%s | acumulado: %s", page, page_count, len(catalog))

        if page >= page_count or not entities:
            break

        page += 1
        time.sleep(config.api_sleep_seconds)

    logger.info("Catálogo de conclusiones cargado: %s", len(catalog))
    return catalog


def get_queue_catalog(config: Config, token: str, logger: logging.Logger) -> Dict[str, str]:
    logger.info("Consultando catalogo de colas para enriquecer la data...")
    catalog: Dict[str, str] = {}
    page = 1

    while True:
        url = f"{config.genesys_api_url}/api/v2/routing/queues?pageSize=100&pageNumber={page}"
        response = request_with_retry("GET", url, config, logger, headers=genesys_headers(token))
        data = response.json() if response.text else {}
        entities = data.get("entities") or []

        for item in entities:
            queue_id = str(item.get("id") or "")
            queue_name = str(item.get("name") or "")
            if queue_id:
                catalog[queue_id] = queue_name

        page_count = int(data.get("pageCount") or page)
        if page >= page_count or not entities:
            break

        page += 1
        if config.api_sleep_seconds > 0:
            time.sleep(config.api_sleep_seconds)

    logger.info("Catalogo de colas cargado: %s", len(catalog))
    return catalog


def extract_wrapup_codes(conversation: Dict[str, Any]) -> List[str]:
    """Extrae todos los wrapUpCode encontrados en participantes, sesiones y segmentos."""
    found: List[str] = []

    def add(value: Any) -> None:
        if value is None:
            return
        text = str(value).strip()
        if text and text not in found:
            found.append(text)

    for participant in conversation.get("participants") or []:
        add(participant.get("wrapUpCode"))
        for session in participant.get("sessions") or []:
            add(session.get("wrapUpCode"))
            for segment in session.get("segments") or []:
                add(segment.get("wrapUpCode"))

    return found


def conversation_was_agent_handled(conversation: Dict[str, Any]) -> bool:
    has_agent_voice = False

    for participant in conversation.get("participants") or []:
        purpose = str(participant.get("purpose") or "").lower()
        if purpose != "agent":
            continue

        for session in participant.get("sessions") or []:
            if str(session.get("mediaType") or "").lower() == "voice":
                has_agent_voice = True
                break

        if has_agent_voice:
            break

    has_wrapup = len(extract_wrapup_codes(conversation)) > 0
    return has_agent_voice and has_wrapup


def extract_queue_info(conversation: Dict[str, Any], queue_catalog: Dict[str, str]) -> Tuple[str, str]:
    queue_ids: List[str] = []
    queue_names: List[str] = []

    def add_unique(target: List[str], value: Any) -> None:
        text = str(value or "").strip()
        if text and text not in target:
            target.append(text)

    for participant in conversation.get("participants") or []:
        purpose = str(participant.get("purpose") or "").lower()
        if purpose == "acd":
            add_unique(queue_names, participant.get("participantName") or participant.get("name"))

        for session in participant.get("sessions") or []:
            for key in ("queueId", "requestedRoutingSkillIds"):
                value = session.get(key)
                if isinstance(value, list):
                    for item in value:
                        add_unique(queue_ids, item)
                else:
                    add_unique(queue_ids, value)

            for segment in session.get("segments") or []:
                add_unique(queue_ids, segment.get("queueId"))

    for queue_id in queue_ids:
        add_unique(queue_names, queue_catalog.get(queue_id))

    return ";".join(queue_ids), ";".join(queue_names)


def wrapup_names_from_ids(wrapup_ids: List[str], wrapup_catalog: Dict[str, str]) -> str:
    names: List[str] = []
    for wrapup_id in wrapup_ids:
        name = wrapup_catalog.get(wrapup_id, "")
        value = name or wrapup_id
        if value and value not in names:
            names.append(value)
    return ";".join(names)


def last_semicolon_value(value: Any) -> str:
    parts = [part.strip() for part in str(value or "").split(";") if part.strip()]
    return parts[-1] if parts else ""

def extract_dialer_attributes(conversation: Dict[str, Any]) -> Dict[str, str]:
    result = {"ContactId": "", "ContactListId": "", "CampaignId": ""}
    for participant in conversation.get("participants") or []:
        attrs = participant.get("attributes") or {}
        ptype = str(participant.get("participantType") or "")
        if ptype.lower() == "dialer" or attrs.get("dialerContactId") or attrs.get("dialerContactListId"):
            result["ContactId"] = str(attrs.get("dialerContactId") or result["ContactId"] or "")
            result["ContactListId"] = str(attrs.get("dialerContactListId") or result["ContactListId"] or "")
            result["CampaignId"] = str(attrs.get("dialerCampaignId") or result["CampaignId"] or "")
    return result


def extract_communication_ids(conversation: Dict[str, Any]) -> List[str]:
    ids: List[str] = []
    for participant in conversation.get("participants") or []:
        for session in participant.get("sessions") or []:
            media_type = str(session.get("mediaType") or "").lower()
            if media_type and media_type != "voice":
                continue
            # En Genesys, el endpoint transcripturl normalmente usa sessionId.
            # Evitamos peerId porque suele generar 404 y retrasa la ejecución.
            value = session.get("sessionId") or session.get("communicationId")
            if value and value not in ids:
                ids.append(str(value))
    return ids


def obtener_session_ids_para_transcript(conversation_details: Dict[str, Any], media_filter: str = "todos") -> List[Dict[str, Any]]:
    candidatos: List[Dict[str, Any]] = []
    seen = set()

    prioridad_por_purpose = {
        "customer": 2,
        "agent": 3,
        "ivr": 4,
        "acd": 5,
    }

    for participant in conversation_details.get("participants") or []:
        purpose = str(participant.get("purpose") or "").lower()
        participant_name = str(participant.get("participantName") or participant.get("name") or "")

        for session in participant.get("sessions") or []:
            media_type = str(session.get("mediaType") or "").lower()
            if media_filter != "todos" and media_type != media_filter:
                continue

            communication_id = session.get("sessionId") or session.get("communicationId")
            if not communication_id:
                continue

            communication_id = str(communication_id)
            if communication_id in seen:
                continue
            seen.add(communication_id)

            recording = session.get("recording") is True
            prioridad = prioridad_por_purpose.get(purpose, 6)
            if purpose == "customer" and recording:
                prioridad = 1

            candidatos.append({
                "communication_id": communication_id,
                "media_type": media_type,
                "purpose": purpose or "desconocido",
                "participant_name": participant_name,
                "recording": recording,
                "prioridad": prioridad,
            })

    return sorted(candidatos, key=lambda item: (item["prioridad"], item["communication_id"]))


def fetch_contact_from_genesys(config: Config, token: str, contact_list_id: str, contact_id: str, logger: logging.Logger) -> Dict[str, Any]:
    if not contact_list_id or not contact_id:
        return {}
    url = f"{config.genesys_api_url}/api/v2/outbound/contactlists/{contact_list_id}/contacts/{contact_id}"
    response = request_with_retry("GET", url, config, logger, headers=genesys_headers(token))
    data = response.json() if response.text else {}
    return data


def flatten_contact_data(contact: Dict[str, Any]) -> Dict[str, Any]:
    if not contact:
        return {}
    out: Dict[str, Any] = {}
    out["contact_callable"] = contact.get("callable")
    out["contact_phone_number_status"] = json.dumps(contact.get("phoneNumberStatus") or {}, ensure_ascii=False)
    data = contact.get("data") or {}
    if isinstance(data, dict):
        for key, value in data.items():
            safe_key = "camp_" + "".join(ch if ch.isalnum() or ch == "_" else "_" for ch in str(key))[:80]
            out[safe_key] = value
    return out


def default_output_path(output_format: str = "xlsx") -> str:
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_dir = env_str("OUTPUT_DIR", "") or os.path.join(os.getcwd(), "output")
    extension = "csv" if output_format == "csv" else "xlsx"
    return os.path.join(output_dir, f"GNS_Transcripciones_{ts}.{extension}")


def resolve_output_path(output_path: str, output_format: str) -> str:
    if not output_path:
        return default_output_path(output_format)

    output_path = str(output_path)
    if os.path.isdir(output_path) or output_path.endswith("/") or output_path.endswith("\\"):
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        extension = "csv" if output_format == "csv" else "xlsx"
        return os.path.join(output_path, f"GNS_Transcripciones_{ts}.{extension}")

    return output_path


# Modelo interno por transcript; STORAGE_TABLES define las columnas físicas.
KEYS = ["CONVERSATION_ID", "COMMUNICATION_ID", "TRANSCRIPT_INSTANCE_ID"]
MAIN = "GNS_API_TRANSCRIPCIONES"
TABLES = {
    MAIN: [(n, "NVARCHAR(250)") for n in KEYS] + [
        ("TRANSCRIPT_ID", "NVARCHAR(250)"), ("RECORDING_ID", "NVARCHAR(250)"),
        ("CONVERSATION_START", "TIMESTAMP"), ("CONVERSATION_END", "TIMESTAMP"),
        ("DURATION_MS", "BIGINT"), ("ORIGINATING_DIRECTION", "NVARCHAR(50)"),
        ("QUEUE_ID", "NVARCHAR(2100)"), ("QUEUE_NAME", "NVARCHAR(2100)"),
        ("WRAP_UP_CODE_ID", "NVARCHAR(2100)"), ("WRAP_UP_CODE_ID_ULTIMA", "NVARCHAR(250)"),
        ("CONCLUSION_ORIGINAL", "NVARCHAR(2100)"), ("CONCLUSION_ULTIMA", "NVARCHAR(500)"),
        ("CONTACT_ID", "NVARCHAR(250)"), ("CONTACT_LIST_ID", "NVARCHAR(250)"), ("CAMPAIGN_ID", "NVARCHAR(250)"),
        ("COMMUNICATION_PURPOSE", "NVARCHAR(100)"), ("TRANSCRIPT_ESTADO", "NVARCHAR(100)"),
        ("TRANSCRIPT_ERROR", "NVARCHAR(2100)"), ("MEDIA_TYPE", "NVARCHAR(100)"),
        ("LANGUAGE", "NVARCHAR(100)"), ("PROGRAM_ID", "NVARCHAR(250)"), ("ENGINE_ID", "NVARCHAR(250)"),
        ("TRANSCRIPT_START_TIME", "TIMESTAMP"), ("TRANSCRIPT_DURATION_MS", "BIGINT"),
        ("SUBJECT", "NVARCHAR(2100)"), ("MESSAGE_TYPE", "NVARCHAR(100)"),
        ("PHRASES_COUNT", "INTEGER"), ("TEXT", "NCLOB"), ("ACOUSTIC_COUNT", "INTEGER"),
        ("ACOUSTIC_JSON", "NCLOB"), ("FECHA_CARGA", "TIMESTAMP")],
}
CHILD_SPECS = {
    "PARTICIPANTES": ("PARTICIPANT_SEQ", {
        "PARTICIPANT_PURPOSE": "NVARCHAR(100)", "PARTICIPANT_NAME": "NVARCHAR(500)",
        **{n: "NVARCHAR(250)" for n in ("USER_ID", "TEAM_ID", "QUEUE_ID", "FLOW_ID", "DIVISION_ID")},
        "FLOW_VERSION": "NVARCHAR(100)", "INITIAL_DIRECTION": "NVARCHAR(100)", "MESSAGE_TYPE": "NVARCHAR(100)",
        "ANI": "NVARCHAR(2100)", "DNIS": "NVARCHAR(2100)", "ADDRESS_TO": "NVARCHAR(2100)",
        "ADDRESS_FROM": "NVARCHAR(2100)", "START_TIME_MS": "BIGINT", "END_TIME_MS": "BIGINT"}),
    "FRASES": ("PHRASE_INDEX", {
        "PARTICIPANT_PURPOSE": "NVARCHAR(100)", "TEXT": "NCLOB", "DECORATED_TEXT": "NCLOB",
        "STABILITY": "DECIMAL(18,8)", "CONFIDENCE": "DECIMAL(18,8)",
        "OFFSET_MS": "BIGINT", "START_TIME_MS": "BIGINT", "DURATION_MS": "BIGINT", "TYPE": "NVARCHAR(100)",
        "WORDS_COUNT": "INTEGER", "WORDS_JSON": "NCLOB", "ALTERNATIVES_COUNT": "INTEGER", "ALTERNATIVES_JSON": "NCLOB"}),
    "SENTIMIENTO": ("SENTIMENT_SEQ", {
        "PARTICIPANT": "NVARCHAR(500)", "PHRASE": "NCLOB", "PHRASE_INDEX": "INTEGER",
        "OFFSET_MS": "BIGINT", "START_TIME_MS": "BIGINT", "DURATION_MS": "BIGINT",
        "SENTIMENT": "NVARCHAR(100)", "TYPE": "NVARCHAR(100)", "SENTIMENT_SCORE": "DECIMAL(18,8)"}),
    "TOPICS": ("TOPIC_SEQ", {
        "PARTICIPANT": "NVARCHAR(500)", "TOPIC_ID": "NVARCHAR(250)", "TOPIC_NAME": "NVARCHAR(500)",
        "TOPIC_PHRASE": "NCLOB", "TRANSCRIPT_PHRASE": "NCLOB", "CONFIDENCE": "DECIMAL(18,8)",
        "OFFSET_MS": "BIGINT", "START_TIME_MS": "BIGINT", "DURATION_MS": "BIGINT", "TYPE": "NVARCHAR(100)"}),
}
for suffix, (seq, fields) in CHILD_SPECS.items():
    TABLES[MAIN + "_" + suffix] = [(n, "NVARCHAR(250)") for n in KEYS] + [(seq, "INTEGER")] + list(fields.items()) + [("FECHA_CARGA", "TIMESTAMP")]

# Modelo interno: conserva metadatos técnicos para agrupar y validar transcripts.
MULTI_COLUMNS = {"COMMUNICATION_ID", "TRANSCRIPT_INSTANCE_ID", "TRANSCRIPT_ID", "RECORDING_ID",
                "COMMUNICATION_PURPOSE", "MEDIA_TYPE", "LANGUAGE", "PROGRAM_ID", "ENGINE_ID",
                "SUBJECT", "MESSAGE_TYPE"}
TABLES[MAIN] = [(name, "NCLOB" if name in MULTI_COLUMNS or name == "TRANSCRIPT_ERROR" else kind)
                for name, kind in TABLES[MAIN]]

# Modelo público V3: exactamente estas 18 columnas en HANA y Excel/CSV.
MAIN_COLUMNS = (
    "CONVERSATION_ID", "CONVERSATION_START", "CONVERSATION_END", "DURATION_MS",
    "ORIGINATING_DIRECTION", "QUEUE_ID", "QUEUE_NAME", "WRAP_UP_CODE_ID",
    "WRAP_UP_CODE_ID_ULTIMA", "CONCLUSION_ORIGINAL", "CONCLUSION_ULTIMA",
    "CONTACT_ID", "CONTACT_LIST_ID", "CAMPAIGN_ID", "MEDIA_TYPE", "PHRASES_COUNT",
    "TEXT", "FECHA_CARGA",
)
STORAGE_TABLES = dict(TABLES)
STORAGE_TABLES[MAIN] = [
    (name, "NVARCHAR(250)" if name in ("ORIGINATING_DIRECTION", "MEDIA_TYPE") else dict(TABLES[MAIN])[name])
    for name in MAIN_COLUMNS
]


def slash_values(values):
    """Valores únicos, sin JSON, manteniendo el orden de aparición."""
    result = []
    def visit(value):
        if isinstance(value, (list, tuple)):
            for item in value:
                visit(item)
        elif value is not None:
            for item in str(value).split("/"):
                item = item.strip()
                if item and item not in result:
                    result.append(item)
    visit(values)
    return "/".join(result) or None


def canonical(data):
    return json.dumps(data, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False)


def array(value):
    if value is None:
        return []
    return value if isinstance(value, list) else [value]


def pick(data, *names):
    for name in names:
        if name in data and data[name] is not None:
            return data[name]
    return None


def milliseconds_value(value):
    """Desenvolver solo unidades explícitas; no adivinar segundos por magnitud."""
    if isinstance(value, dict) and "milliseconds" in value:
        return value["milliseconds"]
    return value


def local_time(value, zone="America/Tegucigalpa"):
    value = milliseconds_value(value)
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        dt = value
    elif isinstance(value, (int, float, Decimal)) or isinstance(value, str) and re.fullmatch(r"[+-]?\d+(?:\.\d+)?", value.strip()):
        ms = numeric_value(value, "BIGINT")
        dt = datetime(1970, 1, 1, tzinfo=timezone.utc) + timedelta(milliseconds=ms)
    else:
        if not isinstance(value, str):
            raise ValueError("Timestamp Genesys requiere ISO-8601 o milisegundos de época")
        dt = parse_genesys_datetime(value)
    if dt is None:
        raise ValueError("Timestamp Genesys inválido")
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(ZoneInfo(zone)).replace(tzinfo=None)


def extract_urls(payload, communication_id):
    """Solo recorre la respuesta de URLs; nunca se usa para extraer texto."""
    found = []
    def visit(item, comm, recording=None):
        if isinstance(item, str):
            if item.startswith("https://"):
                found.append({"url": item, "communicationId": comm, "recordingId": recording})
        elif isinstance(item, list):
            for child in item:
                visit(child, comm, recording)
        elif isinstance(item, dict):
            comm = item.get("communicationId") or comm
            recording = item.get("recordingId") or (item.get("recording") or {}).get("id") or recording
            url = pick(item, "url", "transcriptUrl", "downloadUrl")
            if url:
                visit(url, comm, recording)
            for key in ("urls", "transcriptUrls", "entities"):
                if key in item:
                    visit(item[key], comm, recording)
    visit(payload, communication_id)
    unique = {}
    for item in found:
        unique[(item["communicationId"], item["recordingId"], item["url"])] = item
    return list(unique.values())


def transcript_urls(config, token, cid, comm, logger):
    base = f"{config.genesys_api_url}/api/v2/speechandtextanalytics/conversations/{quote(cid, safe='')}/communications/{quote(comm, safe='')}"
    for endpoint in ("transcripturls", "transcripturl"):
        try:
            response = request_with_retry("GET", base + "/" + endpoint, config, logger,
                                          headers=genesys_headers(token), no_retry_statuses={400, 401, 403, 404},
                                          quiet_statuses={404})
            urls = extract_urls(response.json(), comm)
            if urls:
                return urls
        except requests.HTTPError as exc:
            if exc.response is None or exc.response.status_code != 404:
                raise
        # Singular solo se intenta cuando plural devuelve 404 o no tiene URLs.
    return []


def download_payload(config, url):
    if urlparse(url).scheme != "https":
        raise ValueError("La URL del transcript debe usar HTTPS")
    # Sin Authorization de Genesys y sin registrar URLs firmadas en logs/errores.
    for attempt in range(config.max_retries):
        response = None
        try:
            response = requests.get(url, timeout=config.request_timeout)
            if response.status_code == 200:
                return response.json()
            if response.status_code != 429 and response.status_code < 500:
                raise ValueError(f"Descarga transcript HTTP {response.status_code}")
        except (requests.ConnectionError, requests.Timeout):
            pass
        if attempt + 1 == config.max_retries:
            raise RuntimeError("Descarga transcript agotó reintentos")
        wait_for = response.headers.get("Retry-After", "") if response is not None else ""
        time.sleep(float(wait_for) if wait_for.replace(".", "", 1).isdigit() else min(60, 5 * (attempt + 1)))
    raise RuntimeError("Sin intentos disponibles para descarga")


def transcript_units(payload):
    """Desenvuelve transcripts sin recorrer words/alternatives/analytics como frases."""
    if isinstance(payload, list):
        for item in payload:
            yield from transcript_units(item)
    elif isinstance(payload, dict):
        if "transcripts" in payload:
            outer = {k: v for k, v in payload.items() if k != "transcripts"}
            for item in array(payload["transcripts"]):
                if not isinstance(item, dict):
                    raise ValueError("Elemento transcripts inválido")
                # Las estructuras propias tienen prioridad, incluso si están vacías.
                yield {**outer, **item}
        elif "transcript" in payload and isinstance(payload["transcript"], dict):
            yield {**{k: v for k, v in payload.items() if k != "transcript"}, **payload["transcript"]}
        elif any(k in payload for k in ("phrases", "participants", "analytics", "transcriptId")):
            yield payload
        elif payload:
            raise ValueError("Estructura de transcript no reconocida")
    elif payload is not None:
        raise ValueError("Payload transcript no es objeto/lista")


def duration_milliseconds(item):
    """Genesys puede entregar durationMs o duration: {milliseconds: N}.

    No interpretar unidades ambiguas ni reemplazar estructuras desconocidas por cero.
    Los formatos desconocidos quedan para la validación con contexto de conversación.
    """
    value = pick(item, "durationMs", "duration")
    return milliseconds_value(value)


def field_row(item, fields):
    aliases = {
        "PARTICIPANT_PURPOSE": ("participantPurpose", "purpose"), "PARTICIPANT_NAME": ("participantName", "name"),
        "OFFSET_MS": ("offsetMs", "offset"), "DURATION_MS": ("durationMs", "duration"),
        "SENTIMENT_SCORE": ("sentimentScore", "score"),
    }
    result = {}
    for column in fields:
        parts = column.lower().split("_")
        camel = parts[0] + "".join(p.title() for p in parts[1:])
        result[column] = duration_milliseconds(item) if column == "DURATION_MS" else pick(item, *aliases.get(column, (camel,)))
        if column in ("OFFSET_MS", "START_TIME_MS", "END_TIME_MS"):
            result[column] = milliseconds_value(result[column])
    return result


def build_bundle(base, transcript, descriptor, purpose):
    transcript = dict(transcript)
    recording = transcript.get("recordingId") or descriptor.get("recordingId")
    tid = transcript.get("transcriptId")
    instance = tid or recording or hashlib.sha256(canonical(transcript).encode("utf-8")).hexdigest()
    comm = transcript.get("communicationId") or descriptor["communicationId"]
    parent = {n: None for n, _ in TABLES[MAIN]}
    parent.update(base)
    parent.update(COMMUNICATION_ID=comm, TRANSCRIPT_INSTANCE_ID=str(instance), TRANSCRIPT_ID=tid,
                  RECORDING_ID=recording, COMMUNICATION_PURPOSE=purpose, TRANSCRIPT_ESTADO="OK",
                  TRANSCRIPT_ERROR=None, MEDIA_TYPE=transcript.get("mediaType"), LANGUAGE=transcript.get("language"),
                  PROGRAM_ID=transcript.get("programId"), ENGINE_ID=transcript.get("engineId"),
                  TRANSCRIPT_START_TIME=local_time(pick(transcript, "startTime", "startTimeMs")),
                  TRANSCRIPT_DURATION_MS=duration_milliseconds(transcript),
                  SUBJECT=transcript.get("subject"), MESSAGE_TYPE=transcript.get("messageType"))
    key = {k: parent[k] for k in KEYS}
    bundle = {table: [] for table in TABLES}
    bundle[MAIN] = [parent]
    phrases = array(transcript.get("phrases"))
    if any(not isinstance(p, dict) for p in phrases):
        raise ValueError("phrases contiene elementos inválidos")
    indexed = []
    used = set()
    for original, phrase in enumerate(phrases):
        try:
            index = numeric_value(phrase.get("phraseIndex"), "INTEGER")
            if index is not None and (index < 0 or index in used):
                raise ValueError("phraseIndex negativo o duplicado")
            if index is not None:
                used.add(index)
            start = numeric_value(milliseconds_value(phrase.get("startTimeMs")), "BIGINT")
        except ValueError as exc:
            raise ValueError(f"Frase {original}: {exc}") from exc
        indexed.append((original, phrase, index, start))
    phrases = sorted(indexed, key=lambda item: (item[2] if item[2] is not None else item[3] if item[3] is not None else item[0], item[0]))
    seen_indices, lines = set(), []
    fields = CHILD_SPECS["FRASES"][1]
    for original, phrase, index, _ in phrases:
        if index is None:
            index = original
            while index in used or index in seen_indices:
                index += 1
        if index in seen_indices:
            raise ValueError("phraseIndex duplicado; no se sobrescribirán frases")
        seen_indices.add(index)
        words, alternatives = array(phrase.get("words")), array(phrase.get("alternatives"))
        row = {**key, **field_row(phrase, fields), "PHRASE_INDEX": index,
               "WORDS_COUNT": len(words), "WORDS_JSON": json.dumps(words, ensure_ascii=False),
               "ALTERNATIVES_COUNT": len(alternatives), "ALTERNATIVES_JSON": json.dumps(alternatives, ensure_ascii=False),
               "FECHA_CARGA": parent["FECHA_CARGA"]}
        if row["TEXT"] is not None and str(row["TEXT"]).strip():
            lines.append(f'{row["PARTICIPANT_PURPOSE"]}: {row["TEXT"]}' if row["PARTICIPANT_PURPOSE"] else str(row["TEXT"]))
        bundle[MAIN + "_FRASES"].append(row)
    parent["TEXT"] = "\n".join(lines) or None
    parent["PHRASES_COUNT"] = len(phrases)
    analytics = transcript.get("analytics") or {}
    if isinstance(analytics, list):
        merged = {}
        for entry in analytics:
            if isinstance(entry, dict):
                for field in ("sentiment", "topics", "acoustic"):
                    merged.setdefault(field, []).extend(array(entry.get(field)))
        analytics = merged
    acoustic = array(analytics.get("acoustic"))
    parent.update(ACOUSTIC_COUNT=len(acoustic), ACOUSTIC_JSON=json.dumps(acoustic, ensure_ascii=False))
    for suffix, source in (("PARTICIPANTES", transcript.get("participants")),
                           ("SENTIMIENTO", analytics.get("sentiment")), ("TOPICS", analytics.get("topics"))):
        seq, fields = CHILD_SPECS[suffix]
        for number, item in enumerate(array(source), 1):
            if not isinstance(item, dict):
                raise ValueError(f"Elemento {suffix} inválido")
            row = {**key, **field_row(item, fields), seq: number, "FECHA_CARGA": parent["FECHA_CARGA"]}
            if suffix == "SENTIMIENTO" and isinstance(item.get("sentiment"), (int, float)):
                row["SENTIMENT_SCORE"] = item.get("sentimentScore", item.get("score", item["sentiment"]))
            bundle[MAIN + "_" + suffix].append(row)
    return bundle


def state_bundle(base, comm, state, detail, purpose=None, recording=None):
    # Marcadores técnicos reservados; no representan IDs reales de Genesys.
    bundle = {table: [] for table in TABLES}
    row = {name: None for name, _ in TABLES[MAIN]}
    row.update(base)
    marker = "__STATUS__:" + hashlib.sha256((comm + ":" + (recording or "")).encode()).hexdigest()
    row.update(COMMUNICATION_ID=comm or "__NO_COMMUNICATION__", TRANSCRIPT_INSTANCE_ID=marker,
               RECORDING_ID=recording, COMMUNICATION_PURPOSE=purpose, TRANSCRIPT_ESTADO=state,
               TRANSCRIPT_ERROR=detail, PHRASES_COUNT=0, ACOUSTIC_COUNT=0)
    bundle[MAIN] = [row]
    return bundle


def safe_error(exc):
    text = re.sub(r"https?://\S+", "[URL omitida]", str(exc))
    for key in ("HPR_PASSWORD", "GENESYS_CLIENT_SECRET"):
        if os.getenv(key):
            text = text.replace(os.environ[key], "***")
    return f"{type(exc).__name__}: {text}"[:2000]


def numeric_value(value, kind):
    """Sin equivalencias inventadas para booleanos, etiquetas o porcentajes."""
    if value is None or isinstance(value, str) and not value.strip():
        return None
    if isinstance(value, (bool, dict, list)):
        raise ValueError("valor no numérico")
    try:
        number = Decimal(str(value).strip())
        if not number.is_finite():
            raise ValueError("valor no finito")
        if kind.startswith("DECIMAL"):
            if abs(number) >= Decimal("10000000000"):
                raise ValueError("decimal fuera de rango")
            with localcontext() as context:
                context.prec = 32
                result = number.quantize(Decimal("0.00000001"))
            if abs(result) >= Decimal("10000000000"):
                raise ValueError("decimal fuera de rango después del redondeo")
            return result
        bound = 2 ** (31 if kind == "INTEGER" else 63)
        if number != number.to_integral_value() or not -bound <= number < bound:
            raise ValueError("entero inválido o fuera de rango")
        return int(number)
    except InvalidOperation as exc:
        raise ValueError("formato numérico inválido") from exc


def normalize_optional_scores(bundle, original, config, logger):
    """Conservar fuente antes de sustituir indicadores opcionales incompatibles."""
    parent = bundle[MAIN][0]
    warnings, changes = [], []
    for table, fields in TABLES.items():
        for row in bundle[table]:
            for column, kind in fields:
                if not kind.startswith("DECIMAL"):
                    continue
                value = row.get(column)
                try:
                    numeric_value(value, kind)
                except ValueError as exc:
                    warnings.append({"table": table, "column": column,
                                     "key": {k: row.get(k) for k in primary_keys(table)},
                                     "value": value, "reason": str(exc)})
                    changes.append((row, column))
    if not warnings:
        return
    directory = Path(config.json_output_dir or env_str("OUTPUT_DIR", "") or Path(__file__).parent / "output") / "normalization_warnings"
    directory.mkdir(parents=True, exist_ok=True)
    identity = [parent[k] for k in KEYS]
    # Hash de fuente: nuevas variantes no sobrescriben diagnósticos anteriores.
    filename = hashlib.sha256(canonical([identity, original]).encode("utf-8")).hexdigest() + ".json"
    path = directory / filename
    temporary = path.with_suffix(".part.json")
    temporary.write_text(json.dumps({"keys": dict(zip(KEYS, identity)), "warnings": warnings,
                                     "transcript": original}, ensure_ascii=False, indent=2), encoding="utf-8")
    os.replace(temporary, path)
    for row, column in changes:
        row[column] = None
    columns = sorted({w["column"] for w in warnings})
    parent["TRANSCRIPT_ERROR"] = f"Advertencia: {len(warnings)} indicadores opcionales no numéricos/fuera de rango guardados como NULL ({', '.join(columns)}). Diagnóstico: {filename}"
    logger.warning("Normalización numérica | conversationId=%s | communicationId=%s | transcript=%s | campos=%s | incidencias=%s | fuente íntegra=%s",
                   *identity, ",".join(columns), len(warnings), path)


def process_conversation_transcripts(conv, config, token, wrapups, queues, logger):
    cid = str(conv.get("conversationId") or conv.get("id") or "")
    if not cid:
        raise ValueError("Conversación sin ID")
    start, end = local_time(conv.get("conversationStart"), config.timezone_name), local_time(conv.get("conversationEnd"), config.timezone_name)
    raw_start = parse_genesys_datetime(conv.get("conversationStart"))
    raw_end = parse_genesys_datetime(conv.get("conversationEnd"))
    duration = round((raw_end - raw_start).total_seconds() * 1000) if raw_start and raw_end else None
    if duration is not None and duration < 0:
        raise ValueError(f"Duración negativa: {cid}")
    dialer = extract_dialer_attributes(conv)
    wrap_ids = extract_wrapup_codes(conv)
    names = wrapup_names_from_ids(wrap_ids, wrapups)
    queue_ids, queue_names = extract_queue_info(conv, queues)
    base = dict(CONVERSATION_ID=cid, CONVERSATION_START=start, CONVERSATION_END=end, DURATION_MS=duration,
                ORIGINATING_DIRECTION=slash_values(conv.get("originatingDirection")), QUEUE_ID=queue_ids, QUEUE_NAME=queue_names,
                WRAP_UP_CODE_ID=";".join(wrap_ids), WRAP_UP_CODE_ID_ULTIMA=wrap_ids[-1] if wrap_ids else None,
                CONCLUSION_ORIGINAL=names, CONCLUSION_ULTIMA=wrapups.get(wrap_ids[-1]) if wrap_ids else None,
                CONTACT_ID=dialer.get("ContactId"), CONTACT_LIST_ID=dialer.get("ContactListId"), CAMPAIGN_ID=dialer.get("CampaignId"),
                FECHA_CARGA=datetime.now(ZoneInfo("America/Tegucigalpa")).replace(tzinfo=None))
    campaign = {}
    if config.output_mode == "transcript_campania":
        try:
            list_id = dialer.get("ContactListId") or (split_filter_values(config.contact_list_id) or [""])[0]
            campaign = flatten_contact_data(fetch_contact_from_genesys(config, token, list_id, dialer.get("ContactId"), logger))
        except Exception as exc:
            logger.warning("Enriquecimiento campaña | %s | %s", cid, safe_error(exc))
    candidates = obtener_session_ids_para_transcript(conv, config.media_type)
    if not candidates:
        return [state_bundle(base, "", "SIN_CANDIDATOS", "Sin sesiones candidatas")], campaign
    bundles, seen = [], {}
    for candidate in candidates:
        comm, purpose = candidate["communication_id"], candidate["purpose"]
        try:
            descriptors = transcript_urls(config, token, cid, comm, logger)
        except Exception as exc:
            bundles.append(state_bundle(base, comm, "ERROR", safe_error(exc), purpose))
            continue
        if not descriptors:
            bundles.append(state_bundle(base, comm, "SIN_TRANSCRIPCION_API", "Sin URLs disponibles", purpose))
        for descriptor in descriptors:
            payload = None
            try:
                payload = download_payload(config, descriptor["url"])
                units = list(transcript_units(payload))
                if not units:
                    bundles.append(state_bundle(base, comm, "SIN_TRANSCRIPCION_API", "Payload sin transcripts", purpose, descriptor.get("recordingId")))
                for unit in units:
                    bundle = build_bundle(base, unit, descriptor, purpose)
                    row = bundle[MAIN][0]
                    row["MEDIA_TYPE"] = slash_values(row["MEDIA_TYPE"] or candidate.get("media_type"))
                    # Una grabación puede incluir varios transcripts sin transcriptId.
                    # Desambiguar determinísticamente sin mezclar textos ni usar la URL.
                    if not unit.get("transcriptId") and row["RECORDING_ID"] and len(units) > 1:
                        row["TRANSCRIPT_INSTANCE_ID"] = hashlib.sha256(
                            canonical([row["RECORDING_ID"], unit]).encode("utf-8")).hexdigest()
                        for table in TABLES:
                            for child in bundle[table]:
                                child["TRANSCRIPT_INSTANCE_ID"] = row["TRANSCRIPT_INSTANCE_ID"]
                    key = tuple(row[k] for k in KEYS)
                    fingerprint = hashlib.sha256(canonical(unit).encode()).hexdigest()
                    if key in seen:
                        if seen[key] != fingerprint:
                            raise ValueError("Dos transcripts distintos comparten clave técnica; no se sobrescribirán")
                        continue
                    normalize_optional_scores(bundle, unit, config, logger)
                    # Aislar problemas de formato por transcript antes de entregarlos
                    # al coordinador. Longitudes se validan después con capacidad HANA real.
                    validate_bundle(bundle, check_lengths=False)
                    seen[key] = fingerprint
                    if config.save_transcript_json or config.json_output_dir:
                        directory = Path(config.json_output_dir or Path(__file__).parent / "output" / "json")
                        directory.mkdir(parents=True, exist_ok=True)
                        filename = hashlib.sha256(canonical(key).encode()).hexdigest() + ".json"
                        (directory / filename).write_text(json.dumps(unit, ensure_ascii=False, indent=2), encoding="utf-8")
                    bundles.append(bundle)
            except Exception as exc:
                logger.warning("Transcript | conversationId=%s | communicationId=%s | %s", cid, comm, safe_error(exc))
                if payload is not None:
                    directory = Path(config.json_output_dir or env_str("OUTPUT_DIR", "") or Path(__file__).parent / "output") / "transcript_errors"
                    directory.mkdir(parents=True, exist_ok=True)
                    filename = hashlib.sha256(canonical([cid, comm, payload]).encode("utf-8")).hexdigest() + ".json"
                    path = directory / filename
                    temporary = path.with_suffix(".part.json")
                    temporary.write_text(json.dumps({"conversationId": cid, "communicationId": comm,
                        "error": safe_error(exc), "payload": payload}, ensure_ascii=False, indent=2), encoding="utf-8")
                    os.replace(temporary, path)
                    logger.warning("Fuente íntegra del transcript con error: %s", path)
                bundles.append(state_bundle(base, comm, "ERROR", safe_error(exc), purpose, descriptor.get("recordingId")))
        if config.api_sleep_seconds:
            time.sleep(config.api_sleep_seconds)
    # Un marcador sin-transcript no debe sobrevivir al éxito en la misma comunicación.
    success = {b[MAIN][0]["COMMUNICATION_ID"] for b in bundles if b[MAIN][0]["TRANSCRIPT_ESTADO"] == "OK"}
    unique = {}
    for b in bundles:
        row = b[MAIN][0]
        if row["TRANSCRIPT_ESTADO"] == "SIN_TRANSCRIPCION_API" and row["COMMUNICATION_ID"] in success:
            continue
        key = tuple(row[k] for k in KEYS)
        if key not in unique or row["TRANSCRIPT_ESTADO"] == "ERROR":
            unique[key] = b
    return list(unique.values()), campaign


def primary_keys(table):
    return ["CONVERSATION_ID"] if table == MAIN else KEYS + [CHILD_SPECS[table[len(MAIN) + 1:]][0]]


def aggregate_conversation(bundles):
    """No deduplicar por texto: una frase repetida puede ser legítima."""
    if not bundles:
        raise ValueError("Interacción sin resultados ni estado")
    ids = {b[MAIN][0]["CONVERSATION_ID"] for b in bundles}
    if len(ids) != 1:
        raise ValueError("No se pueden mezclar conversaciones")
    ordered = sorted(bundles, key=lambda b: tuple(str(b[MAIN][0].get(k) or "") for k in KEYS))
    selected, seen = [], {}
    for bundle in ordered:
        row = bundle[MAIN][0]
        identity = row.get("TRANSCRIPT_ID") or (row["COMMUNICATION_ID"], row["TRANSCRIPT_INSTANCE_ID"])
        if row["TRANSCRIPT_ESTADO"] == "OK":
            content = {table: [{k: str(v) if isinstance(v, (Decimal, datetime)) else v for k, v in child.items()
                               if k not in KEYS and k != "FECHA_CARGA"} for child in bundle[table]]
                       for table in TABLES if table != MAIN}
            content["acoustic"] = json.loads(row.get("ACOUSTIC_JSON") or "[]")
            fingerprint = canonical(content)
            if identity in seen:
                if seen[identity] != fingerprint:
                    raise ValueError("Un transcriptId tiene contenidos distintos; no se descartará ninguno silenciosamente")
                continue
            seen[identity] = fingerprint
        selected.append(bundle)
    success = [b for b in selected if b[MAIN][0]["TRANSCRIPT_ESTADO"] == "OK"]
    parent = dict(ordered[0][MAIN][0])
    result = {table: [] for table in TABLES}
    for column in MULTI_COLUMNS:
        values = []
        for b in ordered:
            value = b[MAIN][0].get(column)
            if value is not None and not str(value).startswith(("__STATUS__:", "__NO_COMMUNICATION__")) and value not in values:
                values.append(value)
        parent[column] = json.dumps(values, ensure_ascii=False)
    for column in ("ORIGINATING_DIRECTION", "MEDIA_TYPE"):
        parent[column] = slash_values([b[MAIN][0].get(column) for b in ordered])
    errors = [b[MAIN][0].get("TRANSCRIPT_ERROR") for b in selected if b[MAIN][0].get("TRANSCRIPT_ERROR")]
    has_error = any(b[MAIN][0]["TRANSCRIPT_ESTADO"] == "ERROR" for b in selected)
    parent["TRANSCRIPT_ESTADO"] = "PARCIAL" if success and has_error else "ERROR" if has_error else "OK" if success else ordered[0][MAIN][0]["TRANSCRIPT_ESTADO"]
    parent["TRANSCRIPT_ERROR"] = "\n".join(dict.fromkeys(errors)) or None
    acoustic, timeline = [], []
    for b in success:
        row = b[MAIN][0]
        acoustic.extend(json.loads(row.get("ACOUSTIC_JSON") or "[]"))
        for table in TABLES:
            if table != MAIN:
                result[table].extend(b[table])
        for phrase in b[MAIN + "_FRASES"]:
            start = phrase.get("START_TIME_MS")
            if start is None and row.get("TRANSCRIPT_START_TIME") is not None:
                anchor = row["TRANSCRIPT_START_TIME"].replace(tzinfo=ZoneInfo("America/Tegucigalpa"))
                start = round(anchor.timestamp() * 1000) + (phrase.get("OFFSET_MS") or 0)
            timeline.append((start is None, start or 0, row["COMMUNICATION_ID"], row["TRANSCRIPT_INSTANCE_ID"], phrase["PHRASE_INDEX"], phrase))
    lines = []
    for *_, phrase in sorted(timeline, key=lambda item: item[:-1]):
        if phrase.get("TEXT") is not None and str(phrase["TEXT"]).strip():
            lines.append(f'{phrase["PARTICIPANT_PURPOSE"]}: {phrase["TEXT"]}' if phrase.get("PARTICIPANT_PURPOSE") else str(phrase["TEXT"]))
    parent["TEXT"] = "\n".join(lines) or None
    parent["PHRASES_COUNT"] = len(result[MAIN + "_FRASES"])
    parent["ACOUSTIC_JSON"] = json.dumps(acoustic, ensure_ascii=False)
    parent["ACOUSTIC_COUNT"] = len(acoustic)
    starts = [b[MAIN][0]["TRANSCRIPT_START_TIME"] for b in success if b[MAIN][0].get("TRANSCRIPT_START_TIME") is not None]
    parent["TRANSCRIPT_START_TIME"] = min(starts) if starts else None
    # No sumar duraciones potencialmente solapadas de grabaciones diferentes.
    parent["TRANSCRIPT_DURATION_MS"] = success[0][MAIN][0].get("TRANSCRIPT_DURATION_MS") if len(success) == 1 else None
    result[MAIN] = [parent]
    return result


def validate_bundle(bundle, capacities=None, check_lengths=True, aggregate=False):
    parent = bundle[MAIN][0]
    if any(not parent.get(k) for k in KEYS):
        raise ValueError("Clave técnica vacía")
    if parent.get("DURATION_MS") is not None and parent["DURATION_MS"] < 0:
        raise ValueError("DURATION_MS negativo")
    phrases = bundle[MAIN + "_FRASES"]
    if parent["PHRASES_COUNT"] != len(phrases):
        raise ValueError("PHRASES_COUNT inconsistente")
    if parent["ACOUSTIC_COUNT"] != len(json.loads(parent.get("ACOUSTIC_JSON") or "[]")):
        raise ValueError("ACOUSTIC_COUNT inconsistente")
    if parent.get("TEXT") and not any(p.get("TEXT") for p in phrases):
        raise ValueError("TEXT sin frases reales")
    for phrase in phrases:
        for count, data in (("WORDS_COUNT", "WORDS_JSON"), ("ALTERNATIVES_COUNT", "ALTERNATIVES_JSON")):
            if phrase[count] != len(json.loads(phrase[data])):
                raise ValueError(f"{count} inconsistente")
    for table, rows in bundle.items():
        keys = set()
        for row in rows:
            key = tuple(row.get(k) for k in primary_keys(table))
            if key in keys:
                raise ValueError(f"Clave duplicada: {table}")
            keys.add(key)
            if any(row.get(k) != parent[k] for k in (["CONVERSATION_ID"] if aggregate else KEYS)):
                raise ValueError("Registro hijo sin relación con su transcript")
            for column, kind in (STORAGE_TABLES if aggregate else TABLES)[table]:
                value = row.get(column)
                if value is None:
                    continue
                if kind.startswith("DECIMAL") or kind in ("INTEGER", "BIGINT"):
                    try:
                        row[column] = numeric_value(value, kind)
                        if column in primary_keys(table) and row[column] is None:
                            raise ValueError("clave numérica vacía")
                    except ValueError as exc:
                        raise ValueError(f"{table}.{column}: {exc}; conversationId={parent['CONVERSATION_ID']}, "
                                         f"communicationId={parent['COMMUNICATION_ID']}, transcript={parent['TRANSCRIPT_INSTANCE_ID']}") from exc
                elif kind != "TIMESTAMP":
                    row[column] = json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else str(value)
                    if check_lengths and kind.startswith("NVARCHAR"):
                        limit = (capacities or {}).get((table, column), int(re.search(r"\d+", kind).group()))
                        size = len(row[column].encode("utf-16-le")) // 2
                        if size > limit:
                            raise ValueError(f"{table}.{column}: longitud={size}, capacidad={limit}, conversationId={parent['CONVERSATION_ID']}. Sin truncar.")


def has_transcript_text(bundle):
    text = bundle[MAIN][0].get("TEXT")
    return isinstance(text, str) and bool(text.strip())


class HanaTranscriptWriter:
    """Solo el hilo coordinador usa esta conexión. Sin DDL ni pandas."""
    def __init__(self, logger):
        from hdbcli import dbapi
        self.logger, self.buffer, self.capacities = logger, [], {}
        self.written_counts = {table: 0 for table in TABLES}
        self.conn = dbapi.connect(address=env_str("HPR_HOST", required=True), port=env_int("HPR_PORT", 30015),
                                 user=env_str("HPR_USER", required=True), password=env_str("HPR_PASSWORD", required=True),
                                 connectTimeout=15000)
        try:
            self.conn.setautocommit(False)
            self.validate_schema()
        except Exception:
            self.conn.close()
            raise

    def validate_schema(self):
        cur = self.conn.cursor()
        try:
            for table, fields in STORAGE_TABLES.items():
                cur.execute("SELECT COLUMN_NAME, DATA_TYPE_NAME, LENGTH, SCALE FROM SYS.TABLE_COLUMNS WHERE SCHEMA_NAME = ? AND TABLE_NAME = ?", ("BI_SS", table))
                actual = {r[0]: r[1:] for r in cur.fetchall()}
                if set(actual) != {n for n, _ in fields}:
                    raise ValueError(f"BI_SS.{table}: faltan tablas/columnas o el modelo no coincide. Use el SQL adjunto.")
                for column, kind in fields:
                    base = kind.split("(")[0]
                    data = actual[column]
                    if data[0] != base:
                        raise ValueError(f"BI_SS.{table}.{column}: se esperaba {kind}")
                    if base == "NVARCHAR":
                        size = int(data[1])
                        if size < int(re.search(r"\d+", kind).group()):
                            raise ValueError(f"BI_SS.{table}.{column}: capacidad inferior a {kind}")
                        self.capacities[table, column] = size
                    if base == "DECIMAL" and (int(data[2]) < 8 or int(data[1]) - int(data[2]) < 10):
                        raise ValueError(f"BI_SS.{table}.{column}: precisión insuficiente")
                cur.execute("SELECT COLUMN_NAME FROM SYS.CONSTRAINTS WHERE SCHEMA_NAME = ? AND TABLE_NAME = ? AND IS_PRIMARY_KEY = 'TRUE'", ("BI_SS", table))
                if {r[0] for r in cur.fetchall()} != set(primary_keys(table)):
                    raise ValueError(f"BI_SS.{table}: clave primaria incompatible")
        finally:
            cur.close()

    def add(self, bundle):
        if not has_transcript_text(bundle):
            return
        validate_bundle(bundle, self.capacities, aggregate=True)
        self.buffer.append(bundle)
        if len(self.buffer) >= 20:
            self.flush()

    def flush(self):
        if not self.buffer:
            return
        cur = self.conn.cursor()
        counts = {}
        try:
            # Reemplazo atómico por interacción, incluyendo todas sus filas hijas.
            instances = {b[MAIN][0]["CONVERSATION_ID"]: b for b in self.buffer}
            incomplete = [cid for cid, b in instances.items() if b[MAIN][0]["TRANSCRIPT_ESTADO"] != "OK"]
            if incomplete:
                marks = ", ".join("?" for _ in incomplete)
                cur.execute(f"SELECT CONVERSATION_ID FROM BI_SS.{MAIN} WHERE CONVERSATION_ID IN ({marks})", tuple(incomplete))
                for (cid,) in cur.fetchall():
                    instances.pop(cid, None)
                    self.logger.warning("HANA | %s: se conserva la versión existente ante descarga incompleta/sin transcript", cid)
            for table in reversed(list(TABLES)):
                keys = list(instances)
                for offset in range(0, len(keys), 100):
                    batch = keys[offset:offset + 100]
                    marks = ", ".join("?" for _ in batch)
                    cur.execute(f"DELETE FROM BI_SS.{table} WHERE CONVERSATION_ID IN ({marks})", tuple(batch))
            for table, fields in STORAGE_TABLES.items():
                columns = [n for n, _ in fields]
                rows = [row for b in instances.values() for row in b[table]]
                sql = f"INSERT INTO BI_SS.{table} ({', '.join(columns)}) VALUES ({', '.join('?' for _ in columns)})"
                for offset in range(0, len(rows), 250):
                    cur.executemany(sql, [tuple(row.get(n) for n in columns) for row in rows[offset:offset + 250]])
                counts[table] = len(rows)
            self.conn.commit()
        except Exception as exc:
            self.conn.rollback()
            raise RuntimeError("Rollback de las cinco tablas; lotes anteriores confirmados. " + safe_error(exc)) from exc
        finally:
            cur.close()
        for table, count in counts.items():
            self.written_counts[table] += count
        self.buffer.clear()

    def close(self):
        try:
            self.conn.close()
        finally:
            self.logger.info("Resumen HANA | filas confirmadas por tabla: %s", self.written_counts)


class OutputWriter:
    """CSV completo como spool. Excel nunca recorta textos silenciosamente."""
    def __init__(self, config, logger):
        self.config, self.logger = config, logger
        self.path = Path(resolve_output_path(config.output_csv, config.output_format))
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self.partial = self.path.with_suffix(".partial.csv")
        self.stream = self.partial.open("w", newline="", encoding="utf-8-sig")
        self.columns = list(MAIN_COLUMNS)
        self.writer = csv.DictWriter(self.stream, fieldnames=self.columns)
        self.writer.writeheader()
        self.campaign_file = None
        self.count = 0

    def add(self, bundle):
        if not has_transcript_text(bundle):
            return
        row = bundle[MAIN][0]
        self.writer.writerow({n: row.get(n) for n in self.columns})
        self.count += 1

    def campaign(self, cid, fields):
        if not fields:
            return
        if self.campaign_file is None:
            self.campaign_file = self.path.with_suffix(".campania.jsonl").open("w", encoding="utf-8")
        self.campaign_file.write(json.dumps({"CONVERSATION_ID": cid, **fields}, ensure_ascii=False) + "\n")

    def finish(self):
        self.stream.close()
        if self.config.output_format == "csv":
            os.replace(self.partial, self.path)
        else:
            from openpyxl import Workbook
            from openpyxl.cell import WriteOnlyCell
            from openpyxl.styles import Font, PatternFill
            if self.count > 1048575:
                csv_path = self.path.with_suffix(".completo.csv")
                os.replace(self.partial, csv_path)
                raise ValueError(f"Excel supera su límite de filas. Datos completos: {csv_path}")
            wb = Workbook(write_only=True)
            ws = wb.create_sheet("Transcripciones")
            ws.freeze_panes = "A2"
            cells = []
            for column in self.columns:
                cell = WriteOnlyCell(ws, column)
                cell.fill = PatternFill("solid", fgColor="1F4E78")
                cell.font = Font(color="FFFFFF", bold=True)
                cells.append(cell)
            ws.append(cells)
            overflow = False
            csv.field_size_limit(2 ** 31 - 1)
            with self.partial.open(newline="", encoding="utf-8-sig") as source:
                for row in csv.DictReader(source):
                    cells = []
                    for column in self.columns:
                        value = row[column]
                        if len(value.encode("utf-16-le")) // 2 > 32767 or re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", value):
                            overflow = True
                            value = "Contenido completo en " + self.path.with_suffix(".completo.csv").name
                        cell = WriteOnlyCell(ws, value)
                        cell.data_type = "s"
                        cell.number_format = "@"
                        cells.append(cell)
                    ws.append(cells)
            from openpyxl.utils import get_column_letter
            ws.auto_filter.ref = f"A1:{get_column_letter(len(self.columns))}{self.count + 1}"
            temporary = self.path.with_suffix(".part.xlsx")
            wb.save(temporary)
            os.replace(temporary, self.path)
            if overflow:
                csv_path = self.path.with_suffix(".completo.csv")
                os.replace(self.partial, csv_path)
                self.logger.warning("Excel no admite algunos textos: contenido íntegro en %s y en HANA.", csv_path)
            else:
                self.partial.unlink()
        self.logger.info("Salida generada: %s | filas=%s", self.path, self.count)

    def close(self):
        self.stream.close()
        if self.campaign_file is not None:
            self.campaign_file.close()


def main():
    parser = argparse.ArgumentParser(description="Extractor de transcripciones Genesys Cloud")
    parser.add_argument("--date", default=env_str("DATE", ""), help="Fecha local específica. Ejemplo: 2026-06-01")
    parser.add_argument("--start-date", default=env_str("START_DATE", ""), help="Fecha inicial local inclusiva")
    parser.add_argument("--end-date", default=env_str("END_DATE", ""), help="Fecha final local inclusiva")
    parser.add_argument("--dry-run", action="store_true", help="No genera CSV")
    args = parser.parse_args()
    logger = setup_logger()
    started = time.monotonic()
    hana = output = None
    error_count = total_rows = skipped_without_text = 0
    try:
        config = load_config()
        config.dry_run = config.dry_run or args.dry_run
        start, end, mode = parse_dates(args, config.timezone_name)
        logger.info("Intervalo Genesys: %s/%s | %s | modo=%s", start, end, mode, config.output_mode)
        # Preflight HANA antes de descargar: credenciales, tablas, tipos y claves.
        if not config.dry_run:
            hana = HanaTranscriptWriter(logger)
            output = OutputWriter(config, logger)
        token = get_access_token(config, logger)
        config = apply_resolved_filters(config, token, logger)
        wrapups = get_wrapup_catalog(config, token, logger)
        queues = get_queue_catalog(config, token, logger)
        validate_filter_safety(config, logger)
        ids = split_filter_values(config.conversation_id)
        if ids:
            conversations = [fetch_conversation_details_by_id(config, token, cid, logger) for cid in ids]
        else:
            job = create_conversation_details_job(config, token, start, end, logger)
            wait_details_job(config, token, job, logger)
            conversations = fetch_details_job_results(config, token, job, logger)
        unique = {}
        for conv in conversations:
            if conversation_matches_post_filters(conv, config):
                cid = conv.get("conversationId") or conv.get("id")
                if not cid:
                    raise ValueError("Analytics devolvió conversación sin ID")
                unique[cid] = conv
        conversations = list(unique.values())
        if config.max_conversations:
            conversations = conversations[:config.max_conversations]
        total = len(conversations)
        logger.info("Conversaciones después de filtros y deduplicación: %s", total)
        completed = 0
        last_progress = time.monotonic()
        iterator = iter(conversations)
        # Ventana acotada: solo workers*2 resultados en vuelo; HANA no cruza hilos.
        with ThreadPoolExecutor(max_workers=config.max_transcript_workers) as executor:
            pending = {}
            def submit_one():
                conv = next(iterator, None)
                if conv is None:
                    return
                future = executor.submit(process_conversation_transcripts, conv, config, token, wrapups, queues, logger)
                pending[future] = conv
            for _ in range(config.max_transcript_workers * 2):
                submit_one()
            while pending:
                done, _ = wait(pending, return_when=FIRST_COMPLETED)
                for future in done:
                    conv = pending.pop(future)
                    # Errores inesperados no se silencian ni se convierten en éxito.
                    bundles, campaign = future.result()
                    bundle = aggregate_conversation(bundles)
                    error_count += bundle[MAIN][0]["TRANSCRIPT_ESTADO"] in ("ERROR", "PARCIAL")
                    if has_transcript_text(bundle):
                        validate_bundle(bundle, hana.capacities if hana else None, aggregate=True)
                        if hana:
                            hana.add(bundle)
                        if output:
                            output.add(bundle)
                            output.campaign(conv.get("conversationId") or conv["id"], campaign)
                        total_rows += 1
                    else:
                        skipped_without_text += 1
                    completed += 1
                    now = time.monotonic()
                    if completed == 1 or completed == total or now - last_progress >= 30:
                        logger.info("Avance %s/%s | con texto=%s | omitidas sin texto=%s | errores=%s",
                                    completed, total, total_rows, skipped_without_text, error_count)
                        print(f"PYFLOW_PROGRESS={int(90 * completed / max(1, total))}", flush=True)
                        last_progress = now
                    submit_one()
        if hana:
            hana.flush()
            output.finish()
        print("PYFLOW_PROGRESS=100", flush=True)
        logger.info("Finalizado | filas con texto=%s | omitidas sin texto=%s | errores=%s | dry_run=%s",
                    total_rows, skipped_without_text, error_count, config.dry_run)
        return 1 if error_count else 0
    except Exception as exc:
        logger.error("Proceso fallido: %s", safe_error(exc))
        return 1
    finally:
        if output:
            output.close()
        if hana:
            hana.close()
        for handler in logger.handlers:
            for log_filter in handler.filters:
                if isinstance(log_filter, CompactConsoleFilter):
                    log_filter.summarize(logger)
        logger.info("Duración total: %.2fs", time.monotonic() - started)


if __name__ == "__main__":
    raise SystemExit(main())
