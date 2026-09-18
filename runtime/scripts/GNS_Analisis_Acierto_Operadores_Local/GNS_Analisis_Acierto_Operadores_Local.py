from __future__ import annotations

import csv
import html
import json
import logging
import math
import os
import re
import sys
import time
import unicodedata
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple


PYFLOW_PARAMS = {
    "ORIGEN_DATOS": {
        "type": "select",
        "label": "Origen de datos",
        "options": ["HANA", "ARCHIVO"],
        "default": "ARCHIVO"
    }
}


CONFIG_COLUMNAS = {
    "SOURCE_FILE": ["SOURCE_FILE", "__SOURCE_FILE", "ARCHIVO_ORIGEN"],
    "CONVERSATION_ID": ["CONVERSATION_ID", "CONVERSATIONID", "ID_CONVERSACION"],
    "COMMUNICATION_ID": ["COMMUNICATION_ID", "COMMUNICATIONID"],
    "FECHA_LLAMADA": ["FECHA_LLAMADA", "FECHA_HORA", "CONVERSATIONSTART", "CONVERSATIONSTART_LOCAL", "FECHA"],
    "OPERADOR": ["OPERADOR", "AGENTNAME", "AGENTE", "NOMBRE_AGENTE", "USUARIO"],
    "USER_ID": ["USER_ID", "USERID", "AGENTID"],
    "COLA": ["COLA", "QUEUENAME", "QUEUE_NAME", "NOMBRE_COLA"],
    "QUEUE_ID": ["QUEUE_ID", "QUEUEID"],
    "WRAPUP": ["WRAPUP", "WRAPUPCODE", "CONCLUSION", "CONCLUSION_ORIGINAL", "CONCLUSION_ULTIMA"],
    "CLIENTE_ID": ["CLIENTE_ID", "EXTERNALTAG", "EXTERNAL_TAG", "DNI"],
    "TELEFONO": ["TELEFONO", "ANI", "DNIS", "REMOTE"],
    "DURACION_SEGUNDOS": ["DURACION_SEGUNDOS", "DURATION_SECONDS", "DURACION"],
    "TEMAS_GENESYS": ["TEMAS_GENESYS", "TOPICS", "TEMAS"],
    "TRANSFERENCIAS": ["TRANSFERENCIAS", "TRANSFERS"],
    "RECURRENCIA": ["RECURRENCIA", "RECURRENTE"],
    "TRANSCRIPCION_TEXTO": ["TRANSCRIPCION_TEXTO", "TRANSCRIPCION", "TRANSCRIPT", "TEXT", "TEXTO", "CONVERSATION_TEXT"],
    "TRANSCRIPCION_JSON": ["TRANSCRIPCION_JSON", "TRANSCRIPT_JSON", "JSON_TRANSCRIPT"]
}


MATRIZ_SCORE = {
    "identificacion_motivo": 20,
    "validacion": 15,
    "aplicacion_proceso": 30,
    "calidad_solucion": 15,
    "empatia": 10,
    "tipificacion": 10
}


TEMAS_NEGOCIO = {
    "ATM_RETIRO_NO_DISPENSADO": {
        "descripcion": "Retiro ATM/cajero debitado sin entrega de efectivo.",
        "palabras_cliente": [
            "cajero no me dio el dinero", "no me dispenso", "retiro en atm", "retiro en cajero",
            "cajero automatico", "me debito", "no salio el efectivo", "no me entrego el efectivo",
            "me hizo el cargo", "retiro fallido", "no me dio efectivo", "cajero no entrego"
        ],
        "frases_operador_gestion": [
            "le voy a generar", "vamos a ingresar el reclamo", "se estara validando",
            "numero de gestion", "numero de reclamo", "proceso de reclamo", "plazo de respuesta",
            "caso queda registrado", "validar la transaccion", "reclamo por atm"
        ],
        "frases_mala_gestion": [
            "eso no se atiende aqui", "debe ir a agencia", "vaya a una agencia", "no puedo ayudarle",
            "llame despues", "aqui no vemos eso", "no tengo acceso"
        ],
        "frases_redireccion": ["vaya a agencia", "debe ir a agencia", "sucursal", "oficina"],
        "acciones_esperadas": ["validar fecha", "validar monto", "registrar reclamo", "dar numero de gestion"],
        "wrapups_relacionados": ["ATM", "RETIRO", "NO_DISPENSADO", "CAJERO"],
        "motivos_raiz_posibles": ["No se evidencia gestion", "Redireccion innecesaria o riesgosa", "Informacion incompleta"]
    },
    "BANCA_DIGITAL_ACCESO": {
        "descripcion": "Problemas de acceso a app, web, contrasena, usuario, OTP o token.",
        "palabras_cliente": [
            "no puedo entrar", "no puedo ingresar", "la app no funciona", "banca en linea",
            "contrasena", "usuario bloqueado", "no me llega el codigo", "otp", "token",
            "usuario no cumple", "no puedo registrarme", "me da error", "desbloqueo"
        ],
        "frases_operador_gestion": [
            "restablecer contrasena", "desbloqueo de usuario", "validar codigo", "actualizar correo",
            "actualizar telefono", "soporte de banca digital", "generar gestion", "vamos a validar su usuario",
            "le indicare los pasos"
        ],
        "frases_mala_gestion": ["intente mas tarde", "vaya a agencia", "no se", "no tenemos sistema", "no le puedo ayudar"],
        "frases_redireccion": ["vaya a agencia", "debe ir a agencia", "presencial", "sucursal"],
        "acciones_esperadas": ["diagnosticar acceso", "validar datos", "guiar recuperacion", "escalar si aplica"],
        "wrapups_relacionados": ["BANCA", "TOKEN", "APP", "CONTRASENA", "PASSWORD", "USUARIO", "OTP"],
        "motivos_raiz_posibles": ["Falta de conocimiento del proceso", "No se evidencia gestion", "Informacion incompleta"]
    },
    "TARJETA_PERDIDA_ROBO": {
        "descripcion": "Tarjeta perdida, robada o extraviada.",
        "palabras_cliente": ["perdi mi tarjeta", "me robaron la tarjeta", "extravio", "tarjeta perdida", "tarjeta robada", "billetera", "monedero"],
        "frases_operador_gestion": ["bloquear tarjeta", "bloqueo preventivo", "por seguridad", "reposicion", "validaremos consumos", "confirmar transacciones"],
        "frases_mala_gestion": ["espere a ver si aparece", "llame despues", "no puedo bloquear", "vaya a agencia primero"],
        "frases_redireccion": ["agencia", "sucursal", "oficina"],
        "acciones_esperadas": ["bloquear tarjeta", "validar identidad", "orientar reposicion", "consultar consumos"],
        "wrapups_relacionados": ["TARJETA", "BLOQUEO", "ROBO", "EXTRAVIO", "PERDIDA"],
        "motivos_raiz_posibles": ["No se evidencia gestion", "Informacion incompleta"]
    },
    "TRANSACCION_NO_RECONOCIDA": {
        "descripcion": "Cargo, compra, retiro o transaccion que el cliente no reconoce.",
        "palabras_cliente": [
            "no reconozco", "compra no reconocida", "transaccion no reconocida", "cargo no reconocido",
            "me debitaron", "no hice esa compra", "fraude", "consumo no reconocido", "me sacaron dinero"
        ],
        "frases_operador_gestion": [
            "bloqueo preventivo", "reclamo", "disputa", "investigacion", "numero de caso",
            "validar transaccion", "fecha y monto", "comercio", "se estara investigando"
        ],
        "frases_mala_gestion": ["llame al comercio", "no se puede hacer nada", "vaya a agencia", "no puedo ayudarle"],
        "frases_redireccion": ["vaya a agencia", "sucursal", "oficina", "presencial"],
        "acciones_esperadas": ["validar transaccion", "bloquear si aplica", "registrar reclamo", "explicar tiempos"],
        "wrapups_relacionados": ["NO_RECONOC", "FRAUDE", "RECLAMO", "DISPUTA", "TRANSACCION"],
        "motivos_raiz_posibles": ["No se evidencia gestion", "Informacion incompleta", "Redireccion innecesaria o riesgosa"]
    },
    "BLOQUEO_TARJETA": {
        "descripcion": "Solicitud de bloqueo de tarjeta por seguridad u otra razon.",
        "palabras_cliente": ["bloquear tarjeta", "bloquee mi tarjeta", "cancelar tarjeta", "desactivar tarjeta", "bloqueo de tarjeta"],
        "frases_operador_gestion": ["se bloquea", "bloqueo exitoso", "por seguridad", "confirmo el bloqueo", "reposicion"],
        "frases_mala_gestion": ["no puedo bloquear", "intente despues", "vaya a agencia"],
        "frases_redireccion": ["agencia", "sucursal"],
        "acciones_esperadas": ["validar identidad", "bloquear", "confirmar estado", "orientar reposicion"],
        "wrapups_relacionados": ["BLOQUEO", "TARJETA"],
        "motivos_raiz_posibles": ["No se evidencia gestion"]
    },
    "CONSULTA_SALDO_MOVIMIENTOS": {
        "descripcion": "Consulta de saldo, movimientos o estado de cuenta.",
        "palabras_cliente": ["saldo", "movimientos", "estado de cuenta", "cuanto tengo", "transacciones", "deposito", "abono"],
        "frases_operador_gestion": ["puede consultarlo", "le indico", "saldo", "movimiento", "estado de cuenta", "banca movil"],
        "frases_mala_gestion": ["no se", "no puedo ver", "llame despues"],
        "frases_redireccion": ["agencia", "sucursal"],
        "acciones_esperadas": ["orientar canales", "validar identidad si se brinda informacion"],
        "wrapups_relacionados": ["SALDO", "MOVIMIENTO", "ESTADO_CUENTA", "CONSULTA"],
        "motivos_raiz_posibles": ["Informacion incompleta"]
    },
    "PRESTAMOS_FINANCIAMIENTOS": {
        "descripcion": "Consulta o gestion relacionada con prestamos, cuotas o financiamientos.",
        "palabras_cliente": ["prestamo", "financiamiento", "cuota", "credito", "mora", "pago de prestamo", "saldo de prestamo"],
        "frases_operador_gestion": ["le indico requisitos", "validar su prestamo", "cuota", "fecha de pago", "gestion de cobro", "solicitud"],
        "frases_mala_gestion": ["no se", "vaya a agencia", "no puedo ayudarle"],
        "frases_redireccion": ["agencia", "sucursal", "oficina"],
        "acciones_esperadas": ["identificar producto", "orientar requisitos", "validar datos", "explicar pasos"],
        "wrapups_relacionados": ["PRESTAMO", "CREDITO", "FINANCIAMIENTO"],
        "motivos_raiz_posibles": ["Informacion incompleta", "Redireccion innecesaria o riesgosa"]
    },
    "ACTUALIZACION_DATOS": {
        "descripcion": "Actualizacion de telefono, correo, direccion o datos personales.",
        "palabras_cliente": ["actualizar datos", "cambiar numero", "cambiar correo", "actualizar telefono", "actualizar direccion", "datos personales"],
        "frases_operador_gestion": ["validar sus datos", "requisitos", "actualizar", "generar gestion", "correo", "telefono"],
        "frases_mala_gestion": ["no se puede", "no se", "vaya a agencia", "llame despues"],
        "frases_redireccion": ["agencia", "sucursal", "presencial"],
        "acciones_esperadas": ["identificar dato", "validar identidad", "indicar canal correcto", "generar gestion si aplica"],
        "wrapups_relacionados": ["ACTUALIZACION", "DATOS", "CORREO", "TELEFONO"],
        "motivos_raiz_posibles": ["Informacion incompleta", "No se evidencia gestion"]
    },
    "RECLAMO_GENERAL": {
        "descripcion": "Reclamo no clasificado en otro tema.",
        "palabras_cliente": ["reclamo", "queja", "inconforme", "problema", "no me resuelven", "necesito solucion"],
        "frases_operador_gestion": ["generar reclamo", "numero de caso", "seguimiento", "validar", "gestion", "plazo"],
        "frases_mala_gestion": ["no puedo ayudarle", "llame despues", "no se", "vaya a agencia"],
        "frases_redireccion": ["agencia", "sucursal"],
        "acciones_esperadas": ["escuchar", "validar", "registrar gestion", "dar seguimiento"],
        "wrapups_relacionados": ["RECLAMO", "QUEJA", "GESTION"],
        "motivos_raiz_posibles": ["No se evidencia gestion", "Informacion incompleta"]
    },
    "CONSULTA_GENERAL": {
        "descripcion": "Consulta informativa general.",
        "palabras_cliente": ["consulta", "informacion", "requisitos", "horario", "como puedo", "quiero saber", "me puede indicar"],
        "frases_operador_gestion": ["con gusto le indico", "los requisitos son", "el horario es", "puede realizarlo", "le explico"],
        "frases_mala_gestion": ["no se", "busque en la pagina", "llame despues", "no tengo informacion"],
        "frases_redireccion": ["agencia", "sucursal", "oficina"],
        "acciones_esperadas": ["responder con claridad", "confirmar necesidad", "tipificar correctamente"],
        "wrapups_relacionados": ["CONSULTA", "INFORMACION", "GENERAL"],
        "motivos_raiz_posibles": ["Informacion incompleta"]
    }
}


NEGATIVAS_CLIENTE = [
    "molesto", "enojado", "inconforme", "queja", "reclamo", "no me resuelven",
    "siempre lo mismo", "nadie me ayuda", "pesimo servicio", "mala atencion",
    "ya llame", "tengo dias", "no me solucionan", "necesito solucion", "me urge",
    "voy a cancelar", "voy a reclamar", "no estoy de acuerdo"
]
POSITIVAS = ["gracias", "perfecto", "excelente", "muy amable", "me ayudo", "resuelto", "esta bien", "de acuerdo", "muchas gracias"]
OPERADOR_POCO_EMPATICO = ["no puedo ayudarle", "ese no es mi problema", "tiene que esperar", "no se", "no tengo informacion", "llame despues", "vaya a agencia", "no se puede"]
OPERADOR_EMPATICO = ["con gusto", "le ayudo", "comprendo", "entiendo", "lamento lo ocurrido", "permitame validar", "vamos a revisar", "con mucho gusto", "gracias por esperar"]
FRASES_VALIDACION = ["validar", "confirmar", "numero de identidad", "identidad", "dni", "fecha", "monto", "cuenta", "tarjeta", "correo", "telefono", "nombre completo"]
FRASES_CALIDAD = ["plazo", "tiempo de respuesta", "numero de gestion", "numero de caso", "requisitos", "pasos a seguir", "notificacion", "seguimiento"]
FRASES_REDIRECCION = ["vaya a agencia", "debe ir a agencia", "tiene que presentarse", "en sucursal", "oficina", "agencia mas cercana", "presencial"]


@dataclass
class OptionalLibs:
    rapidfuzz_ratio: Any = None
    sentiment_pipeline: Any = None


@dataclass
class Config:
    fecha_inicio: str
    fecha_fin: str
    origen_datos: str
    tabla_hana: str
    ruta_archivo_entrada: str
    output_dir: Path
    limite_filas: int
    batch_size: int
    reanalizar_existentes: bool
    incluir_transcripcion_completa: bool
    usar_rapidfuzz: bool
    usar_sentimiento_transformers: bool
    modelo_sentimiento_local: str
    chunksize_hana: int


def pyflow_progress(value: int) -> None:
    print(f"PYFLOW_PROGRESS={max(0, min(100, int(value)))}", flush=True)


def get_param(name: str, default: Any = None) -> Any:
    value = os.environ.get(name)
    if value is None or str(value).strip() == "":
        return default
    return value


def env_bool(name: str, default: bool = False) -> bool:
    value = str(get_param(name, "SI" if default else "NO")).strip().upper()
    return value in {"SI", "S", "YES", "Y", "TRUE", "1"}


def env_int(name: str, default: int) -> int:
    try:
        return int(float(str(get_param(name, default)).strip()))
    except Exception:
        return default


def setup_logger() -> logging.Logger:
    logger = logging.getLogger("acierto_local")
    logger.handlers.clear()
    logger.setLevel(logging.INFO)
    handler = logging.StreamHandler(sys.stdout)
    handler.setFormatter(logging.Formatter("%(asctime)s | %(levelname)s | %(message)s", "%Y-%m-%d %H:%M:%S"))
    logger.addHandler(handler)
    return logger


def resolve_path(raw: str, default_name: str) -> Path:
    script_dir = Path(__file__).resolve().parent
    value = str(raw or default_name).strip()
    path = Path(value)
    if not path.is_absolute():
        path = script_dir / path
    path.mkdir(parents=True, exist_ok=True)
    return path


def load_config() -> Config:
    return Config(
        fecha_inicio=str(get_param("FECHA_INICIO", "") or "").strip(),
        fecha_fin=str(get_param("FECHA_FIN", "") or "").strip(),
        origen_datos=str(get_param("ORIGEN_DATOS", "ARCHIVO")).strip().upper(),
        tabla_hana=str(get_param("TABLA_HANA", "ESQUEMA.GNS_TRANSCRIPCIONES_GENESYS")).strip(),
        ruta_archivo_entrada=str(get_param("RUTA_ARCHIVO_ENTRADA", "") or "").strip(),
        output_dir=resolve_path(str(get_param("OUTPUT_DIR", "./output")), "output"),
        limite_filas=env_int("LIMITE_FILAS", 0),
        batch_size=max(1, env_int("BATCH_SIZE", 5000)),
        reanalizar_existentes=env_bool("REANALIZAR_EXISTENTES", False),
        incluir_transcripcion_completa=env_bool("INCLUIR_TRANSCRIPCION_COMPLETA", False),
        usar_rapidfuzz=env_bool("USAR_RAPIDFUZZ", True),
        usar_sentimiento_transformers=env_bool("USAR_SENTIMIENTO_LOCAL_TRANSFORMERS", False),
        modelo_sentimiento_local=str(get_param("MODELO_SENTIMIENTO_LOCAL", "") or "").strip(),
        chunksize_hana=max(1000, env_int("CHUNKSIZE_HANA", 50000))
    )


def normalizar_nombre_columna(col: Any) -> str:
    text = normalizar_texto(str(col or ""))
    text = re.sub(r"[^a-z0-9]+", "_", text)
    return text.strip("_").upper()


def normalizar_texto(texto: Any) -> str:
    if texto is None:
        return ""
    if isinstance(texto, float) and math.isnan(texto):
        return ""
    text = str(texto).replace("\r", "\n")
    text = "".join(ch for ch in text if unicodedata.category(ch)[0] != "C" or ch in "\n\t")
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = text.lower()
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def separar_hablantes(transcripcion: str) -> Tuple[str, str, str, List[str]]:
    text = str(transcripcion or "")
    flags: List[str] = []
    cliente_parts: List[str] = []
    operador_parts: List[str] = []
    pattern = re.compile(
        r"(?im)^\s*(cliente|customer|usuario|caller|operador|agente|agent|representative)\s*:\s*(.*)$"
    )
    for match in pattern.finditer(text):
        speaker = normalizar_texto(match.group(1))
        chunk = match.group(2).strip()
        if speaker in {"cliente", "customer", "usuario", "caller"}:
            cliente_parts.append(chunk)
        else:
            operador_parts.append(chunk)
    general = normalizar_texto(text)
    cliente = normalizar_texto(" ".join(cliente_parts))
    operador = normalizar_texto(" ".join(operador_parts))
    if not cliente or not operador:
        flags.append("HABLANTES_NO_IDENTIFICADOS")
        cliente = cliente or general
        operador = operador or general
    return cliente, operador, general, flags


def cargar_librerias_opcionales(config: Config, logger: logging.Logger) -> OptionalLibs:
    libs = OptionalLibs()
    if config.usar_rapidfuzz:
        try:
            from rapidfuzz import fuzz
            libs.rapidfuzz_ratio = fuzz.partial_ratio
            logger.info("RapidFuzz disponible para coincidencia aproximada.")
        except Exception:
            logger.info("RapidFuzz no instalado. Se usaran coincidencias exactas.")
    if config.usar_sentimiento_transformers:
        try:
            from transformers import pipeline
            if not config.modelo_sentimiento_local:
                logger.warning("USAR_SENTIMIENTO_LOCAL_TRANSFORMERS=SI, pero MODELO_SENTIMIENTO_LOCAL esta vacio. Se usan reglas.")
            else:
                libs.sentiment_pipeline = pipeline(
                    "sentiment-analysis",
                    model=config.modelo_sentimiento_local,
                    local_files_only=True
                )
                logger.info("Modelo local de sentimiento cargado: %s", config.modelo_sentimiento_local)
        except Exception as exc:
            logger.warning("No se pudo cargar Transformers local. Se usan reglas. Detalle: %s", exc)
    return libs


def conectar_hana():
    host = str(get_param("HPR_HOST_ESPEJO", get_param("HPR_HOST", get_param("HANA_HOST", ""))) or "").strip()
    port = int(get_param("HPR_PORT", get_param("HANA_PORT", 30015)) or 30015)
    user = str(get_param("HPR_USER", get_param("HANA_USER", "")) or "").strip()
    password = str(get_param("HPR_PASSWORD", get_param("HANA_PASSWORD", "")) or "").strip()
    missing = [name for name, value in [("HPR_HOST_ESPEJO", host), ("HPR_USER", user), ("HPR_PASSWORD", password)] if not value]
    if missing:
        raise ValueError("Faltan variables globales para HANA: " + ", ".join(missing))
    from hdbcli import dbapi
    return dbapi.connect(address=host, port=port, user=user, password=password)


def parse_date(value: str, name: str) -> Optional[datetime]:
    value = str(value or "").strip()
    if not value:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(value, fmt)
        except ValueError:
            pass
    raise ValueError(f"Fecha invalida en {name}: {value}")


def validate_table_name(table_name: str) -> str:
    clean = str(table_name or "").strip().replace('"', "")
    if not re.fullmatch(r"[A-Za-z0-9_]+(\.[A-Za-z0-9_]+)?", clean):
        raise ValueError(f"Nombre de tabla HANA no permitido: {table_name}")
    return ".".join(f'"{part}"' for part in clean.split("."))


def default_date_range_for_query(config: Config) -> Tuple[str, str]:
    start = parse_date(config.fecha_inicio, "FECHA_INICIO")
    end = parse_date(config.fecha_fin, "FECHA_FIN")
    if not start:
        start = datetime(1900, 1, 1)
    if not end:
        end = datetime.now() + timedelta(days=1)
    else:
        end = end + timedelta(days=1)
    return start.strftime("%Y-%m-%d"), end.strftime("%Y-%m-%d")


def cargar_query_hana(config: Config) -> str:
    query_path = Path(__file__).resolve().parent / "config" / "query_llamadas.sql"
    if query_path.exists():
        query = query_path.read_text(encoding="utf-8-sig")
    else:
        table = validate_table_name(config.tabla_hana)
        query = f"""
            SELECT
                CONVERSATION_ID,
                COMMUNICATION_ID,
                FECHA_LLAMADA,
                OPERADOR,
                USER_ID,
                COLA,
                QUEUE_ID,
                WRAPUP,
                CLIENTE_ID,
                TELEFONO,
                DURACION_SEGUNDOS,
                TEMAS_GENESYS,
                TRANSFERENCIAS,
                RECURRENCIA,
                TRANSCRIPCION_TEXTO,
                TRANSCRIPCION_JSON
            FROM {table}
            WHERE TRANSCRIPCION_TEXTO IS NOT NULL
        """
    start_text, end_text = default_date_range_for_query(config)
    query = query.replace("{{FECHA_INICIO}}", start_text)
    query = query.replace("{{FECHA_FIN}}", end_text)
    if config.limite_filas > 0 and " limit " not in f" {query.lower()} ":
        query = query.rstrip().rstrip(";") + f"\nLIMIT {int(config.limite_filas)}"
    return query


def leer_llamadas_hana(config: Config, logger: logging.Logger):
    import pandas as pd

    query = cargar_query_hana(config)
    logger.info("Leyendo HANA usando config/query_llamadas.sql")
    conn = conectar_hana()
    try:
        return pd.read_sql(query, conn)
    finally:
        conn.close()


def detectar_tipo_archivo(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xls"}:
        return "excel"
    if suffix == ".csv":
        return "csv"
    raise ValueError(f"Tipo de archivo no soportado: {path}")


def sniff_csv_separator(path: Path) -> str:
    sample = path.read_text(encoding="utf-8-sig", errors="ignore")[:4096]
    candidates = [",", ";", "\t"]
    counts = {sep: sample.count(sep) for sep in candidates}
    return max(counts, key=counts.get) if max(counts.values()) > 0 else ","


def read_one_file(path: Path):
    import pandas as pd

    kind = detectar_tipo_archivo(path)
    if kind == "excel":
        df = pd.read_excel(path, dtype=str)
    else:
        sep = sniff_csv_separator(path)
        try:
            df = pd.read_csv(path, sep=sep, dtype=str, low_memory=False, encoding="utf-8-sig")
        except UnicodeDecodeError:
            df = pd.read_csv(path, sep=sep, dtype=str, low_memory=False, encoding="latin-1")
    df["__SOURCE_FILE"] = path.name
    return df


def leer_llamadas_archivo(config: Config, logger: logging.Logger):
    import pandas as pd

    if config.ruta_archivo_entrada:
        path = Path(config.ruta_archivo_entrada)
        if not path.is_absolute():
            path = Path(__file__).resolve().parent / path
        paths = [path]
    else:
        input_dir = Path(__file__).resolve().parent / "input"
        input_dir.mkdir(parents=True, exist_ok=True)
        paths = sorted([p for p in input_dir.iterdir() if p.suffix.lower() in {".csv", ".xlsx", ".xls"}])
    if not paths:
        raise FileNotFoundError("No se encontraron archivos de entrada. Use RUTA_ARCHIVO_ENTRADA o coloque archivos en la carpeta input.")
    frames = []
    for idx, path in enumerate(paths, start=1):
        if not path.exists():
            raise FileNotFoundError(f"No existe archivo: {path}")
        logger.info("Leyendo archivo %s/%s: %s", idx, len(paths), path)
        frames.append(read_one_file(path))
    return pd.concat(frames, ignore_index=True)


def mapear_columnas(df):
    import pandas as pd

    rename = {col: normalizar_nombre_columna(col) for col in df.columns}
    df = df.rename(columns=rename)
    lower_to_real = {col.upper(): col for col in df.columns}
    result = pd.DataFrame()
    for canonical, aliases in CONFIG_COLUMNAS.items():
        source = None
        for alias in aliases:
            key = normalizar_nombre_columna(alias)
            if key in lower_to_real:
                source = lower_to_real[key]
                break
        result[canonical] = df[source] if source else ""
    if result["TRANSCRIPCION_TEXTO"].fillna("").astype(str).str.strip().eq("").all():
        raise ValueError("No se encontro la columna TRANSCRIPCION_TEXTO o equivalente con texto util.")
    return result


def preparar_catalogo_temas(config: Config, libs: OptionalLibs) -> Dict[str, Any]:
    catalog = {}
    for tema, data in TEMAS_NEGOCIO.items():
        normalized = {}
        for key, value in data.items():
            if isinstance(value, list):
                normalized[key] = [normalizar_texto(item) for item in value]
            else:
                normalized[key] = value
        catalog[tema] = normalized
    return catalog


def phrase_hits(text: str, phrases: Sequence[str]) -> List[str]:
    return [phrase for phrase in phrases if phrase and phrase in text]


def fuzzy_hits(text: str, phrases: Sequence[str], ratio_func: Any, threshold: int = 88) -> List[str]:
    if not ratio_func or not text:
        return []
    hits = []
    for phrase in phrases:
        if len(phrase) < 6:
            continue
        try:
            if ratio_func(phrase, text) >= threshold:
                hits.append(phrase)
        except Exception:
            pass
    return hits


def detectar_tema_probable(texto_cliente: str, texto_general: str, temas_genesys: str, catalogo: Dict[str, Any], libs: OptionalLibs):
    scores = []
    genesys = normalizar_texto(temas_genesys)
    for tema, data in catalogo.items():
        cliente_hits = phrase_hits(texto_cliente, data.get("palabras_cliente", []))
        general_hits = phrase_hits(texto_general, data.get("palabras_cliente", []))
        fuzzy = fuzzy_hits(texto_cliente, data.get("palabras_cliente", []), libs.rapidfuzz_ratio)
        wrap_hits = [w for w in data.get("wrapups_relacionados", []) if normalizar_texto(w) in genesys or normalizar_texto(w) in texto_general]
        score = len(cliente_hits) * 5 + len(general_hits) * 2 + len(fuzzy) * 3 + len(wrap_hits) * 2
        evidences = list(dict.fromkeys(cliente_hits + fuzzy + wrap_hits))[:6]
        if score:
            scores.append((tema, score, evidences))
    if not scores:
        return "CONSULTA_GENERAL", 0, "", ""
    scores.sort(key=lambda item: item[1], reverse=True)
    principal = scores[0]
    secundarios = [item[0] for item in scores[1:4]]
    return principal[0], principal[1], " | ".join(principal[2]), " | ".join(secundarios)


def extraer_evidencia(text: str, phrases: Sequence[str], limit: int = 250) -> str:
    if not text:
        return ""
    for phrase in phrases:
        phrase = normalizar_texto(phrase)
        if phrase and phrase in text:
            idx = max(0, text.find(phrase) - 60)
            return text[idx: idx + limit].strip()
    return text[:limit].strip()


def calcular_sentimiento_local_reglas(texto_cliente: str, texto_operador: str, texto_general: str):
    neg_hits = phrase_hits(texto_cliente + " " + texto_general, [normalizar_texto(x) for x in NEGATIVAS_CLIENTE])
    pos_hits = phrase_hits(texto_cliente + " " + texto_general, [normalizar_texto(x) for x in POSITIVAS])
    low_empathy_hits = phrase_hits(texto_operador, [normalizar_texto(x) for x in OPERADOR_POCO_EMPATICO])
    empathy_hits = phrase_hits(texto_operador, [normalizar_texto(x) for x in OPERADOR_EMPATICO])
    score = (len(pos_hits) * 18 + len(empathy_hits) * 8) - (len(neg_hits) * 22 + len(low_empathy_hits) * 18)
    score = max(-100, min(100, score))
    if neg_hits and pos_hits:
        sentiment = "MIXTO"
    elif score <= -25:
        sentiment = "NEGATIVO"
    elif score >= 25:
        sentiment = "POSITIVO"
    else:
        sentiment = "NEUTRO"
    motive = []
    if neg_hits:
        motive.append("Senales negativas cliente: " + ", ".join(neg_hits[:4]))
    if low_empathy_hits:
        motive.append("Baja empatia operador: " + ", ".join(low_empathy_hits[:4]))
    if pos_hits or empathy_hits:
        motive.append("Senales positivas/empatia: " + ", ".join((pos_hits + empathy_hits)[:4]))
    evidence = extraer_evidencia(texto_cliente + " " + texto_operador, neg_hits + low_empathy_hits + pos_hits + empathy_hits)
    return {
        "sentimiento_local": sentiment,
        "score_sentimiento_local": score,
        "motivo_sentimiento_local": " | ".join(motive)[:500],
        "baja_empatia": "SI" if low_empathy_hits and not empathy_hits else "NO",
        "evidencia_sentimiento": evidence
    }


def calcular_sentimiento_local_transformers_opcional(texto_general: str, libs: OptionalLibs, fallback: Dict[str, Any]) -> Dict[str, Any]:
    if not libs.sentiment_pipeline:
        return fallback
    try:
        text = texto_general[:2000]
        result = libs.sentiment_pipeline(text)[0]
        label = normalizar_texto(result.get("label", ""))
        score = float(result.get("score", 0))
        if "neg" in label:
            fallback["sentimiento_local"] = "NEGATIVO"
            fallback["score_sentimiento_local"] = int(-100 * score)
        elif "pos" in label:
            fallback["sentimiento_local"] = "POSITIVO"
            fallback["score_sentimiento_local"] = int(100 * score)
        fallback["motivo_sentimiento_local"] = (fallback.get("motivo_sentimiento_local", "") + f" | Modelo local: {label} {score:.2f}").strip(" |")
    except Exception as exc:
        fallback["motivo_sentimiento_local"] = (fallback.get("motivo_sentimiento_local", "") + f" | Error modelo local: {exc}").strip(" |")
    return fallback


def detectar_redireccion(texto_operador: str) -> Tuple[bool, List[str]]:
    hits = phrase_hits(texto_operador, [normalizar_texto(x) for x in FRASES_REDIRECCION])
    return bool(hits), hits


def detectar_validacion(texto_operador: str) -> Tuple[int, List[str]]:
    hits = phrase_hits(texto_operador, [normalizar_texto(x) for x in FRASES_VALIDACION])
    score = min(15, len(set(hits)) * 3)
    return score, hits


def detectar_gestion(texto_operador: str, tema_data: Dict[str, Any], redireccion: bool) -> Tuple[int, bool, List[str], List[str]]:
    good = phrase_hits(texto_operador, tema_data.get("frases_operador_gestion", []))
    bad = phrase_hits(texto_operador, tema_data.get("frases_mala_gestion", []))
    if good:
        score = min(30, 12 + len(set(good)) * 6)
    elif redireccion:
        score = 10
    else:
        score = 0
    if bad:
        score = min(score, 8)
    return score, bool(good), good, bad


def detectar_calidad_solucion(texto_operador: str) -> Tuple[int, List[str]]:
    hits = phrase_hits(texto_operador, [normalizar_texto(x) for x in FRASES_CALIDAD])
    score = min(15, len(set(hits)) * 4)
    return score, hits


def detectar_tipificacion(wrapup: str, tema_detectado: str, tema_data: Dict[str, Any]) -> Tuple[int, bool]:
    wrap = normalizar_texto(wrapup)
    if not wrap:
        return 0, True
    related = [normalizar_texto(x) for x in tema_data.get("wrapups_relacionados", [])]
    if any(item and item in wrap for item in related):
        return 10, False
    tema_parts = [part for part in normalizar_texto(tema_detectado).split("_") if len(part) > 3]
    if any(part in wrap for part in tema_parts):
        return 5, False
    return 0, True


def calcular_score_acierto(score_tema: int, score_validacion: int, score_gestion: int, score_calidad: int, score_empatia: int, score_tipificacion: int) -> Tuple[int, int]:
    if score_tema >= 12:
        score_identificacion = 20
    elif score_tema >= 6:
        score_identificacion = 12
    elif score_tema > 0:
        score_identificacion = 6
    else:
        score_identificacion = 0
    total = score_identificacion + score_validacion + score_gestion + score_calidad + score_empatia + score_tipificacion
    return score_identificacion, max(0, min(100, total))


def clasificar_score(score: int) -> str:
    if score >= 90:
        return "Acierto alto"
    if score >= 75:
        return "Acierto"
    if score >= 60:
        return "Acierto parcial"
    if score > 0:
        return "No acierto"
    return "No evaluable"


def determinar_motivo_raiz(flags: Sequence[str], score_gestion: int, score_calidad: int, score_tipificacion: int, baja_empatia: bool) -> str:
    if "TRANSCRIPCION_INSUFICIENTE" in flags:
        return "Transcripcion insuficiente"
    if "REDIRECCION_RIESGOSA" in flags or "REDIRECCION_INCORRECTA" in flags:
        return "Redireccion innecesaria o riesgosa"
    if score_gestion == 0 or "NO_GESTIONO" in flags:
        return "No se evidencia gestion"
    if score_tipificacion == 0 or "MALA_TIPIFICACION" in flags:
        return "Mala tipificacion"
    if score_calidad < 7 or "INFORMACION_INCOMPLETA" in flags:
        return "Informacion incompleta"
    if baja_empatia or "BAJA_EMPATIA" in flags:
        return "Baja empatia"
    return "No aplica"


def is_truthy_text(value: Any) -> bool:
    text = normalizar_texto(value)
    return text in {"si", "s", "true", "1", "yes", "recurrente"} or "recurrent" in text


def evaluar_llamada(row: Dict[str, Any], catalogo: Dict[str, Any], libs: OptionalLibs, config: Config) -> Dict[str, Any]:
    try:
        transcript_raw = str(row.get("TRANSCRIPCION_TEXTO", "") or "")
        texto_cliente, texto_operador, texto_general, flags = separar_hablantes(transcript_raw)
        if len(texto_general) < 100:
            flags.append("TRANSCRIPCION_INSUFICIENTE")
            sentiment = calcular_sentimiento_local_reglas(texto_cliente, texto_operador, texto_general)
            return build_result(row, "CONSULTA_GENERAL", 0, "", "", sentiment, flags, 0, 0, 0, 0, 0, 0, "Transcripcion insuficiente", "", "", config, "OK", "")

        tema, score_tema, evidencias_tema, temas_sec = detectar_tema_probable(
            texto_cliente, texto_general, str(row.get("TEMAS_GENESYS", "") or ""), catalogo, libs
        )
        tema_data = catalogo.get(tema, catalogo["CONSULTA_GENERAL"])
        sentiment = calcular_sentimiento_local_reglas(texto_cliente, texto_operador, texto_general)
        sentiment = calcular_sentimiento_local_transformers_opcional(texto_general, libs, sentiment)
        redireccion, redir_hits = detectar_redireccion(texto_operador)
        score_validacion, valid_hits = detectar_validacion(texto_operador)
        score_gestion, gestiono, gestion_hits, mala_hits = detectar_gestion(texto_operador, tema_data, redireccion)
        score_calidad, calidad_hits = detectar_calidad_solucion(texto_operador)
        score_tipificacion, mala_tipificacion = detectar_tipificacion(str(row.get("WRAPUP", "") or ""), tema, tema_data)
        score_empatia = 10
        if sentiment["baja_empatia"] == "SI":
            score_empatia = 2
        elif phrase_hits(texto_operador, [normalizar_texto(x) for x in OPERADOR_EMPATICO]):
            score_empatia = 10
        else:
            score_empatia = 6

        redireccion_riesgosa = (
            redireccion
            and not gestiono
            and (sentiment["sentimiento_local"] == "NEGATIVO" or is_truthy_text(row.get("RECURRENCIA", "")))
            and tema not in {"CONSULTA_GENERAL", "CONSULTA_SALDO_MOVIMIENTOS"}
        )
        if redireccion_riesgosa:
            score_gestion = min(score_gestion, 5)

        score_identificacion, total = calcular_score_acierto(score_tema, score_validacion, score_gestion, score_calidad, score_empatia, score_tipificacion)
        if not gestiono and not redireccion:
            flags.append("NO_GESTIONO")
        if redireccion_riesgosa:
            flags.append("REDIRECCION_RIESGOSA")
        if score_calidad < 7:
            flags.append("INFORMACION_INCOMPLETA")
        if mala_tipificacion:
            flags.append("MALA_TIPIFICACION")
        if sentiment["baja_empatia"] == "SI":
            flags.append("BAJA_EMPATIA")
        if is_truthy_text(row.get("RECURRENCIA", "")):
            flags.append("CLIENTE_RECURRENTE")
        if sentiment["sentimiento_local"] == "NEGATIVO":
            flags.append("SENTIMIENTO_NEGATIVO")
        if total < 60 or redireccion_riesgosa or "INFORMACION_INCOMPLETA" in flags or "MALA_TIPIFICACION" in flags or "BAJA_EMPATIA" in flags:
            flags.append("REQUIERE_REVISION_SUPERVISOR")
        motivo_raiz = determinar_motivo_raiz(flags, score_gestion, score_calidad, score_tipificacion, sentiment["baja_empatia"] == "SI")
        evidencia_cliente = extraer_evidencia(texto_cliente, evidencias_tema.split(" | ") + NEGATIVAS_CLIENTE)
        evidencia_operador = extraer_evidencia(texto_operador, gestion_hits + mala_hits + redir_hits + valid_hits + calidad_hits)
        return build_result(
            row, tema, score_tema, evidencias_tema, temas_sec, sentiment, flags,
            score_identificacion, score_validacion, score_gestion, score_calidad, score_empatia,
            score_tipificacion, motivo_raiz, evidencia_cliente, evidencia_operador, config, "OK", ""
        )
    except Exception as exc:
        flags = ["REQUIERE_REVISION_SUPERVISOR"]
        sentiment = {"sentimiento_local": "NEUTRO", "score_sentimiento_local": 0, "motivo_sentimiento_local": "", "evidencia_sentimiento": "", "baja_empatia": "NO"}
        return build_result(row, "CONSULTA_GENERAL", 0, "", "", sentiment, flags, 0, 0, 0, 0, 0, 0, "No aplica", "", "", config, "ERROR", str(exc))


def build_result(row, tema, score_tema, evidencias_tema, temas_sec, sentiment, flags, score_identificacion, score_validacion, score_gestion, score_calidad, score_empatia, score_tipificacion, motivo_raiz, evidencia_cliente, evidencia_operador, config, estado, error):
    flags = list(dict.fromkeys([flag for flag in flags if flag]))
    total = score_identificacion + score_validacion + score_gestion + score_calidad + score_empatia + score_tipificacion
    resultado = clasificar_score(total)
    redireccion_agencia = any(flag in flags for flag in ["REDIRECCION_RIESGOSA", "REDIRECCION_INCORRECTA"])
    recomendacion = "Validar con supervisor" if "REQUIERE_REVISION_SUPERVISOR" in flags else "Sin accion critica"
    output = {
        "conversation_id": row.get("CONVERSATION_ID", ""),
        "communication_id": row.get("COMMUNICATION_ID", ""),
        "fecha_llamada": row.get("FECHA_LLAMADA", ""),
        "operador": row.get("OPERADOR", ""),
        "user_id": row.get("USER_ID", ""),
        "cola": row.get("COLA", ""),
        "queue_id": row.get("QUEUE_ID", ""),
        "cliente_id": row.get("CLIENTE_ID", ""),
        "telefono": row.get("TELEFONO", ""),
        "wrapup_genesys": row.get("WRAPUP", ""),
        "tema_detectado": tema,
        "score_tema": score_tema,
        "temas_secundarios": temas_sec,
        "motivo_probable": evidencias_tema,
        "sentimiento_local": sentiment.get("sentimiento_local", "NEUTRO"),
        "score_sentimiento_local": sentiment.get("score_sentimiento_local", 0),
        "motivo_sentimiento_local": sentiment.get("motivo_sentimiento_local", ""),
        "operador_gestiono": "SI" if score_gestion >= 18 else "NO",
        "redireccion_agencia": "SI" if redireccion_agencia else "NO",
        "redireccion_riesgosa": "SI" if "REDIRECCION_RIESGOSA" in flags else "NO",
        "mala_tipificacion": "SI" if "MALA_TIPIFICACION" in flags else "NO",
        "informacion_incompleta": "SI" if "INFORMACION_INCOMPLETA" in flags else "NO",
        "baja_empatia": sentiment.get("baja_empatia", "NO"),
        "cliente_recurrente": "SI" if is_truthy_text(row.get("RECURRENCIA", "")) else "NO",
        "score_identificacion_motivo": score_identificacion,
        "score_validacion": score_validacion,
        "score_aplicacion_proceso": score_gestion,
        "score_calidad_solucion": score_calidad,
        "score_empatia": score_empatia,
        "score_tipificacion": score_tipificacion,
        "score_total": total,
        "resultado_acierto": resultado,
        "motivo_raiz_error": motivo_raiz,
        "banderas": " | ".join(flags),
        "evidencia_cliente": evidencia_cliente[:250],
        "evidencia_operador": evidencia_operador[:250],
        "evidencia_sentimiento": str(sentiment.get("evidencia_sentimiento", ""))[:250],
        "recomendacion": recomendacion,
        "requiere_revision_supervisor": "SI" if "REQUIERE_REVISION_SUPERVISOR" in flags else "NO",
        "duracion_segundos": row.get("DURACION_SEGUNDOS", ""),
        "temas_genesys": row.get("TEMAS_GENESYS", ""),
        "transferencias": row.get("TRANSFERENCIAS", ""),
        "recurrencia": row.get("RECURRENCIA", ""),
        "estado_proceso": estado,
        "error_proceso": error
    }
    if config.incluir_transcripcion_completa:
        output["transcripcion_texto"] = row.get("TRANSCRIPCION_TEXTO", "")
    return output


def checkpoint_paths(output_dir: Path) -> Tuple[Path, Path]:
    return output_dir / "checkpoint_acierto_operadores.csv", output_dir / "checkpoint_acierto_operadores.jsonl"


def replace_file_safely(temp_path: Path, final_path: Path, logger: logging.Logger) -> bool:
    try:
        os.replace(str(temp_path), str(final_path))
        return True
    except OSError as exc:
        logger.warning("No se pudo reemplazar checkpoint %s. Detalle: %s", final_path.name, exc)
        return False
    finally:
        try:
            if temp_path.exists():
                temp_path.unlink()
        except OSError:
            pass


def guardar_checkpoint(rows: List[Dict[str, Any]], output_dir: Path, logger: logging.Logger) -> None:
    if not rows:
        return
    try:
        output_dir.mkdir(parents=True, exist_ok=True)
        csv_path, jsonl_path = checkpoint_paths(output_dir)
        stamp = f"{os.getpid()}_{int(time.time() * 1000)}"
        tmp_csv_path = output_dir / f".checkpoint_acierto_operadores_{stamp}.csv.tmp"
        tmp_jsonl_path = output_dir / f".checkpoint_acierto_operadores_{stamp}.jsonl.tmp"
        fieldnames = sorted({key for row in rows for key in row})

        with tmp_csv_path.open("w", newline="", encoding="utf-8-sig") as fh:
            writer = csv.DictWriter(fh, fieldnames=fieldnames, extrasaction="ignore")
            writer.writeheader()
            writer.writerows(rows)
        csv_ok = replace_file_safely(tmp_csv_path, csv_path, logger)
    except Exception as exc:
        logger.warning("No se pudo guardar checkpoint CSV; se continuara sin detener el proceso. Detalle: %s", exc)
        csv_ok = False
        tmp_jsonl_path = output_dir / f".checkpoint_acierto_operadores_{os.getpid()}_{int(time.time() * 1000)}.jsonl.tmp"
        jsonl_path = output_dir / "checkpoint_acierto_operadores.jsonl"

    try:
        with tmp_jsonl_path.open("w", encoding="utf-8") as fh:
            for row in rows:
                fh.write(json.dumps(row, ensure_ascii=False, default=str) + "\n")
        replace_file_safely(tmp_jsonl_path, jsonl_path, logger)
    except Exception as exc:
        logger.warning(
            "No se pudo guardar checkpoint JSONL; se continuara con el avance en memoria. Detalle: %s",
            exc,
        )
    logger.info("Checkpoint %s: %s registros", "guardado" if csv_ok else "omitido temporalmente", len(rows))


def cargar_checkpoint(output_dir: Path, logger: logging.Logger) -> List[Dict[str, Any]]:
    csv_path, jsonl_path = checkpoint_paths(output_dir)
    if not csv_path.exists() and not jsonl_path.exists():
        return []
    rows = []
    def cargar_csv() -> List[Dict[str, Any]]:
        import pandas as pd
        return pd.read_csv(csv_path, dtype=str, low_memory=False).fillna("").to_dict("records")

    try:
        usar_jsonl = jsonl_path.exists() and (
            not csv_path.exists() or jsonl_path.stat().st_mtime >= csv_path.stat().st_mtime
        )
        if usar_jsonl:
            try:
                with jsonl_path.open("r", encoding="utf-8") as fh:
                    for line in fh:
                        line = line.strip()
                        if line:
                            rows.append(json.loads(line))
            except Exception as exc:
                logger.warning("No se pudo cargar checkpoint JSONL; se intentara con CSV. Detalle: %s", exc)
                rows = cargar_csv() if csv_path.exists() else []
        else:
            rows = cargar_csv()
        logger.info("Checkpoint cargado: %s registros", len(rows))
        return rows
    except Exception as exc:
        logger.warning("No se pudo cargar checkpoint: %s", exc)
        return []


def pct(part: float, total: float) -> float:
    return round((float(part) / float(total)) * 100, 1) if total else 0.0


def mode_value(series) -> str:
    values = [str(v) for v in series.dropna().tolist() if str(v).strip()]
    return Counter(values).most_common(1)[0][0] if values else ""


def generar_resumenes(df):
    import pandas as pd

    total = len(df)
    evaluadas = int((df["resultado_acierto"] != "No evaluable").sum()) if total else 0
    acierto_alto = int((df["resultado_acierto"] == "Acierto alto").sum()) if total else 0
    acierto = int((df["resultado_acierto"] == "Acierto").sum()) if total else 0
    no_acierto = int((df["resultado_acierto"] == "No acierto").sum()) if total else 0
    score_promedio = round(pd.to_numeric(df["score_total"], errors="coerce").fillna(0).mean(), 2) if total else 0
    resumen = pd.DataFrame([
        {"Indicador": "Total llamadas", "Valor": total},
        {"Indicador": "Llamadas evaluadas", "Valor": evaluadas},
        {"Indicador": "Llamadas no evaluables", "Valor": int((df["resultado_acierto"] == "No evaluable").sum()) if total else 0},
        {"Indicador": "Acierto alto", "Valor": acierto_alto},
        {"Indicador": "Acierto", "Valor": acierto},
        {"Indicador": "Acierto parcial", "Valor": int((df["resultado_acierto"] == "Acierto parcial").sum()) if total else 0},
        {"Indicador": "No acierto", "Valor": no_acierto},
        {"Indicador": "Porcentaje acierto", "Valor": pct(acierto_alto + acierto, evaluadas)},
        {"Indicador": "Porcentaje no acierto", "Valor": pct(no_acierto, evaluadas)},
        {"Indicador": "Score promedio", "Valor": score_promedio},
        {"Indicador": "Sentimiento negativo", "Valor": int((df["sentimiento_local"] == "NEGATIVO").sum()) if total else 0},
        {"Indicador": "Sentimiento neutro", "Valor": int((df["sentimiento_local"] == "NEUTRO").sum()) if total else 0},
        {"Indicador": "Sentimiento positivo", "Valor": int((df["sentimiento_local"] == "POSITIVO").sum()) if total else 0},
        {"Indicador": "Llamadas revision supervisor", "Valor": int((df["requiere_revision_supervisor"] == "SI").sum()) if total else 0}
    ])
    work = df.copy()
    work["score_total"] = pd.to_numeric(work["score_total"], errors="coerce").fillna(0)
    ranking = work.groupby("operador", dropna=False).agg(
        llamadas_evaluadas=("conversation_id", "count"),
        score_promedio=("score_total", "mean"),
        no_aciertos=("resultado_acierto", lambda s: int((s == "No acierto").sum())),
        sentimiento_negativo=("sentimiento_local", lambda s: int((s == "NEGATIVO").sum())),
        principal_motivo_error=("motivo_raiz_error", mode_value),
        principal_tema=("tema_detectado", mode_value),
        llamadas_revision_supervisor=("requiere_revision_supervisor", lambda s: int((s == "SI").sum()))
    ).reset_index()
    ranking["porcentaje_acierto"] = ranking.apply(lambda r: pct(max(0, r["llamadas_evaluadas"] - r["no_aciertos"]), r["llamadas_evaluadas"]), axis=1)
    ranking["score_promedio"] = ranking["score_promedio"].round(2)
    ranking = ranking.sort_values(["score_promedio", "llamadas_evaluadas"], ascending=[True, False])
    motivos = work.groupby("motivo_raiz_error", dropna=False).agg(cantidad=("conversation_id", "count"), score_promedio=("score_total", "mean")).reset_index().sort_values("cantidad", ascending=False)
    motivos["porcentaje"] = motivos["cantidad"].apply(lambda x: pct(x, total))
    motivos["score_promedio"] = motivos["score_promedio"].round(2)
    temas = work.groupby("tema_detectado", dropna=False).agg(
        cantidad=("conversation_id", "count"),
        score_promedio=("score_total", "mean"),
        no_aciertos=("resultado_acierto", lambda s: int((s == "No acierto").sum())),
        sentimiento_negativo=("sentimiento_local", lambda s: int((s == "NEGATIVO").sum()))
    ).reset_index().sort_values("cantidad", ascending=False)
    temas["porcentaje_no_acierto"] = temas.apply(lambda r: pct(r["no_aciertos"], r["cantidad"]), axis=1)
    temas["score_promedio"] = temas["score_promedio"].round(2)
    sentimiento = work.groupby("sentimiento_local", dropna=False).agg(cantidad=("conversation_id", "count"), score_promedio=("score_total", "mean")).reset_index().sort_values("cantidad", ascending=False)
    sentimiento["porcentaje"] = sentimiento["cantidad"].apply(lambda x: pct(x, total))
    sentimiento["score_promedio"] = sentimiento["score_promedio"].round(2)
    revision = work[
        (work["resultado_acierto"] == "No acierto")
        | (work["score_total"] < 60)
        | (work["requiere_revision_supervisor"] == "SI")
        | work["banderas"].fillna("").str.contains("REDIRECCION_RIESGOSA|INFORMACION_INCOMPLETA|MALA_TIPIFICACION|BAJA_EMPATIA", regex=True)
        | (work["sentimiento_local"] == "NEGATIVO")
    ].copy()
    no_eval = work[work["resultado_acierto"] == "No evaluable"].copy()
    return resumen, ranking, motivos, temas, sentimiento, revision, no_eval


def autosize_excel(path: Path) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import Border, Font, PatternFill, Side

    wb = load_workbook(path)
    fill = PatternFill("solid", fgColor="DA282D")
    font = Font(color="FFFFFF", bold=True)
    border = Border(left=Side(style="thin", color="D9D9D9"), right=Side(style="thin", color="D9D9D9"), top=Side(style="thin", color="D9D9D9"), bottom=Side(style="thin", color="D9D9D9"))
    for ws in wb.worksheets:
        if ws.max_row >= 1:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = fill
            cell.font = font
            cell.border = border
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.border = border
                header = str(ws.cell(row=1, column=cell.column).value or "").lower()
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "0.0" if "porcentaje" in header else "#,##0.00" if isinstance(cell.value, float) else "#,##0"
        for col in ws.columns:
            letter = col[0].column_letter
            width = max(10, min(60, max(len(str(c.value or "")) for c in col) + 2))
            ws.column_dimensions[letter].width = width
    wb.save(path)


def generar_excel(rows: List[Dict[str, Any]], output_dir: Path, logger: logging.Logger) -> Path:
    import pandas as pd

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = output_dir / f"Analisis_Acierto_Operadores_Local_{timestamp}.xlsx"
    df = pd.DataFrame(rows)
    resumen, ranking, motivos, temas, sentimiento, revision, no_eval = generar_resumenes(df)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        resumen.to_excel(writer, sheet_name="Resumen Ejecutivo", index=False)
        df.to_excel(writer, sheet_name="Detalle llamadas", index=False)
        ranking.to_excel(writer, sheet_name="Ranking operadores", index=False)
        motivos.to_excel(writer, sheet_name="Motivos raiz", index=False)
        temas.to_excel(writer, sheet_name="Temas detectados", index=False)
        sentimiento.to_excel(writer, sheet_name="Sentimiento local", index=False)
        revision.to_excel(writer, sheet_name="Revision supervisor", index=False)
        no_eval.to_excel(writer, sheet_name="No evaluables", index=False)
    autosize_excel(path)
    logger.info("Excel generado: %s", path)
    return path


def html_table(title: str, rows: List[Tuple[str, Any]]) -> str:
    html_rows = "".join(
        f"<tr><td>{html.escape(str(a))}</td><td style='text-align:right;font-weight:700'>{html.escape(str(b))}</td></tr>"
        for a, b in rows
    )
    return f"<div class='panel'><h3>{html.escape(title)}</h3><table>{html_rows}</table></div>"


def generar_html(rows: List[Dict[str, Any]], output_dir: Path, logger: logging.Logger) -> Optional[Path]:
    try:
        import pandas as pd
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        path = output_dir / f"Analisis_Acierto_Operadores_Local_{timestamp}.html"
        df = pd.DataFrame(rows)
        resumen, ranking, motivos, temas, sentimiento, revision, _ = generar_resumenes(df)
        res = {r["Indicador"]: r["Valor"] for _, r in resumen.iterrows()}
        top_motivos = [(r["motivo_raiz_error"], f'{int(r["cantidad"]):,}') for _, r in motivos.head(10).iterrows()]
        top_temas = [(r["tema_detectado"], f'{int(r["no_aciertos"]):,}') for _, r in temas.sort_values("no_aciertos", ascending=False).head(10).iterrows()]
        top_ops = [(r["operador"] or "Sin operador", f'{r["score_promedio"]:.1f}') for _, r in ranking.head(10).iterrows()]
        criticas = [(r["conversation_id"], f'{r["resultado_acierto"]} / {r["motivo_raiz_error"]}') for _, r in revision.head(10).iterrows()]
        content = f"""<!doctype html><html><head><meta charset='utf-8'><title>Analisis de Acierto Operativo</title>
<style>
body{{font-family:Arial,sans-serif;background:#F7F7F7;color:#333;margin:0;padding:24px}}
.wrap{{max-width:1100px;margin:auto;background:white;border:1px solid #ddd}}
.head{{background:#DA282D;color:white;padding:24px}}
.cards{{display:grid;grid-template-columns:repeat(5,1fr);gap:12px;padding:18px}}
.card{{background:#fafafa;border:1px solid #e5e5e5;padding:14px}}
.card b{{display:block;font-size:24px;color:#DA282D;margin-top:6px}}
.grid{{display:grid;grid-template-columns:1fr 1fr;gap:14px;padding:18px}}
.panel{{border:1px solid #e5e5e5;padding:14px;background:#fff}}
h3{{margin:0 0 10px;border-bottom:2px solid #DA282D;padding-bottom:6px}}
table{{width:100%;border-collapse:collapse;font-size:13px}}td{{border-bottom:1px solid #eee;padding:7px}}
.note{{padding:18px;color:#555;font-size:13px;border-top:1px solid #eee}}
</style></head><body><div class='wrap'>
<div class='head'><h1>Analisis de Acierto Operativo en Llamadas</h1><p>Preevaluador local sin uso de APIs pagadas.</p></div>
<div class='cards'>
<div class='card'>Total llamadas<b>{int(res.get("Total llamadas", 0)):,}</b></div>
<div class='card'>Score promedio<b>{res.get("Score promedio", 0)}</b></div>
<div class='card'>% acierto<b>{res.get("Porcentaje acierto", 0)}%</b></div>
<div class='card'>No aciertos<b>{int(res.get("No acierto", 0)):,}</b></div>
<div class='card'>Sent. negativo<b>{int(res.get("Sentimiento negativo", 0)):,}</b></div>
</div>
<div class='grid'>
{html_table("Top 10 motivos raiz", top_motivos)}
{html_table("Top 10 temas con mas no acierto", top_temas)}
{html_table("Top 10 operadores menor score", top_ops)}
{html_table("Llamadas criticas", criticas)}
</div>
<div class='note'>Este reporte utiliza un preevaluador automatico basado en reglas locales, diccionarios de negocio y sentimiento estimado. No sustituye la validacion del supervisor.</div>
</div></body></html>"""
        path.write_text(content, encoding="utf-8")
        logger.info("HTML generado: %s", path)
        return path
    except Exception as exc:
        logger.warning("No se pudo generar HTML final: %s", exc)
        return None


def print_config(config: Config, logger: logging.Logger) -> None:
    logger.info("Parametros aplicados:")
    for key, value in config.__dict__.items():
        logger.info("- %s: %s", key, value)


def main() -> int:
    logger = setup_logger()
    start = time.time()
    pyflow_progress(1)
    logger.info("=" * 90)
    logger.info("INICIO ANALISIS LOCAL DE ACIERTO OPERATIVO")
    logger.info("=" * 90)
    try:
        config = load_config()
        print_config(config, logger)
        libs = cargar_librerias_opcionales(config, logger)
        catalogo = preparar_catalogo_temas(config, libs)

        if config.origen_datos == "HANA":
            raw_df = leer_llamadas_hana(config, logger)
        else:
            raw_df = leer_llamadas_archivo(config, logger)
        pyflow_progress(8)
        df = mapear_columnas(raw_df)
        df["CONVERSATION_ID"] = df["CONVERSATION_ID"].astype(str).str.strip()
        before = len(df)
        df = df.drop_duplicates(subset=["CONVERSATION_ID"], keep="last")
        if before != len(df):
            logger.info("Duplicados removidos por CONVERSATION_ID: %s", before - len(df))
        if config.limite_filas > 0:
            df = df.head(config.limite_filas)
        total = len(df)
        logger.info("Llamadas cargadas para evaluar: %s", total)
        if total == 0:
            logger.warning("No hay llamadas para procesar.")
            return 0

        existing = cargar_checkpoint(config.output_dir, logger)
        done_ids = {str(r.get("conversation_id", "")).strip() for r in existing if r.get("conversation_id")}
        rows = [] if config.reanalizar_existentes else list(existing)
        if config.reanalizar_existentes:
            done_ids.clear()

        for idx, (_, pd_row) in enumerate(df.iterrows(), start=1):
            row = {key: ("" if value is None else value) for key, value in pd_row.to_dict().items()}
            cid = str(row.get("CONVERSATION_ID", "") or "").strip()
            if cid and cid in done_ids:
                continue
            result = evaluar_llamada(row, catalogo, libs, config)
            rows.append(result)
            if cid:
                done_ids.add(cid)
            if idx % 1000 == 0:
                logger.info("Avance evaluacion: %s/%s | acumulado=%s", idx, total, len(rows))
            if len(rows) % config.batch_size == 0:
                guardar_checkpoint(rows, config.output_dir, logger)
            pyflow_progress(8 + int((idx / total) * 84))

        guardar_checkpoint(rows, config.output_dir, logger)
        pyflow_progress(94)
        excel = generar_excel(rows, config.output_dir, logger)
        html_path = generar_html(rows, config.output_dir, logger)
        pyflow_progress(100)
        logger.info("=" * 90)
        logger.info("RESUMEN FINAL")
        logger.info("Excel: %s", excel)
        logger.info("HTML: %s", html_path or "No generado")
        logger.info("Registros evaluados: %s", len(rows))
        logger.info("Duracion total: %.2f segundos", time.time() - start)
        logger.info("=" * 90)
        return 0
    except Exception as exc:
        logger.exception("Error general: %s", exc)
        pyflow_progress(100)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
